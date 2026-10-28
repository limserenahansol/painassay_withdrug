"""make_pi2_summary.py - 10-28 (user: 'if I need to show only 2 slides to PI to explain this mini1p pain assay data
(FINAL_20261025_3_mixmarks_mixwin\\ALL_FULL_EN.pptx) could you make this very compact 2 slides'): two summary slides, new figures
drawn from the same data as that deck (analysis\\_variants\\mix4 = mix marks, window per event), numbers re-computed here.

  slide 1  behaviour + overall activity : pin flinch per session; activity of the 185 tracked neurons (no stimulus vs stimulus blocks)
  slide 2  stimulus responses           : population dAUC (after - same-day baseline session) for 4 stimuli; pin UP % of the 185 neurons

  python make_pi2_summary.py -> OneDrive\\Desktop\\mini1p_SBI553_PI_2slides.pptx (+ PNG previews in the same folder name)
"""
import json
import os

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = open(os.path.join(HERE, "make_pi5_deck.py"), encoding="utf-8").read().split("\n")


def find(prefix, start=1):
    for i in range(start - 1, len(SRC)):
        if SRC[i].startswith(prefix):
            return i + 1
    raise KeyError(prefix)


G = {"__name__": "make_pi2_summary", "__file__": os.path.join(HERE, "make_pi5_deck.py")}
exec(compile("\n".join(SRC[:find("prs = Presentation()") - 1]), "make_pi5_deck.py", "exec"), G)
exec(compile("\n".join(SRC[find("SEGT = ") - 1:find("def fig_w2") - 1]), "make_pi5_deck.py", "exec"), G)
globals().update({k: v for k, v in G.items() if not k.startswith("__")})

from scipy.stats import binomtest  # noqa: E402

OUTD = r"C:\Users\hsollim\OneDrive\Desktop\mini1p_SBI553_PI_2slides"
os.makedirs(OUTD, exist_ok=True)
plt.rcParams.update({"font.size": 12, "axes.titlesize": 13, "xtick.labelsize": 11.5, "ytick.labelsize": 11.5})
BN = json.load(open(os.path.join(FIG, "s3a_auc_basenorm.json")))
NOTE = {}


def mlines(a, xs, vals_by_mouse, lw=1.0):
    for m_ in MICE:
        ys = [v.get(m_, np.nan) for v in vals_by_mouse]
        a.plot(xs, ys, color=MC[m_], lw=lw, alpha=0.8, zorder=3)
        for x, y in zip(xs, ys):
            a.scatter(x, y, s=34, marker=MK[m_], color=MC[m_], edgecolor="white", lw=0.6, zorder=5)


def brk(a, x1, x2, y, txt, color="k", h=None, fs=11.5):
    h = h if h is not None else 0.02 * (a.get_ylim()[1] - a.get_ylim()[0])
    a.plot([x1, x1, x2, x2], [y, y + h, y + h, y], color=color, lw=1.2, clip_on=False)
    a.text((x1 + x2) / 2, y + h, txt, ha="center", va="bottom", color=color, fontsize=fs if txt != "ns" else fs - 0.5,
           fontweight="bold" if txt != "ns" else "normal", clip_on=False)


def xy_in(fig, a, x, y, pic_x, pic_y, pic_w):
    """data point -> slide inches (picture placed at pic_x, pic_y with width pic_w)"""
    px, py = a.transData.transform((x, y))
    W, H = fig.get_size_inches()
    s = pic_w / W
    return pic_x + px / fig.dpi * s, pic_y + (H - py / fig.dpi) * s


def legends(fig):
    fig.legend(handles=[Patch(color=CC[c], label=CL[c]) for c in ("vb", "vi", "sb", "si")], ncol=4, loc="upper center",
               bbox_to_anchor=(0.5, 1.0), frameon=False, fontsize=11.5, handlelength=1.2, columnspacing=1.5)
    fig.legend(handles=[Line2D([], [], ls="", marker=MK[m_], color=MC[m_], markeredgecolor="white", markersize=8,
                               label=f"{m_} ({SEX[m_]}, SBI-553 Day {SBIDAY[m_]})") for m_ in MICE], ncol=4, loc="upper center",
               bbox_to_anchor=(0.5, 0.935), frameon=False, fontsize=11, handlelength=1.0, columnspacing=1.5)


# ============================================================================================ figure 1
def fig1(fn):
    fig, (a, b) = plt.subplots(1, 2, figsize=(12.3, 4.35), gridspec_kw=dict(width_ratios=[1, 1.25], wspace=0.30, left=0.075, right=0.99,
                                                                             top=0.80, bottom=0.22))
    # A pin flinch
    pm = {c: beh_pct("flinch", "Pin prick", c)[0] for c in ("vb", "vi", "sb", "si")}
    mu = {c: np.mean(list(pm[c].values())) for c in pm}
    for k, c in enumerate(("vb", "vi", "sb", "si")):
        a.bar(k + (0.4 if k >= 2 else 0), mu[c], color=CC[c], width=0.72, zorder=2)
    X = [0, 1, 2.4, 3.4]
    mlines(a, X[:2], [pm["vb"], pm["vi"]]); mlines(a, X[2:], [pm["sb"], pm["si"]])
    a.set_ylim(0, 128)
    pv, ps, pa = cmh("flinch", "Pin prick", "vehicle vs its baseline"), cmh("flinch", "Pin prick", "SBI-553 vs its baseline"), \
        cmh("flinch", "Pin prick", "SBI-553 vs vehicle")
    brk(a, 0, 1, 104, st(pv), CC["vi"]); brk(a, 2.4, 3.4, 104, st(ps), CC["si"])
    brk(a, 1, 3.4, 116, f"{st(pa)} ({pf(pa)})")
    a.set_xticks(X, ["baseline", "after\nvehicle", "baseline", "after\nSBI-553"])
    a.text(0.5, -0.20, "vehicle day", transform=a.get_xaxis_transform(), ha="center", va="top", color=CC["vi"], fontsize=11.5)
    a.text(2.9, -0.20, "SBI-553 day", transform=a.get_xaxis_transform(), ha="center", va="top", color=CC["si"], fontsize=11.5)
    a.set_ylabel("pin pricks followed by\na flinch (% of deliveries)")
    a.set_title("A  Behaviour: pin prick → flinch", loc="left")
    NOTE["flinch"] = dict(mean={c: round(mu[c], 1) for c in mu}, mice={c: {m_: round(v, 1) for m_, v in pm[c].items()} for c in pm},
                          p_veh=pv, p_sbi=ps, p_after=pa)
    # B activity of the 185 tracked neurons
    D = tracked_change()
    R = {}
    for i, sg in enumerate(("no stimulus", "blocks")):
        x0 = i * 2.6
        for k, dr in enumerate(("veh", "sbi")):
            v = D[f"{dr}_{sg}"].dropna()
            mu_, se_ = v.mean(), v.std(ddof=1) / np.sqrt(len(v))
            b.bar(x0 + k, mu_, yerr=se_, color=CC["vi" if dr == "veh" else "si"], width=0.75, capsize=3, error_kw=dict(lw=1), zorder=2)
            mm = {m_: g[f"{dr}_{sg}"].mean() for m_, g in D.groupby("mouse")}
            for m_ in MICE:
                b.scatter(x0 + k + OFF[m_] * 0.8, mm[m_], s=30, marker=MK[m_], color=MC[m_], edgecolor="white", lw=0.6, zorder=5)
            R[(sg, dr)] = dict(mean=mu_, mm=mm, p0=wilcoxon(v)[1])
        R[(sg, "pair")] = wilcoxon(D[f"sbi_{sg}"], D[f"veh_{sg}"])[1]
        R[(sg, "lower")] = sum(R[(sg, "sbi")]["mm"][m_] < R[(sg, "veh")]["mm"][m_] for m_ in MICE)
    b.axhline(0, color="0.4", lw=0.8)
    allv = [R[(sg, dr)]["mean"] for sg in ("no stimulus", "blocks") for dr in ("veh", "sbi")] + \
        [v for sg in ("no stimulus", "blocks") for dr in ("veh", "sbi") for v in R[(sg, dr)]["mm"].values()]
    top_, bot_ = max(allv + [0]), min(allv + [0])
    sp_ = top_ - bot_
    b.set_ylim(bot_ - 0.08 * sp_, top_ + 0.62 * sp_)
    for i, sg in enumerate(("no stimulus", "blocks")):
        p = R[(sg, "pair")]
        brk(b, i * 2.6, i * 2.6 + 1, top_ + 0.10 * sp_, f"{st(p)}" + (f" ({pf(p)})" if p < 0.05 else ""), h=0.025 * sp_)
        b.text(i * 2.6 + 0.5, top_ + 0.30 * sp_, f"lower on the SBI-553 day\nin {R[(sg, 'lower')]}/4 mice", ha="center", va="bottom", fontsize=10.5,
               color="0.3")
    b.set_xticks([0, 1, 2.6, 3.6], ["vehicle\nday", "SBI-553\nday", "vehicle\nday", "SBI-553\nday"])
    for lab, c in zip(b.get_xticklabels(), ("vi", "si", "vi", "si")):
        lab.set_color(CC[c])
    b.text(0.5, -0.20, "no stimulus (15-180 s)", transform=b.get_xaxis_transform(), ha="center", va="top", fontsize=11.5)
    b.text(3.1, -0.20, "during the 4 stimulus blocks", transform=b.get_xaxis_transform(), ha="center", va="top", fontsize=11.5)
    b.set_ylabel("activity change, same neuron\n(mean z: after − baseline session)")
    b.set_title(f"B  Activity of the same {len(D)} tracked neurons", loc="left")
    NOTE["activity"] = {f"{sg}_{dr}": round(R[(sg, dr)]["mean"], 3) for sg in ("no stimulus", "blocks") for dr in ("veh", "sbi")}
    NOTE["activity"].update({f"p_{sg}": R[(sg, "pair")] for sg in ("no stimulus", "blocks")},
                            lower={sg: R[(sg, "lower")] for sg in ("no stimulus", "blocks")}, n=len(D))
    legends(fig)
    fig.canvas.draw()
    arr = dict(A=xy_in(fig, a, 3.4, mu["si"], 0.5, 1.55, 12.3), B=xy_in(fig, b, 3.6, R[("blocks", "sbi")]["mean"], 0.5, 1.55, 12.3))
    fig.savefig(fn); plt.close(fig)
    return arr


# ============================================================================================ figure 2
def fig2(fn):
    fig, (a, b) = plt.subplots(1, 2, figsize=(12.3, 4.35), gridspec_kw=dict(width_ratios=[1.55, 1], wspace=0.28, left=0.075, right=0.99,
                                                                             top=0.80, bottom=0.22))
    vals = []
    for i, ev in enumerate(STIM):
        r = BN[ev]
        for k, dr in enumerate(("veh", "sbi")):
            x = i * 2.6 + k
            a.bar(x, r[dr + "_mean"], yerr=r[dr + "_sem"], color=CC["vi" if dr == "veh" else "si"], width=0.75, capsize=3,
                  error_kw=dict(lw=1), zorder=2)
            for m_ in MICE:
                a.scatter(x + OFF[m_] * 0.8, r["mouse_" + dr][m_], s=28, marker=MK[m_], color=MC[m_], edgecolor="white", lw=0.6, zorder=5)
            vals += [r[dr + "_mean"] + r[dr + "_sem"], r[dr + "_mean"] - r[dr + "_sem"]] + list(r["mouse_" + dr].values())
    hi, lo = max(vals), min(vals)
    span = hi - lo
    a.set_ylim(lo - 0.08 * span, hi + 0.42 * span)
    a.axhline(0, color="0.4", lw=0.8)
    for i, ev in enumerate(STIM):
        p = BN[ev]["p_between"]
        brk(a, i * 2.6, i * 2.6 + 1, hi + 0.10 * span, f"{st(p)} ({pf(p)})", h=0.03 * span, fs=11)
    a.set_xticks([i * 2.6 + 0.5 for i in range(4)], [f"{ev}\n({BN[ev]['w'][1]:g}-{BN[ev]['w'][2]:g} s)" for ev in STIM])
    a.set_ylabel("response change, same neuron\n(ΔAUC, z·s: after − baseline session)")
    a.set_title("A  Stimulus responses, all neurons (216 / 246)", loc="left")
    # B pin UP, tracked
    t = TRK[TRK.event == "Pin prick"]
    pct = {c: {m_: 100 * g[f"cls_{c}"].eq("activated").mean() for m_, g in t.groupby("mouse")} for c in ("vb", "vi", "sb", "si")}
    tot = {c: 100 * t[f"cls_{c}"].eq("activated").mean() for c in pct}
    cnt = {c: {m_: int(g[f"cls_{c}"].eq("activated").sum()) for m_, g in t.groupby("mouse")} for c in pct}
    X = [0, 1, 2.4, 3.4]
    for x, c in zip(X, ("vb", "vi", "sb", "si")):
        b.bar(x, tot[c], color=CC[c], width=0.72, zorder=2)
    mlines(b, X[:2], [pct["vb"], pct["vi"]]); mlines(b, X[2:], [pct["sb"], pct["si"]])

    def mcn(c0, c1):
        u0, u1 = t[f"cls_{c0}"].eq("activated"), t[f"cls_{c1}"].eq("activated")
        b01, b10 = int((~u0 & u1).sum()), int((u0 & ~u1).sum())
        return binomtest(b01, b01 + b10).pvalue if b01 + b10 else 1.0
    pv, ps, pa = mcn("vb", "vi"), mcn("sb", "si"), mcn("vi", "si")
    b.set_ylim(0, 58)
    brk(b, 0, 1, 46, st(pv), CC["vi"]); brk(b, 2.4, 3.4, 46, st(ps), CC["si"]); brk(b, 1, 3.4, 52, f"{st(pa)} ({pf(pa)})")
    b.set_xticks(X, ["baseline", "after\nvehicle", "baseline", "after\nSBI-553"])
    b.text(0.5, -0.20, "vehicle day", transform=b.get_xaxis_transform(), ha="center", va="top", color=CC["vi"], fontsize=11.5)
    b.text(2.9, -0.20, "SBI-553 day", transform=b.get_xaxis_transform(), ha="center", va="top", color=CC["si"], fontsize=11.5)
    b.set_ylabel("pin-prick UP neurons\n(% of each mouse's tracked neurons)")
    b.set_title(f"B  Pin-prick UP neurons, same {len(t)} neurons", loc="left")
    NOTE["pinUP"] = dict(total={c: round(v, 1) for c, v in tot.items()}, counts=cnt, p_veh=pv, p_sbi=ps, p_after=pa,
                         n_mouse={m_: int(len(g)) for m_, g in t.groupby("mouse")})
    legends(fig)
    fig.canvas.draw()
    arr = dict(A=xy_in(fig, a, 2 * 2.6 + 0.5, hi + 0.13 * span, 0.5, 1.55, 12.3),
               B=xy_in(fig, b, 1, pct["vi"]["5939"], 0.5, 1.55, 12.3))
    fig.savefig(fn); plt.close(fig)
    return arr


F1, F2 = os.path.join(OUTD, "fig1_behaviour_activity.png"), os.path.join(OUTD, "fig2_responses.png")
A1, A2 = fig1(F1), fig2(F2)
json.dump(NOTE, open(os.path.join(OUTD, "numbers.json"), "w"), indent=1, default=float)
print(json.dumps(NOTE, indent=1, default=lambda x: round(float(x), 4)))
print(A1, A2)

# ============================================================================================ slides
from lxml import etree  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402

RED = (0xC0, 0x39, 0x2B)


def tb(s, x, y, w, h, text, size, bold=False, color=(0x22, 0x22, 0x22), italic=False, align=None):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = "Arial"; r.font.color.rgb = RGBColor(*color)
        if align:
            p.alignment = align
    return t


def arrow(s, x0, y0, x1, y1):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = RGBColor(*RED); c.line.width = Pt(3.5)
    te = etree.SubElement(c.line._get_or_add_ln(), "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")


FOOT = "CeA-Ntsr1 mini1p  ·  SBI-553 vs vehicle crossover  ·  n = 4 mice  ·  exploratory, uncorrected  ·  * p < 0.05, ** p < 0.01, *** p < 0.001"
fl, ac, pu = NOTE["flinch"], NOTE["activity"], NOTE["pinUP"]
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

s = prs.slides.add_slide(prs.slide_layouts[6])
tb(s, 0.45, 0.12, 12.4, 0.3, "CeA-Ntsr1 mini1p  ·  SBI-553 vs vehicle  ·  summary 1 / 2: behaviour and overall activity", 12, color=(0x55, 0x55, 0x55))
tb(s, 0.45, 0.36, 12.4, 0.6, "SBI-553 lowered pin flinches and neuronal activity during stimuli", 26)
tb(s, 0.45, 0.98, 12.4, 0.45, f"After SBI-553 vs vehicle: pin flinch {fl['mean']['si']:.0f} vs {fl['mean']['vi']:.0f}%; activity during stimuli "
   f"{ac['blocks_sbi']:.2f} vs {ac['blocks_veh']:.2f} z (lower in {ac['lower']['blocks']}/4 mice)", 15, bold=True, color=RED)
s.shapes.add_picture(F1, Inches(0.5), Inches(1.55), Inches(12.3))
tb(s, 0.45, 5.95, 12.4, 0.75,
   "Design: 4 mice, crossover; each day = baseline session → vehicle or SBI-553 → same session again (light, mild, pin prick, heat).\n"
   "A: bars = mean of 4 mice, lines = mice; CMH over deliveries. B: 185 neurons tracked on both days; paired Wilcoxon.", 12)
tb(s, 0.45, 7.1, 11.6, 0.3, FOOT, 9.5, color=(0x88, 0x88, 0x88), italic=True)
tb(s, 12.2, 7.1, 0.9, 0.3, "1 / 2", 10, color=(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT)
(xa, ya), (xb, yb) = A1["A"], A1["B"]
arrow(s, xa + 0.55, ya - 0.45, xa + 0.22, ya - 0.08)
arrow(s, xb + 0.60, yb + 0.35, xb + 0.25, yb + 0.06)
s.notes_slide.notes_text_frame.text = (
    "Numbers (make_pi2_summary.py, data analysis\\_variants\\mix4 = the ALL_FULL mixmarks deck): pin flinch % (mean of mice) "
    f"{fl['mean']}; per mouse {fl['mice']}; CMH vehicle day p = {fl['p_veh']:.4f}, SBI-553 day p = {fl['p_sbi']:.4f}, after vs after "
    f"p = {fl['p_after']:.4f}. Pin flinch after SBI-553 was lower in 3 mice and equal in 5939. Tracked neurons (n = {ac['n']}): change "
    f"no stimulus vehicle {ac['no stimulus_veh']:+.3f} / SBI {ac['no stimulus_sbi']:+.3f} z (paired p = {ac['p_no stimulus']:.3f}); "
    f"stimulus blocks {ac['blocks_veh']:+.3f} / {ac['blocks_sbi']:+.3f} z (paired p = {ac['p_blocks']:.2g}), lower on the SBI-553 day "
    f"in {ac['lower']['blocks']}/4 mice (both drug orders, so not just day order). SBI-553 dose / route / timing: not in the analysis files.")

s = prs.slides.add_slide(prs.slide_layouts[6])
tb(s, 0.45, 0.12, 12.4, 0.3, "CeA-Ntsr1 mini1p  ·  SBI-553 vs vehicle  ·  summary 2 / 2: stimulus responses", 12, color=(0x55, 0x55, 0x55))
tb(s, 0.45, 0.36, 12.4, 0.6, "Stimulus-evoked responses: no firm SBI-553 effect yet", 26)
tb(s, 0.45, 0.98, 12.4, 0.45, "No stimulus response changed differently on the two days; the pin UP rise after vehicle is mostly one mouse",
   15, bold=True, color=RED)
s.shapes.add_picture(F2, Inches(0.5), Inches(1.55), Inches(12.3))
tb(s, 0.45, 5.95, 12.4, 0.5,
   "A: each neuron's AUC after the injection − its AUC in the same day's baseline session (z − 0.5 s before), mean ± SEM, markers = mice; "
   "bracket = Mann-Whitney. B: % UP (shift test) of the 185 tracked neurons, lines = mice; McNemar (same neurons).", 12)
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(6.50), Inches(12.4), Inches(0.52))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xFE, 0xF9, 0xE7); box.line.color.rgb = RGBColor(0xF3, 0x9C, 0x12); box.shadow.inherit = False
tb(s, 0.6, 6.53, 12.1, 0.46, "Limits: 4 mice; every mouse had more pin UP neurons after its Day-2 injection (day order not ruled out); "
   "neurons pooled, uncorrected. Same result with baseline-z and ΔF/F baselines.", 12.5, color=(0x33, 0x33, 0x33))
tb(s, 0.45, 7.1, 11.6, 0.3, FOOT, 9.5, color=(0x88, 0x88, 0x88), italic=True)
tb(s, 12.2, 7.1, 0.9, 0.3, "2 / 2", 10, color=(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT)
(xa, ya), (xb, yb) = A2["A"], A2["B"]
arrow(s, xa + 0.15, ya - 0.50, xa + 0.02, ya - 0.10)
arrow(s, xb - 0.80, yb + 0.30, xb - 0.12, yb + 0.03)
s.notes_slide.notes_text_frame.text = (
    "A (all neurons, change from each day's baseline session): " + "; ".join(
        f"{ev}: vehicle day {BN[ev]['veh_mean']:+.4f}, SBI-553 day {BN[ev]['sbi_mean']:+.4f} z·s, Mann-Whitney p = {BN[ev]['p_between']:.2f}"
        for ev in STIM) + ". Without the day's baseline session, only pin differed after vs after (p = 0.041), because the SBI-day pin "
    f"response was already ~0 at baseline. B (tracked): pin UP % {pu['total']}; per mouse counts {pu['counts']} of {pu['n_mouse']}; "
    f"McNemar vehicle day p = {pu['p_veh']:.4f}, SBI-553 day p = {pu['p_sbi']:.2f}, after vs after p = {pu['p_after']:.2g}. 5939 alone: "
    "2 -> 19 after vehicle. Checks: the same conclusions with baseline z (-2..0 s), dF/F - mean -2..0 s and a -5..-3 s gap "
    "(analysis\\baseline_methods).")
OUT = os.path.join(OUTD, "mini1p_SBI553_PI_2slides.pptx")
prs.save(OUT)
print("->", OUT)
