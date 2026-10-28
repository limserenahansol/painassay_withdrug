"""apply_v3.py - 10-26: the user's edited deck (compactversio2) -> compactversion3: simpler, checked take-home lines (slides 1-24 titles /
conclusions; mini1p claims fixed against the data), arrows / boxes re-pointed, normalized-AUC slide inserted after the population AUC
slide, legacy labels restyled, page numbers redone. The user's file is not touched; every replaced text is kept in the slide notes."""
import os
import shutil
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deckops import A, E, RED, add_arrow, add_text, geom, get, import_slide, remove, reorder, set_arrow, set_text  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
USER = r"C:\Users\hsollim\OneDrive\Desktop\1111112121212_FINAL_compactversio2.pptx"
SRC = os.path.join(HERE, "v2_user.pptx")
OUT = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\hsollim\OneDrive\Desktop\1111112121212_FINAL_compactversion3.pptx"
NEW9B = os.path.join(HERE, "new9b.pptx")       # copy of ...compact_added_slides\01_NEW_after_slide36_... (that file is open in PowerPoint)
if os.path.getmtime(USER) > os.path.getmtime(SRC) + 1:          # the user saved again: start from the newest file
    shutil.copy2(USER, SRC)
    print("re-copied the user's latest file")

prs = Presentation(SRC)
S = {i: s for i, s in enumerate(prs.slides, 1)}
FOOT = "CeA-Ntsr1 mini1p  ·  SBI-553 vs vehicle crossover  ·  n = 4 mice  ·  exploratory, uncorrected (details in the notes)"
DARK = RGBColor(0x22, 0x22, 0x22)
LOG = []


def keep(s, label, old):
    """only-add rule: the replaced on-slide text goes to the notes"""
    tf = s.notes_slide.notes_text_frame
    tf.text = f"[10-26 edit] previous {label}: {old}\n" + tf.text


def T(si, name, new, near=None, contains=None, **kw):
    s = S[si]
    sh = get(s, name, near=near, contains=contains)
    old = sh.text_frame.text
    if old.strip() == new.strip():
        return sh
    keep(s, name, old.replace("\n", " / "))
    set_text(sh, new, **kw)
    LOG.append(f"slide {si}: {name}: {new.splitlines()[0][:90]}")
    return sh


def replace_in_runs(si, name, old, new):
    sh = get(S[si], name)
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if old in r.text:
                keep(S[si], name, sh.text_frame.text)
                r.text = r.text.replace(old, new)
                LOG.append(f"slide {si}: {name}: ...{new[:80]}")
                return
    raise KeyError(f"slide {si} {name}: {old!r} not found")


def restyle(sh, x, y, w, h, lines):
    """rebuild a text box: lines = [(text, pt, bold, RGBColor)] with explicit Arial runs"""
    geom(sh, x, y, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, (text, pt, bold, col) in enumerate(lines):
        p = p0 if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = text
        r.font.size = Pt(pt); r.font.bold = bold; r.font.name = "Arial"; r.font.color.rgb = col
        p.space_after = Pt(6)


def divider(si, lines):
    s = S[si]
    sh = get(s, "TextBox 1")
    keep(s, "TextBox 1", sh.text_frame.text.replace("\n", " / "))
    sty = [(lines[0], 34, True, DARK), ("", 12, False, DARK)] + [(t, 20, True, RED) for t in lines[1:]]
    restyle(sh, 0.9, 1.4, 11.6, 5.0, sty)
    LOG.append(f"slide {si}: divider restyled: {lines[0]}")


def footer(s, k, n=None, foot=True):
    """set the page number box (any 'k / 14', 'F119', 'F167' box) or add one; add the standard footer if missing"""
    num, fbox = None, None
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if sh.top is not None and sh.top / E > 6.95 and (t.replace(" ", "").replace("/", "").isdigit() or t in ("F119", "F167")
                                                        or (t.endswith("/ 14") and len(t) < 9)):
            num = sh
        if t.startswith("CeA-Ntsr1 mini1p  ·") or t.startswith("Exploratory: neurons pooled"):
            fbox = sh
    lab = f"{k} / {n}"
    if num is None:
        add_text(s, 12.2, 7.1, 0.9, 0.3, lab, 10, color=(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT)
    else:
        set_text(num, lab, size=10, color=RGBColor(0x88, 0x88, 0x88))
        geom(num, 12.2, 7.1, 0.9, 0.3)
    if foot:
        if fbox is None:
            add_text(s, 0.45, 7.1, 11.3, 0.3, FOOT, 9, color=(0x88, 0x88, 0x88), italic=True)
        elif not fbox.text_frame.text.startswith("CeA-Ntsr1"):
            keep(s, "footer", fbox.text_frame.text)
            set_text(fbox, FOOT, size=9, color=RGBColor(0x88, 0x88, 0x88))
            geom(fbox, 0.45, 7.1, 11.3, 0.3)


def red_arrow_style(sh):
    sh.line.color.rgb = RED; sh.line.width = Pt(3.5)
    ln = sh.line._get_or_add_ln()
    for te in ln.findall(A + "tailEnd"):
        ln.remove(te)
    te = etree.SubElement(ln, A + "tailEnd"); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")


def conn(si, near):
    return get(S[si], next(sh.name for sh in S[si].shapes if sh.shape_type == 9 and abs(sh.left / E - near[0]) < 0.03
                           and abs(sh.top / E - near[1]) < 0.03), near=near)


# ====================================================================================== slides 1-24: other projects
T(2, "Title 1", "TRAP whole-brain mapping")
T(2, "Content Placeholder 2", "Which brain regions have more TRAPed cells in Active than in Passive mice?")
T(3, "TextBox 1", "How ORBm and BMAp were picked: Active > Passive in all 4 phases")
T(4, "TextBox 1", "Post, counted as in Smith et al. 2024: ~2x more cells in Active")
T(4, "TextBox 5", "Active ~2x higher at Post, but significant only with cells as n (mice as n: p = 0.33)")
T(5, "TextBox 1", "Pennington 2026 test, mice as n: ORBm and BMAp ~1.4x higher in Active")
T(5, "TextBox 3", "•  Negative binomial per region, 20 mice as n, phase as covariate\n"
                  "•  ORBm 1.43× (p = 0.026), BMAp 1.35× (p = 0.031) higher in Active\n"
                  "•  Not significant after FDR within each cluster: q = 0.16 (ORBm), 0.071 (BMAp)")
T(6, "Title 1", "TRAP conclusion: ORBm and BMAp are the top candidates")
T(6, "Content Placeholder 2", "ORBm and BMAp came out on top at every step: k-means (no labels), group filters, per-mouse check\n"
                              "With mice as n (negative binomial): both ~1.4x higher in Active (p ≈ 0.03)\n"
                              "Not significant after FDR (q = 0.16, 0.071), so still exploratory\n"
                              "Next: add the Xenium Post mice to increase n")
T(7, "Title 1", "Next: TRAP x Xenium in Post mice")
T(7, "Content Placeholder 2", "2 Active + 2 Passive Post mice will get TRAP x Xenium anyway\n"
                              "Their TRAP counts could be added to the Post comparison to increase n\n"
                              "Panel: 297 mouse genes + tdTomato + Ai14 STOP (custom) = 299 targets")
T(8, "TextBox 4", "One panel finds all 20 target cell types (~0.90 recall at Xenium depth) and also reads the morphine state")
T(9, "TextBox 4", "All 32,285 genes scored; the 299 kept separate all 94 same-region type pairs and find all 20 types (~0.90 recall)")
T(11, "TextBox 1", "Addiction behaviour modelling")
T(12, "TextBox 1", "OPIOID SEEKING · BEHAVIOUR")
T(12, "TextBox 4", "Active mice work harder for morphine; on morphine days reward rate predicts pauses less well")
T(13, "TextBox 4", "After a morphine reward, Active mice keep seeking but Passive mice stop")
sh = T(14, "TextBox 1", "Conclusion: addiction behaviour\n\nActive mice work harder for morphine than Passive mice\n"
                        "On morphine days, reward rate predicts pauses less well\n"
                        "Right after a morphine reward, both groups stop more than after water\n"
                        "Then Active mice keep seeking, while Passive mice keep stopping")
geom(sh, w=11.0, h=4.6)
replace_in_runs(16, "TextBox 19", " identified neurons respond to pain (4 up, 2 down, 2 both); none of the 9 checked by HCR has Vipr1 or Ntsr1.",
                " matched neurons respond to pain (4 up, 2 down, 2 both); none of the 9 HCR-checked cells has Vipr1, and the Ntsr1 "
                "probe has not worked yet.")
T(17, "TextBox 7", "Same six WT mice. SBI-553 given 10 min before the Day 2 assay.")
T(18, "TextBox 1", "Fewer pain behaviours on the SBI-553 day (same 6 mice)")
T(18, "TextBox 4", "On the SBI-553 day, escape/rearing dropped at all 4 stimuli and withdrawal at 3 of 4 (paired Wilcoxon p < 0.05, n = 6 mice)")
S[18].notes_slide.notes_text_frame.text += "\n[10-26] Caution: SBI-553 was always given on Day 2, so a day / habituation effect cannot be ruled out."
sh = T(19, "TextBox 1", "Conclusion: pain\n\nTRUFACT works in CeA through a small GRIN lens: 13 imaged cells found again in the slice\n"
                        "One mini2P mouse: pin prick lowered overall CeA-Ntsr1 activity (p = 0.006)\n"
                        "WT mice (n = 6): fewer pain behaviours on the SBI-553 day than on the vehicle day")
geom(sh, w=11.0, h=4.4)
T(20, "TextBox 1", "Face poke excites somatosensory cortex most; more inhibition posteriorly")
T(20, "TextBox 3", "Face poke excited somatosensory cortex the most (SSs 2.01× the cortex-wide rate).\n"
                   "Inhibited cells were more common in the posterior half (+7.2 pp, 18 of 18 recordings, p = 7.6e-06).\n"
                   "n = 4 mice, 18 recordings, 43,156 cells. Areas: binomial test vs cortex-wide rate, BH-FDR. Halves: Wilcoxon signed-rank.")
T(21, "TextBox 1", "Melittin drives medial visual / parietal and motor cortex most")
T(21, "TextBox 3", "Melittin drove medial visual / parietal areas the most (VISam 1.74×), plus motor cortex (MOp 1.42×).\n"
                   "'Up' and 'down' maps are mirror images (ρ = −0.98), so they are one map, not two.\n"
                   "n = 4 mice, 81,834 cells (melittin only). Areas: binomial test vs cortex-wide rate, BH-FDR.")
T(22, "TextBox 1", "Face poke and melittin activate different cortical areas")
T(22, "TextBox 3", "Face poke peaks in somatosensory cortex (SSs 2.01×); melittin peaks in medial visual / parietal areas (VISam 1.74×).\n"
                   "The two area maps are not correlated (Spearman ρ = −0.22, p = 0.36); only motor cortex (MOp, MOs) is high in both.\n"
                   "Face poke: 4 mice, 43,156 cells; melittin: 4 other mice, 81,834 cells. Each area ÷ its own cortex-wide rate.")
sh = T(23, "TextBox 1", "Conclusion: face poke vs melittin\n\nFace poke and melittin activate different cortical areas\n"
                        "Face poke: somatosensory cortex (SSs)\nMelittin: medial visual / parietal areas (VISam, VISpm)\n"
                        "Both: motor cortex (MOp, MOs)")
geom(sh, w=11.0, h=4.2)
T(24, "TextBox 1", "Melittin + morphine: response slightly larger, not smaller (2 mice)")

# ====================================================================================== mini1p (25-46)
T(25, "TextBox 3", "After vehicle, pin flinches and pin UP neurons rose; after SBI-553 they did not")
T(25, "TextBox 7", "design, colour code,\none example session")
T(25, "TextBox 22", "population AUC\n→ UP / DOWN neurons\n→ pin-prick close-up\n→ PCA, overlap")
T(25, "TextBox 27", "summary, three messages\n+ limits")
T(25, "TextBox 28", "4 mice, crossover: each day = a baseline session, then an injection session. SBI-553 on Day 1 in 5941 (F) and 5939 (M); "
                    "on Day 2 in 5937 (M) and 1191 (F). Every comparison is made within a day (vs that day's baseline, same neurons) and "
                    "between days (after vehicle vs after SBI-553).")
# 27 design slide: stray note and empty box out, bottom line in the deck style
s = S[27]
sh = get(s, "TextBox 7"); keep(s, "TextBox 7", sh.text_frame.text); remove(sh)
remove(get(s, "TextBox 5"))
sh = get(s, "TextBox 6"); set_text(sh, sh.text_frame.text, size=11); geom(sh, 0.45, 6.55, 12.4, 0.45)
T(28, "TextBox 3", "Shade = 0-0.5 s after each pin prick (the pin-prick response window)")
T(29, "TextBox 2", "Pin prick: flinches rose after vehicle, not after SBI-553")
sh = T(29, "TextBox 13", "no rise on the SBI-553 day"); geom(sh, w=3.0)
sh = T(30, "TextBox 13", "lower after SBI-553"); geom(sh, 7.89, 2.79, 2.3, 0.31)
set_arrow(conn(30, (5.44, 2.94)), 7.95, 3.02, 5.60, 3.48)          # long arrow re-routed under the pin vehicle-day bracket
divider(31, ["3. Overall neuronal activity", "Activity dropped after both injections, more after SBI-553",
             "Same 185 neurons, stimulus blocks: bigger drop on the SBI-553 day (−0.12 vs −0.03 z)"])
T(32, "TextBox 3", "Overall, all neuronal activity dropped more after SBI-553 (after vs after p = 0.012; baselines equal, p = 0.64)")
T(33, "TextBox 2", "Same 185 neurons, whole session: stimulus blocks lower after SBI-553")
geom(get(S[33], "Rectangle 8"), 0.26, 5.08, 2.2, 1.0)
T(34, "TextBox 2", "Same 185 tracked neurons: activity fell more on the SBI-553 day")
T(34, "TextBox 3", "Stimulus blocks: activity dropped more on the SBI-553 day (−0.12 vs −0.03 z, p < 0.001, 4/4 mice); heat the reverse", size=14)
divider(35, ["4. Stimulus responses", "Population AUC first, then single UP / DOWN neurons",
             "Touch responders rose on the SBI-553 day", "Pin-prick UP neurons rose after vehicle, not after SBI-553"])
s = S[36]
T(36, "TextBox 3", "After vs after only pin differed (p = 0.041), but SBI-day responses were already weak at baseline")
remove(get(s, "Camera 10"))                                        # PowerPoint Cameo live-camera circle (covered the heat panel)
remove(conn(36, (8.57, 2.15)))                                     # long unlabelled arrow onto the pin trace
geom(get(s, "Rectangle 12"), 8.62, 2.70, 1.10, 0.92)               # frame the SBI-day baseline + after bars (weak already at baseline)
divider(37, ["From the population to single neurons", "AUC: only pin differed after SBI-553 vs after vehicle (p = 0.041)",
             "But the baselines were already different; normalized to baseline, no change (p = 0.40)",
             "So we looked at responsive (UP / DOWN) neurons: the same 185 tracked neurons, then all neurons"])
s = S[38]
sh = T(38, "TextBox 2", "Example: UP / DOWN = response beats 1000 time-shifted event trains"); geom(sh, 0.45, 0.32, 12.4, 0.6)
geom(get(s, "TextBox 3"), 0.45, 0.92, 12.4, 0.45)
geom(get(s, "TextBox 5"), 0.45, 6.55, 12.4, 0.42)
s = S[39]
T(39, "TextBox 2", "Same 185 neurons: fewer pin-prick UP neurons after SBI-553 (9% vs 23%)", size=24)
T(39, "TextBox 3", "Pin UP rose after vehicle (13 → 23%), not after SBI-553 (11 → 9%); light-touch UP rose only after SBI-553")
sh = T(39, "TextBox 12", "fewer pin UP (16 vs 43); taller trace = only the strongest pass")
geom(sh, 7.0, 4.03, 6.25, 0.32)
remove(conn(39, (8.45, 2.72)))                                     # arrow onto the UP-trace peak (selection-biased size)
set_arrow(conn(39, (3.82, 1.32)), 4.05, 2.20, 3.80, 2.80)          # short arrow onto the light-touch after-SBI bar
T(40, "TextBox 2", "Responder numbers changed; response size did not")
T(40, "TextBox 3", "Touch responders rose on the SBI-553 day; pin-prick responders rose only on the vehicle day")
s = S[40]
geom(get(s, "Rectangle 6"), 0.55, 4.23, 6.33, 1.82)
geom(get(s, "Rectangle 10"), 7.00, 4.23, 2.85, 1.82)
set_arrow(get(s, "Straight Arrow Connector 8"), 2.75, 5.95, 2.45, 5.40)
set_arrow(get(s, "Straight Arrow Connector 11"), 8.75, 4.30, 8.05, 4.66)
s = S[41]
T(41, "TextBox 1", "4  STIMULUS RESPONSES  ·  pin prick: UP, DOWN and the rest (all neurons)")
T(41, "TextBox 2", "Pin prick: fewer UP neurons after SBI-553 (24 vs 47)")
T(41, "TextBox 3", "Their average looks bigger only because fewer neurons pass the test; neurons UP at baseline changed alike (p = 0.11)", size=15)
remove(get(s, "Rectangle 14"))
remove(conn(41, (7.73, 1.29)))
pic = get(s, "Picture 4"); pic.crop_bottom = 0.0; geom(pic, 0.52, 1.50, 12.30, 4.75)
s = S[42]
add_text(s, 0.45, 0.92, 12.4, 0.45, "After SBI-553, most baseline pin UP neurons went non-responsive (71% vs 28% after vehicle, p = 0.002)",
         15, bold=True, color=(0xC0, 0x39, 0x2B))
set_arrow(conn(42, (7.90, 3.25)), 9.30, 3.10, 8.02, 3.95)
s = S[43]
sh = get(s, "TextBox 1"); keep(s, "TextBox 1", sh.text_frame.text.replace("\n", " / "))
restyle(sh, 0.9, 1.2, 11.6, 5.5, [("Summary: stimulus responses", 34, True, DARK), ("", 12, False, DARK),
                                   ("Pin prick: UP neurons rose after vehicle, not after SBI-553 (same 185 neurons: 23 vs 9%)", 20, True, RED),
                                   ("After SBI-553, most baseline pin UP neurons went non-responsive (71% vs 28%)", 20, True, RED),
                                   ("Touch responders rose on the SBI-553 day", 20, True, RED),
                                   ("Response size per neuron did not change; only the number of responders did", 20, True, RED)])
LOG.append("slide 43: summary rewritten")
s = S[44]
T(44, "TextBox 1", "4  STIMULUS RESPONSES  ·  population PCA (same 185 neurons)", size=11, color=RGBColor(0x55, 0x55, 0x55))
geom(T(44, "TextBox 2", "PCA, same 185 neurons: all stimuli raise PC1; no clear drug pattern", size=24, color=DARK), h=0.6)
sh = T(44, "TextBox 3", "Pin-prick PC1 was lowest after SBI-553; touch and heat changed in both directions (descriptive, no test)",
       size=15, bold=True, color=RED)
geom(sh, 0.45, 0.92, 12.4, 0.45)
sh = T(44, "TextBox 5", "PCA of trial-mean z, 185 tracked neurons (4 mice), 4 sessions x 4 stimuli; PC1 response 0-0.5 s. Descriptive, no test; "
                        "split-half r = 0.33; 5939 carries 54% of the PC1 weights.", size=11, color=DARK)
geom(sh, 0.45, 6.64, 12.4, 0.42)
remove(get(s, "Straight Arrow Connector 8"))                       # heat 'rise' only vs the outlier negative SBI-day baseline
s = S[45]
T(45, "TextBox 1", "4  STIMULUS RESPONSES  ·  touch / noxious overlap (all neurons)")
T(45, "TextBox 2", "Touch and noxious neurons overlap more than chance, after both injections")
sh = T(45, "TextBox 3", "After vehicle: more pain-only, fewer touch-only neurons; after SBI-553: more touch-only neurons\n"
                        "Neurons responding to both: above chance in every session; same after vehicle and after SBI-553")
geom(sh, 0.45, 0.92, 12.4, 0.75)
geom(get(s, "Picture 4"), 0.80, 1.72, 11.73, 4.93)
add_arrow(s, 3.95, 4.05, 3.40, 4.55)                               # after vehicle: the noxious-only (pink) part grew
add_arrow(s, 6.45, 4.55, 5.82, 5.05)                               # after SBI-553: the touch-only (light blue) part grew
for nm in ("TextBox 5", "TextBox 6"):                              # parked below the slide edge: text to the notes
    sh = get(s, nm); keep(s, nm, sh.text_frame.text.replace("\n", " / ")); remove(sh)
add_text(s, 0.45, 6.68, 12.4, 0.4, "Activated neurons, all neurons (vehicle day 216, SBI-553 day 246), 4 mice. Brackets test 'both': McNemar "
         "within day, Fisher after vs after; dashed = expected if independent (observed above it in all 8 bars, hypergeometric p ≤ 0.003).", 10.5)
S[45].notes_slide.notes_text_frame.text += (
    "\n[10-26, recomputed from cells_classes.csv, not drawn on the figure] touch-only neurons: vehicle day 38 -> 4 (McNemar p < 0.001), "
    "SBI-553 day 9 -> 37 (p < 0.001), after vs after 2 vs 15% (Fisher p < 0.001). Noxious-only: vehicle day 22 -> 41 (p < 0.001), "
    "SBI-553 day 25 -> 20 (ns), after vs after 19 vs 8% (p < 0.001). 'Both' (touch & noxious): vehicle 18 -> 21 ns, SBI-553 8 -> 25 "
    "(p = 0.0015), after vs after 10 vs 10% (p = 1.0). Stimulus & reflex 'both' rose after both injections (16 -> 49, 26 -> 56, p < 0.001), "
    "after vs after p = 1.0.")
s = S[46]
T(46, "TextBox 3", "After vehicle, pin flinches and pin UP neurons rose; after SBI-553 they did not")
T(46, "TextBox 6", "Fewer pin flinches and light-touch movements after SBI-553 than after vehicle")
T(46, "TextBox 7", "79 vs 57% and 61 vs 39% (p < 0.01); partly a rise after vehicle")
T(46, "TextBox 10", "Activity fell after both injections, more on the SBI-553 day")
T(46, "TextBox 11", "Same 185 neurons: −0.12 vs −0.03 z (p < 0.001), 4/4 mice")
T(46, "TextBox 14", "Pin UP neurons rose after vehicle, not after SBI-553")
T(46, "TextBox 15", "Same 185 neurons: 23 vs 9% UP (p < 0.001); touch responders rose on the SBI-553 day")
T(46, "TextBox 17", "Limits: n = 4 mice, p values pool neurons or deliveries (exploratory, uncorrected). Drug order = day order: in all 4 mice the "
                    "Day-2 injection session had more pin UP neurons. Population AUC did not differ once each day's baseline was subtracted. "
                    "Pin UP results lean on 5939 and 5941.")
# hidden backups: same corrections
T(52, "TextBox 3", "Fewer pin UP neurons after SBI-553 (24 vs 47); their bigger average is selection: neurons UP at baseline changed alike (p = 0.11)")
T(53, "TextBox 3", "Light / mild DOWN rose on the SBI-553 day; pin DOWN rose only on the vehicle day; heat DOWN rose on both days")
T(54, "TextBox 2", "Population trajectories: every stimulus moves CeA activity along one axis")

# ---------------------------------------------------- arrows: one style (3.5 pt, deck red, triangle head); highlight boxes easier to see
for si in range(25, 47):
    for sh in S[si].shapes:
        if sh.shape_type == 9:
            red_arrow_style(sh)
        elif sh.shape_type == 1 and sh.name.startswith("Rectangle"):
            for al in sh._element.spPr.xpath(".//a:solidFill/a:srgbClr[@val='FFFF00']/a:alpha"):
                al.set("val", "25000")

# ---------------------------------------------------- the normalized-AUC slide after the population AUC slide
nb = import_slide(prs, NEW9B)
for sh in nb.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("Change from each day's baseline session"):
        set_text(sh, "Normalized to each day's baseline: vehicle and SBI-553 days do not differ")
    elif sh.has_text_frame and sh.text_frame.text.startswith("Pin after-vs-after gap"):
        set_text(sh, "The pin gap (p = 0.041) is gone once each neuron's own baseline session is subtracted (p = 0.40)")
LOG.append("NEW slide after 36: population AUC normalized to each day's baseline session")

# ---------------------------------------------------- order + page numbers
order = [S[i] for i in range(1, 37)] + [nb] + [S[37], S[38], S[39], S[40], S[41], S[42], S[44], S[45], S[43], S[46]] + \
        [S[i] for i in range(47, 58)]
reorder(prs, order)
main = [S[25], S[26], S[27], S[28], S[29], S[30], S[31], S[32], S[33], S[34], S[35], S[36], nb, S[37], S[38], S[39], S[40], S[41],
        S[42], S[44], S[45], S[43], S[46]]
for k, s in enumerate(main, 1):
    footer(s, k, len(main), foot=s not in (S[31], S[35], S[37], S[43]))
prs.save(OUT)
open(os.path.join(HERE, "apply_v3_log.txt"), "w", encoding="utf-8").write("\n".join(LOG))
print(len(LOG), "edits ->", OUT, len(prs.slides), "slides")
