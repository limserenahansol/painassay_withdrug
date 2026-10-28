"""deckops.py - small python-pptx helpers to edit the user's deck in place (formatting kept)"""
import copy
import io

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches, Pt

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
E = 914400
RED = RGBColor(0xC0, 0x39, 0x2B)


def shapes_named(slide, name):
    return [sh for sh in slide.shapes if sh.name == name]


def get(slide, name, near=None, contains=None):
    """the shape called name; if several, the one whose text contains `contains` or whose (left, top) is nearest `near` (inches)"""
    c = shapes_named(slide, name)
    if contains is not None:
        c = [sh for sh in c if sh.has_text_frame and contains in sh.text_frame.text]
    if not c:
        raise KeyError(f"{name!r} (contains={contains!r}) not on slide")
    if len(c) > 1 and near is not None:
        c.sort(key=lambda sh: (sh.left / E - near[0]) ** 2 + (sh.top / E - near[1]) ** 2)
    elif len(c) > 1 and contains is None:
        raise KeyError(f"{name!r} is ambiguous ({len(c)} shapes); pass near= or contains=")
    return c[0]


def set_text(sh, text, size=None, bold=None, color=None):
    """replace the text, keeping the first paragraph's and first run's formatting; '\\n' -> new paragraphs with the same formatting"""
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:                                  # first paragraph empty: move the first run found anywhere to the front
        src = next((r for p in tf.paragraphs for r in p.runs), None)
        if src is not None:
            end = p0._p.find(A + "endParaRPr")
            if end is not None:
                end.addprevious(copy.deepcopy(src._r))
            else:
                p0._p.append(copy.deepcopy(src._r))
        else:
            p0.add_run()
    r0 = p0.runs[0]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs)[1:]:
        r._r.getparent().remove(r._r)
    for br in p0._p.findall(A + "br"):
        p0._p.remove(br)
    lines = text.split("\n")
    r0.text = lines[0]
    for ln in lines[1:]:
        np_ = copy.deepcopy(p0._p)
        p0._p.getparent().append(np_)
        rs = np_.findall(A + "r")
        rs[0].find(A + "t").text = ln
    for p in tf.paragraphs:
        for r in p.runs:
            if size:
                r.font.size = Pt(size)
            if bold is not None:
                r.font.bold = bold
            if color is not None:
                r.font.color.rgb = color


def geom(sh, x=None, y=None, w=None, h=None):
    if x is not None:
        sh.left = Inches(x)
    if y is not None:
        sh.top = Inches(y)
    if w is not None:
        sh.width = Inches(w)
    if h is not None:
        sh.height = Inches(h)


def set_arrow(sh, x0, y0, x1, y1):
    """connector from (x0, y0) to (x1, y1); the arrow head (tailEnd) stays at the end point"""
    sh.begin_x, sh.begin_y, sh.end_x, sh.end_y = Inches(x0), Inches(y0), Inches(x1), Inches(y1)


def add_arrow(slide, x0, y0, x1, y1, color=RED, width=3.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = color; c.line.width = Pt(width)
    te = etree.SubElement(c.line._get_or_add_ln(), A + "tailEnd")
    te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def add_text(slide, x, y, w, h, text, size, bold=False, color=(0x22, 0x22, 0x22), italic=False, align=None):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = "Arial"
        r.font.color.rgb = RGBColor(*color) if isinstance(color, tuple) else color
        if align:
            p.alignment = align
    return t


def remove(sh):
    sh._element.getparent().remove(sh._element)


def import_slide(prs, path, idx=0, layout=None):
    """append a copy of slide idx of another deck (pictures + other shapes + notes) to prs; returns it"""
    src = Presentation(path).slides[idx]
    lay = layout or next((l for l in prs.slide_layouts if l.name.lower() == "blank"), prs.slide_layouts[-1])
    s = prs.slides.add_slide(lay)
    for ph in list(s.placeholders):
        remove(ph)
    for sh in src.shapes:
        if sh.shape_type == 13:
            s.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
        else:
            s.shapes._spTree.insert_element_before(copy.deepcopy(sh._element), "p:extLst")
    s.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text if src.has_notes_slide else ""
    return s


def reorder(prs, order):
    """order = list of slide objects in the wanted order (must contain every slide exactly once)"""
    lst = prs.slides._sldIdLst
    els = list(lst)
    idx = [list(prs.slides).index(s) for s in order]
    assert sorted(idx) == list(range(len(els))), "order must contain every slide once"
    for el in els:
        lst.remove(el)
    for i in idx:
        lst.append(els[i])


def add_note(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = text + ("\n\n" + tf.text if tf.text.strip() else "")
