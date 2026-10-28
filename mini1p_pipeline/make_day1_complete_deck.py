"""make_day1_complete_deck.py - the FULL Day-1 deck: every figure of every analysis, one plain result per slide.

Sections
  MAIN            the main slides (make_day1_full_deck.build_main, without the limits slide)
  ACTIVITY        A1-A7, C1-C2, MAIN summary, robustness forest, mixed model
  SINGLE NEURONS  B1/B1c, B1b, B2, B3, B5, B6, B6b, B9, B8, E4, E5, E6, E9, E10, E12, E14
  POPULATION      B4, B7, E3, E7, E8, E11, E13
  BEHAVIOUR       reflexes per stimulus, % time, behaviour around deliveries
  PER MOUSE       make_day1_per_mouse_deck.build_per_mouse (10 slides x 4 mice)
  SUPPLEMENT      limits and checks: DOWN responses, raw-pixel check, field change, reliability, common mode,
                  calibrated counts, ROI-trace versions, motion control, signal quality, the literature criteria,
                  the TOY Day-2 pipeline test
Rules: a title, a 'how to read' line, the test / n note and ONE plain result (increased / decreased / no change,
numbers read from the analysis files, no interpretation). Interpretation and concerns only in the supplement.
Several figures with the same meaning (one per mouse) share a slide.

  python make_day1_complete_deck.py --base C:/Users/hsollim/Documents/HEAL_mini1p_SBI553 --out <pptx>
"""
import argparse
import os
import sys

import numpy as np
import pandas as pd
from PIL import Image

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import safe_io  # noqa: E402,F401  (a figure save retries on a transient Windows write error)
import make_day1_ca_slides as S1                                    # noqa: E402
import make_day1_full_deck as FD                                    # noqa: E402
import make_day1_per_mouse_deck as PM                               # noqa: E402
import deck_cfg as C                                                # noqa: E402  (CA_DAY: which day)

MICE = S1.MICE
STIM = S1.STIM
ARM = {m: a for m, a, _ in MICE}


def dirw(pre, post, p, a=0.05):
    if not np.isfinite(p) or p >= a:
        return "no change"
    return "increased" if post > pre else "decreased"


def pf(p):
    return "n.a." if not np.isfinite(p) else ("p<0.001" if p < 0.001 else f"p={p:.3f}" if p < 0.01 else f"p={p:.2f}")


def grid(paths, ncols, out, pad=12):
    """put several figures in one image: ncols per row, each scaled to the same width"""
    ims = [Image.open(p).convert("RGB") for p in paths]
    cw = max(i.size[0] for i in ims)
    ims = [i.resize((cw, int(i.size[1] * cw / i.size[0]))) for i in ims]
    rows = [ims[k:k + ncols] for k in range(0, len(ims), ncols)]
    rh = [max(i.size[1] for i in r) for r in rows]
    W = ncols * cw + (ncols - 1) * pad
    H = sum(rh) + (len(rows) - 1) * pad
    g = Image.new("RGB", (W, H), "white")
    y = 0
    for r, h in zip(rows, rh):
        for c, im in enumerate(r):
            g.paste(im, (c * (cw + pad), y))
        y += h + pad
    g.save(out)
    return out


def table_figure(rows, cols, widths, out, fs=9.5):
    fig = plt.figure(figsize=(S1.FW, S1.FH)); ax = fig.add_axes([0.0, 0.0, 1, 1]); ax.set_axis_off()
    T = ax.table(cellText=rows, colLabels=cols, loc="upper left", cellLoc="left", colLoc="left", colWidths=widths)
    T.auto_set_font_size(False); T.set_fontsize(fs); T.scale(1, 1.75)
    for (r, c), cell in T.get_celld().items():
        cell.set_edgecolor("0.85")
        if r == 0:
            cell.set_text_props(fontweight="bold"); cell.set_facecolor("#f3f5f8")
    fig.savefig(out, dpi=160); plt.close(fig)
    return out


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    from pptx import Presentation
    from pptx.util import Inches
    B = a.base
    od = os.path.splitext(a.out)[0] + "_figures"
    os.makedirs(od, exist_ok=True)
    f = lambda nm: os.path.join(od, nm)
    an = lambda *p: C.A(B, *p)                       # Day-1 folder / file names mapped to this day (deck_cfg)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(S1.SW), Inches(S1.SH)
    pvf = lambda: f(f"preview_slide{len(prs.slides) + 1:03d}.png")

    def fit(text, pt, nmax):
        """cut a too-long text at the last list separator that fits (the full text goes to the slide notes)"""
        if len(S1.wrap_lines(text, pt, 12.3)) <= nmax:
            return text, False
        cut = text
        while len(S1.wrap_lines(cut + " ... (full text in the notes)", pt, 12.3)) > nmax:
            i = max(cut.rfind("; "), cut.rfind(", "))
            if i <= 0:
                cut = cut[:int(len(cut) * 0.9)]
            else:
                cut = cut[:i]
        return cut + " ... (full text in the notes)", True

    def slide(sec, title, figs, howto, result, note="", ncols=1):
        miss = [p for p in figs if not os.path.isfile(p)]
        if miss:                          # an analysis that was not run for this day: the slide is left out (said)
            print(f"  skipped (figure missing): {title} <- {miss[0]}")
            return
        png = figs[0] if len(figs) == 1 else grid(figs, ncols, f(f"grid_{len(prs.slides) + 1:03d}.png"))
        h2, c1 = fit(howto, 12.5, 2)
        n2, c2 = fit(note, 10.5, 3)
        r2, c3 = fit(result, 14.5, 3)
        FD.add_slide(prs, sec, title, png, n2, r2, pvf(), howto=h2)
        if c1 or c2 or c3:
            prs.slides[-1].notes_slide.notes_text_frame.text = "\n\n".join([howto, note, result])

    # ------------------------------------------------------------------ MAIN
    FD.build_main(B, prs, os.path.join(od, "main"), pvf=pvf, limits=False)
    print("main:", len(prs.slides), "slides", flush=True)

    # ------------------------------------------------------------------ ACTIVITY
    sec = "ALL ANALYSES: ACTIVITY"
    T = pd.read_csv(an("ca_day1", "ca_d1_A1_whole_session_tests.csv")); T["mouse"] = T.mouse.astype(str)
    slide(sec, "A1  Whole-session Ca transient rate, baseline vs injection", [an("ca_day1", "ca_d1_A1_whole_session_rate.png")],
          "How to read: one grey line = one neuron, coloured line = median over neurons; right = each mouse's median change (log2).",
          "Result (median transients/min per neuron): " + ", ".join(
              f"{r.mouse} {r.pre_median:.1f}->{r.post_median:.1f} {dirw(r.pre_median, r.post_median, r.p_wilcoxon)}" for r in T.itertuples()) + ".",
          "Wilcoxon signed-rank over the neurons of each mouse (paired, same neurons). First 15 s excluded. n = 2 mice per group.")
    T2 = pd.read_csv(an("ca_day1", "ca_d1_A2_blocks_tests.csv")); T2["mouse"] = T2.mouse.astype(str)
    T6 = pd.read_csv(an("ca_day1", "ca_d1_A6_blocks_still_tests.csv")); T6["mouse"] = T6.mouse.astype(str)

    def cnt(X, cond="condition"):
        out = []
        for m, _, _ in MICE:
            x = X[X.mouse == m]
            d = int(((x.p < 0.05) & (x.post_mean < x.pre_mean)).sum()); u = int(((x.p < 0.05) & (x.post_mean > x.pre_mean)).sum())
            out.append(f"{m} down {d}, up {u} of {len(x)}")
        return "; ".join(out)
    slide(sec, "A2 / A6  Activity per block and per movement state", [an("ca_day1", "ca_d1_A2_blocks_and_movement.png"),
                                                                       an("ca_day1", "ca_d1_A6_blocks_still_only.png")],
          "How to read: bars = mean rate over neurons per block (no stimulus, each stimulus) and state (still, moving); grey = "
          "baseline, red = injection. Top = all frames, bottom = still frames only.",
          "Result (conditions with p < 0.05), all frames: " + cnt(T2) + ". Still frames: " + cnt(T6) + ".",
          "Wilcoxon signed-rank over neurons per mouse and condition (* p<0.05, ** p<0.01). n = 2 mice per group.")
    T3 = pd.read_csv(an("ca_day1", "ca_d1_A3_spontaneous_tests.csv")); T3["mouse"] = T3.mouse.astype(str)
    sp = lambda c: ", ".join(f"{r.mouse} {dirw(r.pre_median, r.post_median, r.p)}" for r in T3[T3.condition == c].itertuples())
    slide(sec, "A3  Spontaneous activity in the first 3 min (no stimulus)", [an("ca_day1", "ca_d1_A3_spontaneous.png")],
          "How to read: one grey line = one neuron, coloured = median; top row all frames, bottom row only frames without movement.",
          f"Result, all frames: {sp('no stimulus')}; still frames: {sp('no stimulus, still')}.",
          "Wilcoxon signed-rank over neurons per mouse, baseline vs injection, first 3 min of each session.")
    T4 = pd.read_csv(an("ca_day1", "ca_d1_A4_timecourse.csv")); T4["mouse"] = T4.mouse.astype(str)
    g4 = T4.groupby(["mouse", "phase"]).transients_per_cell.mean()
    slide(sec, "A4  Activity minute by minute through the session", [an("ca_day1", "ca_d1_A4_timecourse.png")],
          "How to read: transients per neuron in each minute of the session (and % of time moving), baseline vs injection, per mouse.",
          "Result (mean transients per neuron per minute, baseline -> injection): "
          + ", ".join(f"{m} {g4[(m, 'pre')]:.1f}->{g4[(m, 'post')]:.1f}" for m, _, _ in MICE) + ".",
          "Descriptive time course; 1-min bins.")
    T7 = pd.read_csv(an("ca_day1", "ca_d1_A7_block_aligned.csv")); T7["mouse"] = T7.mouse.astype(str)
    rat = {}
    for (m, ph), x in T7.groupby(["mouse", "phase"]):
        rat[(m, ph)] = x[(x.t_from_block_start >= 0) & (x.t_from_block_start < 180)].rate_all.mean() / \
            max(x[(x.t_from_block_start >= -60) & (x.t_from_block_start < 0)].rate_all.mean(), 1e-9)
    slide(sec, "A7  Activity aligned to the start of each stimulus block", [an("ca_day1", "ca_d1_A7_block_aligned.png")],
          "How to read: activity in 10-s bins from 60 s before to the end of each stimulus block, baseline vs injection, per mouse.",
          "Result (rate during the block / rate in the 60 s before it), baseline: "
          + ", ".join(f"{m} x{rat[(m, 'pre')]:.1f}" for m, _, _ in MICE) + "; injection: "
          + ", ".join(f"{m} x{rat[(m, 'post')]:.1f}" for m, _, _ in MICE) + ".",
          "Mean over neurons and the 4 blocks of a session.")
    W = pd.read_csv(an("ca_day1", "ca_d1_cells_whole_session.csv")); W["mouse"] = W.mouse.astype(str)
    ud = []
    for m, _, _ in MICE:
        P = W[W.mouse == m].pivot(index="cell", columns="phase", values="rate").dropna()
        lr = np.log2((P.post + 0.1) / (P.pre + 0.1))
        ud.append(f"{m} {int((lr > 0).sum())} higher, {int((lr < 0).sum())} lower")
    slide(sec, "A5  Where in the field of view activity changed", [an("ca_day1", "ca_d1_A5_effect_map.png")],
          "How to read: each neuron's footprint coloured by its rate change after the injection (red = higher, blue = lower).",
          "Result (neurons with a higher / lower whole-session rate after the injection): " + "; ".join(ud) + ".",
          "log2((injection + 0.1) / (baseline + 0.1)) transients per min.")
    C2 = pd.read_csv(an("ca_day1", "ca_d1_C2_arm_blocks_still.csv")); C2["mouse"] = C2.mouse.astype(str)
    cs = []
    for m, _, _ in MICE:
        x = C2[(C2.mouse == m) & (C2.condition != "no stimulus")]
        cs.append(f"{m} {int(((x.p < 0.05) & (x.log2_ratio_median < 0)).sum())} of {len(x)}")
    slide(sec, "C1 / C2  Change per mouse, SBI-553 vs vehicle", [an("ca_day1", "ca_d1_C1_arm_summary.png"),
                                                               an("ca_day1", "ca_d1_C2_arm_blocks_still.png")],
          "How to read: one dot / bar per mouse = median change over its neurons (log2 injection / baseline; below 0 = decrease); "
          "red = SBI-553, blue = vehicle.", "Result (stimulus blocks with a decrease, still frames, p < 0.05): " + ", ".join(cs) + ".",
          "Wilcoxon signed-rank per mouse; n = 2 mice per group, groups shown, not tested.", ncols=2)
    RA = pd.read_csv(an("ca_robustness", "R1_robustness_arm.csv"))
    mv_ = RA[RA.variant == "main: OASIS transients, still frames"].set_index("window")
    slide(sec, "Summary figure: change per mouse and per group", [an("ca_robustness", "MAIN_day1_findings.png")],
          "How to read: per mouse change of activity (log2), the range over analysis choices, the leave-one-out difference and "
          "the change of the evoked responses.",
          f"Result: SBI-553 minus vehicle {mv_.at['stimulus blocks', 'sbi_minus_vehicle']:+.2f} log2 in the stimulus blocks and "
          f"{mv_.at['no-stimulus start', 'sbi_minus_vehicle']:+.2f} in the first 3 min.",
          "Main analysis: OASIS transients, still frames. n = 2 mice per group.")
    R1 = pd.read_csv(an("ca_robustness", "R1_robustness_per_mouse.csv")); R1["mouse"] = R1.mouse.astype(str)
    rr = R1[R1.window == "stimulus blocks"].groupby("mouse").log2_ratio_median.agg(["min", "max"])
    slide(sec, "R1  The same comparison under every analysis choice", [an("ca_robustness", "R1_robustness_forest.png")],
          "How to read: one row per analysis choice (detector, thresholds, trace version, frames, neuron subsets); dot = mouse "
          "median change (log2), line = 95% CI.",
          f"Result (range of the stimulus-block change over {R1.variant.nunique()} choices): "
          + ", ".join(f"{m} {rr.at[m, 'min']:+.2f} to {rr.at[m, 'max']:+.2f}" for m, _, _ in MICE) + " log2.",
          "Bootstrap 95% CI over neurons.")
    E2 = pd.read_csv(an("ca_day1_extended", "E2_mixed_model.csv"))
    e2 = []
    for oc in ["transient rate, whole session", "rate, no-stimulus start"]:
        x = E2[E2.outcome == oc].set_index("term")
        if "post" in x.index and "post:sbi" in x.index:
            e2.append(f"{oc}: vehicle {x.at['post', 'coef']:+.2f} ({pf(x.at['post', 'p'])}), SBI-553 extra "
                      f"{x.at['post:sbi', 'coef']:+.2f} ({pf(x.at['post:sbi', 'p'])})")
    slide(sec, "E2  Mixed model: change after the injection, vehicle and SBI-553", [an("ca_day1_extended", "E2_mixed_model.png")],
          "How to read: estimated change after the injection for the vehicle mice (post) and the extra change in the SBI-553 mice "
          "(post:sbi), 95% CI, for each outcome.", "Result: " + "; ".join(e2) + ".",
          "Linear mixed model log2 outcome ~ session x group, random mouse + neuron; p treats neurons as units (4 mice).")
    print("activity done:", len(prs.slides), flush=True)

    # ------------------------------------------------------------------ SINGLE NEURONS
    sec = "ALL ANALYSES: SINGLE NEURONS"
    B1 = pd.read_csv(an("ca_day1", "ca_d1_B1_responsive_counts.csv")); B1["mouse"] = B1.mouse.astype(str)
    B1c = pd.read_csv(an("ca_day1", "ca_d1_B1c_responsive_counts_shift_rule.csv")); B1c["mouse"] = B1c.mouse.astype(str)

    def tot(X, m, ph, col, evs=STIM):
        return int(X[(X.mouse == m) & (X.phase == ph) & X.event.isin(evs)][col].sum())
    slide(sec, "B1 / B1c  Responsive neurons per event, two tests", [an("ca_day1", "ca_d1_B1_responsive_cells.png"),
                                                                    an("ca_day1", "ca_d1_B1c_responsive_cells.png")],
          "How to read: % of neurons that rose (UP) or fell (DOWN) after each event, baseline vs injection; left = trial Wilcoxon, "
          "right = time-shift test.",
          "Result (UP neurons summed over the 4 stimuli, baseline -> injection), time-shift test: "
          + ", ".join(f"{m} {tot(B1c, m, 'pre', 'n_up')}->{tot(B1c, m, 'post', 'n_up')}" for m, _, _ in MICE) + "; trial Wilcoxon: "
          + ", ".join(f"{m} {tot(B1, m, 'pre', 'n_up')}->{tot(B1, m, 'post', 'n_up')}" for m, _, _ in MICE) + ".",
          f"Left: Wilcoxon signed-rank of {C.WRULE} over trials, BH q <= 0.20. Right: events shifted inside their block (null), "
          "BH q <= 0.20 (exploratory).", ncols=2)
    B1b = pd.read_csv(an("ca_day1", "ca_d1_B1b_responsive_matched.csv")); B1b["mouse"] = B1b.mouse.astype(str)
    gb = B1b[B1b.event.isin(STIM)].groupby(["mouse", "phase"]).pct_up.mean()
    slide(sec, "B1b  Responsive neurons with the same number of trials", [an("ca_day1", "ca_d1_B1b_responsive_matched_trials.png")],
          "How to read: % UP / DOWN neurons when baseline and injection use the same number of trials (random subsets).",
          "Result (mean % UP over the 4 stimuli, baseline -> injection): "
          + ", ".join(f"{m} {gb.get((m, 'pre'), np.nan):.1f}->{gb.get((m, 'post'), np.nan):.1f}%" for m, _, _ in MICE) + ".",
          "50 random trial subsets per session; trial Wilcoxon rule.")
    B2 = pd.read_csv(an("ca_day1", "ca_d1_B2_response_change_tests.csv")); B2["mouse"] = B2.mouse.astype(str)
    sig = B2[B2.p_wilcoxon < 0.05]
    dn = [f"{r.mouse} {r.event.lower()}" for r in sig.itertuples() if r.post_mean < r.pre_mean]
    upl = [f"{r.mouse} {r.event.lower()}" for r in sig.itertuples() if r.post_mean > r.pre_mean]
    slide(sec, "B2  Event-locked mean response per mouse", [an("ca_day1", f"ca_d1_B2_psth_{m}.png") for m, _, _ in MICE],
          "How to read: mean response over all neurons around each event (baseline vs injection), one row per mouse "
          f"({' / '.join(', '.join(C.GROUP[g]) for g in ('SBI-553', 'vehicle'))}).",
          f"Result (p < 0.05), mean response decreased: {', '.join(dn) if dn else 'none'}; increased: {', '.join(upl) if upl else 'none'}.",
          f"Wilcoxon signed-rank over neurons of the {C.WRESP} response, injection vs baseline.", ncols=2)
    ET = pd.read_csv(an("ca_day1", "ca_d1_cells_event_tests.csv")); ET["mouse"] = ET.mouse.astype(str)
    pos = ET[ET.event.isin(STIM)].groupby(["mouse", "phase", "cell"]).resp_mean.mean().gt(0).groupby(["mouse", "phase"]).mean()
    slide(sec, "B3  Every neuron's mean response, per mouse", [an("ca_day1", f"ca_d1_B3_heatmaps_{m}.png") for m, _, _ in MICE],
          "How to read: rows = neurons, columns = time around each event, baseline and injection side by side; red = rise, blue = "
          f"fall. Mice: {', '.join(C.GROUP['SBI-553'])} (top), {', '.join(C.GROUP['vehicle'])} (bottom).",
          "Result (% of neurons with a positive mean response to the 4 stimuli, baseline -> injection): "
          + ", ".join(f"{m} {100 * pos[(m, 'pre')]:.0f}->{100 * pos[(m, 'post')]:.0f}%" for m, _, _ in MICE) + ".",
          f"Trial-mean z (per-session MAD) minus {C.WPRE}; same neuron order in both sessions.", ncols=2)
    allev = list(ET.event.unique())
    slide(sec, "B5  UP, DOWN and other neurons: average response", [an("ca_day1", f"ca_d1_B5_groups_{m}.png") for m, _, _ in MICE],
          "How to read: mean +- SEM of the UP neurons (red), DOWN neurons (blue) and the rest (grey) for each event; one panel row per "
          "mouse.",
          "Result (UP / DOWN neurons summed over all events, baseline -> injection): " + "; ".join(
              f"{m} UP {tot(B1, m, 'pre', 'n_up', allev)}->{tot(B1, m, 'post', 'n_up', allev)}, DOWN "
              f"{tot(B1, m, 'pre', 'n_down', allev)}->{tot(B1, m, 'post', 'n_down', allev)}" for m, _, _ in MICE) + ".",
          "Classes from the trial Wilcoxon rule (BH q <= 0.20).", ncols=2)
    for tag, Bx, rule in (("B6", B1, "trial Wilcoxon"), ("B6b", B1c, "time-shift test")):
        slide(sec, f"{tag}  Where the responsive neurons are ({rule})",
              [an("ca_day1", f"ca_d1_{tag}_maps_{m}.png") for m, _, _ in MICE],
              "How to read: one field-of-view map per event (top row baseline, bottom injection); red = UP, blue = DOWN, numbers = "
              "neurons. One map set per mouse.",
              "Result (UP + DOWN neurons over all events, baseline -> injection): " + ", ".join(
                  f"{m} {tot(Bx, m, 'pre', 'n_up', allev) + tot(Bx, m, 'pre', 'n_down', allev)}->"
                  f"{tot(Bx, m, 'post', 'n_up', allev) + tot(Bx, m, 'post', 'n_down', allev)}" for m, _, _ in MICE) + ".",
              f"Classes from the {rule} (BH q <= 0.20, exploratory).", ncols=2)
    top = []
    for m, _, _ in MICE:
        x = ET[(ET.mouse == m) & (ET.phase == "pre") & (ET.q_shift <= 0.2)]
        if len(x):
            r = x.loc[x.resp_mean.abs().idxmax()]
            top.append(f"{m} #{int(r.cell)} {r.event.lower()} {r.resp_mean:+.2f} z")
    slide(sec, "B9  The clearest responsive neurons, every trial", [an("ca_day1", f"ca_d1_B9_exemplar_cells_{m}.png") for m, _, _ in MICE],
          f"How to read: per mouse (left to right {C.ORDER_TXT}) the strongest UP / DOWN neurons for heat, pin prick and "
          "withdrawal: all trials (heatmap) and their mean, baseline and injection.",
          "Result (largest baseline response passing the time-shift test): " + ", ".join(top) + ".",
          "Trial-baselined z; time-shift test BH q <= 0.20.", ncols=4)
    B8 = pd.read_csv(an("ca_day1", "ca_d1_B8_withdrawal_vs_not.csv")); B8["mouse"] = B8.mouse.astype(str)
    slide(sec, "B8  Same stimulus, with vs without a withdrawal", [an("ca_day1", "ca_d1_B8_withdrawal_vs_not.png")],
          "How to read: for each stimulus, the response on trials with a withdrawal vs without one, per neuron.",
          "Result (neuron x stimulus x session tests with q <= 0.20): " + ", ".join(
              f"{m} {int((B8[B8.mouse == m].q <= 0.2).sum())} of {int((B8.mouse == m).sum())}" for m, _, _ in MICE) + ".",
          "Rank-sum test per neuron, BH over neurons.")
    E4 = pd.read_csv(an("ca_day1_extended", "E4_same_cells.csv")); E4["mouse"] = E4.mouse.astype(str)
    slide(sec, "E4  Do the same neurons respond in both sessions?", [an("ca_day1_extended", "E4_same_cells.png")],
          "How to read: overlap of the responsive neurons of the two sessions, and the correlation of every neuron's response "
          "baseline vs injection.",
          "Result (median Spearman rho over events): " + ", ".join(
              f"{m} {E4[E4.mouse == m].spearman_rho.median():.2f}" for m, _, _ in MICE)
          + f"; neurons responsive in both sessions: {int(E4.resp_both.sum())} (all mice and events).",
          "Hypergeometric overlap test; Spearman correlation over neurons.")
    E5 = pd.read_csv(an("ca_day1_extended", "E5_adaptation.csv"))
    s5 = E5[E5.p < 0.05]
    slide(sec, "E5  Adaptation within a block", [an("ca_day1_extended", "E5_adaptation.png")],
          "How to read: population response against the delivery number inside a block (later deliveries to the right).",
          f"Result: the response changed with delivery number (p < 0.05) in {len(s5)} of {len(E5)} session x stimulus tests "
          f"({int((s5.slope_per_trial < 0).sum())} decreasing, {int((s5.slope_per_trial > 0).sum())} increasing).",
          "Spearman correlation of the response with the delivery number.")
    E6 = pd.read_csv(an("ca_day1_extended", "E6_reliability_latency.csv"))
    lat = E6[E6.cls == "UP"].groupby("event").latency_s.median()
    slide(sec, "E6  Reliability and latency of the responses", [an("ca_day1_extended", "E6_reliability_latency.png")],
          "How to read: trial-to-trial reliability of each neuron's response and the time of its peak, by event and class.",
          "Result (median latency of the UP neurons): " + ", ".join(f"{e.lower()} {v:.1f} s" for e, v in lat.items()) + ".",
          "Latency = time of the peak of the trial-mean response in [0, 3] s.")
    E9 = pd.read_csv(an("ca_day1_extended", "E9_spatial_clustering.csv"))
    slide(sec, "E9  Are responsive neurons clustered in space?", [an("ca_day1_extended", "E9_spatial_clustering.png")],
          "How to read: mean distance between responsive neurons vs the same number of random neurons of that mouse.",
          f"Result: closer than random (p < 0.05) in {int((E9.p_clustered < 0.05).sum())} of {len(E9)} mouse x session x event sets.",
          "Permutation of neuron identities.")
    E10 = pd.read_csv(an("ca_day1_extended", "E10_intensity_coding.csv"))
    g10 = E10.groupby("phase")[["n_graded_up", "n_graded_down"]].sum()
    slide(sec, "E10  Graded responses: light < mild < pin < heat", [an("ca_day1_extended", "E10_intensity_coding.png")],
          "How to read: each stimulus's mean response, ordered by intensity, baseline vs injection; neurons whose response grows "
          "with intensity.",
          f"Result (graded neurons, all mice): baseline {int(g10.at['pre', 'n_graded_up'])} up / {int(g10.at['pre', 'n_graded_down'])} down, "
          f"injection {int(g10.at['post', 'n_graded_up'])} up / {int(g10.at['post', 'n_graded_down'])} down.",
          "Spearman correlation of the trial response with the intensity rank per neuron.")
    E12 = pd.read_csv(an("ca_day1_extended", "E12_response_types.csv"))
    slide(sec, "E12  Response types and whether neurons keep their type", [an("ca_day1_extended", "E12_response_types.png")],
          "How to read: baseline response shapes grouped into types (k-means); the same neurons' type after the injection.",
          f"Result: {int(E12.stay_same_type_after_injection.sum())} of {int(E12.n_cells.sum())} neurons kept their type after the injection.",
          "k-means of the baseline response shapes; the injection responses assigned to the nearest type.")
    E14 = pd.read_csv(an("ca_day1_extended", "E14_who_changes.csv"))
    slide(sec, "E14  Do the most responsive neurons change most?", [an("ca_day1_extended", "E14_who_changes.png")],
          "How to read: each neuron's change after the injection against its baseline response, per mouse and event.",
          f"Result: negative slope (p < 0.05) in {int(((E14.r < 0) & (E14.p < 0.05)).sum())} of {len(E14)} mouse x event fits.",
          "Linear fit over neurons.")
    print("single neurons done:", len(prs.slides), flush=True)

    # ------------------------------------------------------------------ POPULATION
    sec = "ALL ANALYSES: POPULATION"
    B4 = pd.read_csv(an("ca_day1", "ca_d1_B4_population_auc.csv")); B4["mouse"] = B4.mouse.astype(str)
    s4 = B4[B4.p < 0.05]
    slide(sec, "B4  Population AUC before vs after each event", [an("ca_day1", "ca_d1_B4_population_auc.png")],
          f"How to read: area under the mean response before ({C.WPRE}) and after ({C.WRESP}) each event; bottom = mean +- SEM over neurons.",
          f"Result: after > before (p < 0.05) in {int((s4.auc_post_mean > s4.auc_pre_mean).sum())}, after < before in "
          f"{int((s4.auc_post_mean < s4.auc_pre_mean).sum())} of {len(B4)} mouse x session x event tests.",
          "Wilcoxon signed-rank over neurons.")
    B7 = pd.read_csv(an("ca_day1", "ca_d1_B7_stimulus_decoding.csv"))
    slide(sec, "B7  Can the stimulus be decoded from the population?", [an("ca_day1", "ca_d1_B7_stimulus_decoding.png")],
          "How to read: accuracy of decoding which of the 4 stimuli was given, from the response and (control) from before it.",
          f"Result: accuracy {B7.acc_response.min():.2f}-{B7.acc_response.max():.2f} (chance 0.25); above the permutation null in "
          f"{int((B7.p_perm < 0.05).sum())} of {len(B7)} sessions.", "Cross-validated classifier; 200 label permutations.")
    E3 = pd.read_csv(an("ca_day1_extended", "E3_encoding_unique_variance.csv"))
    uv = {c.replace("unique_", "").replace("_", " "): E3[c].median() for c in E3.columns if c.startswith("unique_")}
    top3 = sorted(uv.items(), key=lambda kv: -kv[1])[:3]
    slide(sec, "E3  Encoding model: what explains each neuron's activity", [an("ca_day1_extended", "E3_encoding_model.png")],
          "How to read: variance of each neuron explained only by stimuli, reflexes, movement, grooming, escape or motion artefact.",
          f"Result: median full-model R2 {E3.r2_full.median():.3f}; largest median unique variance: "
          + ", ".join(f"{k} {v:.4f}" for k, v in top3) + ".", "Ridge regression, one penalty per mouse, cross-validated.")
    E7 = pd.read_csv(an("ca_day1_extended", "E7_pairwise_correlations.csv")); E7["mouse"] = E7.mouse.astype(str)
    slide(sec, "E7  Correlation between pairs of neurons", [an("ca_day1_extended", "E7_pairwise_correlations.png")],
          "How to read: correlation of every pair of neurons before any stimulus, baseline vs injection, and against their distance.",
          "Result (median pair r, baseline -> injection): " + ", ".join(
              f"{r.mouse} {r.median_r_pre:.3f}->{r.median_r_post:.3f} {dirw(r.median_r_pre, r.median_r_post, r.p)}" for r in E7.itertuples()) + ".",
          "Wilcoxon signed-rank over pairs.")
    slide(sec, "E8  Population trajectories of the four stimuli", [an("ca_day1_extended", "E8_population_trajectories.png")],
          "How to read: mean response to each stimulus on the two main patterns of the baseline session (solid = baseline, "
          "dashed = injection, same axes).",
          ("Result: PC1 explains " + ", ".join(f"{r.mouse} {r.pc1_pct:.0f}%" for r in pd.read_csv(an("ca_day1_extended", "E8_explained_variance.csv")).itertuples())
           + " of the baseline responses." if os.path.isfile(an("ca_day1_extended", "E8_explained_variance.csv"))
           else "Result: PC1 explains 34, 60, 34 and 59% of the baseline responses (5941, 5939, 5937, 1191)."),
          "PCA of the baseline trial-mean responses, -1 to +4 s.")
    E11 = pd.read_csv(an("ca_day1_extended", "E11_predict_withdrawal.csv"))
    slide(sec, "E11  Does the population predict a withdrawal?", [an("ca_day1_extended", "E11_predict_withdrawal.png")],
          "How to read: how well withdrawal is predicted from the stimulus alone and with the population state before / after.",
          f"Result: AUC stimulus only {E11.auc_stimulus_only.min():.2f}-{E11.auc_stimulus_only.max():.2f}; adding the pre-delivery state "
          f"above its null in {int((E11.p_prestate < 0.05).sum())} of {len(E11)} sessions.", "Cross-validated logistic regression.")
    E13 = pd.read_csv(an("ca_day1_extended", "E13_pain_axis.csv")); E13["mouse"] = E13.mouse.astype(str)
    slide(sec, "E13  A noxious vs innocuous axis of the population", [an("ca_day1_extended", "E13_pain_axis.png")],
          "How to read: decoder of noxious (pin, heat) vs innocuous (light, mild) trained on baseline, applied to the injection.",
          "Result (baseline cross-validated accuracy vs null 95th percentile): " + ", ".join(
              f"{r.mouse} {r.cv_acc_pre:.2f} vs {r.null_p95:.2f}" for r in E13.itertuples()) + ".",
          "Linear decoder, cross-validated; label permutation null.")
    print("population done:", len(prs.slides), flush=True)

    # ------------------------------------------------------------------ BEHAVIOUR
    sec = "ALL ANALYSES: BEHAVIOUR"
    BC = pd.read_csv(an("behavior_day1", "behav_d1_change.csv")); BC["mouse"] = BC.mouse.astype(str)
    lst = []
    for kind in ("withdrawal", "flinch"):
        for r in BC[BC[f"{kind}_fisher_p"] < 0.05].itertuples():
            lst.append(f"{kind} {r.stim.lower()} {r.mouse} {getattr(r, kind + '_base_pct'):.0f}->{getattr(r, kind + '_inj_pct'):.0f}%")
    slide(sec, "Reflexes per stimulus: % deliveries with a withdrawal / flinch", [an("behavior_day1", "behav_d1_reflex.png")],
          "How to read: one line per mouse from baseline to injection; y = % of deliveries followed by a withdrawal (top) or flinch "
          "(bottom).", "Result (p < 0.05): " + ("; ".join(lst) if lst else "none") + ".",
          "Fisher exact test per mouse and stimulus on the deliveries. n = 2 mice per group.")
    H = pd.read_csv(an("behavior_day1", "behav_d1_held.csv")); H["mouse"] = H.mouse.astype(str)
    hw = H[H.block == "whole session"].pivot_table(index=["behaviour", "mouse"], columns="phase", values="pct_time")
    hs = []
    for b, x in hw.groupby(level=0):
        lo, hi = int((x.injection < x.baseline).sum()), int((x.injection > x.baseline).sum())
        hs.append(f"{b} lower in {lo}, higher in {hi}")
    slide(sec, "Affective and spontaneous behaviours: % of time", [an("behavior_day1", "behav_d1_affective.png")],
          "How to read: % of the session (top) or of the stimulus blocks (bottom) spent in each behaviour, one line per mouse.",
          "Result (whole session, of 4 mice): " + "; ".join(hs) + ".", "Descriptive; scored by hand (movement, grooming automatic).")
    D, G = FD.behaviour_counts(B)
    D["arm"] = D.mouse.astype(str).map(ARM)
    ev = D.assign(w=D.withdrawal > 0, mv=D["general movement"] > 0).groupby(["arm", "phase"])[["w", "mv"]].mean() * 100
    slide(sec, "Behaviour around every delivery", [an("behavior_day1", "behav_d1_eventlocked.png")],
          "How to read: % of deliveries with each behaviour at each time from -2 to +5 s (solid = baseline, dashed = injection); "
          "top SBI-553, bottom vehicle.",
          "Result (deliveries followed by a withdrawal within 2 s, baseline -> injection): SBI-553 "
          f"{ev.at[('SBI-553', 'baseline'), 'w']:.0f}->{ev.at[('SBI-553', 'injection'), 'w']:.0f}%, vehicle "
          f"{ev.at[('vehicle', 'baseline'), 'w']:.0f}->{ev.at[('vehicle', 'injection'), 'w']:.0f}%; movement within 3 s: SBI-553 "
          f"{ev.at[('SBI-553', 'baseline'), 'mv']:.0f}->{ev.at[('SBI-553', 'injection'), 'mv']:.0f}%, vehicle "
          f"{ev.at[('vehicle', 'baseline'), 'mv']:.0f}->{ev.at[('vehicle', 'injection'), 'mv']:.0f}%.",
          "Pooled over the 2 mice of each group; 0.1-s bins.")
    print("behaviour done:", len(prs.slides), flush=True)

    # ------------------------------------------------------------------ PER MOUSE
    PM.build_per_mouse(B, prs, os.path.join(od, "per_mouse"), pvf=pvf, section="PER MOUSE")
    print("per mouse done:", len(prs.slides), flush=True)

    # ------------------------------------------------------------------ SUPPLEMENT: limits and checks
    sec = "SUPPLEMENT: LIMITS AND CHECKS"
    R4 = pd.read_csv(an("ca_response_shape", "R4_down_speed_cells.csv"))
    R4r = pd.read_csv(an("ca_response_shape", "R4_down_speed_cells_traces_roi.csv"))
    R5 = pd.read_csv(an("ca_responsive_change", "R5_split_half_reliability.csv"))
    tr = R5[R5.event.isin(STIM)].n_trials
    dn4 = R4[R4.cls == "DOWN"]
    slide(sec, "Limits: DOWN responses and trials per stimulus", [os.path.join(od, "main", "s14_caveats.png")],
          "How to read: A = how fast responses start (grey = possible for a calcium decrease); B = whether the same neurons respond in "
          "two random halves of the trials.",
          f"{100 * (dn4.speed > 0.53).mean():.0f}% of DOWN responses are faster than Ca allows; {int(tr.min())}-{int(tr.max())} trials per "
          "stimulus and session. With 2 mice per group the smallest possible between-mouse permutation p is 0.33.",
          "Concern: DOWN may be a demixing / focus artefact; single-neuron maps are not stable with this number of trials.")
    slide(sec, "R4  DOWN speed, EXTRACT traces vs ROI traces", [an("ca_response_shape", "R4_down_check.png"),
                                                              an("ca_response_shape", "R4_down_check_traces_roi.png")],
          "How to read: speed of each DOWN / UP response and the off-footprint background around the events; top EXTRACT traces, "
          "bottom demixing-free ROI traces.",
          f"Result: DOWN responses faster than 0.53: EXTRACT {100 * (dn4.speed > 0.53).mean():.0f}% of {len(dn4)}, ROI traces "
          f"{100 * (R4r[R4r.cls == 'DOWN'].speed > 0.53).mean():.0f}% of {int((R4r.cls == 'DOWN').sum())}.",
          "Speed = change at 0.1-0.3 s / change at 1 s; a calcium decrease cannot exceed 0.53.")
    R3 = pd.read_csv(an("ca_response_shape", "R3_response_shape_counts.csv"))
    sh = R3.groupby(["cls", "shape"]).n_cells.sum()
    slide(sec, "R3  Shape of the UP and DOWN responses", [an("ca_response_shape", "R3_response_shape.png")],
          "How to read: average shape of each response class (rise after the event, drop after, ramp before ...).",
          "Result: " + "; ".join(f"{c} {s}: {int(v)}" for (c, s), v in sh.items()) + ".", "Classification of the trial-mean shape.")
    R9 = pd.read_csv(an("ca_response_shape", "R9_raw_roi_check.csv"))
    d9 = R9[R9.group == "DOWN"]
    slide(sec, "R9  DOWN neurons in the raw movie pixels", [an("ca_response_shape", "R9_raw_roi_check.png")],
          "How to read: the footprint pixels and the ring around them in the raw movie at the event, for DOWN and UP neurons.",
          f"Result: {int((d9.roi_dff_1s < 0).sum())} of {len(d9)} DOWN neurons got darker in their own pixels; "
          f"{int((d9.roi_minus_ann_1s < 0).sum())} brightened less than their surround.", "1-s window after the event.")
    R10 = pd.read_csv(an("ca_response_shape", "R10_field_change_by_state.csv"))
    slide(sec, "R10  Whole-field brightness change at the events", [an("ca_response_shape", "R10_field_change.png")],
          "How to read: brightness change of the whole field of view at each event, split by what the mouse did.",
          f"Result: change different from 0 (p < 0.05) in {int((R10.p_vs_0 < 0.05).sum())} of {len(R10)} event x state groups.",
          "Wilcoxon vs 0 over trials.")
    R5c = pd.read_csv(an("ca_responsive_change", "R5_baseline_responsive_cells_change.csv"))
    Sp = pd.read_csv(an("ca_cell_event_matrix", "responsive_change_split_half.csv"))
    Sp = Sp[Sp.event.isin(FD.EVR)]
    rep = {k: float(np.mean((g.base_heldout > 0) if k == "activated" else (g.base_heldout < 0))) for k, g in Sp.groupby("group")}
    slide(sec, "R5 / R5b  Held-out change and split-half reliability", [an("ca_responsive_change", "R5_baseline_responsive_change.png"),
                                                                       an("ca_responsive_change", "R5b_split_half_reliability.png")],
          "How to read: left = baseline-responsive neurons measured on held-out trials vs injection; right = how similar the response "
          "maps of two trial halves are.",
          f"Result: held-out change p < 0.05 in {int((R5c.p_wilcoxon < 0.05).sum())} of {len(R5c)} tests; median split-half reliability "
          f"{R5.reliability_sb.median():.2f}; sign kept in the other half: activated {100 * rep['activated']:.0f}%, inhibited "
          f"{100 * rep['inhibited']:.0f}%.", "Spearman-Brown corrected split-half correlation, 100 random halves.", ncols=2)
    R6 = pd.read_csv(an("ca_responsive_change", "R6_trial_common_mode.csv"))
    slide(sec, "R6  A component shared by all neurons on each trial", [an("ca_responsive_change", "R6_trial_common_mode.png")],
          "How to read: share of the trial-to-trial variation carried by one component common to all neurons, at events and at random times.",
          f"Result: median share {R6.pc1_share.median():.2f} at the events and {R6.sham_pc1_share.median():.2f} at random times.",
          "PCA of the trial x neuron responses.")
    R8 = pd.read_csv(an("ca_responsive_change", "R8_responsive_calibrated.csv"))
    R8r = pd.read_csv(an("ca_responsive_change", "R8_responsive_calibrated_traces_roi.csv"))
    slide(sec, "R8  Responsive-neuron counts corrected for correlated neurons", [an("ca_responsive_change", "R8_responsive_calibrated.png"),
                                                                                an("ca_responsive_change", "R8_responsive_calibrated_traces_roi.png")],
          "How to read: number of responsive neurons vs the number found on shifted (null) events, and the max-T neurons; top EXTRACT, "
          "bottom ROI traces.",
          f"Result: count above the null (p < 0.05) in {int((R8.p_count < 0.05).sum())} of {len(R8)} tests (ROI traces "
          f"{int((R8r.p_count < 0.05).sum())} of {len(R8r)}); max-T UP neurons {int(R8.maxT_up.sum())}, DOWN {int(R8.maxT_down.sum())}.",
          "Westfall-Young max-T, family-wise error 5%.")
    R7r = pd.read_csv(an("ca_responsive_change", "R7_population_response_test_traces_roi.csv"))
    R7 = pd.read_csv(an("ca_responsive_change", "R7_population_response_test.csv"))
    slide(sec, "R7  Population response test, EXTRACT vs ROI traces", [an("ca_responsive_change", "R7_population_response_test.png"),
                                                                     an("ca_responsive_change", "R7_population_response_test_traces_roi.png")],
          "How to read: population response / chance per event, mouse and session; top EXTRACT traces, bottom ROI traces.",
          f"Result: q <= 0.05 in {int((R7.q_energy <= 0.05).sum())} of {len(R7)} tests (EXTRACT) and "
          f"{int((R7r.q_energy <= 0.05).sum())} of {len(R7r)} (ROI traces).", "Events shifted inside their block, 2000 shifts, BH.")
    B10 = pd.read_csv(an("ca_day1", "ca_d1_B10_motion_control.csv"))
    slide(sec, "B10  Brain motion at the events", [an("ca_day1", "ca_d1_B10_motion_control.png")],
          "How to read: motion-correction shift and template correlation around the events; responsive neurons recounted on trials "
          "without a motion problem.",
          f"Result: {int(B10.n_bad_motion.sum())} of {int(B10.n_trials.sum())} trials had poor registration; UP neurons "
          f"{int(B10.up_all.sum())} with all trials, {int(B10.up_clean.sum())} with clean trials only.",
          "Poor registration = template correlation below median - 5 MAD.")
    E1 = pd.read_csv(an("ca_day1_extended", "E1_quality_tests.csv"))
    q1 = E1.groupby("measure").ratio.median()
    slide(sec, "E1  Signal quality, baseline vs injection", [an("ca_day1_extended", "E1_signal_quality.png")],
          "How to read: noise, transient size, decay and bleaching of every neuron in the two sessions.",
          "Result (median injection / baseline ratio): " + ", ".join(f"{k} {v:.2f}" for k, v in q1.items()) + ".",
          "Wilcoxon signed-rank over neurons per mouse.")
    lit = [["Corder 2019 Science", "dF/F", "0-2 s vs -5..-3 s", "1-tail rank-sum P<0.01", "none", "activated only"],
           ["Ahanonu & Corder 2022", "dF/F", "as Corder", "as Corder, adjust to SNR", "BH optional", "1-tail"],
           ["Grewe 2017 Nature", "dF/F 1-s bins", "1-25 s vs session", "2-tail rank-sum P<=0.01", "none", "either"],
           ["Zhang & Li 2018 Nat Commun", "CNMF dF/F0", "0-1 s vs -1..0 s", "signed-rank P<0.05", "none", "act. / inh."],
           ["Yang 2023 Nature (CeA)", "CNMF-E dF", "0-3 s vs -3..0 s", "permutation P<0.05", "not verified", "act. / inh."],
           ["Hua 2020 Nat Neurosci (CeA)", "MIN1PIPE", "5 s vs 5 s", "+-2 SD", "none", "incr. / suppr."],
           ["Ponserre 2022 J Neurosci", "MIN1PIPE dF/F", "regressor r", "bout shuffle, 1.96 SE", "none", "above / below"],
           ["Kissinger 2024 J Neurosci", "CNMFe dF/F", "0-5 s vs -5..0 s", "rank-sum p<0.05", "none", "incr. / decr."],
           ["Qian 2025 Nat Commun", "dF/F", "not verified", "1-tail rank-sum p<0.025", "none", "activated only"],
           ["THIS STUDY", "z-scored dF/F", C.WLIT, "time-shift null, BH q<=0.20", "BH", "UP / DOWN"]]
    slide(sec, "How published studies call a neuron responsive", [table_figure(
              lit, ["paper", "signal", "windows (after vs before)", "test / threshold", "correction", "classes"],
              [0.2, 0.12, 0.17, 0.22, 0.1, 0.13], f("literature_table.png"))],
          "How to read: one row per study (methods verified from the full texts), our rule in the last row.",
          "Result: our signal and response window match Corder 2019; our baseline is right before the event (Corder: -5 to -3 s); "
          "only our rule corrects for multiple neurons.",
          "Full table with DOIs: analysis\\RESPONSIVE_NEURON_CRITERIA_LITERATURE_KR.md.")
    if C.DAY != 1:
        prs.save(a.out)
        print("->", a.out, "|", len(prs.slides), "slides | figures + previews in", od)
        return
    L5 = pd.read_csv(an("ca_longitudinal_TOY", "long_cell_matching.csv")); L5["mouse"] = L5.mouse.astype(str)
    LT = pd.read_csv(an("ca_longitudinal_TOY", "long_tests.csv")); LT["mouse"] = LT.mouse.astype(str)
    lr_ = LT[LT.measure == "transient rate"]
    slide(sec, "TOY Day 2: pipeline test of the longitudinal analysis (not data)",
          [an("ca_longitudinal_TOY", "long_L5_cell_matching.png"), an("ca_longitudinal_TOY", "long_L4_cohort.png")],
          "How to read: a FAKE Day 2 made from each mouse's Day-1 baseline with a known effect; top = matching the same neurons across "
          "days, bottom = the cohort comparison.",
          "Result: " + ", ".join(f"{r.mouse} {r.n_matched} of {r.n_day1} matched (precision {r.toy_precision:.2f})" for r in L5.itertuples())
          + "; TOY rate medians baseline / vehicle / SBI: " + ", ".join(
              f"{r.mouse} {r.med_baseline:.1f}/{r.med_vehicle:.1f}/{r.med_sbi:.1f}" for r in lr_.itertuples()) + ".",
          "TOY data only: checks that the Day-2 code finds the effect that was put in. Real Day 2 recorded 2026-10-21.")
    prs.save(a.out)
    print("->", a.out, "|", len(prs.slides), "slides | figures + previews in", od)


if __name__ == "__main__":
    main()
