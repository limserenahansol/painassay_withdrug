"""make_day1_ca_slides.py - a short deck (6 slides) of the Day-1 mini1p Ca results for the PI.

Rules (PI, 2026-10-16): few polished slides; every graph with a title, labelled axes, legend, n per
group and the statistical test; a title at the top and one or two sentences at the bottom saying what
we did and what we found. Every number is read from the analysis files, nothing is typed by hand.

Slides
  1  Design and selected neurons      session timeline + the 4 fields of view (max dF/F, kept ROIs)
  2  Representative session           behaviour + Ca dF/F movie frame + traces (video file named)
  3  Activity per mouse               stimulus-block activity, baseline vs injection, every cell paired
  4  SBI-553 vs vehicle               per-mouse change, the arm difference and how robust it is
  5  Evoked responses                 population response vs the same events shifted in their block
  6  What the data can and cannot say caveats with their numbers (DOWN artefacts, trials per stimulus)

Layout per slide (13.33 x 7.5 in): title 0.25-1.0 in, figure 1.0-5.55 in (every figure is drawn at
12.4 x 4.55 in, so nothing is rescaled), the graph note (tests, n, definitions) 5.6-6.15 in at 11 pt,
the one or two sentences 6.25-7.3 in at 15 pt. Text widths are measured; a title that does not fit
at 22 pt stops the build.

  python make_day1_ca_slides.py --base C:/Users/hsollim/Documents/HEAL_mini1p_SBI553 --out <pptx>
Writes the pptx, one figure PNG per slide and a preview PNG per slide (same layout) next to it.
"""
import argparse
import os
import sys
import textwrap

import numpy as np
import pandas as pd
import scipy.io as sio

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.transforms as mtransforms
from matplotlib.patches import Patch
from matplotlib.lines import Line2D
from matplotlib.textpath import TextPath
from matplotlib.font_manager import FontProperties

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import safe_io  # noqa: E402,F401  (a figure save retries on a transient Windows write error)
import deck_cfg as C                                                # noqa: E402  (CA_DAY: which day)

MICE = C.MICE                   # Day 1: 5941, 5939 SBI-553 / 5937, 1191 vehicle; Day 2 the crossover
ARMC = {"SBI-553": "#c0392b", "vehicle": "#2471a3"}
STIM = ["Light touch", "Mild touch", "Pin prick", "Heat"]
SW, SH = 13.333, 7.5
FW, FH = 12.4, 4.55                                    # every figure, inches
X0 = 0.45                                              # left margin, inches
plt.rcParams.update({"font.size": 11, "axes.titlesize": 12, "axes.labelsize": 11, "font.family": "Arial"})


def pstr(p):
    return "n.a." if not np.isfinite(p) else ("p < 0.001" if p < 0.001 else f"p = {p:.3f}" if p < 0.01 else f"p = {p:.2f}")


def text_width_in(s, pt, bold=False):
    fp = FontProperties(family="Arial", weight="bold" if bold else "normal", size=pt)
    return TextPath((0, 0), s, prop=fp).get_extents().width / 72.0


def wrap_lines(s, pt, width_in):
    """Greedy wrap by measured width (Arial, the pptx font, so the preview wraps like the slide)."""
    out, cur = [], ""
    for w in s.split():
        t = (cur + " " + w).strip()
        if cur and text_width_in(t, pt) > width_in:
            out.append(cur); cur = w
        else:
            cur = t
    return out + ([cur] if cur else [])


def clean(a):
    for sp in ("top", "right"):
        a.spines[sp].set_visible(False)


# ------------------------------------------------------------------------------------------ data
def whole(B):
    W = pd.read_csv(C.A(B, "ca_day1", "ca_d1_cells_whole_session.csv"))
    W["mouse"] = W.mouse.astype(str)
    num = sum(W[f"still rate | {s}"].fillna(0) * W[f"still min | {s}"].fillna(0) for s in STIM)
    den = sum(W[f"still min | {s}"].fillna(0) for s in STIM)
    W["stim_still"] = num / den.replace(0, np.nan)
    return W


def kept_counts(B):
    OUT = os.path.join(B, "mini1p_out", C.DATE)
    n = {}
    for m, _, _ in MICE:
        R = pd.read_csv(os.path.join(OUT, f"{m}{C.T}", "4_cell_review", f"{m}{C.T}_cell_review.csv"))
        att = C.one_day_flags(B, m, R)
        k = R.keep == 1
        n[m] = (int(k.sum()), int((k & (att != "")).sum()))
    return n


# ------------------------------------------------------------------------------------- figures
def fig_design(B, out):
    OUT = os.path.join(B, "mini1p_out", C.DATE)
    fig = plt.figure(figsize=(FW, FH))
    ax = fig.add_axes([0.035, 0.80, 0.60, 0.08])
    blocks = [("no stimulus", 0, 180, "#bbbbbb")] + [(f"block {i + 1}", s0, s1, "#f5b041")
                                                    for i, (s0, s1) in enumerate([(180, 360), (390, 570), (600, 780), (810, 990)])]
    for nm, s0, s1, c in blocks:
        ax.barh(0, (s1 - s0) / 60, left=s0 / 60, height=0.8, color=c, edgecolor="k", lw=0.5)
        ax.text((s0 + s1) / 120, 0, nm, ha="center", va="center", fontsize=9)
    ax.set_xlim(0, 17); ax.set_ylim(-0.5, 0.5); ax.set_yticks([])
    ax.set_xlabel("time in session (min)", fontsize=10, labelpad=2)
    ax.tick_params(labelsize=9)
    ax.set_title("A  one session: 3 min without stimulus, then 4 stimulus blocks of 3 min", fontsize=11, loc="left")
    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    fig.legend(handles=[Patch(color="#bbbbbb", label="no stimulus (spontaneous)"),
                        Patch(color="#f5b041", label="light touch, mild touch, pin prick or heat"),
                        Line2D([], [], color="#2ecc40", lw=2, label=C.GREEN),
                        Line2D([], [], color=C.ONE_COL, lw=2, label=C.ORANGE)],
               loc="upper left", bbox_to_anchor=(0.655, 0.985), frameon=False, fontsize=9, ncol=1, handlelength=1.6)
    fig.text(0.035, 0.635, "B  fields of view: max dF/F projection of the whole day (both sessions), selected neurons outlined",
             fontsize=11)
    wd, gap = 0.235, 0.012
    for k, (m, arm, sex) in enumerate(MICE):
        a = fig.add_axes([0.035 + k * (wd + gap), 0.0, wd, 0.55])
        tag = f"{m}{C.T}"
        try:
            S = np.asarray(sio.loadmat(os.path.join(OUT, tag, "3_extract", f"{tag}_extract.mat"), variable_names=["spatial_weights"])["spatial_weights"], np.float32)
            R = pd.read_csv(os.path.join(OUT, tag, "4_cell_review", f"{tag}_cell_review.csv"))
            bg = np.asarray(sio.loadmat(os.path.join(OUT, tag, "3_extract", f"{tag}_projections.mat"), variable_names=["max_dff"])["max_dff"], float)
        except Exception:                                                      # noqa: BLE001
            a.set_axis_off(); a.set_title(f"{m}: not ready", fontsize=10); continue
        lo, hi = np.percentile(bg, [1, 99.8])
        a.imshow(np.clip((bg - lo) / (hi - lo), 0, 1) ** 0.7, cmap="gray")
        att = C.one_day_flags(B, m, R)
        nk = na = 0
        for j in range(S.shape[2]):
            if R.keep.iloc[j] != 1:
                continue
            f = S[:, :, j]
            if f.max() <= 0:
                continue
            col = C.ONE_COL if att.iloc[j] else "#2ecc40"
            na += bool(att.iloc[j]); nk += 1
            a.contour((f >= 0.3 * f.max()).astype(float), levels=[0.5], colors=[col], linewidths=0.8)
        a.set_axis_off()
        a.set_title(f"{m} ({arm}, {sex}): {nk} neurons ({nk - na} + {na})", fontsize=10, color=ARMC[arm])
    fig.savefig(out, dpi=160); plt.close(fig)
    return ("Order of the 4 stimuli differs between the two sessions; gaps of 30 s between blocks. Neurons per mouse: "
            + "on both days (green) + " + C.ONE_LAB + (" (magenta)." if C.DAY == 1 else " (cyan).")
            + " Same neurons in both sessions (one field of view, one detection per day).")


def fig_session(B, frame_png, out):
    """Slide 2: the Ca frame of a representative video (middle frame) + the same 8 neurons' traces and the
    behaviour, redrawn at slide size. Window, stimulus and cells are those of ca_behaviour_video.py."""
    import re as _re
    import cv2
    import sys
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from ca_behaviour_video import scoring_for, BLOCKS, FPS, SCOL, TRC
    base = os.path.basename(frame_png)
    mt = _re.match(r"(\d+)" + C.T + r"_(pre|post)_behaviour_ca_([A-Za-z]+)_(\d+)s_frame\.png", base)
    mouse, phase, tgt, t0 = mt.group(1), mt.group(2), mt.group(3), float(mt.group(4))
    rid, tag, dur = f"{mouse}{C.T}_{phase}", f"{mouse}{C.T}", 40.0
    t1 = t0 + dur
    day = os.path.join(B, "mini1p_out", C.DATE, tag)
    S = scoring_for(os.path.join(B, C.SCORE), mouse, phase)
    stim = [str(x) for x in np.atleast_1d(S["stimNames"])]
    target = next(s for s in stim if s.replace(" ", "") == tgt)
    dF = np.atleast_1d(S["dFrames"]).astype(int); dT = np.atleast_1d(S["dTypes"]).astype(int)
    rex = np.atleast_2d(S["reflexEvents"]).astype(int).reshape(-1, 2)
    move = np.asarray(S["moveVec"]).astype(bool).ravel() if "moveVec" in S else None
    f0, f1 = int(round(t0 * FPS)), int(round(t1 * FPS))
    T = dict(np.load(os.path.join(day, "5_analysis", "traces", f"{rid}_traces.npz")))
    fr, ids, D = T["frame"].astype(int), T["cell_ids"].astype(int), T["dff"]
    rows = np.clip(np.searchsorted(fr, np.arange(f0, f1)), 0, len(fr) - 1)
    Dw = D[rows]
    Ev = pd.read_csv(os.path.join(day, "5_analysis", "traces", f"{rid}_events.csv"))
    nev = Ev[(Ev.onset_s >= t0) & (Ev.onset_s < t1)].groupby("cell").size().reindex(ids, fill_value=0)
    R = pd.read_csv(C.A(B, "ca_day1", "ca_d1_cells_event_tests.csv"))
    R = R[(R["mouse"].astype(str) == mouse) & (R.phase == phase) & (R.event == target)]
    R = R[R.cell.isin(ids) & (R.cls_shift == "UP")].sort_values("resp_mean", ascending=False)
    up = [int(c) for c in R.cell.values[:8]]
    rest = [int(c) for c in nev.sort_values(ascending=False).index if int(c) not in up]
    pick = (up + rest)[:8]
    cj = [int(np.where(ids == c)[0][0]) for c in pick]
    im = cv2.imread(frame_png)[40:600, 1280:1920, ::-1]
    tmid = (f0 + (f1 - f0) // 2) / FPS                                   # the frame written by the video script
    fig = plt.figure(figsize=(FW, FH))
    a = fig.add_axes([0.0, 0.06, 0.315, 0.82])
    a.imshow(im); a.set_axis_off()
    a.set_title(f"A  Ca dF/F movie, {mouse} {('baseline' if phase == 'pre' else 'injection')} session, t = {tmid:.0f} s",
                loc="left", fontsize=11)
    ax = fig.add_axes([0.40, 0.38, 0.58, 0.50])
    axb = fig.add_axes([0.40, 0.12, 0.58, 0.21], sharex=ax)
    tt = np.arange(f0, f1) / FPS
    sc = np.nanpercentile(np.abs(D[:, cj]), 99.5, axis=0); sc[~(sc > 0)] = 1
    for i, j in enumerate(cj):
        ax.plot(tt, np.nan_to_num(Dw[:, j] / sc[i]) * 0.9 + (len(cj) - i), color=TRC[i % len(TRC)], lw=1)
        ax.text(t0 - 0.4, len(cj) - i + 0.2, f"#{ids[j]}" + (" *" if int(ids[j]) in up else ""), color=TRC[i % len(TRC)],
                ha="right", va="center", fontsize=9.5, fontweight="bold")
    for f, k in zip(dF, dT):
        if f0 <= f < f1 and 1 <= k <= len(stim):
            ax.axvline(f / FPS, color=SCOL.get(stim[k - 1], "k"), lw=0.9, alpha=0.7)
            axb.plot([f / FPS] * 2, [2.6, 3.4], color=SCOL.get(stim[k - 1], "k"), lw=2)
    ax.axvline(tmid, color="0.2", lw=1, ls=":")
    ax.set_yticks([]); ax.set_xlim(t0, t1); ax.set_ylim(0.4, len(cj) + 1.1)
    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    ax.tick_params(labelbottom=False)
    ax.set_title(f"B  8 neurons, dF/F (each scaled to its own maximum), {target} block", loc="left", fontsize=11)
    for typ, y, col in ((1, 2, "#6c3483"), (2, 1, "#1e8449")):
        x = rex[(rex[:, 1] == typ) & (rex[:, 0] >= f0) & (rex[:, 0] < f1), 0] / FPS
        axb.plot(x, np.full(len(x), y), "|", color=col, ms=11, mew=2)
    if move is not None:
        seg = move[f0 - 1:f1 - 1]
        axb.fill_between(tt[:len(seg)], -0.35, 0.35, where=seg, color="#17a2b8", lw=0, step="mid")
    axb.set_yticks([3, 2, 1, 0], [f"{target} delivery", "withdrawal", "flinch", "movement"], fontsize=9.5)
    axb.set_ylim(-0.6, 3.6)
    axb.set_xlabel("time in session (s)")
    for sp in ("top", "right"):
        axb.spines[sp].set_visible(False)
    fig.legend(handles=[Line2D([], [], color="0.2", ls=":", label="time of the frame in A"),
                        Line2D([], [], color=SCOL.get(target, "k"), lw=1.5, label=f"{target} delivery")],
               loc="upper right", bbox_to_anchor=(0.985, 1.01), ncol=2, frameon=False, fontsize=9.5)
    fig.savefig(out, dpi=160); plt.close(fig)
    nup = len(R)                                                         # all heat responders, not only those shown
    return (rid, target, nup, len(ids),
            f"A: one frame of the movie (motion-corrected dF/F, F0 = day mean, haze removed). White outlines = all {len(ids)} "
            f"selected neurons; coloured = the 8 traces in B. B: * = activity rose after {target.lower()} deliveries more than at "
            f"shifted times ({len(up)} of the {nup} such neurons shown); the others are the most active in these {dur:.0f} s.")


def fig_activity(W, out):
    fig, ax = plt.subplots(1, 2, figsize=(FW, FH), gridspec_kw=dict(width_ratios=[1.35, 1]))
    fig.subplots_adjust(left=0.065, right=0.99, top=0.83, bottom=0.17, wspace=0.18)
    rng = np.random.default_rng(0)
    from scipy.stats import wilcoxon
    for k, (col, ttl, panel) in enumerate((("stim_still", "stimulus blocks (4 blocks pooled)", "A"),
                                           ("still rate | no stimulus", "first 3 min, no stimulus", "B"))):
        a = ax[k]
        xt, xl, ymax = [], [], 0
        lab = []
        for i, (m, arm, sex) in enumerate(MICE):
            P = W[W.mouse == m].pivot(index="cell", columns="phase", values=col).dropna()
            if not {"pre", "post"} <= set(P.columns) or not len(P):
                continue
            x0, x1 = 3 * i, 3 * i + 1
            for _, r in P.iterrows():
                a.plot([x0, x1], [r.pre, r.post], color="0.78", lw=0.5, zorder=1)
            for x, ph in ((x0, "pre"), (x1, "post")):
                a.scatter(np.full(len(P), x) + rng.uniform(-0.13, 0.13, len(P)), P[ph], s=6, color=ARMC[arm],
                          alpha=0.35 if ph == "pre" else 0.8, zorder=2, lw=0)
                a.plot([x - 0.32, x + 0.32], [P[ph].median()] * 2, color="k", lw=2.5, zorder=3)
            p = wilcoxon(P.post, P.pre).pvalue if (P.post != P.pre).any() else np.nan
            ymax = max(ymax, np.nanmax(np.r_[P.pre, P.post]))
            lab.append(((x0 + x1) / 2, f"n = {len(P)} cells\n{pstr(p)}"))
            xt += [x0, x1]; xl += ["base", "inj."]
            a.text((x0 + x1) / 2, -0.13, f"{m} ({arm})", ha="center", va="top", fontsize=9, color=ARMC[arm],
                   transform=mtransforms.blended_transform_factory(a.transData, a.transAxes))
        a.set_ylim(0, ymax * 1.22)
        tr = mtransforms.blended_transform_factory(a.transData, a.transAxes)
        for x, s in lab:
            a.text(x, 0.985, s, ha="center", va="top", fontsize=8.5, transform=tr)
        a.set_xticks(xt, xl, fontsize=9)
        a.set_xlim(-0.7, 10.7)
        a.set_ylabel("Ca transients per min (still frames)")
        a.set_title(f"{panel}  {ttl}", loc="left", pad=8)
        clean(a)
    fig.legend(handles=[Patch(color=ARMC["SBI-553"], label=f"SBI-553 (n = {len(C.GROUP['SBI-553'])} mice)"),
                        Patch(color=ARMC["vehicle"], label=f"vehicle (n = {len(C.GROUP['vehicle'])} mice)"),
                        Line2D([], [], color="k", lw=2.5, label="median over cells"), Line2D([], [], color="0.75", lw=1, label="one neuron, baseline to injection")],
               loc="upper center", bbox_to_anchor=(0.5, 1.0), ncol=4, frameon=False, fontsize=9.5)
    fig.savefig(out, dpi=160); plt.close(fig)
    return ("Dots = neurons; base = baseline session, inj. = after the injection. Test: Wilcoxon signed-rank, injection vs "
            "baseline, over the neurons of each mouse (within-mouse, not a test between groups). Still frames = no movement "
            "(movement changes Ca activity and SBI-553 changed movement).")


def fig_arm(B, out):
    rb = C.A(B, "ca_robustness")
    R1 = pd.read_csv(os.path.join(rb, "R1_robustness_per_mouse.csv")); R1["mouse"] = R1.mouse.astype(str)
    RL = pd.read_csv(os.path.join(rb, "R2_leave_one_mouse_out.csv"))
    RA = pd.read_csv(os.path.join(rb, "R1_robustness_arm.csv"))
    E2 = pd.read_csv(C.A(B, "ca_day1_extended", "E15_block_position_model.csv"))
    main = "main: OASIS transients, still frames"
    fig, ax = plt.subplots(1, 3, figsize=(FW, FH), gridspec_kw=dict(width_ratios=[1.25, 1.1, 1.15]))
    fig.subplots_adjust(left=0.075, right=0.985, top=0.9, bottom=0.2, wspace=0.42)
    V = R1[R1.window == "stimulus blocks"]
    for i, (m, arm, sex) in enumerate(MICE):
        X = V[V.mouse == m]
        if not len(X):
            continue
        r = X[X.variant == main].iloc[0]
        ax[0].plot([i, i], [X.log2_ratio_median.min(), X.log2_ratio_median.max()], color="0.82", lw=11, solid_capstyle="butt", zorder=1)
        ax[0].errorbar(i, r.log2_ratio_median, yerr=[[r.log2_ratio_median - r.ci_lo], [r.ci_hi - r.log2_ratio_median]],
                       fmt="o", color=ARMC[arm], ms=8, capsize=4, zorder=2)
    ax[0].axhline(0, color="0.4", lw=0.8)
    nc = {m: int(V[(V.mouse == m) & (V.variant == main)].n_cells.iloc[0]) for m, _, _ in MICE if len(V[(V.mouse == m) & (V.variant == main)])}
    ax[0].set_xticks(range(4), [f"{m}\n{arm}\n{nc.get(m, 0)} cells" for m, arm, _ in MICE], fontsize=9)
    ax[0].set_xlim(-0.6, 3.6)
    ax[0].set_ylabel("log2(injection / baseline)\nstimulus-block transient rate")
    ax[0].set_title("A  change per mouse", loc="left")
    ax[0].legend(handles=[Line2D([], [], marker="o", color="k", ls="", label="median over cells, 95% bootstrap CI"),
                          Patch(color="0.82", label=f"range over {V.variant.nunique()} analysis choices")], fontsize=8.5,
                 frameon=False, loc="lower right")
    lo0 = min(V.log2_ratio_median.min(), V[V.variant == main].ci_lo.min())
    ax[0].set_ylim(lo0 - 0.15, 0.35)
    Q = RL[RL.window == "stimulus blocks"].reset_index(drop=True)
    cols = ["0.15"] + [ARMC[dict((m, a) for m, a, _ in MICE)[d]] for d in Q.left_out[1:]]
    ax[1].bar(range(len(Q)), Q.sbi_minus_vehicle, color=cols, alpha=0.85, width=0.7)
    for i, v in enumerate(Q.sbi_minus_vehicle):
        ax[1].text(i, v + (0.03 if v >= 0 else -0.03), f"{v:+.2f}", ha="center", va="bottom" if v >= 0 else "top", fontsize=9)
    ax[1].axhline(0, color="0.4", lw=0.8)
    ax[1].set_ylim(min(-0.2, Q.sbi_minus_vehicle.min() * 1.22), max(0.2, Q.sbi_minus_vehicle.max() + 0.15))
    ax[1].set_xticks(range(len(Q)), ["all 4\nmice"] + [f"without\n{d}" for d in Q.left_out[1:]], fontsize=9)
    ax[1].set_ylabel("SBI-553 minus vehicle\n(log2 change, mean of mouse medians)")
    ax[1].set_title("B  leave one mouse out", loc="left")
    y = E2[(E2.term.isin(["post", "post:sbi"])) & E2.model.str.contains("position")].reset_index(drop=True)
    lab = {"post": "vehicle mice:\nchange at retest", "post:sbi": "SBI-553 mice:\nextra change"}
    for i, r in y.iterrows():
        ax[2].errorbar(r.coef, i, xerr=[[r.coef - r.ci_lo], [r.ci_hi - r.coef]], fmt="s", color="k", ms=7, capsize=4)
        ax[2].text(r.ci_lo, i + 0.17, f"{r.coef:+.2f}  (cell-level {pstr(r.p)})", ha="left", va="bottom", fontsize=9)
    ax[2].axvline(0, color="0.4", lw=0.8)
    ax[2].set_yticks(range(len(y)), [lab.get(t, t) for t in y.term], fontsize=9.5)
    ax[2].set_ylim(-0.6, len(y) - 0.25)
    ax[2].set_xlabel("effect on log2 block activity (95% CI)")
    ax[2].set_title("C  mixed model", loc="left")
    ax[2].set_xlim(min(-1.3, y.ci_lo.min() - 0.15), max(0.7, y.ci_lo.max() + 1.2))
    for a in ax:
        clean(a)
    fig.savefig(out, dpi=160); plt.close(fig)
    q = RA[RA.window == "stimulus blocks"]
    return (f"n = 2 mice per group (SBI-553: {', '.join(C.GROUP['SBI-553'])}; vehicle: {', '.join(C.GROUP['vehicle'])}). A: grey = the mouse median under each analysis choice "
            f"(detector, thresholds, trace version, frames, cell subsets); SBI-553 < vehicle in {int((q.sbi_minus_vehicle < 0).sum())}/{len(q)} choices. "
            "C: linear mixed model log2(rate) ~ session x group + block position + stimulus, random mouse + cell; its p treats cells as "
            "units. Between mice (2 vs 2) no test can reach p < 0.33 (permutation).")


def fig_evoked(B, out):
    R7 = pd.read_csv(C.A(B, "ca_responsive_change", "R7_population_response_test.csv"))
    R7["mouse"] = R7.mouse.astype(str)
    evs = ["Light touch", "Mild touch", "Pin prick", "Heat", "withdrawal", "movement onset"]
    fig, ax = plt.subplots(1, 2, figsize=(FW, FH), sharey=True)
    fig.subplots_adjust(left=0.075, right=0.99, top=0.83, bottom=0.2, wspace=0.06)
    ymax = 0
    for k, ph in enumerate(("pre", "post")):
        a = ax[k]
        for i, (m, arm, sex) in enumerate(MICE):
            X = R7[(R7.mouse == m) & (R7.phase == ph)].set_index("event")
            for j, ev in enumerate(evs):
                if ev not in X.index:
                    continue
                r = X.loc[ev]
                v = r.energy / r.energy_null_median
                ymax = max(ymax, v)
                x = j + (i - 1.5) * 0.19
                a.bar(x, v, 0.17, color=ARMC[arm], alpha=0.95 if i % 2 == 0 else 0.5)
                if r.q_energy <= 0.05:
                    a.text(x, v + 0.03, "*", ha="center", va="bottom", fontsize=13)
        a.axhline(1, color="0.3", lw=1, ls="--")
        a.set_xticks(range(len(evs)), [e if e != "movement onset" else "movement\nonset" for e in evs], fontsize=9.5)
        a.set_title(("A  baseline" if ph == "pre" else "B  after injection") + " session", loc="left")
        if k == 0:
            a.set_ylabel("population response / chance\n(1 = no response)")
        clean(a)
    ax[0].set_ylim(0, ymax * 1.15)
    fig.legend(handles=[Patch(color=ARMC[arm], alpha=0.95 if i % 2 == 0 else 0.5, label=f"{m} ({arm})") for i, (m, arm, _) in enumerate(MICE)]
                       + [Line2D([], [], color="0.3", ls="--", label="chance"), Line2D([], [], color="k", ls="", marker="$*$", ms=9, label="* = that bar vs chance (q ≤ 0.05)")],
               loc="upper center", bbox_to_anchor=(0.5, 1.0), ncol=6, frameon=False, fontsize=9.5)
    fig.savefig(out, dpi=160); plt.close(fig)
    sig = R7[R7.q_energy <= 0.05]
    return ("One bar = one mouse and session (n = 2 mice per group). Response = mean z[0, 2 s] - z[-2, 0 s] after each event, "
            "RMS over neurons; chance = the same events shifted in time inside their own block (2000 shifts). "
            f"* = Benjamini-Hochberg q ≤ 0.05 over all {len(R7)} tests ({len(sig)} significant).")


def fig_caveats(B, out):
    R4 = pd.read_csv(C.A(B, "ca_response_shape", "R4_down_speed_cells.csv"))
    R5 = pd.read_csv(C.A(B, "ca_responsive_change", "R5_split_half_reliability.csv"))
    fig, ax = plt.subplots(1, 2, figsize=(FW, FH))
    fig.subplots_adjust(left=0.07, right=0.99, top=0.9, bottom=0.2, wspace=0.25)
    dn, up = R4[R4.cls == "DOWN"].speed.clip(-1, 3), R4[R4.cls == "UP"].speed.clip(-1, 3)
    bins = np.linspace(-1, 3, 33)
    ax[0].axvspan(-1, 0.53, color="0.9", zorder=0, label="speed possible for a Ca decrease (≤ 0.53)")
    h1 = ax[0].hist(up, bins, color="#c0392b", alpha=0.6, label=f"UP responses (n = {len(up)})")[0]
    h2 = ax[0].hist(dn, bins, color="#2471a3", alpha=0.6, label=f"DOWN responses (n = {len(dn)})")[0]
    ax[0].set_ylim(0, max(h1.max(), h2.max()) * 1.45)
    ax[0].set_xlim(-1, 3)
    ax[0].set_xlabel("response speed: change at 0.1-0.3 s / change at 1 s")
    ax[0].set_ylabel("responses (cell x event x session)")
    ax[0].set_title("A  most DOWN responses fall faster than Ca can", loc="left")
    ax[0].legend(fontsize=9, frameon=False, loc="upper right")
    g = R5.groupby("event").agg(n=("n_trials", "mean"), r=("reliability_sb", "mean"), s=("reliability_sb", "std"))
    g = g.reindex([e for e in ["Light touch", "Mild touch", "Pin prick", "Heat", "withdrawal", "movement onset"] if e in g.index])
    ax[1].bar(range(len(g)), g.r, yerr=g.s, color="0.55", capsize=3, width=0.7)
    for i, (e, r) in enumerate(g.iterrows()):
        ax[1].text(i, max(r.r, 0) + (r.s if np.isfinite(r.s) else 0) + 0.04, f"{r.n:.0f} trials", ha="center", fontsize=9)
    ax[1].axhline(0, color="0.4", lw=0.8)
    ax[1].set_xticks(range(len(g)), [e if e != "movement onset" else "movement\nonset" for e in g.index], fontsize=9.5)
    ax[1].set_ylabel("split-half reliability of which\nneurons respond (Spearman-Brown)")
    ax[1].set_title("B  single-neuron response maps are not stable yet", loc="left")
    ax[1].set_ylim(min(-1.05, (g.r - g.s).min() - 0.05), 1.05)
    for a in ax:
        clean(a)
    fig.savefig(out, dpi=160); plt.close(fig)
    return ("A: GCaMP fluorescence cannot fall faster than its decay (tau ~1.3 s); step-like DOWN responses come with a brightening "
            "of the whole field and are treated as artefacts. B: mean ± SD over the 8 sessions (4 mice x 2 sessions), 100 random "
            "halves of the trials per session; 1 = the same neurons respond in both halves.")


# ---------------------------------------------------------------------------------- slides
def add_slide(prs, title, fig_png, note, bottom, preview_png):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image
    tpt = next((p for p in (26, 24, 22) if text_width_in(title, p, bold=True) <= 12.3), None)
    if tpt is None:
        raise SystemExit(f"title too long for one line: {title!r}")
    nl = wrap_lines(note, 11, 12.3) if note else []
    bl = wrap_lines(bottom, 15, 12.3)
    if len(nl) > 3 or len(bl) > 3:
        raise SystemExit(f"text too long (note {len(nl)} lines, bottom {len(bl)} lines) on: {title!r}")
    s = prs.slides.add_slide(prs.slide_layouts[6])

    def box(x, y, w, h, text, pt, bold=False, rgb=(0x33, 0x33, 0x33)):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
        p = tf.paragraphs[0]; p.text = text; p.font.size = Pt(pt); p.font.bold = bold; p.font.name = "Arial"
        p.font.color.rgb = RGBColor(*rgb)
    box(X0, 0.3, 12.4, 0.65, title, tpt, True, (0x1F, 0x2D, 0x3D))
    if C.PROV_MARK:
        box(8.3, 0.04, 4.6, 0.3, C.PROV_MARK, 10.5, True, (0xC0, 0x39, 0x2B))
    im = Image.open(fig_png)
    boxw, boxh, top = 12.4, 4.55, 1.0
    w = boxw; h = w * im.size[1] / im.size[0]
    if h > boxh:
        h = boxh; w = h * im.size[0] / im.size[1]
    s.shapes.add_picture(fig_png, Inches(X0 + (boxw - w) / 2), Inches(top), Inches(w), Inches(h))
    if note:
        box(X0, 5.6, 12.4, 0.6, note, 11, False, (0x55, 0x55, 0x55))
    box(X0, 6.3, 12.4, 1.0, bottom, 15)
    # preview of the same layout (Arial, wrapped exactly as measured)
    fig = plt.figure(figsize=(SW, SH), dpi=100)
    fig.text((X0 + 0.05) / SW, 1 - 0.62 / SH, title, fontsize=tpt, fontweight="bold", color="#1f2d3d", va="center")
    ax = fig.add_axes([(X0 + (boxw - w) / 2) / SW, 1 - (top + h) / SH, w / SW, h / SH])
    ax.imshow(im); ax.set_axis_off()
    if nl:
        fig.text((X0 + 0.05) / SW, 1 - 5.63 / SH, "\n".join(nl), fontsize=11, color="#555555", va="top", linespacing=1.15)
    fig.text((X0 + 0.05) / SW, 1 - 6.33 / SH, "\n".join(bl), fontsize=15, color="#333333", va="top", linespacing=1.15)
    fig.savefig(preview_png, dpi=100); plt.close(fig)


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--video-frame", default="")
    ap.add_argument("--layout-test", action="store_true", help="skip the selection = analysis check (old data, layout only)")
    a = ap.parse_args(argv)
    from pptx import Presentation
    from pptx.util import Inches
    B = a.base
    od = os.path.splitext(a.out)[0] + "_figures"
    os.makedirs(od, exist_ok=True)
    W = whole(B)
    figs, notes = {}, {}
    for k, fn, nm in ((1, lambda o: fig_design(B, o), "s1_design_cells.png"), (3, lambda o: fig_activity(W, o), "s3_activity_per_mouse.png"),
                      (4, lambda o: fig_arm(B, o), "s4_sbi_vs_vehicle.png"), (5, lambda o: fig_evoked(B, o), "s5_evoked.png"),
                      (6, lambda o: fig_caveats(B, o), "s6_caveats.png")):
        figs[k] = os.path.join(od, nm); notes[k] = fn(figs[k])
    vf = a.video_frame
    ET = pd.read_csv(C.A(B, "ca_day1", "ca_d1_cells_event_tests.csv"))
    ET = ET[ET.event == "Heat"]
    nup_all = ET.assign(up=ET.cls_shift == "UP").groupby([ET["mouse"].astype(str), "phase"]).up.sum()
    if vf == "auto":        # the example session = the one with the most heat responders (said on the slide)
        import glob
        bm, bp = nup_all.idxmax()
        c = sorted(glob.glob(C.A(B, "ca_videos", f"{bm}{C.T}_{bp}_behaviour_ca_Heat_*s_frame.png")), key=os.path.getmtime)
        vf = c[-1] if c else ""
    # numbers for the sentences
    rb = C.A(B, "ca_robustness")
    R1 = pd.read_csv(os.path.join(rb, "R1_robustness_per_mouse.csv")); R1["mouse"] = R1.mouse.astype(str)
    RL = pd.read_csv(os.path.join(rb, "R2_leave_one_mouse_out.csv"))
    V = R1[(R1.window == "stimulus blocks") & (R1.variant == "main: OASIS transients, still frames")].set_index("mouse")
    ch = {m: 100 * (2 ** V.at[m, "log2_ratio_median"] - 1) for m, _, _ in MICE if m in V.index}
    V0 = R1[(R1.window == "no-stimulus start") & (R1.variant == "main: OASIS transients, still frames")].set_index("mouse")
    ch0 = {m: 100 * (2 ** V0.at[m, "log2_ratio_median"] - 1) for m, _, _ in MICE if m in V0.index}
    R5 = pd.read_csv(C.A(B, "ca_responsive_change", "R5_split_half_reliability.csv"))
    tr = R5[R5.event.isin(STIM)].n_trials
    big0 = min(ch, key=ch.get)                       # the mouse with the largest drop (Day 1: 5941)
    l5941 = RL[(RL.window == "stimulus blocks") & (RL.left_out == big0)].sbi_minus_vehicle
    R7 = pd.read_csv(C.A(B, "ca_responsive_change", "R7_population_response_test.csv"))
    st = R7[~R7.event.isin(["movement onset", "flinch"])]; mv = R7[R7.event == "movement onset"]
    R4 = pd.read_csv(C.A(B, "ca_response_shape", "R4_down_speed_cells.csv"))
    nk = kept_counts(B)
    ncell = {m: int(W[(W.mouse == m) & (W.phase == "pre")].cell.nunique()) for m, _, _ in MICE}
    bad = {m: (ncell[m], nk[m][0]) for m, _, _ in MICE if ncell[m] != nk[m][0]}
    if bad and not a.layout_test:
        raise SystemExit(f"analysis cells differ from the confirmed selection (analysis, selection): {bad}")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    pv = lambda k: os.path.join(od, f"preview_slide{k}.png")
    add_slide(prs, f"{C.LABEL}: Ca imaging of 4 mice before and after SBI-553 or vehicle", figs[1], notes[1],
              "Each mouse was imaged in a baseline session and again after the injection (same day, same neurons, same stimuli). "
              f"Neurons were chosen on the max projection: {', '.join(f'{m} {nk[m][0]}' for m, _, _ in MICE)}.", pv(1))
    n = 2
    if vf and os.path.isfile(vf):
        f2 = os.path.join(od, "s2_representative_session.png")
        rid, target, nup, nall, note2 = fig_session(B, vf, f2)
        oth = nup_all.drop(nup_all.idxmax())
        add_slide(prs, f"Example session: Ca activity of the neurons in the {target.lower()} block", f2, note2,
                  f"{nup} of {nall} neurons responded time-locked to {target.lower()} here, the most of the 8 sessions (others "
                  f"{int(oth.min())}-{int(oth.max())}). Video with both cameras: analysis\\{C.an('ca_videos')}\\{os.path.basename(vf).replace('_frame.png', '.mp4')}.", pv(2))
        n = 3
    big = min(ch, key=ch.get)
    arm_of = {m: arm for m, arm, _ in MICE}
    t3 = (f"Stimulus-block activity after the injection: largest drop in {big} ({arm_of[big]})" if ch[big] < 0
          else "Stimulus-block activity after the injection, per mouse")
    add_slide(prs, t3, figs[3], notes[3],
              "We counted Ca transients per minute of still frames for every neuron. Median change per neuron in the stimulus blocks: "
              + ", ".join(f"{m} {ch[m]:+.0f}%" for m in ch) + "; in the first 3 min without stimulus: "
              + ", ".join(f"{m} {ch0[m]:+.0f}%" for m in ch0) + ".", pv(n)); n += 1
    if C.DAY == 1:
        t4 = ("SBI-553 vs vehicle: a larger drop with SBI-553, carried by one mouse (5941)" if len(l5941) and l5941.iloc[0] > -0.15
              else "SBI-553 vs vehicle: a larger drop with SBI-553 than with vehicle")
        s4 = ("The SBI-553 mice fell more than the vehicle mice in most analysis choices, but without 5941 the group difference is "
              f"{l5941.iloc[0]:+.2f} log2; with 2 mice per group it cannot be tested between mice. Day 2 (crossover) is needed.")
    else:
        RA = pd.read_csv(os.path.join(rb, "R1_robustness_arm.csv")); q4 = RA[RA.window == "stimulus blocks"]
        all4 = RL[(RL.window == "stimulus blocks") & (RL.left_out == "none")].sbi_minus_vehicle
        t4 = "SBI-553 vs vehicle: change per mouse and the group difference"
        s4 = (f"SBI-553 minus vehicle: {all4.iloc[0]:+.2f} log2 with all 4 mice, {l5941.iloc[0]:+.2f} without {big0}; SBI-553 lower in "
              f"{int((q4.sbi_minus_vehicle < 0).sum())}/{len(q4)} analysis choices. 2 mice per group: no test between mice.")
    add_slide(prs, t4, figs[4], notes[4], s4, pv(n)); n += 1
    t5 = ("Evoked responses: clear for movement onset, weak for the stimuli" if C.DAY == 1
          else "Evoked population responses per event, baseline and after injection")
    add_slide(prs, t5, figs[5], notes[5],
              "We compared each event's population response with the same events shifted inside their block. "
              f"Movement onset beat chance in {int((mv.q_energy <= 0.05).sum())}/{len(mv)} sessions, the stimuli and withdrawal in "
              f"{int((st.q_energy <= 0.05).sum())}/{len(st)}; a drug effect on evoked responses cannot be measured yet.", pv(n)); n += 1
    dn = R4[R4.cls == "DOWN"]
    add_slide(prs, "Limits of the current data", figs[6], notes[6],
              f"{100 * (dn.speed > 0.53).mean():.0f}% of the DOWN responses are faster than Ca allows, so we report UP responses only; "
              f"with {int(tr.min())}-{int(tr.max())} trials per stimulus and session the single-neuron response maps are not stable. "
              + ("More trials and Day 2 will address both." if C.DAY == 1 else "More trials will address both."),
              pv(n))
    prs.save(a.out)
    print("->", a.out, "| figures + previews in", od)
    print("cells per mouse (selection = analysis):", {m: nk[m] for m, _, _ in MICE})


if __name__ == "__main__":
    main()
