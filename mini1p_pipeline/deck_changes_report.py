"""deck_changes_report.py - a readable table of how each slide of the morning PI deck (FINAL_20261023, before the timing fix)
changed in the new PI deck (FINAL_20261023_timefix): new slide number, what changed (figure / numbers / text), the red take-home
line before and after; plus the list of new slides. Uses the diff of deck_diff.py (DIFF_PI_EN.csv).

  python deck_changes_report.py
Writes FINAL_20261023_timefix\\CHANGES_PI_EN.md
"""
import os

import pandas as pd
from pptx import Presentation
from pptx.util import Inches

V = r"C:\Users\hsollim\Documents\HEAL_mini1p_SBI553\lab_meeting\decks_Day1_Day2\v2_20261022"
OLD, NEW = os.path.join(V, "FINAL_20261023", "ALL_PI_EN.pptx"), os.path.join(V, "FINAL_20261023_timefix", "ALL_PI_EN.pptx")


def keys(path):
    """per slide: (title, red key line, hidden)"""
    out = {}
    for i, s in enumerate(Presentation(path).slides, 1):
        title, key = "", ""
        for sh in s.shapes:
            if not (sh.has_text_frame and sh.text_frame.text.strip()):
                continue
            t = (sh.top or 0) / Inches(1)
            if 0.25 <= t < 0.6 and not title:
                title = sh.text_frame.text.strip()
            elif 0.8 <= t < 1.2 and not key:
                key = sh.text_frame.text.strip().replace("\n", " ")
        out[i] = (title, key, s._element.get("show") == "0")
    return out


def main():
    D = pd.read_csv(os.path.join(V, "FINAL_20261023_timefix", "DIFF_PI_EN.csv"))
    KO, KN = keys(OLD), keys(NEW)
    rows = ["# PI deck: before vs after the timing fix (10-23)", "",
            f"Morning deck {len(KO)} slides -> new deck {len(KN)} slides ({sum(v[2] for v in KN.values())} hidden backup).", "",
            "## The slides of the morning deck", "", "| old | new | what changed | take-home before | take-home after |", "|---|---|---|---|---|"]
    M = D[D.old_no.notna()].sort_values("old_no")
    for r in M.itertuples():
        o = int(r.old_no)
        if pd.isna(r.new_no):
            rows.append(f"| {o} | - | removed | {KO[o][1]} | |"); continue
        n = int(r.new_no)
        what = ", ".join(w for w, f in (("figure", r.figure_changed), ("numbers", r.numbers_changed), ("text", r.text_changed and not r.numbers_changed),
                                         ("notes", r.notes_changed)) if f is True or f == "True") or "same"
        ko, kn = KO[o][1], KN[n][1]
        rows.append(f"| {o} | {n} | {what} | {ko if ko != kn else '(same)'} | {kn} |")
    rows += ["", "## New slides", "", "| new | hidden | title | take-home |", "|---|---|---|---|"]
    for r in D[D.old_no.isna()].sort_values("new_no").itertuples():
        n = int(r.new_no)
        rows.append(f"| {n} | {'yes' if KN[n][2] else ''} | {KN[n][0]} | {KN[n][1]} |")
    fn = os.path.join(V, "FINAL_20261023_timefix", "CHANGES_PI_EN.md")
    open(fn, "w", encoding="utf-8").write("\n".join(rows).replace("\n| ", "\n| ") + "\n")
    print("->", fn, len(M), "old slides,", int(D.old_no.isna().sum()), "new")


if __name__ == "__main__":
    main()
