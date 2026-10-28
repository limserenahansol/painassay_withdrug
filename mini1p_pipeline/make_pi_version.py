"""make_pi_version.py - 10-26 (user: 'I need to send this ppt to my PI ... clear label, intuitive, no wordy, readable font size ...
check ppt to fit this and if it did not meet make a new version for PI from this ppt (do not overwrite)').

Input : the user's latest deck  OneDrive\\Desktop\\1111112121212_FINAL_compactversion3.pptx (read through a scratch copy; never written)
Output: OneDrive\\Desktop\\1111112121212_FINAL_PI_version.pptx

What it does (each replaced text is kept in that slide's notes as '[PI version] previous ...'):
  * the verified fixes of the 5-reviewer check (claims vs data, layout, story): wrong 'pin dropped after SBI-553' lines (35, 43, 47),
    slide 44's red line (heat / pin DOWN), slide 46's unfinished line, the deleted red line of 30, the removed Limits (48), arrows
    that point at gaps or cross labels, yellow fills that recolour bars (-> gold outlines), the cropped pin figure (re-drawn with its
    time axis), the behaviour figure with 7-8 pt text (re-drawn, 11-13 pt), two-line titles, page labels
  * readability: one short bottom line per mini1p slide at 12 pt (the full legend goes to the notes; star key moved to the footer),
    section labels 12 pt
  * order = the 4 parts of the e-mail: TRAP + Xenium -> addiction model -> crystal skull + mini2P / TRUFACT -> SBI-553 (WT
    behaviour, then mini1p); redundant slides (gene list at 7 pt, 2nd design table, the all-neuron response-change slide) go to the
    hidden backup; nothing is deleted.

  python make_pi_version.py
"""
import copy
import io
import os
import re
import shutil
import sys

from lxml import etree
from matplotlib.font_manager import FontProperties
from matplotlib.textpath import TextPath
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from deckops import add_arrow, add_note, add_text, geom, get, remove, reorder, set_arrow, set_text  # noqa: E402

USER = r"C:\Users\hsollim\OneDrive\Desktop\1111112121212_FINAL_compactversion3.pptx"
SCR = r"C:\Users\hsollim\AppData\Local\Temp\claude\C--Users-hsollim\e78b98bc-45eb-4eba-b6ff-ba7ac412ccce\scratchpad\hsedit"
SRC = os.path.join(SCR, "v3_2155.pptx")
OUT = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\hsollim\OneDrive\Desktop\1111112121212_FINAL_PI_version.pptx"
FIG = r"C:\Users\hsollim\Documents\HEAL_mini1p_SBI553\lab_meeting\decks_Day1_Day2\v2_20261022\FINAL_20261026_PI5_mixmarks\figures"
if os.path.getmtime(USER) > os.path.getmtime(SRC) + 1:          # the user saved again: start from the newest file
    shutil.copy2(USER, SRC)
assert os.path.abspath(OUT) != os.path.abspath(USER)

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
E = 914400
RED = RGBColor(0xC0, 0x39, 0x2B)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x55, 0x55, 0x55)
GOLD = RGBColor(0xF1, 0xC4, 0x0F)
STARS = "* p < 0.05, ** p < 0.01, *** p < 0.001, ns p ≥ 0.05"
LOG = []

prs = Presentation(SRC)
S = list(prs.slides)
# numbers below = the 60-slide deck that was reviewed (21:55). At 21:56 the user deleted slide 43 ('Responder numbers changed; response
# size did not'), so if it is gone every reviewed number above 43 is one lower in the file.
HAS43 = any(sh.has_text_frame and sh.text_frame.text.startswith("Responder numbers changed") for sh in S[42].shapes)
NREV = len(S) + (0 if HAS43 else 1)                 # 60


def sl(n):
    if n == 43 and not HAS43:
        return None
    return S[n - 1] if (HAS43 or n < 43) else S[n - 2]


def width_in(text, pt, bold=False):
    fp = FontProperties(family="Arial", weight="bold" if bold else "normal")
    return TextPath((0, 0), text, size=pt, prop=fp).get_extents().width / 72


def keep_old(n, sh, what=None):
    old = sh.text_frame.text.replace("\n", " / ") if sh.has_text_frame else ""
    add_note(sl(n), f"[PI version] previous {what or sh.name}: {old}")


def T(n, name, new, near=None, contains=None, size=None, bold=None, color=None, fit=None):
    """replace a shape's text (formatting of its first run kept); old text -> notes; fit = (width in, pt, bold) to warn on overflow"""
    sh = get(sl(n), name, near=near, contains=contains)
    keep_old(n, sh)
    set_text(sh, new, size=size, bold=bold, color=color)
    if fit:
        for line in new.split("\n"):
            w = width_in(line, fit[1], fit[2])
            if w > fit[0]:
                LOG.append(f"WARNING slide {n} {name}: '{line[:60]}' {w:.2f} in > {fit[0]} in")
    LOG.append(f"slide {n} {name}: {new[:90]!r}")
    return sh


def para_set(sh, idx, text):
    p = sh.text_frame.paragraphs[idx]
    rs = p.runs
    if not rs:
        r = p.add_run()
        rs = [r]
    rs[0].text = text
    for r in rs[1:]:
        r._r.getparent().remove(r._r)


def para_del(sh, idx):
    p = sh.text_frame.paragraphs[idx]._p
    p.getparent().remove(p)


def outline(sh, color=GOLD, w=2.25):
    """yellow translucent fill recolours the bars under it -> no fill, gold outline (same place = same highlight)"""
    sh.fill.background()
    sh.line.color.rgb = color
    sh.line.width = Pt(w)


def dup_shape(slide, sh):
    el = copy.deepcopy(sh._element)
    sh._element.addnext(el)
    return [x for x in slide.shapes if x._element is el][0]


def replace_picture(slide, old, path, x, y, w):
    new = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w))
    old._element.addprevious(new._element)              # same z-order as the old picture (arrows / boxes stay on top)
    remove(old)
    return new


def notes_replace(n, pairs):
    tf = sl(n).notes_slide.notes_text_frame
    t = tf.text
    t2 = t
    for a, b in pairs:
        t2 = t2.replace(a, b)
    if t2 != t:
        tf.text = t2


def white_box(slide, x, y, w, h):
    from pptx.enum.shapes import MSO_SHAPE
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.line.fill.background()
    r.shadow.inherit = False
    return r


# =============================================================================================== slides 1-23 (other projects)
for n in (1, 15):                                   # '+ With pain model' ran off the slide and sat on the dashed border
    sh = get(sl(n), "TextBox 1")
    geom(sh, 9.94, 6.31, 2.90, 0.22)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    bp = tf._txBody.find(A + "bodyPr")
    for k in list(bp):
        if k.tag in (A + "spAutoFit", A + "normAutofit", A + "noAutofit"):
            bp.remove(k)
    etree.SubElement(bp, A + "noAutofit")
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
r1 = dup_shape(sl(1), get(sl(1), "Rectangle 11"))  # slide 1: also mark the behaviour box (the addiction model, slides 11-14)
geom(r1, 0.17, 1.83, 3.03, 0.95)
add_note(sl(1), "[PI version] First the addiction aim: behaviour modelling, TRAP and Xenium (yellow boxes); then the pain aim.")
geom(get(sl(15), "Rectangle 3"), 6.85, 1.12, 6.40, 3.23)
add_note(sl(15), "[PI version] Now the pain aim: crystal-skull widefield, mini2P + TRUFACT, then SBI-553 (WT behaviour, then mini1p).")

sh = get(sl(4), "TextBox 4")                        # repeated the grey bottom line (mice as n: p = 0.33)
keep_old(4, sh); remove(sh)

sh = get(sl(6), "Content Placeholder 2")
keep_old(6, sh)
para_set(sh, 0, "ORBm and BMAp were the 2 regions left after all 3 filters")       # k-means does not rank regions
para_set(sh, 1, "Mice as n: both ~1.4x higher in Active (p ≈ 0.03)")
sh = get(sl(7), "Content Placeholder 2")
keep_old(7, sh)
para_set(sh, 1, "Their TRAP counts can be added to the Post comparison")
para_set(sh, 2, "Panel: 297 genes + tdTomato + Ai14 STOP = 299 targets")
T(8, "TextBox 1", "XENIUM")                         # old section number '4 —' from another deck
T(9, "TextBox 1", "XENIUM  ·  HOW THE PANEL WAS SEARCHED")
pic9 = get(sl(9), "Picture 6")                      # a picture entirely off the slide (x = -9.23 in): to a hidden backup slide
blank = next((l for l in prs.slide_layouts if l.name.lower() == "blank"), prs.slide_layouts[-1])
S9b = prs.slides.add_slide(blank)
for ph in list(S9b.placeholders):
    remove(ph)
S9b.shapes.add_picture(io.BytesIO(pic9.image.blob), Inches(2.2), Inches(1.2), pic9.width, pic9.height)
add_text(S9b, 0.45, 0.25, 12, 0.5, "Backup: picture that sat outside slide 'How thoroughly the 299 genes were searched' (moved here)", 16,
         color=(0x55, 0x55, 0x55))
remove(pic9)
add_note(sl(9), "[PI version] the picture that sat off the slide (left of it) is now on a hidden backup slide at the end.")
T(10, "TextBox 12", "Druggable  (46)")             # wrapped onto the rule; same name as slide 8
add_note(sl(10), "[PI version] hidden (7 pt gene list is not readable when projected); same 299 genes as slide 'One 299-gene panel', "
                 "grouped a little differently (here the 13 TFs are counted under cell identity).")
sh = get(sl(11), "TextBox 1")
geom(sh, 0.92, 2.12, 11.50, 0.64)
for p in sh.text_frame.paragraphs:
    p.alignment = PP_ALIGN.CENTER
T(12, "TextBox 6", "n = 6 Active, 8 Passive mice (7 Passive after day 13)  ·  A: Mann-Whitney U per phase  ·  B–C: Wilcoxon signed-rank")
T(13, "TextBox 6", "n = 6 Active, 8 Passive mice  ·  A: 95% bootstrap band over mice  ·  B: Wilcoxon signed-rank (details in notes)")

# conclusion slides 14 / 19 / 23: same place, bold heading, same bullets, one line per statement
bul = copy.deepcopy(get(sl(14), "TextBox 1").text_frame.paragraphs[2]._p.find(A + "pPr"))
for n, head, lines in ((14, "Conclusion: addiction behaviour", {4: "Right after a morphine reward, both stop more than after water",
                                                                5: "Then Active mice keep seeking; Passive mice keep stopping"}),
                       (19, "Conclusion: mini2P / TRUFACT and SBI-553 behaviour",
                        {2: "TRUFACT works in CeA: 13 imaged cells found again in the slice",
                         3: "1 mini2P mouse: pin prick lowered CeA-Ntsr1 activity (p = 0.006)",
                         4: "6 WT mice: fewer pain behaviours after SBI-553 (always Day 2)"}),
                       (23, "Conclusion: face poke vs melittin", {})):
    sh = get(sl(n), "TextBox 1")
    keep_old(n, sh)
    geom(sh, 0.92, 0.94, 11.50)
    para_set(sh, 0, head)
    for r in sh.text_frame.paragraphs[0].runs:
        r.font.bold = True; r.font.size = Pt(32)
    for k, t in lines.items():
        para_set(sh, k, t)
    if n != 14:
        for p in sh.text_frame.paragraphs[2:]:
            old = p._p.find(A + "pPr")
            if old is not None:
                p._p.remove(old)
            p._p.insert(0, copy.deepcopy(bul))

sh = get(sl(16), "TextBox 19")
keep_old(16, sh)
rs = sh.text_frame.paragraphs[0].runs
rs[2].text = " matched neurons respond to pain; none of 9 HCR-checked cells has Vipr1"
notes_replace(16, [("pain-session cell", "in vivo cell"), ("pain 세포가 있는", "in vivo 세포가 있는")])
for p in get(sl(17), "TextBox 5").text_frame.paragraphs:
    for r in p.runs:
        r.text = r.text.replace("Day 1 Vehicle", "Day 1   Vehicle")
s18 = sl(18)
geom(get(s18, "Rectangle 9"), 4.08, 4.26, 0.56, 1.44)
for nm in ("Rectangle 5", "Rectangle 7", "Rectangle 9", "Rectangle 11"):
    outline(get(s18, nm))
geom(get(s18, "TextBox 1"), x=0.55)
T(18, "TextBox 2", "Escape/rearing dropped at all 4 stimuli and withdrawal at 3 of 4 on the SBI-553 day", size=16, bold=True, color=RED)
T(18, "TextBox 4", "n = 6 WT mice  ·  paired Wilcoxon per behaviour × stimulus (exact, floor p = 0.031)\n"
                   "SBI-553 was always given on Day 2, so a day / habituation effect is not ruled out", size=13, bold=False)
for n, nm in ((3, "TextBox 4"), (5, "TextBox 4"), (16, "TextBox 20")):   # 11 pt grey notes under the figures -> 12 pt
    for p in get(sl(n), nm).text_frame.paragraphs:
        for r in p.runs:
            if r.font.size and r.font.size.pt < 12:
                r.font.size = Pt(12)
white_box(sl(20), 2.35, 0.90, 0.88, 0.23)           # 'Slide 2 |' printed in the image header (from another deck)
white_box(sl(21), 2.30, 0.95, 0.85, 0.22)
white_box(sl(22), 2.02, 0.95, 0.84, 0.22)
white_box(sl(22), 8.12, 0.95, 3.22, 0.22)           # '- this is what registration buys us'

# =============================================================================================== mini1p
MINI = [n for n in range(24, 49) if sl(n) is not None]
for n in MINI:                                      # section labels 11 -> 12 pt; footer carries the star key
    s = sl(n)
    for sh in s.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        t = sh.text_frame.text
        if sh.top / E < 0.2 and sh.left / E < 1 and sh.width / E < 12.2:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.font.size.pt <= 11:
                        r.font.size = Pt(12)
            sh.top = Inches(0.10)
        if t.startswith("CeA-Ntsr1 mini1p  ·  SBI-553 vs vehicle crossover"):
            set_text(sh, "CeA-Ntsr1 mini1p  ·  SBI-553 vs vehicle crossover  ·  n = 4 mice  ·  exploratory, uncorrected  ·  " + STARS,
                     size=9.5)
            sh.width = Inches(11.5)

# 24 overview: question title, caveat, roadmap matches the order
T(24, "TextBox 2", "Does SBI-553 change how CeA-Ntsr1 neurons respond to touch and pain?", fit=(12.2, 24, False))
T(24, "TextBox 7", "one example session\ndesign, colour code")
T(24, "TextBox 17", "all neurons\n→ tracked neurons")
T(24, "TextBox 22", "population AUC\n→ pin-prick close-up\n→ UP / DOWN neurons\n→ PCA, overlap")
T(24, "TextBox 27", "3 messages + limits")
sh = get(sl(24), "TextBox 28")
keep_old(24, sh)
p2 = copy.deepcopy(sh.text_frame.paragraphs[0]._p)
sh.text_frame._txBody.append(p2)
para_set(sh, len(sh.text_frame.paragraphs) - 1,
         "Caveat: in all 4 mice the Day-2 injection session had more pin UP neurons, so day order is not ruled out.")
sh.height = Inches(1.40)
T(25, "TextBox 1", "1  THE DATA  ·  design and colour code")
T(27, "TextBox 3", "Pin-prick UP neurons respond in the 0-0.5 s after a prick (shade)", fit=(12.2, 15, True))

# 28 behaviour divider (new, the user's)
sh = get(sl(28), "TextBox 1")
keep_old(28, sh)
para_set(sh, 0, "2. Behaviour")
para_set(sh, 2, "Fewer pin flinches and light-touch movements after SBI-553 than after vehicle")
sl(28).notes_slide.notes_text_frame.text = ("Section 2: behaviour (next 2 slides). [PI version] the notes copied from the overall-activity "
                                           "divider were replaced; previous divider text: 2. behavior / pin flinches and light-touch "
                                           "movements dropped after SBI-553 than after vehicle")

# 29 every behaviour: readable figure, no tint, arrow off the 'heat' label
s = sl(29)
T(29, "TextBox 2", "Pin prick: flinches rose after vehicle, not after SBI-553", fit=(12.2, 24, False))
T(29, "TextBox 3", "Only pin flinch (p < 0.001) and pin escape (p = 0.031) changed differently on the two days", fit=(12.2, 15, True))
replace_picture(s, get(s, "Picture 4"), os.path.join(FIG, "sA1_behaviour_per_stimulus_stats_PI.png"), 0.50, 1.45, 12.30)
r9 = get(s, "Rectangle 9")
outline(r9); geom(r9, 7.10, 2.40, 0.80, 1.62)       # flinch - pin (new figure geometry)
r9b = dup_shape(s, r9); geom(r9b, 7.10, 4.52, 0.80, 1.62)  # escape/rearing - pin
set_arrow(get(s, "Connector 12"), 7.97, 4.12, 7.70, 3.47)
geom(get(s, "TextBox 13"), 8.00, 3.97)
T(29, "TextBox 5", "Bars = mean of 4 mice (after − same-day baseline); star = CMH vs baseline; bracket = day difference (logistic)",
  size=12, fit=(12.2, 12, False))

# 30 behaviours that differed: red line back (numbers), arrows shortened
s = sl(30)
t30 = add_text(s, 0.45, 0.92, 12.40, 0.40, "Pin flinch 57 vs 79 %, light-touch movement 39 vs 61 % (after SBI-553 vs after vehicle, p < 0.01)",
               15, bold=True, color=(0xC0, 0x39, 0x2B))
LOG.append(f"slide 30 red line added ({width_in(t30.text_frame.text, 15, True):.2f} in)")
lng = [c for c in s.shapes if c.shape_type == 9 and abs(c.begin_x / E - 7.95) < 0.02][0]   # 2.4-in arrow across the pin group
set_arrow(lng, 5.20, 3.06, 5.20, 3.36)
add_text(s, 4.30, 2.74, 2.0, 0.30, "lower after SBI-553", 12.5, bold=True, color=(0xC0, 0x39, 0x2B))
geom(get(s, "TextBox 13"), x=7.75)
set_arrow([c for c in s.shapes if c.shape_type == 9 and abs(c.begin_x / E - 7.89) < 0.02][0], 7.75, 2.92, 7.10, 2.92)
T(30, "TextBox 5", "n = 4 mice (markers); CMH over deliveries: blue / red = vs same-day baseline, black = after vehicle vs after SBI-553",
  size=12, fit=(12.2, 12, False))

# 31 activity divider: one line
sh = get(sl(31), "TextBox 1")
keep_old(31, sh)
para_set(sh, 0, "3. Overall activity")
para_set(sh, 2, "Activity dropped after both injections, more on the SBI-553 day (same 185 neurons: −0.12 vs −0.03 z)")
para_del(sh, 3)
T(32, "TextBox 2", "Ca transients fell after both injections, more after SBI-553", fit=(12.2, 24, False))
T(32, "TextBox 5", "Dot = one neuron (216 / 246), 4 mice; Wilcoxon within day (same neurons), Mann-Whitney between days", size=12,
  fit=(12.2, 12, False))
s = sl(33)
T(33, "TextBox 3", "Stimulus blocks got darker after SBI-553 (0.87 → 0.76 z) than after vehicle (0.90 → 0.86 z)", fit=(12.2, 15, True))
outline(get(s, "Rectangle 8"))
T(33, "TextBox 5", "Mean z of the 185 tracked neurons per mouse; bottom row of each group = mean of 4 mice; descriptive", size=12,
  fit=(12.2, 12, False))
s = sl(34)
T(34, "TextBox 3", "Stimulus blocks: activity dropped more on the SBI-553 day (−0.12 vs −0.03 z, p < 0.001, 4/4 mice), except heat",
  fit=(12.2, 14, True))
c34 = [c for c in s.shapes if c.shape_type == 9]
set_arrow(min(c34, key=lambda c: abs(c.begin_x / E - 12.65)), 12.70, 5.20, 12.25, 4.66)
set_arrow(min(c34, key=lambda c: abs(c.begin_x / E - 9.06)), 9.00, 5.25, 8.56, 4.72)
T(34, "TextBox 5", "Per tracked neuron: after − same-day baseline (z); star = Wilcoxon vs 0; bracket = paired Wilcoxon; 'mice' = 4 mice",
  size=12, fit=(12.2, 12, False))

# 35 responses divider: the wrong 'dropped' line
sh = get(sl(35), "TextBox 1")
keep_old(35, sh)
para_set(sh, 2, "Touch responders rose on the SBI-553 day")
para_set(sh, 3, "Pin-prick UP neurons rose after vehicle, not after SBI-553")
para_del(sh, 4)
outline(get(sl(36), "Rectangle 12"))
T(36, "TextBox 5", "AUC per neuron (216 / 246), mean ± SEM, 4 mice; star = Wilcoxon vs 0; black bracket = Mann-Whitney after vs after",
  size=12, fit=(12.2, 12, False))
sh = T(37, "TextBox 5", "ΔAUC per neuron = after − same-day baseline session (z·s); 216 / 246 neurons, 4 mice; bracket = Mann-Whitney; "
                        "vs 0 = Wilcoxon", size=12, fit=(24.4, 12, False))
geom(sh, 0.45, 6.25, 12.40, 0.70)
sh = get(sl(38), "TextBox 1")
keep_old(38, sh)
para_set(sh, 2, "All-neuron AUC: no day difference once each day's baseline is subtracted (pin p = 0.041 → 0.40)")
para_set(sh, 3, "But UP and DOWN neurons cancel in that average")
para_set(sh, 4, "So next: single UP / DOWN neurons, pin prick first, then all 4 stimuli")
T(39, "TextBox 3", "In this example session, 21 of 47 neurons were UP and 9 were DOWN", fit=(12.2, 15, True))
T(39, "TextBox 5", "Example: 5939, Day 2 after vehicle, pin prick; top = every trial, middle = mean ± SEM, bottom: grey = 1000 shifted trains",
  size=12, fit=(12.2, 12, False))

# 40 pin UP / DOWN / rest: re-drawn with its time axis; counts per day in the title
s = sl(40)
T(40, "TextBox 2", "Pin prick: UP neurons rose after vehicle (25 → 47), not after SBI-553 (28 → 24)", fit=(12.2, 24, False))
T(40, "TextBox 3", "The after-SBI-553 UP trace looks bigger only because fewer neurons pass the test (baseline-UP neurons: p = 0.11)",
  fit=(12.2, 15, True))
replace_picture(s, get(s, "Picture 4"), os.path.join(FIG, "sF119_pin_only_PI.png"), 0.50, 1.55, 12.30)
set_arrow(get(s, "Connector 8"), 10.40, 2.80, 10.98, 2.97)
sh = T(40, "TextBox 5", "Mean ± SEM of z − its 0.5 s before, for the neurons called UP, DOWN or neither in each session; McNemar within day, "
                        "Fisher after vs after; 4 mice, neurons pooled", size=12, fit=(24.4, 12, False))
geom(sh, 0.45, 5.65, 12.40, 0.70)
add_note(s, "[PI version] neurons UP at the day's baseline changed alike on both days (-0.099 vs -0.057 z, p = 0.11). Per mouse, fewer "
            "pin UP after SBI-553 only in 5941 and 5939 (SBI-553 on Day 1); 5937 7 vs 2 and 1191 6 vs 0 go the other way (day order).")

# 41 pin close-up
T(41, "TextBox 2", "Pin prick: baseline UP neurons lost their response after SBI-553", fit=(12.2, 24, False))
T(41, "TextBox 5", "Each baseline class = 100 %; all neurons (216 / 246), 4 mice; Fisher exact; markers = mice", size=12, fit=(12.2, 12, False))
add_note(sl(41), "[PI version] vehicle side: 22 of the 25 vehicle-day baseline-UP neurons are from 5941 (within 5941: stay UP 15/22 vs 3/10).")

# 42 tracked % UP: one-line title, red line back under it, bar arrow on the after-SBI-553 bar
s = sl(42)
sh = T(42, "TextBox 2", "Same 185 tracked neurons: fewer pin UP neurons after SBI-553 (9% vs 23%)", fit=(12.2, 24, False))
geom(sh, 0.45, 0.32, 12.40, 0.60)
geom(get(s, "TextBox 3"), 0.45, 0.92, 12.40, 0.35)
c42 = [c for c in s.shapes if c.shape_type == 9]
set_arrow(min(c42, key=lambda c: abs(c.begin_x / E - 8.59) + abs(c.begin_y / E - 4.12)), 9.95, 2.62, 9.60, 2.98)
T(42, "TextBox 5", "185 neurons found on both days, 4 mice (markers); brackets = McNemar (same neurons); bottom = mean trace of UP neurons",
  size=12, fit=(12.2, 12, False))

# 43 -> backup (top row = slide 37's quantity, bottom row = slide 44's bottom row); fixed anyway. (The user deleted it at 21:56.)
if HAS43:
    T(43, "TextBox 2", "Responder numbers changed; the all-neuron average did not")
    T(43, "TextBox 3", "Touch responders rose on the SBI-553 day; pin responders rose only on the vehicle day (p = 0.001)")
    for nm in ("Rectangle 6", "Rectangle 10"):
        outline(get(sl(43), nm))

# 44 UP / DOWN AUC: correct red line, picture back in place, gold outlines, arrow on data, legend off the footer
s = sl(44)
T(44, "TextBox 3", "Compared with the vehicle day, touch responses grew and pin responses shrank on the SBI-553 day; heat did not differ",
  size=15, fit=(12.2, 15, True))
geom(get(s, "TextBox 3"), 0.45, 0.92, 12.40, 0.40)
geom(get(s, "Picture 4"), 0.50, 1.45)
r10, r14 = get(s, "Rectangle 10"), get(s, "Rectangle 14")
outline(r10); outline(r14)
geom(r10, 7.02, 1.95, 2.93, 1.95)
geom(r14, 7.02, 4.02, 2.93, 2.13)
set_arrow(get(s, "Connector 8"), 8.75, 3.02, 8.30, 3.40)
sh = T(44, "TextBox 5", "Top: AUC change (after − baseline session) of the tracked UP / DOWN neurons, same neurons on both days; star = Wilcoxon "
                        "vs 0, bracket = paired Wilcoxon. Bottom: change in % UP / DOWN, all neurons (216 / 246). 4 mice.", size=12,
       fit=(24.4, 12, False))
geom(sh, 0.45, 6.22, 12.40, 0.70)
notes_replace(44, [("(slide 17)", "(the earlier 'each neuron vs its own baseline' slide, now removed)"), ("slide 17", "the earlier (now removed) slide 'each neuron vs "
                                                                                                  "its own baseline'")])

# 45 PCA
T(45, "TextBox 2", "PCA, same 185 neurons: stimuli raise PC1; no clear drug pattern", fit=(12.2, 24, False))
T(45, "TextBox 3", "Pin-prick PC1 was lowest after SBI-553 (descriptive, no test)")
T(45, "TextBox 5", "PCA of trial-mean z, 185 tracked neurons, 4 mice; descriptive, no test (split-half r = 0.33)", size=12, fit=(12.2, 12, False))

# 46 overlap: the unfinished red line, in the figure's words
sh = T(46, "TextBox 3", "After vehicle, noxious-only neurons rose (10 → 19%); after SBI-553, touch-only (4 → 15%) and ‘both’ (3 → 10%) rose",
       size=14.5, fit=(12.2, 14.5, True))
geom(sh, 0.45, 0.92, 12.40, 0.40)
add_note(sl(46), "[PI version] 'fewer touch-only after vehicle' (38 -> 4) is mostly one mouse (1191: 28 -> 0) and 'fewer noxious-only "
                 "after SBI-553' is ns within the day (25 -> 20), so they are not on the slide.")
T(46, "TextBox 11", "All neurons (216 / 246), 4 mice; brackets test ‘both’: McNemar within day, Fisher after vs after; dashed = chance level",
  size=12, fit=(12.2, 12, False))

# 47 summary divider
sh = get(sl(47), "TextBox 1")
keep_old(47, sh)
para_set(sh, 2, "Pin UP neurons rose after vehicle (13 → 23%), not after SBI-553 (11 → 9%)")
para_set(sh, 4, "Touch responders rose on the SBI-553 day; pin responders only on the vehicle day")
for p in sh.text_frame.paragraphs[2:]:
    p.space_before = Pt(12)
    for r in p.runs:
        r.font.size = Pt(20)
sh.height = Inches(4.2)

# 48 take-home: red line under the title, cards say which number is which day, limits back (short)
s = sl(48)
geom(get(s, "TextBox 3"), 0.45, 0.92, 12.40, 0.40)
for nm in ("Rounded Rectangle 4", "TextBox 5", "TextBox 6", "TextBox 7", "Rounded Rectangle 8", "TextBox 9", "TextBox 10", "TextBox 11",
           "Rounded Rectangle 12", "TextBox 13", "TextBox 14", "TextBox 15"):
    sh = get(s, nm)
    sh.top = sh.top + Inches(0.30)
T(48, "TextBox 5", "2  Behaviour"); T(48, "TextBox 9", "3  Overall activity"); T(48, "TextBox 13", "4  Stimulus responses")
T(48, "TextBox 7", "After vehicle vs after SBI-553: pin flinch 79 vs 57%, light-touch movement 61 vs 39% (p < 0.01); partly a rise after vehicle")
T(48, "TextBox 11", "Same 185 neurons: SBI-553 day −0.12 vs vehicle day −0.03 z (p < 0.001), 4/4 mice")
T(48, "TextBox 15", "Same 185 neurons: 13 → 23% after vehicle (p < 0.01), 11 → 9% after SBI-553 (ns); touch responders rose on the SBI-553 day")
for nm in ("TextBox 7", "TextBox 11", "TextBox 15"):
    get(s, nm).height = Inches(1.10)
lim_box = dup_shape(s, get(s, "Rounded Rectangle 4"))
geom(lim_box, 0.45, 5.20, 12.40, 1.10)
lim = add_text(s, 0.65, 5.30, 12.00, 0.95,
               "Limits: 4 mice, exploratory (neurons pooled, uncorrected).\n"
               "Every mouse had more pin UP neurons after its Day-2 injection, so day order is not ruled out for the pin UP result; "
               "the pin-flinch and activity results held in all 4 mice.", 14, color=(0x33, 0x33, 0x33))
lim.text_frame.paragraphs[0].runs[0].font.bold = True

# =============================================================================================== final QA round (10-26, verified)
# The vehicle-day pin UP rise is mostly one mouse (5939: 2 -> 21 of 22 extra) and follows day order; pin flinch was lower after SBI-553
# in 3 mice and equal in 1; stimulus-block activity was lower on the SBI-553 day in 4/4 -> lead with the robust results.
HEAD = "After SBI-553: fewer pin flinches (57 vs 79%) and a bigger drop in neuronal activity during stimuli than after vehicle"
T(48, "TextBox 3", HEAD, fit=(12.2, 14.5, True))
T(24, "TextBox 3", HEAD, fit=(12.2, 15, True))
T(48, "TextBox 10", "Activity during stimuli fell after both injections, more on the SBI-553 day")
T(48, "TextBox 15", "Same 185 neurons: 13 → 23% after vehicle (p < 0.01; mostly one mouse), 11 → 9% after SBI-553 (ns); touch responders "
                    "rose on the SBI-553 day")
set_text(lim, "Limits: 4 mice, exploratory (neurons pooled, uncorrected).\nPin UP neurons follow day order: every mouse had more after its "
              "Day-2 injection, and the vehicle-day rise is mostly one mouse (5939). Pin flinch (lower in 3 mice, equal in 1) and "
              "stimulus-block activity (lower in 4/4) do not follow day order.")
for r in lim.text_frame.paragraphs[1].runs:
    r.font.bold = False
sh = get(sl(24), "TextBox 28")
para_set(sh, len(sh.text_frame.paragraphs) - 1, "Caveat: every mouse had more pin UP neurons after its Day-2 injection (day order not ruled "
                                               "out); the vehicle-day rise is mostly one mouse (5939).")
T(24, "TextBox 22", "population AUC\n→ UP / DOWN neurons\n→ pin close-up, tracked\n→ PCA, overlap")
T(42, "TextBox 3", "Pin UP rose after vehicle (13 → 23%; 17 of 19 extra in mouse 5939), not after SBI-553 (11 → 9%)", fit=(12.2, 15, True))
c = [c for c in sl(42).shapes if c.shape_type == 9 and abs(c.begin_x / E - 4.05) < 0.03]
if c:
    remove(c[0])                                     # light-touch arrow: its clause left the red line (pooled, 1191 15 -> 0)
sh = T(40, "TextBox 3", "Rise after vehicle is mostly one mouse (5939: 2 → 21); taller after-SBI-553 trace = fewer, stronger neurons pass "
                        "(baseline-UP: p = 0.11)", fit=(24.4, 15, True))
sh.height = Inches(0.60)
set_arrow(get(sl(40), "Connector 8"), 10.40, 2.80, 11.10, 2.92)
T(41, "TextBox 5", "Each baseline class = 100 %; all neurons (216 / 246), 4 mice; Fisher exact; markers = mice; vehicle-day baseline UP: 22 "
                   "of 25 from 5941", size=12, fit=(12.2, 12, False))
sh = get(sl(31), "TextBox 1")
para_set(sh, 2, "Activity during stimuli dropped after both injections, more on the SBI-553 day (−0.12 vs −0.03 z); no change without stimuli")
T(34, "TextBox 2", "Same 185 neurons: activity during stimuli fell more on the SBI-553 day", fit=(12.2, 24, False))
T(44, "TextBox 2", "UP / DOWN neurons: touch grew, pin shrank from baseline on the SBI-553 day", fit=(12.2, 24, False))
T(44, "TextBox 3", "Change from each day's baseline: touch grew, pin shrank on the SBI-553 day; heat did not differ", fit=(12.2, 15, True))
sh = get(sl(44), "TextBox 5")
set_text(sh, sh.text_frame.text + " Touch-UP and pin-DOWN gaps come from different baselines (notes).")
T(45, "TextBox 3", "Noxious stimuli moved PC1 more than touch (3 of 4 sessions); no consistent SBI-553 change (split-half r = 0.33)",
  fit=(12.2, 15, True))
for c in [c for c in sl(45).shapes if c.shape_type == 9]:
    remove(c)
add_note(sl(45), "[PI version] arrow on the pin bar removed (the red line no longer singles out pin).")
T(29, "TextBox 3", "Pin flinch: +20 points after vehicle vs −15 after SBI-553 (day difference p < 0.001)", fit=(12.2, 15, True))
T(36, "TextBox 2", "Population responses: weaker on the whole SBI-553 day, already at baseline", fit=(12.2, 24, False))
T(36, "TextBox 3", "After vs after only pin differed (p = 0.041), but the SBI-day pin response was already ~0 at baseline", fit=(12.2, 15, True))
sh = get(sl(47), "TextBox 1")
para_set(sh, 2, "Pin UP neurons rose after vehicle (13 → 23%), not after SBI-553 (11 → 9%), but this follows day order")
para_set(sh, 3, "After SBI-553, most baseline pin UP neurons went non-responsive (71% vs 28%; vehicle side mostly one mouse)")
para_set(sh, 4, "Touch responders rose on the SBI-553 day (light-touch UP 1 → 16%, 3 of 4 mice)")
T(32, "TextBox 3", "Fewer Ca transients after SBI-553 than after vehicle (median 1.3 vs 2.7 /min, p = 0.012); baselines equal (p = 0.64)",
  fit=(12.2, 15, True))
T(37, "TextBox 2", "Normalized to each day's baseline: no day difference (all neurons)", fit=(12.2, 24, False))
geom(get(sl(37), "TextBox 5"), 0.45, 6.55, 12.40, 0.50)
T(33, "TextBox 3", "Stimulus-block activity fell more after SBI-553 (0.87 → 0.76 z) than after vehicle (0.90 → 0.86 z)", fit=(12.2, 15, True))
for p in get(sl(27), "TextBox 5").text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(12)
for n in (27, 33, 41, 45):                          # run-level Arial so these match the other mini1p slides
    for nm in ("TextBox 1", "TextBox 2", "TextBox 3", "TextBox 5", "TextBox 10"):
        try:
            sh = get(sl(n), nm)
        except KeyError:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
for n, cb, h in ((41, 0.045, 4.97), (46, 0.05, 4.68)):   # crop the in-figure star key (the footer carries it)
    pic = get(sl(n), "Picture 4")
    pic.crop_bottom = cb; pic.height = Inches(h)

# other projects
add_text(sl(1), 6.85, 0.55, 6.40, 0.33, "Today: TRAP → Xenium → behaviour model, then the pain aim", 14, bold=True, color=(0x26, 0x26, 0x26))
add_text(sl(15), 6.85, 0.55, 6.40, 0.33, "Today: crystal skull → mini2P → SBI-553 (WT, then mini1p)", 14, bold=True, color=(0x26, 0x26, 0x26))
notes_replace(1, [("Today is mostly the addiction aim.", "")]); notes_replace(15, [("Today is mostly the addiction aim.", "")])
for n, tag in ((20, "CRYSTAL SKULL MESOSCOPE  ·  FACE POKE"), (21, "CRYSTAL SKULL MESOSCOPE  ·  MELITTIN"),
               (22, "CRYSTAL SKULL MESOSCOPE  ·  FACE POKE vs MELITTIN")):
    geom(get(sl(n), "TextBox 1"), y=0.28)
    add_text(sl(n), 0.45, 0.02, 8.00, 0.25, tag, 12, bold=True, color=(0x55, 0x55, 0x55))
    p0 = get(sl(n), "TextBox 3").text_frame.paragraphs[0]
    for r in p0.runs:
        r.font.color.rgb = RED; r.font.bold = True; r.font.size = Pt(14)
sh = get(sl(23), "TextBox 1"); para_set(sh, 0, "Conclusion: crystal skull, face poke vs melittin")
for r in sh.text_frame.paragraphs[0].runs:
    r.font.bold = True; r.font.size = Pt(32)


def red_para(n, name, idx=None):
    ps = get(sl(n), name).text_frame.paragraphs
    for p in (ps if idx is None else [ps[idx]]):
        for r in p.runs:
            r.font.color.rgb = RED; r.font.bold = True


red_para(3, "TextBox 3", 2); red_para(5, "TextBox 3", 1); red_para(8, "TextBox 4"); red_para(9, "TextBox 4")
T(4, "TextBox 5", "n = 2 Active + 2 Passive mice at Post  ·  p < 0.001 only with cells as n (mice as n: p = 0.33)")
sh = get(sl(4), "TextBox 3"); para_set(sh, 1, "•  ORBm 2.1×, BMAp 1.9× more cells in Active")
for r in sh.text_frame.paragraphs[1].runs:
    r.font.color.rgb = RED; r.font.bold = True; r.font.size = Pt(16)
rs = get(sl(16), "TextBox 19").text_frame.paragraphs[0].runs
rs[1].text = "8 of the 12"; rs[2].text = " matched neurons with an in vivo trace respond to pain; none of 9 HCR-checked cells has Vipr1"
T(17, "TextBox 1", "SBI-553 in WT mice: acute pain behaviour")
T(17, "TextBox 3", "Behaviour only, no imaging  ·  assay 2026-09-03")

# =============================================================================================== order, backup, page labels
HID = [10, 26] + ([43] if HAS43 else [])
main_other = [n for n in range(1, 15) if n not in HID] + [15, 20, 21, 22, 23, 16, 17, 18, 19]
main_mini = [24, 27, 25, 28, 29, 30, 31, 32, 33, 34, 35, 36, 37, 38, 39, 40, 41, 42, 44, 45, 46, 47, 48]
backup = [49] + HID + list(range(50, NREV + 1))
order = [sl(n) for n in main_other + main_mini + backup] + [S9b]
assert len(order) == len(prs.slides), (len(order), len(prs.slides))
reorder(prs, order)
for n in HID:
    sl(n)._element.set("show", "0")
S9b._element.set("show", "0")
N = len(main_mini)
for k, n in enumerate(main_mini, 1):
    lab = [sh for sh in sl(n).shapes if sh.has_text_frame and abs(sh.left / E - 12.20) < 0.05 and abs(sh.top / E - 7.10) < 0.05]
    if lab:
        set_text(lab[0], f"{k} / {N}")
    else:
        add_text(sl(n), 12.20, 7.10, 0.90, 0.30, f"{k} / {N}", 10, color=(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT)
for n in HID:
    lab = [sh for sh in sl(n).shapes if sh.has_text_frame and abs(sh.left / E - 12.20) < 0.05 and abs(sh.top / E - 7.10) < 0.05]
    if lab:
        set_text(lab[0], "backup")
    add_note(sl(n), "[PI version] moved to the hidden backup.")

prs.save(OUT)
pos = {id(s): i for i, s in enumerate(prs.slides, 1)}
LOG.append("FINAL ORDER (new# = user's #): " + ", ".join(f"{pos[id(sl(n))]}={n}" for n in main_other + main_mini + backup))
open(os.path.join(SCR, "pi_version_log.txt"), "w", encoding="utf-8").write("\n".join(LOG))
print("\n".join(x for x in LOG if x.startswith("WARNING") or x.startswith("FINAL") or "red line added" in x))
print("->", OUT, len(prs.slides), "slides")
