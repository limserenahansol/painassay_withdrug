"""make_auc_updown.py - 10-26 (user, on slide 41 'Responder numbers changed; response size did not', image of its figure: 'can you make
AUC of up and down cell so in this upper panel, instead of response change z among either up and down cell in 185 neuron?').

Same figure as slide 41 (make_pi5_deck.fig_single_norm), top row replaced:
    top    = AUC change (after the injection - the same day's baseline session, z*s) of the UP neurons and of the DOWN neurons among the
             185 neurons tracked on both days. UP = called UP in ANY of the 4 sessions (both baselines, both injection sessions);
             DOWN likewise. The selection is the same for all 4 sessions, so it does not push one day or one session up (no regression
             to the mean toward one side) and the SAME neurons are compared on the two days (paired test).
    bottom = unchanged: change in % UP / % DOWN, all neurons of each day (McNemar within day, Mann-Whitney between days).
AUC = area of (trial-mean z - its mean in the 0.5 s before the event) over the response window (pin 0-0.5 s; light, mild, heat 0-1 s),
the same AUC as the population-AUC slides (make_auc_basenorm.auc_session).

  python make_auc_updown.py  -> <PI5 figures>\\s2c_auc_updown.png + .json, and the 1-slide pptx + png in
                                OneDrive\\Desktop\\1111112121212_FINAL_compact_added_slides
"""
import json
import os

import numpy as np
import pandas as pd
from scipy.stats import binomtest, mannwhitneyu, wilcoxon

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402

B = r"C:\Users\hsollim\Documents\HEAL_mini1p_SBI553"
A = B + r"\analysis\_variants\mix4"
FIG = B + r"\lab_meeting\decks_Day1_Day2\v2_20261022\FINAL_20261026_PI5_mixmarks\figures"
DST = r"C:\Users\hsollim\OneDrive\Desktop\1111112121212_FINAL_compact_added_slides"
os.makedirs(DST, exist_ok=True)

plt.rcParams.update({"font.family": "Arial", "font.size": 12, "axes.spines.top": False, "axes.spines.right": False, "savefig.dpi": 200,
                     "hatch.color": "white"})
MICE = ["5941", "5939", "5937", "1191"]
SBI = {1: {"5941", "5939"}, 2: {"5937", "1191"}}
SBIDAY = {"5941": 1, "5939": 1, "5937": 2, "1191": 2}
SEX = {"5941": "F", "5939": "M", "5937": "M", "1191": "F"}
MC = {"5941": "#e67e22", "5939": "#8e44ad", "5937": "#16a085", "1191": "#2c3e50"}
MK = {"5941": "o", "5939": "s", "5937": "^", "1191": "D"}
OFF = {"5941": -0.21, "5939": -0.07, "5937": 0.07, "1191": 0.21}
CC = {"vi": "#2471a3", "si": "#c0392b"}
STIM = ["Light touch", "Mild touch", "Pin prick", "Heat"]
SES = ("cls_vb", "cls_vi", "cls_sb", "cls_si")
STARS = "* p < 0.05, ** p < 0.01, *** p < 0.001, ns p ≥ 0.05"


def st(p):
    return "ns" if p is None or not np.isfinite(p) or p >= 0.05 else "***" if p < 0.001 else "**" if p < 0.01 else "*"


def pf(p):
    if p is None or not np.isfinite(p):
        return "n.a."
    return "p < 0.001" if p < 0.001 else f"p = {p:.4f}" if p < 0.01 else f"p = {p:.3f}" if p < 0.1 else f"p = {p:.2f}"


cc = pd.read_csv(A + r"\event_analysis\cells_classes.csv")
cc["mouse"] = cc.mouse.astype(str)
cc["drug"] = [("SBI-553" if m in SBI[d] else "vehicle") for m, d in zip(cc.mouse, cc.day)]
cc["UP"] = cc.cls.eq("activated"); cc["DOWN"] = cc.cls.eq("inhibited")
PSTH = np.load(B + r"\analysis\new_decks\psth_cache_plainz_mix4.npz", allow_pickle=True)["D"].item()
TRK = pd.read_csv(A + r"\tracked_crossover\tracked_responses.csv"); TRK["mouse"] = TRK.mouse.astype(str)


def win(ev):
    e = cc[cc.event == ev].iloc[0]
    return float(e.pre), float(e.a), float(e.b)


def auc_session(m, d, ph, ev):
    """per neuron of one session: AUC (z*s) of (trial-mean z - its mean in the 0.5 s before) over the response window"""
    v = PSTH[(m, d, ph, ev)]
    tt, Y = v["tt"], v["mean"]
    pre, a, b = win(ev)
    a0 = min(a, 0)
    mb = (tt >= a0 - pre - 1e-9) & (tt < a0 - 1e-9)
    Yb = Y - np.nanmean(Y[mb], 0, keepdims=True)
    k0, k1 = np.searchsorted(tt, a - 1e-9), np.searchsorted(tt, b - 1e-9)
    return pd.Series(np.trapezoid(Yb[k0:k1 + 1], tt[k0:k1 + 1], axis=0), index=[int(i) for i in v["ids"]])


# ------------------------------------------------------------------------------------------------ top: UP / DOWN neurons, same 185
def mouse_paired(g, a, b):
    """mouse means of a and b (cells of each mouse), paired Wilcoxon over mice (smallest possible p 0.125 with 4 mice)"""
    mm = g.groupby("mouse")[[a, b]].mean().reindex(MICE)
    ok = mm.dropna()
    return mm, (wilcoxon(ok[b], ok[a])[1] if len(ok) >= 4 else np.nan), int((ok[b] > ok[a]).sum()), len(ok)


def updown(ev):
    t = TRK[TRK.event == ev].copy()
    A_ = {(m, d, ph): auc_session(m, d, ph, ev) for m in MICE for d in (1, 2) for ph in ("pre", "post")}
    for col, dk, ph, ik in (("avb", "veh_day", "pre", "veh_id"), ("avi", "veh_day", "post", "veh_id"),
                            ("asb", "sbi_day", "pre", "sbi_id"), ("asi", "sbi_day", "post", "sbi_id")):
        t[col] = [A_[(r.mouse, int(getattr(r, dk)), ph)].get(int(getattr(r, ik)), np.nan) for r in t.itertuples()]
    assert t[["avb", "avi", "asb", "asi"]].notna().all().all(), ev
    t["dv"] = t.avi - t.avb; t["ds"] = t.asi - t.asb
    R = dict(ev=ev, w=win(ev), n_trk=len(t))
    for lab, cls in (("UP", "activated"), ("DOWN", "inhibited")):
        sel = t[list(SES)].eq(cls).any(axis=1)
        g = t[sel]
        mm, pm, nm_up, nm = mouse_paired(g, "dv", "ds")
        sv = t.cls_vb.eq(cls) | t.cls_vi.eq(cls); ss = t.cls_sb.eq(cls) | t.cls_si.eq(cls)       # robustness: per-day selection
        other = "inhibited" if cls == "activated" else "activated"
        g1 = g[~g[list(SES)].eq(other).any(axis=1)]                                               # robustness: never in the other group
        R[lab] = dict(n=int(sel.sum()), n_mouse={m: int((g.mouse == m).sum()) for m in MICE},
                      veh=g.dv.mean(), veh_sem=g.dv.std(ddof=1) / np.sqrt(len(g)), sbi=g.ds.mean(), sbi_sem=g.ds.std(ddof=1) / np.sqrt(len(g)),
                      p_veh=wilcoxon(g.dv)[1], p_sbi=wilcoxon(g.ds)[1], p_pair=wilcoxon(g.ds, g.dv)[1],
                      mouse_veh={m: float(mm.dv[m]) for m in MICE}, mouse_sbi={m: float(mm.ds[m]) for m in MICE},
                      p_mice=pm, mice_sbi_higher=nm_up, mice_n=nm,
                      lv={k: float(g[k].mean()) for k in ("avb", "avi", "asb", "asi")},
                      # raw levels, same neurons: is the day difference in the change carried by the baselines or by the after sessions?
                      p_base=wilcoxon(g.asb, g.avb)[1], p_after=wilcoxon(g.asi, g.avi)[1],
                      perday=dict(n_veh=int(sv.sum()), n_sbi=int(ss.sum()), veh=t.dv[sv].mean(), sbi=t.ds[ss].mean(),
                                  p_veh=wilcoxon(t.dv[sv])[1], p_sbi=wilcoxon(t.ds[ss])[1], p_between=mannwhitneyu(t.ds[ss], t.dv[sv])[1]),
                      pure=dict(n=len(g1), p_pair=wilcoxon(g1.ds, g1.dv)[1], p_veh=wilcoxon(g1.dv)[1], p_sbi=wilcoxon(g1.ds)[1]))
    R["n_both"] = int((t[list(SES)].eq("activated").any(axis=1) & t[list(SES)].eq("inhibited").any(axis=1)).sum())
    R["n_rest"] = int((~t[list(SES)].isin(["activated", "inhibited"]).any(axis=1)).sum())
    return R


UD = {ev: updown(ev) for ev in STIM}


# ------------------------------------------------------------------------------------------------ bottom: as make_pi5_deck (unchanged)
def pairs(x):
    return x[x.phase == "pre"].set_index(["mouse", "day", "cell"]).join(
        x[x.phase == "post"].set_index(["mouse", "day", "cell"]), lsuffix="_b", rsuffix="_a", how="inner")


def mcnemar(x, col):
    j = pairs(x)
    b10 = int((j[col + "_b"] & ~j[col + "_a"]).sum()); b01 = int((~j[col + "_b"] & j[col + "_a"]).sum())
    return binomtest(b01, b01 + b10).pvalue if b01 + b10 else 1.0


def cls_change(ev):
    e = cc[cc.event == ev]
    out = {}
    for dr in ("vehicle", "SBI-553"):
        j = pairs(e[e.drug == dr]).reset_index()
        for col in ("UP", "DOWN"):
            j["d" + col] = j[col + "_a"].astype(int) - j[col + "_b"].astype(int)
        out[dr] = j
        for col in ("UP", "DOWN"):
            out[(col, dr)] = mcnemar(e[e.drug == dr], col)
    return out


CLS = {ev: cls_change(ev) for ev in STIM}


def bracket(a, x1, x2, y, txt, color="k", fs=10, h=None):
    h = h if h is not None else 0.015 * (a.get_ylim()[1] - a.get_ylim()[0])
    a.plot([x1, x1, x2, x2], [y, y + h, y + h, y], color=color, lw=1.2, clip_on=False)
    a.text((x1 + x2) / 2, y + h, txt, ha="center", va="bottom", color=color, fontsize=fs if txt != "ns" else fs - 0.5,
           fontweight="bold" if txt != "ns" else "normal", clip_on=False)


BARS = (("UP", "vehicle", 0.0), ("UP", "SBI-553", 1.0), ("DOWN", "vehicle", 2.5), ("DOWN", "SBI-553", 3.5))


def fig(fn):
    fig, ax = plt.subplots(2, 4, figsize=(12.3, 4.75), gridspec_kw=dict(wspace=0.22, hspace=0.70, left=0.075, right=0.995, top=0.80,
                                                                         bottom=0.07))
    POS = {}
    for j, ev in enumerate(STIM):
        R = UD[ev]
        a = ax[0, j]
        vals = []
        for col, dr, x in BARS:
            r = R[col]; k = "veh" if dr == "vehicle" else "sbi"
            mu, se = r[k], r[k + "_sem"]
            a.bar(x, mu, yerr=se, color=CC["vi"] if dr == "vehicle" else CC["si"], width=0.8, hatch=None if col == "UP" else "///", lw=0,
                  capsize=3, error_kw=dict(lw=1), zorder=2)
            mm = r["mouse_" + k]
            for m in MICE:
                if np.isfinite(mm[m]):
                    a.scatter(x + OFF[m] * 0.9, mm[m], s=22, marker=MK[m], color=MC[m], edgecolor="white", lw=0.6, zorder=4)
            vals += [mu + se, mu - se] + [v for v in mm.values() if np.isfinite(v)]
            POS[(ev, col, dr)] = (x, mu - se if mu < 0 else mu + se)
        hi, lo = max(vals + [0]), min(vals + [0])
        span = hi - lo
        a.set_ylim(lo - 0.08 * span, hi + 0.75 * span)
        for col, dr, x in BARS:
            p = R[col]["p_veh" if dr == "vehicle" else "p_sbi"]
            a.text(x, hi + 0.08 * span, st(p), ha="center", va="bottom", fontsize=9.5, color=CC["vi"] if dr == "vehicle" else CC["si"],
                   fontweight="normal" if p >= 0.05 else "bold")
        for col, x0 in (("UP", 0.0), ("DOWN", 2.5)):
            bracket(a, x0, x0 + 1, hi + 0.34 * span, st(R[col]["p_pair"]), fs=9.5, h=0.04 * span)
            POS[(ev, col, "bracket")] = (x0 + 0.5, hi + 0.38 * span)
        a.axhline(0, color="0.4", lw=0.8)
        a.set_xticks([0.5, 3.0], ["UP neurons", "DOWN neurons"], fontsize=10.5)
        a.tick_params(axis="y", labelsize=10)
        pre, a_, b_ = win(ev)
        a.set_title(f"{ev}  ({a_:g}-{b_:g} s" + (", original marks)" if ev == "Heat" else ")") +
                    f"\nsame {R['n_trk']} neurons: UP {R['UP']['n']}, DOWN {R['DOWN']['n']}", fontsize=11)
        # bottom: change in % UP and % DOWN (as on slide 41)
        b = ax[1, j]
        C = CLS[ev]
        vals, Sx = [], {}
        for col, dr, x in BARS:
            d = C[dr]
            v = 100 * d["d" + col].mean()
            b.bar(x, v, color=CC["vi"] if dr == "vehicle" else CC["si"], width=0.8, hatch=None if col == "UP" else "///", lw=0, zorder=2)
            mm = {m: 100 * g["d" + col].mean() for m, g in d.groupby(d.mouse.astype(str))}
            for m, w in mm.items():
                b.scatter(x + OFF[m] * 0.9, w, s=22, marker=MK[m], color=MC[m], edgecolor="white", lw=0.6, zorder=4)
            Sx[(col, dr)] = dict(p=C[(col, dr)], x=x)
            vals += [v] + list(mm.values())
        hi, lo = max(vals + [0]), min(vals + [0])
        span = hi - lo
        b.set_ylim(lo - 0.08 * span, hi + 0.75 * span)
        for (col, dr), r in Sx.items():
            b.text(r["x"], hi + 0.08 * span, st(r["p"]), ha="center", va="bottom", fontsize=9.5,
                   color=CC["vi"] if dr == "vehicle" else CC["si"], fontweight="normal" if r["p"] >= 0.05 else "bold")
        for col, x0 in (("UP", 0.0), ("DOWN", 2.5)):
            bracket(b, x0, x0 + 1, hi + 0.34 * span, st(mannwhitneyu(C["SBI-553"]["d" + col], C["vehicle"]["d" + col])[1]), fs=9.5,
                    h=0.04 * span)
        b.axhline(0, color="0.4", lw=0.8)
        b.set_xticks([0.5, 3.0], ["% UP", "% DOWN"], fontsize=10.5)
        b.tick_params(axis="y", labelsize=10)
        b.set_title(f"all neurons ({len(C['vehicle'])} / {len(C['SBI-553'])})", fontsize=9.5, color="0.35")
    ax[0, 0].set_ylabel("AUC change (z·s)\n(after - baseline)", fontsize=10.5)
    ax[1, 0].set_ylabel("change in % neurons\n(after - baseline)", fontsize=10.5)
    h = [Patch(color=CC["vi"], label="vehicle day"), Patch(color=CC["si"], label="SBI-553 day"),
         Patch(facecolor="0.55", hatch="///", edgecolor="white", label="hatched = DOWN")] + \
        [Line2D([], [], ls="", marker=MK[m], color=MC[m], markeredgecolor="white", markersize=8,
                label=f"{m} ({'female' if SEX[m] == 'F' else 'male'}, SBI-553 Day {SBIDAY[m]})") for m in MICE]
    fig.legend(handles=h, ncol=7, loc="upper center", bbox_to_anchor=(0.5, 1.0), frameon=False, fontsize=9.5, handlelength=1.2,
               columnspacing=1.0)
    fig.canvas.draw()
    # slide coordinates (inches) of chosen points, for the arrows: picture placed at (0.5, 1.45), 12.3 in wide = 1:1
    XY = {}
    for (ev, col, what), (x, y) in POS.items():
        px, py = ax[0, STIM.index(ev)].transData.transform((x, y))
        XY[(ev, col, what)] = (0.5 + px / fig.dpi, 1.45 + 4.75 - py / fig.dpi)
    fig.savefig(fn); plt.close(fig)
    return XY


F = FIG + r"\s2c_auc_updown.png"
XY = fig(F)
json.dump({ev: UD[ev] for ev in STIM}, open(F.replace(".png", ".json"), "w"), indent=1, default=float)
for ev in STIM:
    R = UD[ev]
    for lab in ("UP", "DOWN"):
        r = R[lab]
        print(f"{ev:12s} {lab:4s} n={r['n']:3d} {r['n_mouse']}  veh {r['veh']:+.4f} ({pf(r['p_veh'])})  sbi {r['sbi']:+.4f} ({pf(r['p_sbi'])})"
              f"  paired {pf(r['p_pair'])} | mice SBI>veh {r['mice_sbi_higher']}/{r['mice_n']} ({pf(r['p_mice'])})"
              f" | per-day n {r['perday']['n_veh']}/{r['perday']['n_sbi']} {r['perday']['veh']:+.4f}/{r['perday']['sbi']:+.4f}"
              f" MW {pf(r['perday']['p_between'])} | pure n={r['pure']['n']} {pf(r['pure']['p_pair'])}"
              f" | levels vb {r['lv']['avb']:+.3f} vi {r['lv']['avi']:+.3f} sb {r['lv']['asb']:+.3f} si {r['lv']['asi']:+.3f}"
              f" base-vs-base {pf(r['p_base'])} after-vs-after {pf(r['p_after'])}")
    print(f"{'':12s} both groups {R['n_both']}, never UP/DOWN {R['n_rest']}")
print("->", F)
print({k: tuple(round(v, 2) for v in xy) for k, xy in XY.items() if k[0] in ("Light touch", "Pin prick")})


# ------------------------------------------------------------------------------------------------ the 1-slide pptx (compact-deck style)
from lxml import etree  # noqa: E402
from matplotlib.font_manager import FontProperties  # noqa: E402
from matplotlib.textpath import TextPath  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

FOOT = "CeA-Ntsr1 mini1p  ·  SBI-553 vs vehicle crossover  ·  n = 4 mice  ·  exploratory, uncorrected (details in the notes)"


def fit(text, pt, bold, width, lo):
    prop = FontProperties(family="Arial", weight="bold" if bold else "normal")
    while pt > lo and TextPath((0, 0), text, size=pt, prop=prop).get_extents().width / 72 > width:
        pt -= 0.5
    if TextPath((0, 0), text, size=pt, prop=prop).get_extents().width / 72 > width:
        print("TOO LONG:", text)
    return pt


def tbox(s, x, y, w, h, text, size, bold=False, color=(0x22, 0x22, 0x22), align=None, italic=False):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = "Arial"; r.font.color.rgb = RGBColor(*color)
        if align:
            p.alignment = align
    return t


def arrow(s, x0, y0, x1, y1):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = RGBColor(0xC0, 0x39, 0x2B); c.line.width = Pt(3.5)
    te = etree.SubElement(c.line._get_or_add_ln(), "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")


L = UD["Light touch"]
SECTION = "4  STIMULUS RESPONSES  ·  UP / DOWN neurons, each vs its own baseline (same 185 neurons)"
TITLE = "UP and DOWN neurons: stronger touch, weaker pin responses on the SBI-553 day"
KEY = f"Light touch got stronger after SBI-553 (UP {pf(L['UP']['p_sbi'])}, DOWN {pf(L['DOWN']['p_sbi'])}); no change after vehicle"
BOTTOM = ("Top: AUC change (after - baseline session, z·s), mean ± SEM, of the tracked neurons called UP (or DOWN) in any of the 4 "
          "sessions, the same neurons on both days; markers = mouse means. Star = Wilcoxon vs 0; bracket = paired Wilcoxon, vehicle vs "
          "SBI-553 day. Bottom: change in % UP / DOWN, all neurons (slide 17). 4 mice; 5937 adds few neurons. Light-touch UP, mild-touch UP "
          "and pin DOWN gaps come mostly from different baseline sessions (notes). " + STARS)


def res_line(ev):
    R = UD[ev]
    out = []
    for lab in ("UP", "DOWN"):
        r = R[lab]
        nm = ", ".join(f"{m} {r['n_mouse'][m]}" for m in MICE)
        out.append(f"{lab} n = {r['n']} ({nm}), change vehicle day {r['veh']:+.3f} ({pf(r['p_veh'])}), SBI-553 day {r['sbi']:+.3f} "
                   f"({pf(r['p_sbi'])}), vehicle vs SBI-553 day paired {pf(r['p_pair'])}; mice with a higher change on the SBI-553 day "
                   f"{r['mice_sbi_higher']}/{r['mice_n']} (mouse-level paired Wilcoxon {pf(r['p_mice'])}, lowest possible 0.125)")
    return f"{ev}: " + "; ".join(out)


def lev_line(ev):
    R = UD[ev]
    out = []
    for lab in ("UP", "DOWN"):
        r = R[lab]; v = r["lv"]
        out.append(f"{lab} baseline vehicle day {v['avb']:+.3f} -> after {v['avi']:+.3f}, baseline SBI-553 day {v['asb']:+.3f} -> after "
                   f"{v['asi']:+.3f} (baseline vs baseline {pf(r['p_base'])}, after vs after {pf(r['p_after'])})")
    return f"{ev}: " + "; ".join(out)


def rob_line(ev):
    R = UD[ev]
    out = []
    for lab in ("UP", "DOWN"):
        r = R[lab]; q = r["perday"]
        out.append(f"{lab} per-day selection n = {q['n_veh']} / {q['n_sbi']}, {q['veh']:+.3f} vs {q['sbi']:+.3f} (Mann-Whitney "
                   f"{pf(q['p_between'])}); without the neurons that were in both groups n = {r['pure']['n']}, paired {pf(r['pure']['p_pair'])}")
    return f"{ev}: " + "; ".join(out)


def md(r, k):
    return ", ".join(f"{m} {r['mouse_' + k][m]:+.3f}" for m in MICE if np.isfinite(r["mouse_" + k][m]))


NB = {ev: UD[ev]["n_both"] for ev in STIM}; NR = {ev: UD[ev]["n_rest"] for ev in STIM}
NOTES = "\n".join([
    "WHY (user 10-26, on slide 17 'Responder numbers changed; response size did not'): the top row of slide 17 shows the response change "
    "of ALL neurons, where UP, DOWN and non-responsive neurons cancel. Here the same change is shown as AUC, separately for the UP and the "
    "DOWN neurons, in the 185 neurons tracked on both days. Bottom row = slide 17's bottom row, unchanged.",
    "MEASURE: per neuron-session AUC = area of (trial-mean z - its mean in the 0.5 s before the event) over the response window (pin 0-0.5 "
    "s; light, mild, heat 0-1 s), z·s - the same AUC as the population-AUC slides. Change = AUC after the injection - AUC in the same "
    "day's baseline session (same neuron). Pin AUC is about half of its mean z because its window is 0.5 s.",
    "WHICH NEURONS: UP = called UP (shift test, BH) in ANY of the 4 sessions (vehicle-day baseline, after vehicle, SBI-553-day baseline, "
    "after SBI-553); DOWN likewise. The same rule for all 4 sessions, so the selection does not favour one day or one session (no "
    "regression to the mean toward one side), and the same neurons are compared on the two days. Neurons UP in one session and DOWN in "
    "another are in both groups (" + ", ".join(f"{ev.split()[0].lower()} {NB[ev]}" for ev in STIM) + "); never UP or DOWN: "
    + ", ".join(f"{ev.split()[0].lower()} {NR[ev]}" for ev in STIM) + " of 185. A neuron that becomes UP after the injection also has a "
    "positive change, so this row is the size version of the bottom row (more neurons passing and stronger responses both raise it).",
    "TESTS: star over a bar = Wilcoxon signed-rank of the change vs 0; bracket = paired Wilcoxon, vehicle-day change vs SBI-553-day change "
    "(same neurons); markers = mouse means of the same neurons; a mouse-level paired Wilcoxon cannot go below 0.125 with 4 mice. "
    "Exploratory, uncorrected.",
    "RESULTS: " + " | ".join(res_line(ev) for ev in STIM) + ".",
    "WHERE THE DAY GAP COMES FROM (raw AUC of the same neurons, paired Wilcoxon): " + " | ".join(lev_line(ev) for ev in STIM) + ". "
    "So: light-touch UP, mild-touch UP and pin DOWN = the baseline sessions differ, the after sessions do not; light-touch DOWN = both "
    "(after SBI-553 more inhibited than after vehicle); mild-touch DOWN and pin UP = the after sessions differ, the baselines do not "
    "(after SBI-553: mild DOWN more inhibited, pin UP smaller). Heat: both differ, the changes do not.",
    "ROBUSTNESS: " + " | ".join(rob_line(ev) for ev in STIM) + ". Pin UP is the weakest: per-day selection ns, without the both-group "
    "neurons ns, mice 2/4.",
    f"DAY ORDER: SBI-553 was on Day 1 for 5941 / 5939 and Day 2 for 5937 / 1191. Light-touch DOWN mouse means, vehicle day "
    f"{md(L['DOWN'], 'veh')}; SBI-553 day {md(L['DOWN'], 'sbi')} - more inhibited on the SBI-553 day in 5941 (SBI-553 Day 1) and 1191 "
    "(SBI-553 Day 2), so not only the day order. 5937 has 0-1 light-touch UP / DOWN neurons.",
    "RELATION TO SLIDE 17: its title 'response size did not [change]' is about the mean of ALL neurons. In the UP / DOWN neurons the size did "
    "change for touch and pin (this slide). Suggested slide-17 title if both slides are kept: 'Responder numbers changed; the all-neuron "
    "average did not'.",
    r"Made by mini1p_pipeline\make_auc_updown.py (numbers in figures\s2c_auc_updown.json)."])

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
tbox(s, 0.45, 0.10, 11, 0.3, SECTION, 11, color=(0x55, 0x55, 0x55))
tbox(s, 0.45, 0.32, 12.4, 0.6, TITLE, fit(TITLE, 24, False, 12.15, 20))
tbox(s, 0.45, 0.92, 12.4, 0.5, KEY, fit(KEY, 15, True, 12.15, 13), bold=True, color=(0xC0, 0x39, 0x2B))
s.shapes.add_picture(F, Inches(0.5), Inches(1.45), Inches(12.3))
tbox(s, 0.45, 6.25, 12.4, 0.8, BOTTOM, 10.5)
tbox(s, 0.45, 7.1, 11.3, 0.3, FOOT, 9, color=(0x88, 0x88, 0x88), italic=True)
tbox(s, 12.2, 7.1, 0.9, 0.3, "17b / 23", 10, color=(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT)
x1, y1 = XY[("Light touch", "DOWN", "SBI-553")]
arrow(s, x1 + 0.55, y1 + 0.45, x1 + 0.12, y1 + 0.06)                      # at the light-touch DOWN bar of the SBI-553 day
s.notes_slide.notes_text_frame.text = NOTES
OUTP = os.path.join(DST, "02_NEW_after_slide41_UP_DOWN_neurons_AUC.pptx")    # slide 41 of the v3 deck = '17 / 23'
prs.save(OUTP)
print("->", OUTP)
