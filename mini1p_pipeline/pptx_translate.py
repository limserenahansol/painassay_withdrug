"""pptx_translate.py - the Korean version of a deck made by the deck builders (python-pptx only, no PowerPoint).

Slide text is translated through TEMPLATES: every number in a paragraph is replaced by a placeholder ({0}, {1}, ...),
the template is looked up in a dictionary (translations_kr.json: English template -> Korean template, placeholders
in any order), and the numbers of THIS deck are put back. So when a neuron is removed tomorrow and the numbers change,
the same dictionary still translates the rebuilt deck. Paragraphs without a template are split into sentences /
clauses and those are looked up; what is still unknown stays English and is listed (--dump) for translation.
Figures stay as they are (their labels are English). Korean text gets the East-Asian font Malgun Gothic and is
shrunk (0.5 pt steps, down to 70 %) where it would overflow its text box.

  python pptx_translate.py --in deck.pptx --out deck_KR.pptx [--dict translations_kr.json] [--dump missing.json]
"""
import argparse
import copy
import json
import os
import re

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
NUM = re.compile(r"(?<![A-Za-z_\d.#])[-+\u2212]?\d+(?:\.\d+)?")
FONT_KR = r"C:\Windows\Fonts\malgun.ttf"


KEEP = ("SBI-553", "GCaMP6", "GCaMP", "nVista3", "mini2p", "mini1p", "Ntsr1", "CeA-Ntsr1", "R2_", "log2", "Day-1", "Day-2")


def template(text):
    """numbers -> {0}, {1}, ... (names that contain digits, like SBI-553, are kept as text)"""
    prot = {}
    for i, w in enumerate(KEEP):
        if w in text:
            tok = "\u0001" + chr(65 + i) + "\u0002"
            text = text.replace(w, tok)
            prot[tok] = w
    nums = NUM.findall(text)
    k = iter(range(len(nums)))
    t = NUM.sub(lambda m: "{" + str(next(k)) + "}", text)
    for tok, w in prot.items():
        t = t.replace(tok, w)
    return t, nums


def fill(tpl, nums):
    out = tpl
    for i in range(len(nums) - 1, -1, -1):
        out = out.replace("{" + str(i) + "}", nums[i])
    return out


SPLIT = re.compile(r"(?<=[.;:])\s+(?=[A-Z(\"'])|(?<=; )")


def translate(text, D, missing):
    if not text.strip():
        return text, True
    t, nums = template(text)
    if t in D and D[t]:
        return fill(D[t], nums), True
    # sentence / clause pieces (the dictionary is kept at this level)
    parts = re.split(r"(?<=[.;])\s+", text)
    if len(parts) > 1:
        out, ok_all = [], True
        for q in parts:
            tq, nq = template(q)
            if tq in D and D[tq]:
                out.append(fill(D[tq], nq))
            else:
                out.append(glossary(q)); ok_all = False
                missing.setdefault(tq, "")
        return " ".join(out), ok_all
    missing.setdefault(t, "")
    return glossary(text), False


def glossary(text):
    """last resort for a sentence without a template (lists that depend on the data): the recurring terms"""
    G = GLOSS
    for en in sorted(G, key=len, reverse=True):
        text = re.sub(r"(?<![A-Za-z])" + re.escape(en) + r"(?![A-Za-z])", G[en], text)
    return text


GLOSS = {}


def set_ea_font(run, name="Malgun Gothic"):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", "Arial" if tag == "a:latin" else name)


def text_lines(text, pt, width_in, font):
    """wrapped line count of `text` at `pt` in a box `width_in` wide (Korean wraps at any character, English at spaces)"""
    from PIL import ImageFont
    f = ImageFont.truetype(font, max(int(round(pt * 4)), 1))           # 4 px per pt -> 288 px per inch
    maxw = width_in * 288
    lines, cur = 1, ""
    tokens = re.findall(r"[A-Za-z0-9.,%+\-()/#:=<>'\"]+\s*|\s+|.", text)
    for tok in tokens:
        t = cur + tok
        if f.getlength(t.rstrip()) > maxw and cur:
            lines += 1
            cur = tok.lstrip()
        else:
            cur = t
    return lines


def shrink_to_fit(shape, tf):
    """reduce the font size until the Korean text fits the box height (line height 1.2 x size)"""
    if not shape.width or not shape.height:
        return
    w = shape.width / 914400 - 0.12
    h = shape.height / 914400 - 0.05
    runs = [r for p in tf.paragraphs for r in p.runs]
    if not runs or runs[0].font.size is None:
        return
    pt0 = runs[0].font.size.pt
    pt = pt0
    while pt >= 0.7 * pt0:
        n = sum(text_lines("".join(r.text for r in p.runs), pt, w, FONT_KR) for p in tf.paragraphs)
        if n * pt * 1.2 / 72 <= h + 0.02:
            break
        pt -= 0.5
    if pt != pt0:
        for r in runs:
            r.font.size = Pt(pt)


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="inp", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--dict", default=os.path.join(HERE, "translations_kr.json"))
    ap.add_argument("--dump", default="")
    a = ap.parse_args(argv)
    D = json.load(open(a.dict, encoding="utf-8")) if os.path.isfile(a.dict) else {}
    gl = os.path.join(os.path.dirname(a.dict), "glossary_kr.json")
    if os.path.isfile(gl):
        GLOSS.update(json.load(open(gl, encoding="utf-8")))
    prs = Presentation(a.inp)
    missing, n_par, n_ok = {}, 0, 0
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            changed = False
            for p in tf.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt.strip():
                    continue
                n_par += 1
                new, ok = translate(txt, D, missing)
                n_ok += ok
                if new != txt:
                    p.runs[0].text = new
                    for r in p.runs[1:]:
                        r._r.getparent().remove(r._r)
                    set_ea_font(p.runs[0])
                    changed = True
            if changed:
                shrink_to_fit(sh, tf)
        if s.has_notes_slide:
            nt = s.notes_slide.notes_text_frame
            for p in nt.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if txt.strip():
                    new, _ = translate(txt, D, missing)
                    if new != txt and p.runs:
                        p.runs[0].text = new
                        for r in p.runs[1:]:
                            r._r.getparent().remove(r._r)
    prs.save(a.out)
    print(f"{n_ok} of {n_par} paragraphs fully translated; {len(missing)} templates unknown -> {a.out}")
    if a.dump:
        old = json.load(open(a.dump, encoding="utf-8")) if os.path.isfile(a.dump) else {}
        old.update({k: v for k, v in missing.items() if k not in old})
        json.dump(old, open(a.dump, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print("missing templates ->", a.dump)


if __name__ == "__main__":
    main()
