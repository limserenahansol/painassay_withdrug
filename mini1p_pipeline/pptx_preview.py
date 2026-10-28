"""pptx_preview.py - approximate PNG render of every slide of a .pptx (no PowerPoint, no LibreOffice).

Draws what python-pptx can read, at its real position: pictures (with their crop), text boxes and
placeholders (wrapped at the box width with the run's font size; Arial metrics), filled shapes and
tables as boxes. Good enough to check layout: overlaps, text running out of its box or off the
slide, empty areas. Not a faithful renderer (no theme fonts, no effects, no charts).

  python pptx_preview.py <deck.pptx> <out_dir> [--slides 1,3,5] [--dpi 80]
Writes <out_dir>/<deck>_slideNN.png and prints a warning for text that overflows its box.
"""
import argparse
import io
import os

import numpy as np
from PIL import Image
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.textpath import TextPath
from matplotlib.font_manager import FontProperties
from pptx import Presentation
from pptx.util import Emu
import sys as _sys, os as _os
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import safe_io  # noqa: E402,F401  (a figure save retries on a transient Windows write error)

EMU = 914400.0


def fam(s):
    """Arial, or Malgun Gothic for Korean text (the East-Asian font of the Korean decks)"""
    return "Malgun Gothic" if any("가" <= ch <= "힣" for ch in s) else "Arial"


def width_in(s, pt, bold):
    if not s:
        return 0.0
    fp = FontProperties(family=fam(s), weight="bold" if bold else "normal", size=pt)
    return TextPath((0, 0), s, prop=fp).get_extents().width / 72.0


def wrap(text, pt, bold, w):
    lines = []
    for para in text.split("\n"):
        cur = ""
        for word in para.split(" "):
            t = (cur + " " + word) if cur else word
            if cur and width_in(t, pt, bold) > w:
                lines.append(cur); cur = word
            else:
                cur = t
        lines.append(cur)
    return lines


def run_style(tf):
    pt, bold, col = None, False, "#222222"
    for p in tf.paragraphs:
        for r in p.runs:
            if r.font.size:
                pt = r.font.size.pt
            bold = bold or bool(r.font.bold)
            try:
                if r.font.color and r.font.color.type is not None and r.font.color.rgb is not None:
                    col = "#" + str(r.font.color.rgb)
            except Exception:                                                  # noqa: BLE001
                pass
            if pt:
                return pt, bold, col
        if p.font.size:
            pt = p.font.size.pt
    return pt or 18.0, bold, col


def render(prs, idx, s, out, dpi):
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    fig = plt.figure(figsize=(W, H), dpi=dpi)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(H, 0); ax.set_axis_off()
    ax.add_patch(plt.Rectangle((0, 0), W, H, color="white", zorder=0))
    warn = []
    for sh in s.shapes:
        if sh.left is None or sh.width is None:
            continue
        x, y, w, h = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
        if sh.shape_type == 13:                                                # picture
            try:
                im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
                cl, cr, ct, cb = sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom
                iw, ih = im.size
                im = im.crop((int(cl * iw), int(ct * ih), int(iw * (1 - cr)), int(ih * (1 - cb))))
                ax.imshow(np.asarray(im), extent=(x, x + w, y + h, y), zorder=2, aspect="auto")
            except Exception:                                                  # noqa: BLE001
                ax.add_patch(plt.Rectangle((x, y), w, h, fc="#dddddd", ec="k", zorder=2))
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            try:
                fill = sh.fill.fore_color.rgb if sh.fill.type == 1 else None
            except Exception:                                                  # noqa: BLE001
                fill = None
            if fill is not None:
                ax.add_patch(plt.Rectangle((x, y), w, h, fc="#" + str(fill), ec="none", zorder=1))
            pt, bold, col = run_style(sh.text_frame)
            lines = wrap(sh.text_frame.text, pt, bold, max(w - 0.2, 0.3))
            lh = pt * 1.2 / 72.0
            for k, ln in enumerate(lines):
                ax.text(x + 0.1, y + 0.05 + k * lh, ln, fontsize=pt, fontweight="bold" if bold else "normal",
                        color=col, va="top", ha="left", family=fam(ln), zorder=3)
            need = 0.1 + len(lines) * lh
            if need > h + 0.05:
                warn.append(f"slide {idx}: text box '{sh.text_frame.text[:40]}' needs {need:.2f} in, box {h:.2f} in")
            if y + need > H or x + w > W + 0.01:
                warn.append(f"slide {idx}: text '{sh.text_frame.text[:40]}' runs off the slide")
        elif sh.shape_type in (1, 17):                                         # autoshape / textbox without text
            try:
                al = sh._element.spPr.xpath(".//a:solidFill/a:srgbClr/a:alpha/@val")          # 10-26: transparency
                alpha = int(al[0]) / 100000 if al else 1.0
                if sh.fill.type == 1:
                    ax.add_patch(plt.Rectangle((x, y), w, h, fc="#" + str(sh.fill.fore_color.rgb), ec="none", alpha=alpha,
                                               zorder=1 if alpha == 1.0 else 4))
                if sh.line.fill.type == 1:
                    ax.add_patch(plt.Rectangle((x, y), w, h, fc="none", ec="#" + str(sh.line.color.rgb), lw=1.2, zorder=4))
            except Exception:                                                  # noqa: BLE001
                pass
        elif sh.shape_type == 9:                                               # 10-26: line / connector (arrows)
            xf = sh._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
            fh, fv = xf is not None and xf.get("flipH") == "1", xf is not None and xf.get("flipV") == "1"
            x0, x1 = (x + w, x) if fh else (x, x + w)
            y0, y1 = (y + h, y) if fv else (y, y + h)
            try:
                col = "#" + str(sh.line.color.rgb)
            except Exception:                                                  # noqa: BLE001
                col = "k"
            head = bool(sh._element.spPr.xpath(".//a:tailEnd[@type!='none']"))
            ax.annotate("", xy=(x1, y1), xytext=(x0, y0), zorder=5,
                        arrowprops=dict(arrowstyle="-|>" if head else "-", color=col, lw=2.5, mutation_scale=18))
        elif sh.has_table:
            ax.add_patch(plt.Rectangle((x, y), w, h, fc="#f0f0f0", ec="k", zorder=1))
    fig.savefig(out, dpi=dpi); plt.close(fig)
    return warn


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("out_dir")
    ap.add_argument("--slides", default="")
    ap.add_argument("--dpi", type=int, default=80)
    a = ap.parse_args(argv)
    os.makedirs(a.out_dir, exist_ok=True)
    prs = Presentation(a.deck)
    want = {int(v) for v in a.slides.split(",") if v}
    stem = os.path.splitext(os.path.basename(a.deck))[0][:40].replace(" ", "_")
    for i, s in enumerate(prs.slides, 1):
        if want and i not in want:
            continue
        for wmsg in render(prs, i, s, os.path.join(a.out_dir, f"{stem}_slide{i:02d}.png"), a.dpi):
            print("WARNING", wmsg)
    print("->", a.out_dir)


if __name__ == "__main__":
    main()
