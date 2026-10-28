"""assemble_master.py - ONE big FULL deck and ONE PI deck with everything made so far (2026-10-22, user: 'put all slides ppt we
made so far in one big PPT ... day1+2 full = day1 (history) + day2 (history) + 1+2 days, and a PI version').

FULL  cover | PART 1 Day 1 + Day 2 (new style, best window) | PART 2 every mouse, Day 1 and Day 2 (best window) | PART 3 Day 1
      alone (new style + the earlier Day-1 slides that do not depend on the response window) | PART 4 Day 2 alone (same) |
      PART 5 the earlier Day 1 + Day 2 (longitudinal) slides that do not depend on the window | ARCHIVE the earlier slides made
      with the OLD 2-s window, each marked 'OLD WINDOW - superseded' (kept so nothing made so far is lost)
PI    cover | the Day 1 + Day 2 PI deck | the Day-1 and the Day-2 summary card

Slides are copied with their pictures, tables, text, notes and movies. Inputs: lab_meeting\\decks_Day1_Day2\\v2_20261022.

  python assemble_master.py --lang en|kr
"""
import argparse
import copy
import io
import os
import re
import sys
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import make_new_deck as ND  # noqa: E402

V = r"C:\Users\hsollim\Documents\HEAL_mini1p_SBI553\lab_meeting\decks_Day1_Day2\v2_20261022"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}

# earlier (old-style) slides that do NOT depend on the event response window (titles; checked by hand on the 10-22 decks)
KEEP_OLD = [r"^Design", r"^From the raw movie", r"^Neurons per mouse", r"^Every neuron over a whole session", r"^Example session",
            r"^Stimulus-block activity after the injection", r"^SBI-553 vs vehicle: change per mouse", r"^What we score",
            r"^Behaviour per delivery, baseline vs injection", r"^Reflexes and affective behaviour per stimulus", r"^A1 ", r"^A2 / A6",
            r"^A3 ", r"^A4 ", r"^A5 ", r"^A7 ", r"^C1 / C2", r"^R1 ", r"^E2 ", r"^E7 ", r"^Reflexes per stimulus",
            r"^Affective and spontaneous behaviours", r"^Behaviour around every delivery", r"^E1 ", r"^How published studies",
            # longitudinal deck
            r"tracking", r"[Tt]racked", r"^The same neurons", r"[Ww]ithin-day change", r"transient rate", r"[Bb]ehaviour",
            r"[Cc]rossover", r"[Mm]ovement and grooming", r"^Scoring quality", r"^How each Day-1 neuron", r"^Within each mouse",
            r"^Cell matching", r"^Annotation quality"]
DEP_OLD = [r"event responses", r"mean response around each event", r"^Day 1 \+ Day 2 \(crossover"]   # old window: archive
DROP_OLD = [r"PER-MOUSE SUPPLEMENT", r"^Day \d per mouse"]                          # replaced by PART 2 (best window)


def title_of(slide):
    tx = sorted(((sh.top or 0, sh.text_frame.text.strip()) for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()))
    t = [x[1] for x in tx]
    return (t[1] if len(t) > 1 else t[0] if t else ""), (t[0] if t else "")


def is_mouse_slide(sec):
    return bool(re.match(r"^MOUSE \d{4}|^마우스 \d{4}", sec))


def clone(dst, slide, tag=None, lang="en", pic=None):
    """copy one slide (shapes, pictures, movies, notes) into dst; tag = a red label at the top right; pic = (png bytes, left, top,
    width, height) in place of the slide's (only) picture"""
    new = dst.slides.add_slide(dst.slide_layouts[6])
    for shp in slide.shapes:
        el = shp._element
        if pic is not None and el.tag.endswith("}pic"):
            new.shapes.add_picture(io.BytesIO(pic[0]), *pic[1:])
            continue
        if el.tag.endswith("}pic"):
            vid = el.findall(".//a:videoFile", NS)
            blip = el.find(".//a:blip", NS)
            poster = None
            if blip is not None:
                rid = blip.get(f"{{{NS['r']}}}embed")
                if rid:
                    poster = slide.part.related_part(rid).blob
            if vid:
                rid = vid[0].get(f"{{{NS['r']}}}link") or vid[0].get(f"{{{NS['r']}}}embed")
                try:
                    blob = slide.part.related_part(rid).blob
                    with tempfile.NamedTemporaryFile(suffix=".mp4", delete=False) as tf:
                        tf.write(blob)
                    new.shapes.add_movie(tf.name, shp.left, shp.top, shp.width, shp.height,
                                         poster_frame_image=io.BytesIO(poster) if poster else None, mime_type="video/mp4")
                    os.unlink(tf.name)
                    continue
                except Exception as e:                                   # fall back to the poster frame
                    print("  movie not copied:", e)
            if poster is not None:
                new.shapes.add_picture(io.BytesIO(poster), shp.left, shp.top, shp.width, shp.height)
            continue
        if el.findall(".//a:blip", NS):                                  # a group or frame with pictures: keep the pictures
            for b in el.findall(".//a:blip", NS):
                rid = b.get(f"{{{NS['r']}}}embed")
                if rid:
                    new.shapes.add_picture(io.BytesIO(slide.part.related_part(rid).blob), shp.left, shp.top, shp.width, shp.height)
            continue
        new.shapes._spTree.insert_element_before(copy.deepcopy(el), "p:extLst")
    if slide.has_notes_slide:
        new.notes_slide.notes_text_frame.text = slide.notes_slide.notes_text_frame.text
    if tag:
        tb = new.shapes.add_textbox(Inches(8.2), Inches(0.05), Inches(5.0), Inches(0.3))
        p = tb.text_frame.paragraphs[0]
        p.text = tag; p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        p.font.name = "Malgun Gothic" if lang == "kr" else "Arial"
    return new


def part_suffix(slide, k, n):
    """append ' (k/n)' to the slide title (the largest text near the top)"""
    best = None
    for sh in slide.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip() or (sh.top or 0) > Inches(1.0):
            continue
        sz = max([r.font.size.pt for pp in sh.text_frame.paragraphs for r in pp.runs if r.font.size] or [0])
        if best is None or sz > best[0]:
            best = (sz, sh)
    if best and best[0] >= 18:
        runs = [r for r in best[1].text_frame.paragraphs[0].runs]
        if runs:
            runs[-1].text = runs[-1].text + f" ({k}/{n})"


def clone_split(dst, slide, tag, lang, min_scale=0.7):
    """an earlier (tagged) slide with one figure shown below min_scale of its drawn size -> one slide per part of the figure"""
    import split_parts as SPL
    from PIL import Image
    from pptx.util import Emu
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    if tag is None or len(pics) != 1 or pics[0]._element.findall(".//a:videoFile", NS):
        return [clone(dst, slide, tag, lang)]
    sh = pics[0]
    try:
        im = Image.open(io.BytesIO(sh.image.blob))
    except Exception:
        return [clone(dst, slide, tag, lang)]
    dpi = im.info.get("dpi", (100, 100))[0] or 100
    if Emu(sh.width).inches / (im.size[0] / dpi) >= min_scale:
        return [clone(dst, slide, tag, lang)]
    top = Emu(sh.top).inches
    below = [Emu(x.top).inches for x in slide.shapes if x.has_text_frame and x.text_frame.text.strip() and Emu(x.top).inches > top + 0.3]
    box_h = (min(below) if below else 6.6) - top - 0.05
    crops, sc = SPL.parts(im, 12.4, box_h, target=min_scale)
    if len(crops) == 1:
        return [clone(dst, slide, tag, lang)]
    out = []
    for k, c in enumerate(crops, 1):
        bio = io.BytesIO()
        c.save(bio, format="PNG", dpi=(dpi, dpi))
        w_in, h_in = c.size[0] / dpi, c.size[1] / dpi
        f_ = min(12.4 / w_in, box_h / h_in)
        new = clone(dst, slide, tag, lang, pic=(bio.getvalue(), Inches(0.45 + (12.4 - w_in * f_) / 2), Inches(top),
                                                Inches(w_in * f_), Inches(h_in * f_)))
        part_suffix(new, k, len(crops))
        out.append(new)
    return out


TESTS = [(r"Wilcoxon", "Wilcoxon signed-rank"), (r"Mann-Whitney", "Mann-Whitney U"), (r"McNemar", "McNemar"),
         (r"Fisher", "Fisher exact"), (r"Mantel|CMH|Cochran", "Cochran-Mantel-Haenszel"), (r"time-shift|shift test|shifted", "time-shift test"),
         (r"max-T", "max-T"), (r"binomial", "binomial test"), (r"Spearman", "Spearman"), (r"GLM", "GLM"), (r"bootstrap", "bootstrap"),
         (r"permutation", "permutation"), (r"[Dd]escriptive|기술 통계", "descriptive")]
NSLIDE = re.compile(r"\bn\s*=\s*\d|\b\d+\s*(mice|mouse)\b|\b2\s*\+\s*2\b|마리|마우스\s*\d|one mouse|한 마리", re.I)
TSLIDE = re.compile("|".join(t for t, _ in TESTS) + "|시간 이동|검정", re.I)   # 10-23: KR '검정: 시간 이동' got an extra English line


def ensure_n_test(slide, lang):
    """new-style slides: if the bottom sentences do not state n animals and the test, add one short line built ONLY from what
    the slide's notes say (nothing invented); nothing is added when the notes do not say it either"""
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    boxes = [sh for sh in slide.shapes if sh.has_text_frame and (sh.top or 0) >= Inches(6.5)]
    if not boxes or not notes:
        return False
    bot = boxes[-1]
    txt = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    need_n, need_t = not NSLIDE.search(txt), not TSLIDE.search(txt)
    if not (need_n or need_t) or len(bot.text_frame.text) > 190:
        return False
    parts = []
    if need_n:
        # only an explicit 'n = K mice' of the notes (a loose 'K mice' took '2 mice each order' as n = 2 on 10-22)
        m = re.search(r"\bn\s*=\s*(\d+)\s*(?:mice|mouse|마리)", notes)
        if m:
            k = m.group(1)
            parts.append(f"n = {k} " + ("마리" if lang == "kr" else ("mouse" if k == "1" else "mice")))
    if need_t:
        ts = [lab for pat, lab in TESTS if re.search(pat, notes)]
        if ts:
            parts.append(", ".join(dict.fromkeys(ts)))
    if not parts:
        return False
    para = bot.text_frame.add_paragraph()
    para.text = "; ".join(parts) + (" (자세한 내용은 노트)" if lang == "kr" else " (details in the notes)")
    para.font.size = Pt(11.5); para.font.name = "Malgun Gothic" if lang == "kr" else "Arial"
    para.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return True


# earlier (old-style) slides: what n and which test apply, by the English title (checked by hand on the 10-22 decks against
# each slide's own sentences; used only when the slide does not say it already). (regex, n EN, n KR, test EN, test KR)
N4 = ("n = 4 mice (2 SBI-553, 2 vehicle)", "n = 4마리 (SBI-553 2, vehicle 2)")
DESC = ("descriptive (no test)", "기술 통계 (검정 없음)")
OLD_NT = [
    (r"^Day \d \(SBI-553 vs vehicle", None, None, "tests: on each result slide (Wilcoxon over the neurons of each mouse; groups of "
     "2 mice shown, not tested)", "검정: 각 결과 슬라이드 (마우스마다 뉴런 Wilcoxon; 2마리 그룹은 보여 주기만, 검정 안 함)"),
    (r"^Design", *N4, *DESC),
    (r"^From the raw movie", "n = 1 example mouse (5941)", "n = 예시 1마리 (5941)", *DESC),
    (r"^Neurons per mouse", *N4, *DESC),
    (r"^Every neuron over a whole session", "n = 1 example mouse (5941)", "n = 예시 1마리 (5941)", *DESC),
    (r"^Example session", "n = 1 example mouse (5941)", "n = 예시 1마리 (5941)", "* = time-shift test (BH q <= 0.20)",
     "* = 시간 이동 검정 (BH q <= 0.20)"),
    (r"^Responsive neurons, activated and inhibited", *N4, None, None),
    (r"^Stimulus-block activity after the injection", *N4, None, None),
    (r"^What we score", *N4, *DESC),
    (r"^Reflexes and affective behaviour per stimulus", *N4, *DESC),
    (r"^A[1-7]\b|^A2 / A6", *N4, *DESC),
    (r"^C1 / C2", *N4, "per mouse: Wilcoxon over neurons; groups of 2 mice: descriptive (not tested)",
     "마우스마다 뉴런 Wilcoxon; 2마리 그룹: 기술 통계 (검정 안 함)"),
    (r"^Summary figure", *N4, "per mouse: Wilcoxon over neurons; groups of 2 mice: descriptive (not tested)",
     "마우스마다 뉴런 Wilcoxon; 2마리 그룹: 기술 통계 (검정 안 함)"),
    (r"^R1 ", *N4, "bootstrap 95% CI over neurons", "뉴런 bootstrap 95% CI"),
    (r"^E2 ", *N4, "linear mixed model (neurons nested in mice)", "선형 혼합 모형 (마우스 안 뉴런)"),
    (r"^B[0-9]+b? |^B1 / B1c", *N4, *DESC),
    (r"^E3 ", *N4, "descriptive (cross-validated R2, no test)", "기술 통계 (교차 검증 R2, 검정 없음)"),
    (r"^E6 |^E8 ", *N4, *DESC),
    (r"^E11 ", "n = 4 mice (8 sessions)", "n = 4마리 (세션 8개)", "cross-validated AUC vs a shuffle null (permutation)",
     "교차 검증 AUC vs 섞은 null (permutation)"),
    (r"^E12 ", *N4, "k-means clustering: descriptive (no test)", "k-means 군집: 기술 통계 (검정 없음)"),
    (r"^E14 ", *N4, "slope of a linear fit per mouse x event", "마우스 x 사건마다 선형 적합의 기울기 검정"),
    (r"^E[0-9]+ ", *N4, *DESC),
    (r"^R[0-9]+ |^R5 / R5b", *N4, *DESC),
    (r"^Behaviour around every delivery|^Reflexes per stimulus|^Affective and spontaneous|^Behaviour per delivery, baseline", *N4, *DESC),
    (r"^How published studies", "n: not applicable (literature)", "n: 해당 없음 (문헌)", "no test (literature table)",
     "검정 없음 (문헌 표)"),
    (r"^TOY Day 2", "n: simulated data (no mice)", "n: 모의 데이터 (마우스 아님)", *DESC),
    (r"^Limits", *N4, *DESC),
    (r"^Each neuron's response|^Single neurons respond", *N4, *DESC),
    (r"^SBI-553 vs vehicle: change per mouse", *N4, "per mouse: Wilcoxon over neurons; groups of 2 mice: descriptive (not tested)",
     "마우스마다 뉴런 Wilcoxon; 2마리 그룹: 기술 통계 (검정 안 함)"),
    # Day 1 + Day 2 (longitudinal) deck
    (r"^Day 1 \+ Day 2 \(crossover", "n = 4 mice (crossover)", "n = 4마리 (crossover)", "tests: on each result slide",
     "검정: 각 결과 슬라이드"),
    (r"^Behaviour per delivery: vehicle day vs SBI-553 day|^Within each mouse", "n = 4 mice (crossover)", "n = 4마리 (crossover)",
     "Poisson GLM (events ~ session x drug)", "Poisson GLM (사건 ~ 세션 x 약물)"),
    (r"^Tracked neurons: event responses", "n = 4 mice (tracked neurons of each mouse)", "n = 4마리 (마우스마다 추적 뉴런)",
     "Friedman + Wilcoxon over tracked neurons", "추적 뉴런 Friedman + Wilcoxon"),
    (r"^Tracked neurons: mean response", "n = 4 mice (crossover)", "n = 4마리 (crossover)",
     "descriptive (the tests are on the previous slides)", "기술 통계 (검정은 앞 슬라이드)"),
    (r"^Tracked neurons: baseline vs vehicle vs SBI-553|^Four mice, each its own control", "n = 4 mice (crossover)", "n = 4마리 (crossover)",
     "Wilcoxon over tracked neurons", "추적 뉴런 Wilcoxon"),
    (r".", "n = 4 mice (crossover)", "n = 4마리 (crossover)", *DESC),          # tracking, matching, quality, design: descriptive
]


def old_nt(slide, en_title, lang):
    """add the n / test of OLD_NT that the slide does not state yet (audit_rules' N_RE / T_RE decide what is stated)"""
    import audit_rules as AR
    row = next((r for r in OLD_NT if re.search(r[0], en_title)), None)
    if row is None:
        return False
    txt = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    kr = lang == "kr"
    parts = []
    if not AR.N_RE.search(txt) and row[1]:
        parts.append(row[2] if kr else row[1])
    if not AR.T_RE.search(txt) and row[3]:
        parts.append(row[4] if kr else row[3])
    if not parts:
        return False
    boxes = sorted([sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip() and (sh.top or 0) >= Inches(5.4)],
                   key=lambda sh: sh.top)
    if boxes:
        para = boxes[-1].text_frame.add_paragraph()
    else:
        tb = slide.shapes.add_textbox(Inches(0.45), Inches(6.95), Inches(12.4), Inches(0.4))
        para = tb.text_frame.paragraphs[0]
    para.text = "; ".join(parts) + "."
    para.font.size = Pt(11.5); para.font.name = "Malgun Gothic" if kr else "Arial"
    para.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return True


# the ARCHIVE index: each earlier OLD-window analysis -> the best-window slides that replace it (PART 1 titles)
ARCH_MAP = [
    ("B1 / B1c, B1b  responsive neurons per event", "How many neurons respond to each stimulus / behaviour (best window); Responsive vs non-responsive",
     "B1 / B1c, B1b  사건별 반응 뉴런", "사건별 반응 뉴런 수 (최적 창); 반응 vs 무반응"),
    ("B2, B4  event-locked mean response, population AUC", "Population response to each stimulus / to each behaviour",
     "B2, B4  사건 정렬 평균 반응, 집단 AUC", "자극별 / 행동별 집단 반응"),
    ("B3  every neuron's mean response", "EVERY EVENT: <event>: every neuron in the three conditions",
     "B3  모든 뉴런의 평균 반응", "EVERY EVENT: <사건>: 세 조건의 모든 뉴런"),
    ("B5  UP, DOWN and other neurons", "EVERY EVENT: <event>: UP, DOWN and the other neurons",
     "B5  UP, DOWN, 그 외 뉴런", "EVERY EVENT: <사건>: UP, DOWN, 그 외 뉴런"),
    ("B6 / B6b  where the responsive neurons are", "Where the <event> neurons are; SUPPLEMENT <mouse>: where the responsive neurons are",
     "B6 / B6b  반응 뉴런의 위치", "<사건> 뉴런의 위치; 부록 <마우스>: 반응 뉴런의 위치"),
    ("B8, E3  with vs without a withdrawal; encoding model", "Stimulus, withdrawal and flinch come almost together; Check 3",
     "B8, E3  withdrawal 있음 vs 없음; encoding model", "자극, withdrawal, flinch는 거의 동시; 확인 3"),
    ("B7, E13  decoding the stimulus", "Other way 4 and Other way 5 (decoding)", "B7, E13  자극 해독", "다른 방법 4, 5 (해독)"),
    ("E4, R5 / R5b, E14  the same neurons after the injection", "Do the same neurons keep responding after the injection?; What happens to "
     "the baseline-responsive neurons after the injection", "E4, R5 / R5b, E14  주사 후 같은 뉴런", "주사 후에도 같은 뉴런이 반응하는가?; 기준 반응 뉴런은 주사 후 어떻게 되나"),
    ("E9  clustered in space", "Are the responsive neurons clustered in the field of view?", "E9  공간 군집", "반응 뉴런이 시야에 모여 있는가?"),
    ("E12, R3  response types and shapes", "Other way 2 (functional types); Other way 3 (response shapes)",
     "E12, R3  반응 유형과 모양", "다른 방법 2 (기능 유형); 다른 방법 3 (반응 모양)"),
    ("R8  max-T count", "Other way 1 (max-T); Check 1 and Check 2", "R8  max-T 수", "다른 방법 1 (max-T); 확인 1, 2"),
    ("B10  brain motion at the events", "DATA QUALITY: Why does '% responsive' jump between sessions?",
     "B10  사건 때 뇌 움직임", "데이터 품질: 왜 '% 반응'이 세션마다 튀는가?"),
    ("Tracked neurons: event responses (Day 1 + 2 deck)", "EVERY EVENT: <event>: every neuron in the three conditions; Population response",
     "추적 뉴런: 사건 반응 (Day 1+2 덱)", "EVERY EVENT: <사건>: 세 조건의 모든 뉴런; 집단 반응"),
    ("NOT YET REDONE: B9, E5, E6, E8, E10, E11, R4, R6, R7, R9, R10", "only the OLD 2-s window so far (being recomputed with the best window)",
     "아직 다시 안 함: B9, E5, E6, E8, E10, E11, R4, R6, R7, R9, R10", "아직 옛 2 s 창만 (최적 창으로 다시 계산 중)"),
]


def archive_index(dk, lang, redone=False):
    """two table slides: old analysis -> replacement; redone = the best-window re-run decks are in PARTS 3 / 4"""
    from pptx.util import Inches as In, Pt as P_
    kr = lang == "kr"
    amap = ARCH_MAP
    if redone:
        amap = ARCH_MAP[:-1] + [("ALL OF THEM (B, E, R, summary figure)", "also re-run with the best window: PART 3 / PART 4, tag 're-run "
                                 "with the BEST window'", "전부 (B, E, R, 요약 그림)", "최적 창으로도 다시 계산: PART 3 / PART 4, 표시 '최적 창으로 다시 계산'")]
    rows = [(r[2], r[3]) if kr else (r[0], r[1]) for r in amap]
    half = (len(rows) + 1) // 2
    for k, chunk in enumerate((rows[:half], rows[half:]), 1):
        s = dk.head("ARCHIVE" if not kr else "기록",
                    (f"Archive index ({k}/2): which best-window slide replaces each OLD-window slide" if not kr else
                     f"기록 안내 ({k}/2): 옛 창 슬라이드마다 대체한 최적 창 슬라이드"),
                    ("The archive keeps every earlier slide (OLD 2-s window, marks before the 10-23 fix); use the PART 1 slides on the right"
                     if not kr else "기록 = 옛 2 s 창·10-23 보정 전 표시의 이전 슬라이드 전부; 오른쪽의 PART 1 슬라이드를 쓰세요"))
        tb = s.shapes.add_table(len(chunk) + 1, 2, In(0.45), In(1.45), In(12.4), In(0.5 * (len(chunk) + 1))).table
        tb.columns[0].width = In(4.9); tb.columns[1].width = In(7.5)
        for j, h in enumerate(("earlier OLD-window slide", "best-window replacement (PART 1)") if not kr else ("옛 창 슬라이드", "최적 창 대체 (PART 1)")):
            c = tb.cell(0, j); c.text = h
            for pp in c.text_frame.paragraphs:
                pp.font.size = P_(13); pp.font.bold = True; pp.font.name = dk.font
        for i, (a_, b_) in enumerate(chunk, 1):
            for j, t_ in enumerate((a_, b_)):
                c = tb.cell(i, j); c.text = t_
                for pp in c.text_frame.paragraphs:
                    pp.font.size = P_(12); pp.font.name = dk.font
                    if t_.startswith(("NOT YET", "아직")) or "being recomputed" in t_ or "다시 계산 중" in t_:
                        pp.font.bold = True; pp.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        dk.box(s, 0.45, 6.68, 12.4, 0.8, ("Index of the archive (no data). n = 4 mice for every analysis listed; descriptive (no test)."
                                          if not kr else "기록의 목록 (데이터 아님). 모든 분석 n = 4마리; 기술 통계 (검정 없음)."), 13.5)
        s.notes_slide.notes_text_frame.text = ("Archive = the earlier Day-1 / Day-2 / Day 1 + 2 slides whose numbers depend on the event "
                                               "response window (old: 0-2 s vs -2-0 s). Window-independent earlier slides are in PARTS 3-5.")


# ------------------------------------------------------------------ story order (10-22 night, user: 'organise the decks well:
# first get attention, then the questions solved in order, overall to details'). Key of a slide = 'SECTION|title' (English).
EVN = [("Light touch", "light touch"), ("Mild touch", "mild touch"), ("Pin prick", "pin prick"), ("Heat", "heat"),
       ("withdrawal", "withdrawal"), ("flinch", "flinch"), ("attending onset", "attending"), ("movement onset", "movement onset"),
       ("grooming onset", "grooming"), ("escape/rearing onset", "escape / rearing")]
_EV = [q for A, b in EVN for q in (r"\|" + re.escape(A) + r": every neuron in the three conditions", r"\|" + re.escape(A) + r": UP, DOWN",
                                    r"\|Where the " + re.escape(A) + r" neurons are", r"\|" + re.escape(b) + r": every neuron by its class")]
STORY = [
    ("cover", [r"^COVER$"]),
    ("attention", [r"^PIPE:video$"]),
    ("question", [r"\|Crossover: vehicle on one day", r"\|The same neurons on Day 1 and Day 2"]),
    ("glance", [r"^GLANCE2$"]),
    ("methods", [r"^PIPE:glance$", r"^PIM\|.*Design: one field of view", r"^PIM\|.*From the raw movie to neurons",
                 r"^PIM\|.*Neurons per mouse", r"^PIPE:logger$", r"^PIPE:beh1$", r"^PIPE:beh2$", r"^PIPE:bvideo$", r"^PIPE:auto$",
                 r"^X-methods\|"]),
    ("Q1 behaviour", [r"\|Behaviour per delivery: every mouse", r"\|Every mouse, every behaviour", r"\|Behaviour per stimulus, all mice",
                      r"\|Behaviour per stimulus", r"\|Affective and reflexive", r"\|Movement before any stimulus",
                      r"\|Time spent in each behaviour", r"^X-Q1 behaviour\|"]),
    ("Q2 activity", [r"^X-Q2 activity\|", r"\|The same, averaged over the 4 mice", r"\|The same average, with its colour scale", r"\|Each injection next to the baseline",
                     r"\|Activity of all neurons", r"\|Every neuron across the three conditions", r"every neuron through all four sessions",
                     r"\|Where in the field of view activity changed"]),
    ("Q3 population", [r"\|Population response to each stimulus", r"\|Population response to each behaviour", r"^X-Q3 population\|",
                       # 10-23 (user): the same neurons on both days, baseline-normalized (deck_tracked_slides.py)
                       r"\|Same neurons on both days", r"\|% UP and DOWN, baseline-normalized", r": the same neurons in all four sessions", r": the same neurons, baseline-normalized",
                       r": each neuron through the four sessions", r"\|Other way 4", r"\|Other way 5"]),
    ("Q3 neurons", [r"^EVERY NEURON\|Heat: every neuron", r"^PIX\|.*\|Pin prick: every neuron", r"\|How many neurons respond to each stimulus",
                    r"\|How many neurons respond to each behaviour", r"\|Responsive vs non-responsive", r"^X-Q3 neurons\|",
                    r"\|Summary: responsive neurons per event",
                    r"RESPONSIVE NEURONS\|UP, DOWN and the other neurons", r"\|Responsive neurons session by session",
                    r"\|Where the Pin prick neurons are, session by session", r"\|Are the responsive neurons clustered",
                    r"\|Do the same neurons keep responding", r"\|The same neurons from baseline to injection", r"\|Stable across days",
                    r"\|What happens to the baseline-responsive"]),
    ("Q3 every event", _EV),
    ("Q4 stimulus or movement", [r"\|Stimulus, withdrawal and flinch", r"\|Do sessions with more behaviour", r"^X-Q4 stimulus or movement\|"]),
    ("Q3 other ways", [r"\|Other way 1", r"\|Other way 2", r"\|Other way 3"]),
    ("Q5 sex", [r"\|Is there a pattern by sex", r"by sex", r"\|What sex can", r"^X-Q5 sex\|"]),
    ("other", []),
    # 10-23 round 2: the window slides and check 3 inside HOW SURE? (checks 1-2-3-4 in order), then the Day-alone cards
    ("how sure", [r"\|Check 1", r"\|Check 2", r"\|Your windows", r"\|Exploring start", r"\|Long windows", r"\|When do the responses happen",
                  r"\|Why does '% responsive'", r"\|Check 3", r"\|Check 4", r"^DAYCARD1$", r"^DAYCARD2$", r"\|Day 1: SBI-553 vs vehicle",
                  r"\|Day 2: SBI-553 vs vehicle"]),
    ("limits", [r"\|What the data can and cannot", r"^LIMITS2$"]),
    ("summary", [r"^SUMMARY1$"]),
    ("supplement", [r"^X-supplement\|", r"where the responsive neurons are \(", r"\|Day 1 \+ Day 2: SBI-553 vs vehicle in the same mice"]),
]
_GNAMES = [g for g, _ in STORY]
SECLAB = {"attention": ("THE DATA", "데이터"), "question": ("QUESTION AND DESIGN", "질문과 설계"), "glance": ("RESULTS AT A GLANCE", "결과 한눈에"),
          "methods": ("HOW WE DID IT", "방법"), "Q1 behaviour": ("Q1  BEHAVIOUR", "Q1  행동"),
          "Q2 activity": ("Q2  BLA ACTIVITY", "Q2  BLA 활동"), "Q3 population": ("Q3  RESPONSES: POPULATION", "Q3  반응: 집단"),
          "Q3 neurons": ("Q3  RESPONSES: SINGLE NEURONS", "Q3  반응: 단일 뉴런"),
          "Q4 stimulus or movement": ("Q4  STIMULUS OR MOVEMENT?", "Q4  자극인가 움직임인가?"),
          "Q3 other ways": ("Q3  OTHER ANALYSES", "Q3  다른 분석"), "Q5 sex": ("Q5  SEX", "Q5  성별"),
          "how sure": ("HOW SURE?", "얼마나 확실한가?"), "limits": ("LIMITS", "한계"), "summary": ("SUMMARY", "요약"),
          "supplement": ("SUPPLEMENT", "부록")}


def story_rank(key):
    if key == "OTHER":
        return None
    for gi, (_, pats) in enumerate(STORY):
        for pi, pat in enumerate(pats):
            if re.search(pat, key):
                return (gi, pi)
    return (_GNAMES.index("other"), 0)


def reorder(prs, keys, lang):
    """story slides first (by STORY group, pattern, then their order), then the rest ('OTHER') in their order"""
    sld = prs.slides._sldIdLst
    items = list(sld)
    assert len(items) == len(keys), (len(items), len(keys))
    rk = [story_rank(k) for k in keys]
    st = sorted([i for i, r in enumerate(rk) if r is not None], key=lambda i: (rk[i], i))
    order = st + [i for i, r in enumerate(rk) if r is None]
    slides = list(prs.slides)
    for el in items:
        sld.remove(el)
    for i in order:
        sld.append(items[i])
    # the section label of every story slide = its question (10-22 story judge: the old labels alternated and did not match Q1-Q5)
    for i in order:
        if rk[i] is None:
            continue
        g = _GNAMES[rk[i][0]]
        lab = SECLAB.get(g)
        for sh in slides[i].shapes:
            if not (sh.has_text_frame and (sh.top or 0) < Inches(0.3) and sh.text_frame.text.strip()):
                continue
            t0 = sh.text_frame.text.strip()
            if (sh.left or 0) > Inches(7):
                continue                                       # the red tags at the top right stay
            r_ = sh.text_frame.paragraphs[0].runs
            if not r_:
                continue
            if g == "Q3 every event":
                new = ("Q3 " + t0) if t0.startswith(("EVERY EVENT", "모든 사건")) else ("Q3 EVERY EVENT" if lang != "kr" else "Q3 모든 사건")
            elif lab:
                new = lab[1] if lang == "kr" else lab[0]
            else:
                continue
            r_[0].text = new
            for x in r_[1:]:
                x.text = ""
            break
    return [_GNAMES[rk[i][0]] if rk[i] is not None else "OTHER" for i in order]


# take-home message of each Day 1 + Day 2 result slide, PER PART of a split figure (10-22 night triple check: the parts had copied
# the first part's line). Numbers checked on the 10-22 data (cells_classes.csv, P5_fate_tests.csv, P6_behaviour_tests.csv,
# P7_correlations.csv, V4_key_findings_per_mouse.csv, summary.json 'overlap', validation.json 'spikein', the sex rate table).
# (title regex on the English title, take-home EN, take-home KR, {title_en, title_kr, bottom_en, bottom_kr})
# bottoms shared by several TAKE rows (10-23)
B_RV_EN = 'Solid = activated, hatched = inhibited, in the condition colour (grey = baseline, blue = vehicle, red = SBI-553); the rest = non-responsive. Stars = neurons pooled (vehicle day 216, SBI-553 day 246; exploratory); per mouse and strict test: check 4. n = 4 mice (crossover); Test: time-shift test (per neuron, BH q <= 0.20), McNemar, Fisher exact.'
B_RV_KR = '채운 막대 = 활성화, 빗금 = 억제, 색 = 조건 (회색 기준, 파랑 vehicle, 빨강 SBI-553); 나머지 = 무반응. 마우스별·엄격한 검정: 검증 4. 별표 = 뉴런 합산 (vehicle 날 216, SBI-553 날 246; 탐색적). n = 4마리 (crossover); 검정: 시간 이동 검정 (뉴런별, BH q <= 0.20), McNemar, Fisher exact.'
B_P5_EN = 'Neurons chosen on half of the baseline trials, measured on the other half (no regression to the mean). n in panels = cells pooled over 4 mice (exploratory){extra}. Test: Wilcoxon signed-rank, Mann-Whitney U.'
B_P5_KR = '기준 trial 절반으로 뉴런을 고르고 나머지 절반으로 측정 (평균 회귀 아님). 패널의 n = 4마리 뉴런 합산 (탐색적){extra}. 검정: Wilcoxon signed-rank, Mann-Whitney U.'
TAKE = [
    (r"^Other way 2: functional types of all neurons", None, None,
     dict(bottom_en="Same neurons, baseline -> injection: 75% stay in their type with vehicle, 63% with SBI-553. n = 4 mice (crossover); "
                    "k-means clustering: descriptive (no test).",
          bottom_kr="같은 뉴런, 기준 -> 주사: vehicle 75%, SBI-553 63%가 같은 유형에 머묾. n = 4마리 (crossover); k-means 군집: 기술 통계 "
                    "(검정 없음).")),
    (r"^Exploring start, end, length and delay of the window", None, None,
     dict(bottom_en="Each cell = one window (rows: baseline length before the event; columns: response window). n = 4 mice (crossover); "
                    "descriptive: real minus a jitter sham (no test).")),
    (r"^Long windows: slow changes, not single responses", None, None,
     dict(bottom_en="Same grid, scored against the block sham (any change, slow ones included) instead of the 2-5 s jitter. n = 4 mice "
                    "(crossover); descriptive (no test).")),
    (r"^Crossover: vehicle on one day",
     "Question: does SBI-553 change pain behaviour and BLA pain responses? Crossover: each mouse is its own control",
     "질문: SBI-553가 통증 행동과 BLA 통증 반응을 바꾸나? crossover: 각 마우스가 자기 대조군",
     dict(notes_add_en="Dose, route and the time from the injection to the injection session are not in the analysis files: add them "
                       "on this slide before the meeting.",
          notes_add_kr="용량, 투여 경로, 주사부터 주사 세션까지의 시간은 분석 파일에 없음: 회의 전에 이 슬라이드에 추가.")),
    (r"^The same, averaged over the 4 mice",
     "Stimuli raise population activity at baseline (+0.10 z) and after vehicle (+0.07 z), barely after SBI-553 (+0.04 z)",
     "자극이 시작되면 집단 활동 증가: 기준 +0.10 z, vehicle 후 +0.07 z, SBI-553 후 +0.04 z로 가장 적음",
     dict(title_en="Whole-session population activity, averaged over the 4 mice", title_kr="세션 전체 집단 활동: 4마리 평균")),
    (r"^Each injection next to the baseline of its own day \(1/2",
     "SBI-553 day: the stimulus-block rise shrank from +0.13 z at baseline to +0.04 z after SBI-553 (descriptive)",
     "SBI-553 날: 자극 블록 증가가 기준 +0.13 z에서 SBI-553 후 +0.04 z로 줄어듦 (기술 통계)", {}),
    (r"^Each injection next to the baseline of its own day \(2/2",
     "Vehicle day: the stimulus-block rise stayed about the same (+0.08 z at baseline, +0.07 z after vehicle; descriptive)",
     "vehicle 날: 자극 블록 증가가 거의 같음 (기준 +0.08 z, vehicle 후 +0.07 z; 기술 통계)", {}),
    (r"^Activity of all neurons: baseline vs vehicle vs SBI-553",
     "Firing fell after both injections; the same-day drop was bigger after SBI-553 in 3/4 mice (B); after SBI-553 vs vehicle 1.3 vs 2.7/min (p = 0.01, Day 2)",
     "두 주사 모두 활동 감소; 같은 날 감소는 SBI-553 후 더 큼 (3/4마리, B); 주사 후 수준 1.3 vs 2.7/분 (p = 0.01, Day 2 결과)",
     dict(notes_add_en="Same-day drop of the median stimulus-block rate: Day 1 SBI-553 mice 7.7 -> 3.5 (-54%) vs vehicle mice 4.4 -> 2.7 "
                       "(-37%); Day 2 SBI-553 mice 0.80 -> 0.26 (-68%) vs vehicle mice 3.9 -> 2.5 (-36%). The level after the injection "
                       "(1.3 vs 2.7 /min, neurons pooled, Mann-Whitney p = 0.01, exploratory) comes from Day 2: on Day 1 it is 3.5 vs 2.7 "
                       "(p = 0.56). Pooled baselines 4.3 vs 4.0 hide this day split.",
          notes_add_kr="같은 날 감소 (자극 블록 rate 중앙값): Day 1 SBI-553 마우스 7.7 -> 3.5 (-54%) vs vehicle 마우스 4.4 -> 2.7 (-37%); "
                       "Day 2 SBI-553 마우스 0.80 -> 0.26 (-68%) vs vehicle 마우스 3.9 -> 2.5 (-36%). 주사 후 수준 (1.3 vs 2.7/분, 뉴런 합산, "
                       "Mann-Whitney p = 0.01, 탐색적)은 Day 2에서 옴: Day 1은 3.5 vs 2.7 (p = 0.56). 합친 기준 4.3 vs 4.0은 이 차이를 가림.")),
    (r"^Where in the field of view activity changed.*\(1/2",
     "5941 and 5939: firing fell more on the SBI-553 day than the vehicle day (whole-session median log2 -1.77 vs -0.52; -0.33 vs +0.01)",
     "5941·5939: SBI-553 날 vehicle 날보다 더 감소 (세션 전체 log2 중앙값 -1.77 vs -0.52; -0.33 vs +0.01)", {}),
    (r"^Where in the field of view activity changed.*\(2/2",
     "1191 fell more on the SBI-553 day (whole-session log2 -0.83 vs -0.13); 5937 did not (0.00 vs -0.53): half its neurons fired in neither session",
     "1191은 SBI-553 날 더 감소 (세션 전체 log2 -0.83 vs -0.13); 5937은 아님 (0.00 vs -0.53): 뉴런 절반이 두 세션 모두 활동 없음", {}),
    (r"^Population response to each stimulus",
     "Only pin prick: smaller after SBI-553 than after vehicle (+0.002 vs +0.033 z, p = 0.04, neurons pooled); every other comparison ns",
     "pin prick만: SBI-553 후가 vehicle 후보다 작음 (+0.002 vs +0.033 z, p = 0.04, 뉴런 합산); 나머지 비교는 모두 ns",
     dict(bottom_en="Response = 0-0.5 s minus the 0.5 s before; except for pin prick after SBI-553 (+0.004 z), that 0.5 s dips below the earlier 3 s (-0.004 to -0.023 z), so part of each "
                    "small response is recovery. n = 4 mice; brackets = the pair tested (neurons). Test: Wilcoxon signed-rank, Mann-Whitney U.",
          bottom_kr="반응 = 0-0.5 s에서 직전 0.5 s를 뺀 값; SBI-553 후 pin prick (+0.004 z)을 빼면 그 0.5 s가 앞선 3 s보다 낮아서 (-0.004 ~ -0.023 z) 작은 반응의 일부는 회복분. "
                    "n = 4마리; 괄호 = 비교한 두 조건 (뉴런 단위). 검정: Wilcoxon signed-rank, Mann-Whitney U.",
          notes_add_en="Against the earlier 3 s (-3 to -0.5 s) the 0-0.5 s responses are only -0.012 to +0.016 z (all stimuli and conditions; "
                       "pin prick: baseline -0.009, vehicle +0.006, SBI-553 +0.007, so the SBI-553 vs vehicle gap comes from the 0.5 s before). Checked 10-23 on the trace cache.",
          notes_add_kr="앞선 3 s (-3 ~ -0.5 s) 대비 0-0.5 s 반응은 -0.012 ~ +0.016 z뿐 (모든 자극·조건; pin prick: 기준 -0.009, vehicle "
                       "+0.006, SBI-553 +0.007; SBI-553 vs vehicle 차이는 직전 0.5 s에서 옴). 10-23 trace 캐시로 확인.")),
    (r"^Population response to each behaviour.*\(1/2",
     "Withdrawal, flinch, attending: no significant change (withdrawal after SBI-553 +0.008 vs +0.035 z after vehicle, p = 0.11)",
     "withdrawal·flinch·attending: 유의한 변화 없음 (withdrawal: SBI-553 후 +0.008 vs vehicle 후 +0.035 z, p = 0.11)",
     dict(          bottom_en="We measured every neuron's response in the best window (panel title). n = 4 mice (attending: 1-2 mice per bar); "
                    "brackets = the pair tested (neurons); no bar = fewer than 5 events in that session. Test: Wilcoxon signed-rank, Mann-Whitney U.",
          bottom_kr="모든 뉴런의 반응을 최적 창(패널 제목)으로 측정. n = 4마리 (attending: 막대당 1-2마리); 괄호 = 비교한 두 조건 (뉴런 "
                    "단위); 막대 없음 = 그 세션의 사건 5개 미만. 검정: Wilcoxon signed-rank, Mann-Whitney U.")),
    (r"^Population response to each behaviour.*\(2/2",
     "Movement onset: SBI-553 vs vehicle ns (+0.022 vs +0.026 z), above its baseline after SBI-553 (*); too few escape / grooming events",
     "movement onset: SBI-553 vs vehicle ns (+0.022 vs +0.026 z), SBI-553 후 기준보다 높음 (*); SBI-553 후 escape·grooming 사건 부족",
     dict(          bottom_en="We measured every neuron's response in the best window (panel title). n = 4 mice (attending: 1-2 mice per bar); "
                    "brackets = the pair tested (neurons); no bar = fewer than 5 events in that session. Test: Wilcoxon signed-rank, Mann-Whitney U.",
          bottom_kr="모든 뉴런의 반응을 최적 창(패널 제목)으로 측정. n = 4마리 (attending: 막대당 1-2마리); 괄호 = 비교한 두 조건 (뉴런 "
                    "단위); 막대 없음 = 그 세션의 사건 5개 미만. 검정: Wilcoxon signed-rank, Mann-Whitney U.")),
    (r"^Pin prick: every neuron in the three conditions",
     "Pin-prick mean response: smaller after SBI-553 (vehicle +0.033 vs SBI-553 +0.002 z, p = 0.04, neurons pooled); fewer neurons respond too",
     "pin prick 평균 반응: SBI-553 후 더 작음 (vehicle 후 +0.033 vs SBI-553 후 +0.002 z, p = 0.04, 뉴런 합산); 반응하는 뉴런도 더 적음",
     dict(notes_add_en="Baseline +0.014 = both baselines pooled (the heatmap's baseline panel); the UP / DOWN counts per stimulus are on the "
                       "'Neurons UP and DOWN after each stimulus' slide.",
          notes_add_kr="기준 +0.014 = 두 기준 세션 합산 (히트맵의 기준 패널); 자극별 UP / DOWN 수는 '자극마다 UP·DOWN 뉴런 비율' 슬라이드.")),
    (r"^Heat: every neuron in the three conditions",
     "Heat responses are small on average (baseline / vehicle / SBI-553: +0.017 / +0.034 / +0.011 z)",
     "heat 반응은 평균적으로 작음 (기준 / vehicle / SBI-553: +0.017 / +0.034 / +0.011 z)", {}),
    (r"^How many neurons respond to each stimulus",
     "Pin-prick UP neurons rose after vehicle (12 -> 22%, ***), not after SBI-553 (11 -> 10%, ns), neurons pooled; the rise is mostly 5939",
     "pin prick UP 뉴런: vehicle 후 증가 (12 -> 22%, ***), SBI-553 후 변화 없음 (11 -> 10%, ns), 뉴런 합산; 증가는 주로 5939",
     dict(title_en="Neurons UP and DOWN after each stimulus (best window)", title_kr="자극마다 UP·DOWN 뉴런 비율 (최적 창)",
          bottom_en="Solid = % of neurons that rose after the event (UP), hatched = % that fell (DOWN); markers = mice. Stars = neurons "
                    "pooled (vehicle day 216, SBI-553 day 246; exploratory); per mouse and max-T: check 4. n = 4 mice (crossover); Test: "
                    "time-shift test (per neuron, BH q <= 0.20), McNemar, Fisher exact.",
          bottom_kr="단색 막대 = 사건 후 상승한 뉴런 %(UP), 빗금 막대 = 하강 %(DOWN); 표식 = 마우스. 별표 = 뉴런 합산 (vehicle 날 216, "
                    "SBI-553 날 246; 탐색적); 마우스별·max-T: 검증 4. n = 4마리 (crossover); 검정: 시간 이동 검정 (뉴런별, BH q <= 0.20), "
                    "McNemar, Fisher exact.",
          notes_add_en="Per mouse (pin UP): vehicle day 5939 2 -> 21 / 47 rose, 5937 1 -> 2 / 24, 5941 22 -> 24, 1191 0 -> 0; SBI-553 day "
                       "5941 10 -> 7 / 67 and 1191 18 -> 6 / 86 fell, 5937 0 -> 7, 5939 0 -> 4 rose. The SBI-553 vs vehicle gap after the injection "
                       "is Day 2 only (42 vs 12%; Day 1 2 vs 8%, reversed). Heat UP rose after both (8 -> 23%, 2 -> 24%); light UP rose after SBI-553 only "
                       "(2 -> 12%; vehicle 10 -> 8%). McNemar: pin vehicle p < 0.001, SBI-553 p = 0.62; heat both p < 0.001; light vehicle p = 0.51, "
                       "SBI-553 p < 0.001. Fisher after vehicle vs after SBI-553: pin 22 vs 10% (p < 0.001), heat 23 vs 24% (ns), light "
                       "8 vs 12% (ns).",
          notes_add_kr="마우스별 (pin UP): vehicle 날 5939 2 -> 21 / 47 증가, 5937 1 -> 2 / 24, 5941 22 -> 24, 1191 0 -> 0; SBI-553 날 5941 "
                       "10 -> 7 / 67, 1191 18 -> 6 / 86 감소, 5937 0 -> 7, 5939 0 -> 4 증가. 주사 후 SBI-553 vs vehicle 차이는 Day 2에만 (42 vs 12%; "
                       "Day 1 2 vs 8%, 반대). heat UP은 둘 다 증가 (8 -> 23%, 2 -> 24%); light UP은 SBI-553 후에만 증가 (2 -> 12%; vehicle 10 -> 8%). McNemar: "
                       "pin vehicle p < 0.001, SBI-553 p = 0.62; heat 둘 다 p < 0.001; light vehicle p = 0.51, SBI-553 p < 0.001. Fisher "
                       "vehicle 후 vs SBI-553 후: pin 22 vs 10% (p < 0.001), heat 23 vs 24% (ns), light 8 vs 12% (ns).")),
    (r"^Responsive vs non-responsive, with the chosen window \(1/3",
     "Vs the same-day baseline: pin prick gained responders after vehicle (16 -> 31% ***), not after SBI-553 (25 -> 21%, ns); light and heat gained after SBI-553",
     "같은 날 기준 대비: pin prick 반응 뉴런은 vehicle 후 증가 (16 -> 31% ***), SBI-553 후 변화 없음 (25 -> 21%, ns); light·heat은 SBI-553 후 증가",
     dict(bottom_en=B_RV_EN, bottom_kr=B_RV_KR,
          notes_add_en="Responsive (UP + DOWN), same neurons (McNemar): SBI-553 day light 3 -> 24%, mild 19 -> 19% (p = 1.0), pin 25 -> 21% (p = 0.22), "
                       "heat 2 -> 44%; vehicle day light 14 -> 12% (p = 0.58), mild 20 -> 8% (fell, p < 0.001), pin 16 -> 31% (mostly 5939), heat "
                       "11 -> 36%. After vehicle vs after SBI-553 (black, Fisher): pin 31 vs 21% fewer after SBI-553 (*), heat 36 vs 44% "
                       "(ns); light 12 vs 24% (***) and mild 8 vs 19% (***) more after SBI-553.",
          notes_add_kr="반응 (UP + DOWN), 같은 뉴런 (McNemar): SBI-553 날 light 3 -> 24%, mild 19 -> 19% (p = 1.0), pin 25 -> 21% (p = 0.22), heat "
                       "2 -> 44%; vehicle 날 light 14 -> 12% (p = 0.58), mild 20 -> 8% (감소, p < 0.001), pin 16 -> 31% (주로 5939), heat 11 -> 36%. "
                       "vehicle 후 vs SBI-553 후 (검은 괄호, Fisher): pin 31 vs 21%는 SBI-553 후 적음 (*), heat 36 vs 44%는 차이 없음 (ns); light 12 vs "
                       "24% (***), mild 8 vs 19% (***)는 SBI-553 후 많음.")),
    (r"^Responsive vs non-responsive, with the chosen window \(2/3",
     "Withdrawal / flinch: more responders after SBI-553 than vehicle (61 vs 45%, 50 vs 36%, **), all from 1191 (Day 1); movement rose after both (***)",
     "withdrawal·flinch: 반응 뉴런이 vehicle 후보다 SBI-553 후 많음 (61 vs 45%, 50 vs 36%, **), 전부 1191 때문 (Day 1); movement는 둘 다 증가 (***)",
     dict(bottom_en=B_RV_EN, bottom_kr=B_RV_KR,
          notes_en="Key numbers (responsive %, UP + DOWN): withdrawal 35 -> 45% (vehicle day), 49 -> 61% (SBI-553 day), after vehicle vs "
                   "after SBI-553 p = 0.0015 (Fisher); flinch 27 -> 36%, 25 -> 50%, p = 0.003; movement onset 42 -> 58%, 44 -> 61%, p = 0.64. After the injection, vehicle vs SBI-553: Day 1 withdrawal 27 vs 65%, flinch 6 vs 50%; Day 2 65 vs 55%, 65 vs 49% (reversed). 1191 alone: after vehicle 25%, 2% vs after SBI-553 60%, 57%; without 1191, SBI-553 vs vehicle: 61 vs 59%, 46 vs 57% (both ns).",
          notes_kr="핵심 숫자 (반응 %, UP + DOWN): withdrawal 35 -> 45% (vehicle 날), 49 -> 61% (SBI-553 날), vehicle 후 vs SBI-553 후 "
                   "p = 0.0015 (Fisher); flinch 27 -> 36%, 25 -> 50%, p = 0.003; movement onset 42 -> 58%, 44 -> 61%, p = 0.64. 주사 후 vehicle vs SBI-553: Day 1 withdrawal 27 vs 65%, flinch 6 vs 50%; Day 2 65 vs 55%, 65 vs 49% (반대). 1191만: vehicle 후 25%, 2% vs SBI-553 후 60%, 57%; 1191 제외 시 SBI-553 vs vehicle: 61 vs 59%, 46 vs 57% (둘 다 ns).")),
    (r"^Responsive vs non-responsive, with the chosen window \(3/3",
     "Escape, grooming and attending responders: no significant change (all ns); too few escape / grooming events after SBI-553",
     "escape·grooming·attending 반응 뉴런: 유의한 변화 없음 (모두 ns); SBI-553 후 escape·grooming 사건은 너무 적음",
     dict(bottom_en=B_RV_EN, bottom_kr=B_RV_KR,
          notes_en="Key numbers (responsive %): escape vehicle day 19 -> 12% (ns); SBI-553 day baseline 49% (3 mice), no escape events after "
                   "SBI-553. Grooming: only 5937 groomed (24 neurons): vehicle day 0 -> 17%, SBI-553 day baseline 4%, no grooming after "
                   "SBI-553. Attending: after vehicle 0% (1 mouse); SBI-553 day 3 -> 0% (2 mice).",
          notes_kr="핵심 숫자 (반응 %): escape vehicle 날 19 -> 12% (ns); SBI-553 날 기준 49% (3마리), SBI-553 후 escape 없음. grooming: "
                   "5937만 (뉴런 24): vehicle 날 0 -> 17%, SBI-553 날 기준 4%, SBI-553 후 grooming 없음. attending: vehicle 후 0% "
                   "(1마리); SBI-553 날 3 -> 0% (2마리).")),
    (r"^Where the Pin prick neurons are, session by session \(1/2",
     "5941: UP 10 -> 7 of 67 after SBI-553, 22 -> 24 of 60 after vehicle; 5939: 2 -> 21 of 47 after vehicle (most of the pooled vehicle rise), 0 -> 4 of 69 after SBI-553",
     "5941: UP 10 -> 7/67 (SBI-553 후), 22 -> 24/60 (vehicle 후); 5939: vehicle 후 2 -> 21/47 (vehicle 증가의 대부분), SBI-553 후 0 -> 4/69",
     dict(bottom_kr="마우스마다 세션 순서대로 (Day 1 기준, 주사, Day 2 기준, 주사). n = 4마리 (crossover); 검정: 시간 이동 검정 "
                    "(뉴런별 UP/DOWN, BH q <= 0.20).")),
    (r"^Where the Pin prick neurons are, session by session \(2/2",
     "5937: UP 1 -> 2 of 24 after vehicle, 0 -> 7 after SBI-553; 1191: 0 -> 0 of 85 after vehicle, 18 -> 6 of 86 after SBI-553",
     "5937: UP 1 -> 2/24 (vehicle 후), 0 -> 7/24 (SBI-553 후); 1191: 0 -> 0/85 (vehicle 후), 18 -> 6/86 (SBI-553 후)",
     dict(bottom_kr="마우스마다 세션 순서대로 (Day 1 기준, 주사, Day 2 기준, 주사). n = 4마리 (crossover); 검정: 시간 이동 검정 "
                    "(뉴런별 UP/DOWN, BH q <= 0.20).")),
    (r"^What happens to the baseline-responsive neurons after the injection \(1/2",
     "Pin-prick activated neurons: no significant change after vehicle (0.28 -> 0.20 z, 11 cells) or SBI-553 (0.16 -> 0.20 z, 17 cells); no difference (ns)",
     "pin prick 활성 뉴런: vehicle 후 (0.28 -> 0.20 z, 11개)도 SBI-553 후 (0.16 -> 0.20 z, 17개)도 유의한 변화 없음; 차이 ns",
     dict(bottom_en=B_P5_EN.format(extra="; no test = fewer than 6 cells (no bracket)"),
          bottom_kr=B_P5_KR.format(extra="; 검정 없음 = 뉴런 6개 미만 (괄호 없음)"))),
    (r"^What happens to the baseline-responsive neurons after the injection \(2/2",
     "Movement-onset neurons weakened after vehicle (0.20 -> 0.15 z, **) but grew after SBI-553 (0.14 -> 0.20 z, *); differ ***",
     "movement onset 활성 뉴런: vehicle 후 약해짐 (0.20 -> 0.15 z, **), SBI-553 후 강해짐 (0.14 -> 0.20 z, *); 차이 ***",
     dict(bottom_en=B_P5_EN.format(extra="; escape, grooming, attending: too few such cells"),
          bottom_kr=B_P5_KR.format(extra="; escape·grooming·attending: 그런 뉴런이 너무 적음"))),
    (r"^Stimulus, withdrawal and flinch come almost together \(1/2",
     "Responses are larger when a withdrawal follows (0.30 vs 0.17 z, ***), yet 55% of neuron-stimulus pairs keep half without one",
     "withdrawal이 뒤따를 때 반응이 더 큼 (0.30 vs 0.17 z, ***); 그래도 뉴런-자극 쌍의 55%는 없어도 반응 절반 이상 유지",
     dict(bottom_en="A: the same stimulus with vs without a withdrawal (neuron x stimulus pairs); B: unique variance = cross-validated R2 "
                    "(full) minus R2 without that family, <= 0 = none. n = 4 mice (crossover); Test: Wilcoxon signed-rank (A; *** p < 0.001), ridge "
                    "encoding model (B).",
          bottom_kr="A: 같은 자극, withdrawal 있음 vs 없음 (뉴런 x 자극 쌍); B: 고유 설명량 = 교차 검증 R2 (전체) - 그 요인을 뺀 R2, "
                    "0 이하 = 없음. n = 4마리 (crossover); 검정: Wilcoxon signed-rank (A; *** p < 0.001), ridge encoding model (B).")),
    (r"^Stimulus, withdrawal and flinch come almost together \(2/2",
     "Movement alone explains the most neurons (38-54%); few follow the stimulus alone (2% Day 1, 7% Day 2)",
     "움직임만으로 설명되는 뉴런이 가장 많음 (38-54%); 자극만 따르는 뉴런은 적음 (Day 1 2%, Day 2 7%)",
     dict(bottom_en="Classes from the encoding model (ridge, cross-validated) + block-shift test per neuron (BH q <= 0.20); shares by day / "
                    "condition: descriptive (neurons nested in 4 mice).",
          bottom_kr="encoding model (ridge, 교차 검증) + 뉴런별 블록 이동 검정 (BH q <= 0.20)으로 분류; 날·조건별 비율: 기술 통계 "
                    "(4마리 안의 뉴런).",
          notes_add_en="Stimulus-only share after vehicle 13% vs after SBI-553 2% (descriptive, no test).",
          notes_add_kr="자극만 따르는 뉴런: vehicle 후 13% vs SBI-553 후 2% (기술 통계, 검정 없음).")),
    (r"^Do sessions with more behaviour have more responsive neurons\? \(1/2",
     "Activated neurons: no behaviour tracks them across sessions (strongest: attending, within mouse rho -0.18, p = 0.51)",
     "활성화 뉴런 %는 어떤 행동과도 관련 없음 (가장 강한 것: attending, 마우스 안 rho -0.18, p = 0.51)",
     dict(title_en="Do sessions with more behaviour have more responsive neurons? (1/2: activated neurons)",
          title_kr="행동이 많은 세션에 반응 뉴런도 많은가? (1/2: 활성화 뉴런)",
          bottom_en="One point = one session (16): behaviour per delivery vs % of neurons activated by any stimulus. n = 4 mice "
                    "(crossover); Test: Spearman (all ns).",
          bottom_kr="점 하나 = 세션 하나 (16개): 자극당 행동 vs 어떤 자극에든 활성화된 뉴런 %. n = 4마리 (crossover); 검정: "
                    "Spearman (모두 ns).")),
    (r"^Do sessions with more behaviour have more responsive neurons\? \(2/2",
     "Inhibited neurons: fewer in sessions with more movement (within mouse rho -0.75, p < 0.001; all sessions -0.60, p = 0.013)",
     "억제 뉴런 %는 movement가 많은 세션일수록 적음 (마우스 안 rho -0.75, p < 0.001; 전체 -0.60, p = 0.013)",
     dict(title_en="Do sessions with more behaviour have more responsive neurons? (2/2: inhibited neurons)",
          title_kr="행동이 많은 세션에 반응 뉴런도 많은가? (2/2: 억제 뉴런)",
          bottom_en="One point = one session (16): behaviour per delivery vs % of neurons inhibited by any stimulus. n = 4 mice "
                    "(crossover); Test: Spearman (movement within mouse p < 0.001, others ns).",
          bottom_kr="점 하나 = 세션 하나 (16개): 자극당 행동 vs 어떤 자극에든 억제된 뉴런 %. n = 4마리 (crossover); 검정: "
                    "Spearman (movement 마우스 안 p < 0.001, 나머지 ns).")),
    (r"^Your windows: 1/1, 2/2, 3/3, 5/5 and 10/10 s \(1/2",
     "Clearest event-locked signal: movement onset (10% at 1/1 s); stimuli score low partly because the jitter sham holds the next delivery",
     "여기서는 movement onset이 가장 뚜렷한 사건 고정 신호 (1/1 s에서 10%); 자극이 낮은 이유 일부는 2-5 s 가짜가 다음 자극을 포함하기 때문",
     dict(title_en="Which response window? 1/1 to 10/10 s (1/2: event-locked neurons)",
          title_kr="어떤 반응 창? 1/1 ~ 10/10 s (1/2: 사건에 고정된 뉴런)",
          bottom_en="Answer: short windows, used for every neuronal number: stimuli 0-0.5 s vs 0.5 s before, behaviours 0-0.5 s vs 1 s before "
                    "(a finer search, scored on a held-out mouse). " + "Points = mean of 16 sessions, n = 4 mice; descriptive. Per-neuron calls: Wilcoxon signed-rank, time-shift test.",
          bottom_kr="답: 짧은 창 (모든 뉴런 수치에 사용): 자극 0-0.5 s vs 직전 0.5 s, 행동 0-0.5 s vs 직전 1 s (더 촘촘한 탐색, 빼 둔 "
                    "마우스로 평가). " + "점 = 16개 세션 평균, n = 4마리; 기술 통계. 뉴런별 판정: Wilcoxon signed-rank, 시간 이동 검정.",
          notes_add_en="Deliveries come about every 3 s, so a fake event shifted by 2-5 s often lands next to another real delivery of "
                       "the same stimulus: the sham then holds real responses and the stimulus values here are low partly by construction. "
                       "The per-neuron counts (Q3) use the block-shift test instead.",
          notes_add_kr="자극은 약 3 s마다 오므로 2-5 s 옮긴 가짜 사건은 같은 자극의 다른 실제 자극 옆에 오기 쉬움: 가짜에 실제 반응이 들어가 "
                       "여기서 자극 값은 구조상 일부 낮음. Q3의 뉴런별 수는 블록 이동 검정을 씀.")),
    (r"^Your windows: 1/1, 2/2, 3/3, 5/5 and 10/10 s \(2/2",
     "Long windows catch slow changes, but at 10/10 s 98-99% of stimulus trials already hold the next delivery (movement 57%)",
     "긴 창은 느린 변화를 잡지만, 10/10 s에서는 자극 시행의 98-99%에 다음 자극이 들어 있음 (movement 57%)",
     dict(title_en="Which response window? 1/1 to 10/10 s (2/2: slow changes and the next stimulus)",
          title_kr="어떤 반응 창? 1/1 ~ 10/10 s (2/2: 느린 변화와 다음 자극)",
          bottom_en="Withdrawal / flinch: the delivery inside even the 1/1-s window is mostly the one that triggered them. " + "Points = mean of 16 sessions, n = 4 mice; descriptive. Per-neuron calls: Wilcoxon signed-rank, time-shift test.",
          bottom_kr="withdrawal / flinch: 1/1 s 창 안의 자극은 대부분 그 반사를 일으킨 자극. " + "점 = 16개 세션 평균, n = 4마리; 기술 통계. 뉴런별 판정: Wilcoxon signed-rank, 시간 이동 검정.")),
    (r"^Check 1: what does the test call on FAKE events",
     "Fake events: 1.6% (BH) / 0.1% (max-T) of neurons called vs 28% / 8% for real events",
     "가짜 사건: 뉴런 1.6% (BH) / 0.1% (max-T) 판정 vs 실제 사건 28% / 8%",
     dict(bottom_en="On 4.6% of the fake runs BH calls >= 10% of the neurons at once; max-T never (0%). 'selective' = the same "
                    "test minus the all-neuron mean (shown, not used). n = 4 mice (crossover); Test: BH q <= 0.20 and max-T 5% on shifted fake events.",
          bottom_kr="가짜 실행의 4.6%에서 BH가 뉴런 10% 이상을 한꺼번에 판정; max-T는 없음 (0%). 'selective' = 전체 뉴런 평균을 뺀 "
                    "같은 검정 (보여 주기만, 사용 안 함). n = 4마리 (crossover); 검정: 이동한 가짜 사건에 BH q <= 0.20, max-T 5%.")),
    (r"^Check 2: responses of known size added",
     "With BH a +0.5 z response is found in 96% of the spiked neurons (old 2-s window 71%); with max-T 82% vs 38%",
     "+0.5 z를 넣은 뉴런 중 BH는 96% 찾음 (옛 2 s 창 71%); max-T는 82% vs 38%",
     dict(bottom_en="Decreases (-0.5 z): 86% vs 65% (BH), 77% vs 39% (max-T); BH also calls ~2% of the other neurons, max-T 0%. "
                    "n = 4 mice (crossover); Test: BH q <= 0.20 and max-T 5%.",
          bottom_kr="감소 (-0.5 z): 86% vs 65% (BH), 77% vs 39% (max-T); BH는 다른 뉴런도 ~2% 판정, max-T는 0%. n = 4마리 "
                    "(crossover); 검정: BH q <= 0.20, max-T 5%.")),
    (r"^Check 3: can the model tell a stimulus response from a withdrawal",
     "Day 2 (clicker): 65% of stimulus-locked and 73% of withdrawal-locked neurons get the right class; Day 1 (taps): 30% / 69%",
     "Day 2 (로거): 자극 고정 뉴런의 65%, withdrawal 고정 뉴런의 73%가 맞는 분류; Day 1 (탭): 30% / 69%",
     dict(bottom_en="On Day 1 47% of the stimulus-locked neurons come out 'cannot separate' (taps sit on the withdrawal), so the stimulus-only "
                    "shares (2% / 7%) are lower bounds. n = 4 mice (crossover); classes: ridge encoding model + per-neuron block-shift test (BH q <= 0.20).",
          bottom_kr="Day 1에서 자극 고정 뉴런의 47%가 '구분 불가' (탭이 withdrawal과 거의 겹침): 자극만 따르는 비율 (2% / 7%)은 하한값. "
                    "n = 4마리 (crossover); 분류: ridge encoding model + 뉴런별 블록 이동 검정 (BH q <= 0.20).")),
    (r"^Check 4: do the key findings hold in every mouse.*\(1/2",
     "Pin: smaller SBI-553-day than vehicle-day change in 3/4 mice (vehicle-day rise * with max-T); heat and movement-onset neurons rose on both days",
     "pin: 4마리 중 3마리에서 SBI-553 날 변화가 vehicle 날보다 작음 (vehicle 날 증가는 max-T로도 *); heat·movement onset 뉴런은 두 날 모두 증가",
     dict(title_en="Check 4: do the key findings hold in every mouse? (1/2: neurons)",
          title_kr="검증 4: 마우스마다 유지되나? (1/2: 뉴런)",
          bottom_en="Marker = one mouse (filled = BH q <= 0.20, open = max-T 5%); stars = pooled-neuron binomial test, baseline vs "
                    "injection of the same day (* p < 0.05, ** p < 0.01, *** p < 0.001, ns). n = 4 mice (crossover).",
          bottom_kr="표식 = 마우스 (채움 = BH q <= 0.20, 빈 = max-T 5%); 별표 = 뉴런 합산 이항 검정, 같은 날 기준 vs 주사 (* p < 0.05, "
                    "** p < 0.01, *** p < 0.001, ns). n = 4마리 (crossover).",
          notes_add_en="Per mouse (BH, points): pin vehicle day 1191 0, 5937 +8.3, 5939 +59.6, 5941 +1.7; SBI-553 day -24.4, +29.2, +13, -9. "
                       "Heat vehicle day +27.1, -20.8, +23.4, +40; SBI-553 day +36, +20.8, +62.3, +34.3. Movement onset: rose on the SBI-553 day "
                       "in 4/4 mice, on the vehicle day in 2/4, larger on the SBI-553 day in 2/4 (one tie). max-T: pin p = 0.021 / 0.63 (vehicle *, SBI-553 ns), "
                       "heat vehicle p = 0.022, SBI-553 p < 0.001.",
          notes_add_kr="마우스별 (BH, %p): pin vehicle 날 1191 0, 5937 +8.3, 5939 +59.6, 5941 +1.7; SBI-553 날 -24.4, +29.2, +13, -9. heat "
                       "vehicle 날 +27.1, -20.8, +23.4, +40; SBI-553 날 +36, +20.8, +62.3, +34.3. movement onset: SBI-553 날 4/4마리, vehicle 날 "
                       "2/4마리 증가, SBI-553 날이 더 큰 마우스 2/4 (1마리 같음). max-T: pin p = 0.021 / 0.63 (vehicle *, SBI-553 ns), heat vehicle p = 0.022, "
                       "SBI-553 p < 0.001.")),
    (r"^Check 4: do the key findings hold in every mouse.*\(2/2",
     "Behaviour drops after SBI-553 hold per mouse: flinch after pin prick 3/4, movement after light touch 4/4, escape after heat 3/4",
     "행동 감소는 마우스별로도 유지: pin prick 후 flinch 3/4, light touch 후 movement 4/4, heat 후 escape 3/4",
     dict(title_en="Check 4: do the key findings hold in every mouse? (2/2: behaviour)",
          title_kr="검증 4: 마우스마다 유지되나? (2/2: 행동)",
          bottom_en="One line = one mouse, its vehicle vs SBI-553 injection session (descriptive; with 4 mice the smallest p is 0.125). "
                    "n = 4 mice (crossover).",
          bottom_kr="선 하나 = 마우스 하나, vehicle vs SBI-553 주사 세션 (기술 통계; 4마리로는 최소 p = 0.125). n = 4마리 (crossover).")),
    (r"^Behaviour per delivery: every mouse",
     "SBI-553 vs vehicle: less movement (0.56 vs 0.76 per delivery, ***; mostly SBI-553-first 5941, 5939), less escape (**); reflexes ns",
     "SBI-553 vs vehicle: 움직임 (자극당 0.56 vs 0.76, ***; 주로 SBI-553 먼저인 5941·5939)과 escape (**) 감소; 반사는 차이 없음",
     dict(notes_add_en="Withdrawal and flinch rose after vehicle (***) but not after SBI-553 (ns); the SBI-553 day started higher, and "
                       "after vehicle vs after SBI-553 is ns. 0.56 vs 0.76 = the two injection sessions (mean of mice), not before -> after. "
                       "SBI-553 on Day 1 = 5941, 5939; on Day 2 = 5937, 1191.",
          notes_add_kr="withdrawal·flinch는 vehicle 후 증가 (***), SBI-553 후는 아님 (ns); SBI-553 날은 처음부터 높았고, vehicle 후 vs "
                       "SBI-553 후는 ns. 0.56 vs 0.76 = 두 주사 세션 (마우스 평균), 전후 비교 아님. Day 1 SBI-553 = 5941, 5939; Day 2 = 5937, 1191.")),
    (r"^Behaviour per stimulus, all mice \(1/3",
     "SBI-553 vs vehicle: only pin-prick flinch differs (57 vs 79%, **), from a vehicle-day rise (59 -> 79%, *); SBI-553 day 72 -> 57% ns",
     "SBI-553 vs vehicle: pin prick flinch만 다름 (57 vs 79%, **); vehicle 날 59 -> 79% (*), SBI-553 날 72 -> 57% (ns)",
     dict(bottom_en="% of the deliveries of each stimulus followed by the behaviour; bars = mean of mice, dots = mice. n = 4 mice (crossover); Test: Cochran-Mantel-Haenszel; ns = tested, not significant; no bracket = not testable.",
          bottom_kr="자극마다 그 행동이 뒤따른 자극 %; 막대 = 마우스 평균, 점 = 마우스. n = 4마리 (crossover); 검정: Cochran-Mantel-Haenszel; ns = 검정했으나 유의하지 않음; 괄호 없음 = 검정 불가.")),
    (r"^Behaviour per stimulus, all mice \(2/3",
     "Attending and lick / bite: no difference between SBI-553 and vehicle",
     "attending, lick / bite: SBI-553와 vehicle 차이 없음",
     dict(bottom_en="% of the deliveries of each stimulus followed by the behaviour; bars = mean of mice, dots = mice. n = 4 mice (crossover); Test: Cochran-Mantel-Haenszel; ns = tested, not significant; no bracket = not testable.", bottom_kr="자극마다 그 행동이 뒤따른 자극 %; 막대 = 마우스 평균, 점 = 마우스. n = 4마리 (crossover); 검정: Cochran-Mantel-Haenszel; ns = 검정했으나 유의하지 않음; 괄호 없음 = 검정 불가.")),
    (r"^Behaviour per stimulus, all mice \(3/3",
     "Less movement after light touch with SBI-553 than vehicle (39 vs 61%, **); escape after heat 1 vs 10%, ns",
     "SBI-553 후 light touch 뒤 움직임이 vehicle보다 적음 (39 vs 61%, **); heat 뒤 escape는 1 vs 10%, ns",
     dict(bottom_en="% of the deliveries of each stimulus followed by the behaviour; bars = mean of mice, dots = mice. n = 4 mice (crossover); Test: Cochran-Mantel-Haenszel; ns = tested, not significant; no bracket = not testable.", bottom_kr="자극마다 그 행동이 뒤따른 자극 %; 막대 = 마우스 평균, 점 = 마우스. n = 4마리 (crossover); 검정: Cochran-Mantel-Haenszel; ns = 검정했으나 유의하지 않음; 괄호 없음 = 검정 불가.")),
    (r"^Activity of every neuron by sex",
     "On the SBI-553 day both females dropped more than both males (stimulus blocks, log2 -1.58, -0.53 vs -0.22, 0.00; 5937 at its floor)",
     "SBI-553 날 두 암컷이 두 수컷보다 더 감소 (자극 블록, log2 -1.58, -0.53 vs -0.22, 0.00; 5937은 바닥값)",
     dict(title_kr="성별로 본 모든 뉴런의 활동",
          bottom_en="Stars = neurons pooled over 2 mice per sex: exploratory (* p < 0.05, ** p < 0.01, *** p < 0.001, ns). Neurons F / M: "
                    "baseline 298 / 164, vehicle 145 / 71, SBI-553 153 / 93; 2 + 2 mice. Test: Mann-Whitney (A), Wilcoxon vs 0 (B).",
          bottom_kr="별표 = 성별 2마리의 뉴런 합산: 탐색적 (* p < 0.05, ** p < 0.01, *** p < 0.001, ns). 뉴런 암컷 / 수컷: 기준 298 / "
                    "164, vehicle 145 / 71, SBI-553 153 / 93; 마우스 2 + 2. 검정: Mann-Whitney (A), 0 대비 Wilcoxon (B).")),
    (r"^What the data can and cannot say yet",
     "With 4 mice no test across mice can reach p < 0.05; 5937 (Day 2) and 5939 (tracking) are the weakest data",
     "4마리로는 마우스 간 검정이 p < 0.05에 닿을 수 없음; 가장 약한 데이터는 5937 (Day 2)과 5939 (추적)", {}),
    (r"^Day 1 \+ Day 2: SBI-553 vs vehicle in the same mice",
     "Movement fell after SBI-553 in 3/4 mice (strong in 2); pin-prick UP neurons 10% vs 22% after vehicle (Day 2 only)",
     "SBI-553 후 움직임 감소 3/4마리 (2마리에서 큼); pin prick UP 뉴런 10% vs vehicle 후 22% (Day 2만)", {}),
    (r"^Day 1: SBI-553 vs vehicle",
     "Day 1: after SBI-553 light-touch UP neurons 16% (vehicle 1%), movement 0.40 vs 0.76 per delivery (2 vs 2 mice)",
     "Day 1: SBI-553 후 light touch UP 뉴런 16% (vehicle 1%), 움직임 자극당 0.40 vs 0.76 (2 vs 2마리)", {}),
    (r"^Day 2: SBI-553 vs vehicle",
     "Day 2: after SBI-553 pin-prick UP neurons 12% (vehicle 42%); attending similar (0.08 vs 0.09; 2 vs 2 mice)",
     "Day 2: SBI-553 후 pin prick UP 뉴런 12% (vehicle 42%); attending 비슷 (0.08 vs 0.09; 2 vs 2마리)", {}),
]

# exact text corrections (whole text frame): (old, new)
FIXTEXT = [
    ("0.76 after vehicle -> 0.40 after SBI-553 (mean of mice); 4 mice: min p = 0.125.",
     "vehicle mice 0.76 vs SBI-553 mice 0.40 (mean of 2 mice each); 2 vs 2 mice: not testable."),
    ("0.09 after vehicle -> 0.08 after SBI-553 (mean of mice); 4 mice: min p = 0.125.",
     "vehicle mice 0.09 vs SBI-553 mice 0.08 (mean of 2 mice each); 2 vs 2 mice: not testable."),
    ("vehicle 후 0.76 -> SBI-553 후 0.40 (마우스 평균); 4마리: 최소 p = 0.125.",
     "vehicle 마우스 0.76 vs SBI-553 마우스 0.40 (각 2마리 평균); 2 vs 2마리: 검정 불가."),
    ("vehicle 후 0.09 -> SBI-553 후 0.08 (마우스 평균); 4마리: 최소 p = 0.125.",
     "vehicle 마우스 0.09 vs SBI-553 마우스 0.08 (각 2마리 평균); 2 vs 2마리: 검정 불가."),
    # sex slide, card 4 (10-23: vs the SAME day's baseline, P6 movement per delivery; the old card used the mean of both baselines)
    ("Only the two mice with SBI-553 on Day 1: 5941 (F) 56 -> 39 %, 5939 (M) 55 -> 33 %; not 1191 (F) 54 -> 56 % or 5937 (M) 66 -> 61 %.",
     "Lower than after vehicle in 3 of 4 mice (not 1191), strongly only in the SBI-553-first mice: 5941 (F) 0.44 vs 0.87, 5939 (M) 0.36 vs "
     "0.63; 1191 (F) 0.69 vs 0.66, 5937 (M) 0.74 vs 0.87 per delivery."),
    ("Day 1에 SBI-553을 받은 두 마리만: 5941 (F) 56 -> 39 %, 5939 (M) 55 -> 33 %; 1191 (F) 54 -> 56 %, 5937 (M) 66 -> 61 %는 아님.",
     "4마리 중 3마리에서 vehicle 후보다 낮음 (1191 제외), 크게 줄어든 것은 SBI-553을 먼저 받은 두 마리: 5941 (F) 0.44 vs 0.87, 5939 (M) 0.36 vs 0.63; "
     "1191 (F) 0.69 vs 0.66, 5937 (M) 0.74 vs 0.87 (자극당)."),
    # sex slide, card 2 (light touch / heat: also after vehicle) and card 3 (withdrawal: same-day baselines)
    ("Males 29 / 0 %; heat the same way (females 39 / 37 %, males 62 / 29 %). The reflex neurons do not split by sex.",
     "Males 29 / 0 %: no clean sex split (after vehicle F 40 / 0 vs M 0 / 4 %). Heat after SBI-553 F 39 / 37 vs M 62 / 29 %, "
     "after vehicle F 43 / 27, M 60 / 0 %."),
    ("수컷 29 / 0 %; heat도 같음 (암컷 39 / 37 %, 수컷 62 / 29 %). 반사 뉴런은 성별로 나뉘지 않음.",
     "수컷 29 / 0 %: 뚜렷한 성별 차이 없음 (vehicle 후 암컷 40 / 0 vs 수컷 0 / 4 %). heat는 SBI-553 후 암컷 39 / 37 vs 수컷 62 / 29 %, "
     "vehicle 후 암컷 43 / 27, 수컷 60 / 0 %."),
    ("Males 54 / 55 %; after SBI-553 both females higher too; at baseline all 4 mice 46-52 %. Not drug-specific.",
     "Males 54 / 55 %. Females rose more after vehicle (+19 / +28 vs +8 / +10 points), not after SBI-553 (+6 / +1 vs +8 / +6). "
     "Not drug-specific."),
    ("수컷 54 / 55 %; SBI-553 후에도 두 암컷이 높음; 기준에서는 4마리 모두 46-52 %. 약물 특이적이지 않음.",
     "수컷 54 / 55 %. 암컷은 vehicle 후 더 증가 (+19 / +28 vs +8 / +10 %p), SBI-553 후는 아님 (+6 / +1 vs +8 / +6). 약물 특이적이지 "
     "않음."),
]
# Korean wording inside a text (substring): (old, new)
KR_SUB = [("실선 = 사건 후 상승한", "단색 막대 = 사건 후 상승한"), ("빗금 = 하강", "빗금 막대 = 하강"),
          ("(B; 검정 = 마우스)", "(B; 검은 표식 = 마우스별 중앙값)"), ("으로 잼.", "으로 측정."), ("나머지 절반으로 잼", "나머지 절반으로 측정"),
          ("확인 4", "검증 4"), ("반응 vs 무반응, 선택한 창으로", "반응 vs 무반응 (UP + DOWN), 최적 창")]
# both languages (substring, run level): the earlier 'chosen window' name and wording addressed to one person
SUB_ALL = [("Responsive vs non-responsive, with the chosen window", "Responsive vs non-responsive (UP + DOWN), best window"),
           ("Neuron selection confirmed by you (10-22)", "Neuron selection checked by eye (10-22)"),
           ("Day-1 selection confirmed by you", "Day-1 selection checked by eye"),
           ("Next: population dynamics (state space) on top of these basic results.", "Next: more mice, then population dynamics (state space)."),
           ("다음: 이 기본 결과 위에 집단 동역학(state space) 분석.", "다음: 마우스 추가, 그다음 집단 동역학(state space)."),
           # 10-24 audit: the Day-1 decks say 'per group' - in the crossover the drugs swap on Day 2
           ("n = 2 mice per group", "Day 1: 2 mice per drug (swapped on Day 2)"), ("군당 n = 2마리", "Day 1: 약물당 2마리 (Day 2에 바뀜)")]
# notes (10-23 round 3): wording addressed to one person and an outdated bracket rule
NOTES_SUB = [("Your rule:", "Scoring rule:"), ("your rule", "the scoring rule"), ("Brackets (only p < 0.05 drawn)",
             "Brackets: every tested pair (ns shown); no bracket = not testable"), ("괄호 (p < 0.05만)", "괄호: 검정한 모든 쌍 (ns 표시); 괄호 없음 = 검정 불가")]

def fit_big_numbers(slide):
    """card numbers (34 pt) longer than 6 characters ran out of their 1.9-in box (preview warning, sex cards): 24 / 20 pt"""
    k = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        L_ = len(sh.text_frame.text.strip())
        for pp in sh.text_frame.paragraphs:                    # the size can sit on the paragraph (Deck.box) or on the run
            big = (pp.font.size and pp.font.size.pt >= 33) or any(r_.font.size and r_.font.size.pt >= 33 for r_ in pp.runs)
            if big and L_ > 6:
                new = Pt(24 if L_ <= 9 else 20)
                pp.font.size = new
                for r_ in pp.runs:
                    r_.font.size = new
                k += 1
    return k


def _set_text(tf, new):
    runs = tf.paragraphs[0].runs
    if not runs:
        return False
    runs[0].text = new
    for x in runs[1:]:
        x.text = ""
    for pp in list(tf.paragraphs[1:]):
        pp._p.getparent().remove(pp._p)
    return True


def fix_texts(slide, lang):
    k = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        for o, n in FIXTEXT:
            if t == o and _set_text(sh.text_frame, n):
                k += 1
        for pp in sh.text_frame.paragraphs:
            for r_ in pp.runs:
                for o, n in SUB_ALL + (KR_SUB if lang == "kr" else []):
                    if o in r_.text:
                        r_.text = r_.text.replace(o, n); k += 1
    return k


def notes_fix(slide):
    """10-23 round 3: NOTES_SUB in the notes (both languages)"""
    if not slide.has_notes_slide:
        return
    nt = slide.notes_slide.notes_text_frame
    t = nt.text
    t2 = t
    for o, n in NOTES_SUB:
        t2 = t2.replace(o, n)
    if t2 != t:
        nt.text = t2


# 10-24 (window per event, WIN_MODE=mix): texts the decks wrote that assume ONE searched window (applied to every final slide
# and its notes; regex, paragraph by paragraph)
MIX_RX = [
    (r"(the 1 s before) \(window search\)", r"\1 (window per event)"),
    (r"(직전 1 s) \(창 탐색\)", r"\1 (사건별 창)"),
    (r"behaviours 0-1 s after vs the 1 s before", "behaviours withdrawal, flinch, movement 0-1 s; escape, grooming, attending 0-0.5 s "
                                                   "(vs the 1 s before)"),
    (r"행동 = 후 0-1 s vs 직전 1 s", "행동 = withdrawal, flinch, movement 후 0-1 s; escape, grooming, attending 후 0-0.5 s (직전 1 s 대비)"),
    (r"행동 사건 후 0-1 s vs 직전 1 s", "행동 = withdrawal, flinch, movement 후 0-1 s; escape, grooming, attending 후 0-0.5 s (직전 1 s 대비)"),
    (r"\s*\((?:nan|[0-9.]+)x the old window(?: in a held-out mouse)?\)", ""),
    (r"\s*\((?:검증 마우스에서 )?기존 창의 (?:nan|[0-9.]+)배\)", ""),
    (r"Chosen: stimuli [^,]+, behaviours [^.(]+", "Used: one window per event, 0-0.5 or 0-1 s (notes)"),
    (r"선택: 자극 [^,]+, 행동 [^.(]+", "사용: 사건마다 창 하나, 0-0.5 또는 0-1 s (노트)"),
    (r"^Best: stimuli .*$", "Window search = exploration only; the windows used are per event, 0-0.5 or 0-1 s (notes)"),
    (r"^최적: 자극 .*$", "창 탐색 = 탐색만; 실제 사용 창은 사건마다 0-0.5 또는 0-1 s (노트)"),
    # 10-24 audit: 'best window' -> the window of each event (not in the archive texts about the 10-22 re-run)
    (r"in the best window \(panel title\)", "in its own window (panel title)"),
    (r"(?<!of )\(best window\)", "(window per event)"),
    (r"best-window response", "response (window per event)"),
    (r"(?<![-\w])[Bb]est windows?(?! of 10-2)(?! re-run)", "window per event"),
    (r"최적 창\(패널 제목\)으로", "사건별 창(패널 제목)으로"),
    (r"(?<!10-22 )최적 창(?!으로 다시)", "사건별 창"),
    # 10-24 audit: leftovers addressed to one person, internal labels, untranslated words
    (r"Ntsr1-deck layout: ", ""),
    (r"median difference \(d\)", "median difference"),
    (r"\(your confirmed tracking\)", "(tracking checked by eye)"),
    (r"검정: descriptive\.", "기술 통계 (검정 없음)."),
    (r"time-shift 검정", "시간 이동 검정"),
    (r"뉴런 선택 확정 \(10-22\)", "뉴런 선택은 눈으로 확인 (10-22)"),
    (r"Day 2 \(로거\)", "Day 2 (클리커)"),
    (r"로거 표시가", "클리커 표시가"),
]


def mix_fix(slide):
    """MIX_RX on every paragraph of the slide (first run keeps its format) and on the notes; True if anything changed"""
    import winwords as WW
    if not WW.mixed():
        return False
    hit = False
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for pp in sh.text_frame.paragraphs:
            t = "".join(r.text for r in pp.runs)
            t2 = t
            for o, n in MIX_RX:
                t2 = re.sub(o, n, t2)
            if t2 != t and pp.runs:
                pp.runs[0].text = t2
                for r_ in pp.runs[1:]:
                    r_.text = ""
                hit = True
    if slide.has_notes_slide:
        nt = slide.notes_slide.notes_text_frame
        t = nt.text
        t2 = t
        for o, n in MIX_RX:
            t2 = re.sub(o, n, t2, flags=re.M)
        if t2 != t:
            nt.text = t2
            hit = True
    return hit

# 10-24 audit: (English title regex, EN, KR) - the test the figure prints, where the deck wrote 'descriptive (no test)'
NO_TEST_FIX = [
    (r"^(Where are the clustered responsive neurons|군집된 반응 뉴런은 어디에)", "nearest-neighbour test vs 500 random draws per map (uncorrected)",
     "최근접 이웃 거리 vs 무작위 500번 (보정 없음)"),
    (r"^(Is the clustering more than chance|군집이 우연 이상인가)", "binomial test, one-sided (tests not independent)", "이항 검정, 단측 (검정들은 독립 아님)"),
    (r"^(Are the responsive neurons clustered|반응 뉴런이 시야에서 모여)", "nearest-neighbour test vs 500 random draws per map (uncorrected)",
     "최근접 이웃 거리 vs 무작위 500번 (보정 없음)"),
    (r"^(Stable across days, shared between events|날짜를 넘어 안정적)", "binomial test (neurons pooled, exploratory)", "이항 검정 (뉴런 합산, 탐색적)"),
    (r"^(5941|5939|5937|1191)\b.*where the responsive neurons are", "classes: time-shift test (BH q <= 0.20)",
     "분류: 시간 이동 검정 (BH q <= 0.20)"),
    (r"^(Summary: responsive neurons per event|요약: 사건별 반응 뉴런)", "classes: time-shift test (BH q <= 0.20); % descriptive",
     "분류: 시간 이동 검정 (BH q <= 0.20); 비율은 기술 통계"),
    (r"UP, DOWN (and the other neurons|, 나머지 뉴런)", "classes: time-shift test (BH q <= 0.20); traces descriptive",
     "분류: 시간 이동 검정 (BH q <= 0.20); trace는 기술 통계"),
]
WINDOW_CAVEAT = ("Analyst choice: the response window was chosen per event (0-0.5 or 0-1 s, whichever showed the clearer vehicle vs "
                 "SBI-553 difference; exploratory, no correction), so the drug p values are optimistic.",
                 "분석 선택: 반응 창을 사건마다 고름 (0-0.5 또는 0-1 s 중 vehicle vs SBI-553 차이가 더 뚜렷한 쪽; 탐색적, 보정 없음) - "
                 "약물 p 값은 낙관적.")


def _para_sub(sh, pat, rep):
    hit = False
    for pp in sh.text_frame.paragraphs:
        t = "".join(r.text for r in pp.runs)
        t2 = re.sub(pat, rep, t)
        if t2 != t and pp.runs:
            pp.runs[0].text = t2
            for r_ in pp.runs[1:]:
                r_.text = ""
            hit = True
    return hit


def face_fix(slide, lang):
    """the slide-face fixes of the 10-24 audit that follow from the title (both variants; True if anything changed)"""
    import winwords as WW
    kr = lang == "kr"
    T = [sh for sh in slide.shapes if sh.has_text_frame]
    tb = [sh for sh in T if Inches(0.25) <= (sh.top or 0) < Inches(0.8)]
    if not tb:
        return False
    title = tb[0].text_frame.text.strip()
    bot = [sh for sh in T if (sh.top or 0) >= Inches(6.5)]
    hit = False
    m = re.match(r"^(5941|5939|5937|1191)\b", title)
    if m and bot:                                                # one mouse on the slide
        for b in bot:
            hit |= _para_sub(b, r"n = 4 mice \(crossover\)", f"n = 1 mouse ({m.group(1)})")
            hit |= _para_sub(b, r"n = 4마리 \(crossover\)", f"n = 1마리 ({m.group(1)})")
            hit |= _para_sub(b, r"마우스 4(?!\d)", f"마우스 1 ({m.group(1)})")
    if re.search(r"(field of view activity changed|시야에서 활동이 바뀐 곳).*\((1|2)/2", title) and bot:
        for b in bot:
            hit |= _para_sub(b, r"n = 4 mice(?: \(crossover\))?", "n = 2 of 4 mice on this slide")
            hit |= _para_sub(b, r"(n = 4마리(?: \(crossover\))?|마우스 4(?!\d))", "이 슬라이드 2마리 (전체 4마리)")
    en_title = title
    for pat, en, ko in NO_TEST_FIX:
        if re.search(pat, en_title):
            for b in bot:
                bt = b.text_frame.text
                en_, ko_ = (("uncorrected" if en.split(" (")[0] in bt else en),
                            ("보정 없음" if ko.split(" (")[0] in bt else ko))       # 10-25: no repeated test phrase
                hit |= _para_sub(b, r"descriptive \(no test\)", en_)
                hit |= _para_sub(b, r"기술 통계 \(검정 없음\)", ko_)
    # behaviour titles start with a capital; one name for escape / rearing
    new = re.sub(r"^(withdrawal|flinch|movement onset)(:| )", lambda q: q.group(1)[0].upper() + q.group(1)[1:] + q.group(2), title)
    new = re.sub(r", escape\)$", ", escape/rearing)", new)
    new = re.sub(r",? ?\(?(window per event|사건별 창)\)?(?= \(|$)", "", new)      # 10-25: window words only in the notes
    if new != title:
        hit |= _para_sub(tb[0], re.escape(title), new.replace("\\", "\\\\"))
    if WW.mixed() and re.match(r"^(What the data can and cannot say yet|지금 데이터로 말할 수 있는 것과 없는 것)", title):
        nt = slide.notes_slide.notes_text_frame
        cav = WINDOW_CAVEAT[1 if kr else 0]
        if cav not in nt.text:
            nt.text = (cav + "\n\n" + nt.text).strip()
            hit = True
    return hit


def _tnorm(t):
    """a title for matching, insensitive to the 10-24 title rewrites (window words, escape / rearing, capitals)"""
    t = t.strip().lower()
    for w in ("(best window)", "(window per event)", "(최적 창)", "(사건별 창)", "best window", "window per event", "최적 창",
              "사건별 창", "/rearing", ","):
        t = t.replace(w, "")
    return " ".join(t.split())


def apply_overrides(prs, lang):
    """10-24 audit: per-variant slide texts checked against that variant's outputs
    (analysis\\_variants\\<VARIANT_KEY>\\slide_overrides.json: [{title_en, title_kr, key_en, key_kr, bottom_en, bottom_kr,
    notes_add_en, notes_add_kr, new_title_en, new_title_kr}]), matched on the exact final slide title; the replaced text is kept
    in the notes (only add, never delete). Returns the number of slides changed."""
    import json as _j
    v = os.environ.get("VARIANT_KEY", "")
    fn = os.path.join("C:/Users/hsollim/Documents/HEAL_mini1p_SBI553/analysis/_variants", v, "slide_overrides.json")
    if not v or not os.path.isfile(fn):
        return 0
    kr = lang == "kr"
    O = {}
    for r in _j.load(open(fn, encoding="utf-8")):
        t = (r.get("title_kr") if kr else r.get("title_en")) or ""
        if t:
            O.setdefault(_tnorm(t), []).append(r)
    k = 0
    for sl in prs.slides:
        T = [sh for sh in sl.shapes if sh.has_text_frame]
        tb = [sh for sh in T if Inches(0.25) <= (sh.top or 0) < Inches(0.8)]
        if not tb:
            continue
        title = tb[0].text_frame.text.strip()
        rows = O.pop(_tnorm(title), None)               # first slide with that title only (Day-alone FULL parts share titles)
        if not rows:
            continue
        nt = sl.notes_slide.notes_text_frame
        done = False
        for r in rows:
            g = lambda f: (r.get(f + ("_kr" if kr else "_en")) or "").strip()
            key, bot, nadd, ntit = g("key"), g("bottom"), g("notes_add"), g("new_title")
            keyb = sorted([sh for sh in T if Inches(0.8) <= (sh.top or 0) <= Inches(1.3) and sh.text_frame.text.strip()],
                          key=lambda sh: sh.top or 0)
            if key and keyb:
                old = keyb[0].text_frame.text.strip()
                if old != key and _set_text(keyb[0].text_frame, key):
                    keyb[0].text_frame.paragraphs[0].runs[0].font.size = Pt(key_pt(key, kr))
                    nt.text = (nt.text + "\n\n" + ("[이전 버전, 대체됨] 이전 핵심 문장: " if kr else "[Superseded, earlier version] key line: ")
                               + old).strip()
                    done = True
            # bottom zone: new-style slides put the bottom at 6.68 in, the earlier decks at ~6.2 in with a grey methods paragraph
            # above it (~5.6 in); the override replaces the lowest box and moves the other long texts of the zone to the notes
            bb = sorted([sh for sh in T if (sh.top or 0) >= Inches(5.5) and sh.text_frame.text.strip()
                         and sh not in keyb[:1] and (sh.width or 0) > Inches(6)              # not the BACKUP tag / banner
                         and sh.text_frame.text.strip() not in ("BACKUP (hidden)", "보조 슬라이드 (숨김)")],
                        key=lambda sh: sh.top or 0)
            if bot and bb:
                old = bb[-1].text_frame.text.strip()
                if old != bot and _set_text(bb[-1].text_frame, bot):
                    for pp in list(bb[-1].text_frame.paragraphs[1:]):
                        pp._p.getparent().remove(pp._p)
                    nt.text = (nt.text + "\n\n" + ("[이전 버전, 대체됨] 아래 문장: " if kr else "[Superseded, earlier version] bottom line: ")
                               + old).strip()
                    done = True
                for sh in bb[:-1]:
                    t_ = sh.text_frame.text.strip()
                    if len(t_) > 120 and (sh.width or 0) > Inches(10) and _set_text(sh.text_frame, " "):   # full-width paragraphs only
                        nt.text = (nt.text + "\n\n" + ("[이전 버전, 노트로 옮김] " if kr else "[Moved from the slide] ") + t_).strip()
            if nadd and nadd not in nt.text:
                nt.text = (nadd + "\n\n" + nt.text).strip()
                done = True
            if ntit and ntit != title and _set_text(tb[0].text_frame, ntit):
                done = True
        k += done
    return k


def fit_title(slide):
    """10-23 round 3: a title that no longer fits one line (after a rename) steps down 26 -> 22 -> 19 pt"""
    for sh in slide.shapes:
        if not (sh.has_text_frame and Inches(0.25) <= (sh.top or 0) < Inches(0.8) and sh.text_frame.text.strip()):
            continue
        runs = [r_ for pp in sh.text_frame.paragraphs for r_ in pp.runs]
        cur = max([r_.font.size.pt for r_ in runs if r_.font.size] or [0])
        if cur < 18:
            continue
        t = sh.text_frame.text.strip()
        em = sum(1.0 if "\uac00" <= ch <= "\ud7a3" else 0.28 if ch == " " else 0.66 if ch.isupper() else 0.53 for ch in t) * 1.05
        for pt in (26, 22, 19):
            if pt <= cur and (em * pt / 72 <= 12.2 or pt == 19):
                if pt < cur:
                    for r_ in runs:
                        r_.font.size = Pt(pt)
                break
        return


def kr_notes_fix(slide):
    nt = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
    if nt is None:
        return
    t = nt.text
    t2 = re.sub(r"검정\s*\((veh-SBI|vehicle vs SBI-553)\)", r"검은 괄호 (\1)", t)
    if t2 != t:
        nt.text = t2


def kr_notes(slide, en_slide):
    """the Korean notes keep the English statistics when the translation is much shorter (10-22 triple check)"""
    if en_slide is None or not en_slide.has_notes_slide:
        return False
    en = en_slide.notes_slide.notes_text_frame.text.strip()
    kr = slide.notes_slide.notes_text_frame.text.strip()
    has_kr = any("가" <= ch <= "힣" for ch in kr)
    if en and (not has_kr or len(kr) < 0.6 * len(en)):
        slide.notes_slide.notes_text_frame.text = (kr + "\n\n" if has_kr else "") + "원문 (EN): " + en
        return True
    return False


def key_pt(t, kr):
    """key line size: one line of 12.2 in at 16 or 14 pt, else 12.5 pt (10-23: KR keys were 12.5 pt where 16 fits)"""
    if not kr:
        n_ = len(t)
        return 16 if n_ <= 115 else 14 if n_ <= 150 else 12.5
    em = sum(1.0 if "\uac00" <= ch <= "\ud7a3" else 0.3 if ch == " " else 0.56 for ch in t)
    for pt in (16, 14):
        if em * pt / 72 <= 12.1:
            return pt
    return 12.5


def resize_key(slide, lang):
    """the red key line of a copied slide at the size key_pt gives (never smaller than it was)"""
    for sh in slide.shapes:
        if not (sh.has_text_frame and Inches(0.8) <= (sh.top or 0) <= Inches(1.3) and sh.text_frame.text.strip()):
            continue
        runs = [r_ for pp in sh.text_frame.paragraphs for r_ in pp.runs]
        if not runs or not runs[0].font.bold:
            continue
        cur = max([r_.font.size.pt for r_ in runs if r_.font.size] or [0])
        new = key_pt(sh.text_frame.text.strip(), lang == "kr")
        if cur and new > cur:
            for r_ in runs:
                r_.font.size = Pt(new)
        return


def apply_take(slide, en_title, lang):
    row = next((r for r in TAKE if re.search(r[0], en_title)), None)
    if row is None:
        return False
    kr = lang == "kr"
    opt = row[3]
    T = sorted([sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()], key=lambda sh: sh.top or 0)
    keyb = [sh for sh in T if Inches(0.8) <= (sh.top or 0) <= Inches(1.3)]
    nt = slide.notes_slide.notes_text_frame
    new = row[2] if kr else row[1]
    if keyb and new:
        tf = keyb[0].text_frame
        old = tf.text.strip()
        if _set_text(tf, new):
            tf.paragraphs[0].runs[0].font.size = Pt(key_pt(new, kr))
            nt.text = (nt.text + "\n\n" + ("[이전 버전, 대체됨] 이전 핵심 문장: " if kr else "[Superseded, earlier version] key line: ")
                       + old).strip()
    elif new:                                                   # no key line on the slide (10-23: the first limits slide)
        tb_ = slide.shapes.add_textbox(Inches(0.45), Inches(0.92), Inches(12.4), Inches(0.45))
        tb_.text_frame.word_wrap = True
        r_ = tb_.text_frame.paragraphs[0].add_run(); r_.text = new
        r_.font.size = Pt(key_pt(new, kr)); r_.font.bold = True; r_.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        r_.font.name = "Malgun Gothic" if kr else "Arial"
    nrep = opt.get("notes_kr" if kr else "notes_en")
    if nrep:                                                    # notes copied from another part: its own numbers first
        nt.text = (nrep + "\n\n" + (("[이전 버전, 대체됨] 다른 부분에서 복사된 노트: " if kr else
                                       "[Superseded, earlier version] notes copied from another part: ") + nt.text)).strip()
    nadd = opt.get("notes_add_kr" if kr else "notes_add_en")
    if nadd:
        nt.text = (nadd + "\n\n" + nt.text).strip()
    tnew = opt.get("title_kr" if kr else "title_en")
    if tnew:
        tb = [sh for sh in T if Inches(0.25) <= (sh.top or 0) < Inches(0.8)]
        if tb and _set_text(tb[0].text_frame, tnew):
            tb[0].text_frame.paragraphs[0].runs[0].font.size = Pt(26 if len(tnew) <= 55 else 22 if len(tnew) <= 68 else 19)
    bnew = opt.get("bottom_kr" if kr else "bottom_en")
    if bnew:
        bb = [sh for sh in T if (sh.top or 0) >= Inches(6.5)]
        if bb:
            oldb = bb[-1].text_frame.text.strip()
            if _set_text(bb[-1].text_frame, bnew):
                nt.text = (nt.text + "\n\n" + (("[이전 버전, 대체됨] 아래 문장: " if kr else "[Superseded, earlier version] bottom line: ")
                                                + oldb)).strip()
    return True


_NT_SENT = re.compile(r"\bn\s*=|\bTest\b|\btest:|[Dd]escriptive|검정|기술 통계|마리", re.I)


def trim_bottom(slide, lang):
    """too wordy (user, 10-22): the bottom keeps its first sentence and the n / test sentence; the rest goes into the notes"""
    boxes = [sh for sh in slide.shapes if sh.has_text_frame and (sh.top or 0) >= Inches(6.5) and sh.text_frame.text.strip()]
    if not boxes:
        return False
    b = boxes[-1]
    full = " ".join(pp.text.strip() for pp in b.text_frame.paragraphs if pp.text.strip())
    sents = [x for x in re.split(r"(?<=\.)\s+", full) if x.strip()]
    if len(sents) <= 2:
        return False
    first = 1 if (re.match(r"^(So|그래서|따라서) ", sents[0]) and len(sents) > 1) else 0     # a 'So ...' link needs the key line above
    keep = [sents[first]] + [x for j, x in enumerate(sents) if j != first and _NT_SENT.search(x)]
    rest = [x for j, x in enumerate(sents) if j != first and not _NT_SENT.search(x)]
    if len(" ".join(keep).split()) > 38:                 # 10-24 (user: details in the notes): of the n / test sentences keep only
        k2 = [keep[0]]                                     # their n / test clauses
        for x in keep[1:]:
            parts = [c for c in re.split(r";\s+", x.rstrip(".")) if c.strip()]
            pk = [c for c in parts if _NT_SENT.search(c)] or parts
            rest += [c + "." for c in parts if c not in pk]
            k2.append("; ".join(pk) + ".")
        keep = k2
    if not rest:
        return False
    tf = b.text_frame
    runs = tf.paragraphs[0].runs
    if not runs:
        return False
    runs[0].text = " ".join(keep)
    for x in runs[1:]:
        x.text = ""
    for pp in list(tf.paragraphs[1:]):
        pp._p.getparent().remove(pp._p)
    nt = slide.notes_slide.notes_text_frame
    nt.text = ((("자세히: " if lang == "kr" else "More: ") + " ".join(rest)) + "\n\n" + nt.text).strip()
    return True


HEAD = ("SBI-553: less movement (3/4 mice, 2 strongly), bigger firing drop (3/4), fewer pin UP neurons than vehicle (Day 2)",
        "SBI-553: 움직임 감소 (3/4마리, 2마리 크게), 활동 감소 더 큼 (3/4), pin UP 뉴런 vehicle보다 적음 (Day 2)")


def cards6(dk, section, title, key, cards, bottom, notes=""):
    """3 x 2 cards (10-23: the cover gives Q4 and Q5 their own card; the second limits slide holds 6 limits)"""
    from pptx.util import Inches as In
    from pptx.dml.color import RGBColor as RGB
    from pptx.enum.shapes import MSO_SHAPE
    s = dk.head(section, title, key)
    for i, (big, head, txt) in enumerate(cards):
        x = 0.45 + (i % 3) * 4.17; y = 1.5 + (i // 3) * 2.36
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, In(x), In(y), In(4.0), In(2.26))
        r.fill.solid(); r.fill.fore_color.rgb = RGB(0xF3, 0xF5, 0xF8); r.line.fill.background()
        dk.box(s, x + 0.1, y + 0.1, 1.3, 0.75, big, 28 if len(big) <= 3 else 22 if len(big) <= 5 else 18, True, (0xC0, 0x39, 0x2B))
        dk.box(s, x + 1.4, y + 0.1, 2.5, 0.8, head, 14, True, (0x1F, 0x2D, 0x3D))
        dk.box(s, x + 1.4, y + 0.88, 2.5, 1.34, txt, 12)
    if bottom:
        dk.box(s, 0.45, 6.3, 12.4, 1.0, bottom, 14)
    s.notes_slide.notes_text_frame.text = notes
    return s


def one_page_summary(dk, lang):
    """the last slide (user, 10-22): what we found, on one page (numbers re-checked on the data 10-23)"""
    from pptx.util import Inches as In, Pt as P_
    kr = lang == "kr"
    L = lambda en, k_: k_ if kr else en
    s = dk.head(L("SUMMARY", "요약"), L("What we found, on one page", "한 페이지로 본 결과"), L(*HEAD))
    for sh in s.shapes:                                          # the headline at the size that keeps it on one line
        if sh.has_text_frame and sh.text_frame.text.strip() == L(*HEAD):
            for r_ in sh.text_frame.paragraphs[0].runs:
                r_.font.size = P_(key_pt(L(*HEAD), kr))
    rows = [(L("Q1  Does SBI-553 change pain behaviour?", "Q1  SBI-553가 통증 행동을 바꾸나?"),
             L("Less movement after the deliveries than after vehicle in 3/4 mice (0.56 vs 0.76 per delivery), strongly only in the two "
               "SBI-553-first mice. Vs vehicle: less movement after light touch (39 vs 61%), fewer flinches after pin prick (57 vs 79%, a "
               "vehicle-day rise), less escape after heat (1 vs 10%, ns).",
               "자극 후 움직임이 vehicle 후보다 적음, 4마리 중 3마리 (자극당 0.56 vs 0.76); 크게 줄어든 것은 SBI-553을 먼저 받은 두 마리. vehicle "
               "대비: light touch 후 움직임 (39 vs 61%), pin prick 후 flinch (57 vs 79%, vehicle 날 증가 때문) 감소; heat 후 escape (1 vs 10%)는 ns."),
             L("3/4 mice (over mice p = 0.25); per delivery: CMH p < 0.05 (heat escape p = 0.06)", "3/4마리 (마우스 단위 p = 0.25); 자극 단위 CMH p < 0.05 (heat escape p = 0.06)")),
            (L("Q2  Does it change BLA activity?", "Q2  BLA 활동을 바꾸나?"),
             L("Firing fell after both injections; the same-day drop was bigger after SBI-553 on both days (Day 1 -54 vs -37%, Day 2 -68 "
               "vs -36%; 3/4 mice). The level gap after the injection (1.3 vs 2.7 /min) is a Day-2 result: those SBI-553 mice fired "
               "little already (0.80 vs 3.9). The stimulus-block rise was smaller after SBI-553 (+0.04 vs +0.07 z).",
               "두 주사 모두 활동 감소; 같은 날 감소는 두 날 모두 SBI-553 후 더 큼 (Day 1 -54 vs -37%, Day 2 -68 vs -36%; 3/4마리). 주사 후 "
               "수준 차이 (분당 1.3 vs 2.7)는 Day 2 결과: 그 SBI-553 마우스는 원래 활동이 적음 (0.80 vs 3.9). 자극 블록 증가도 SBI-553 후 더 "
               "작음 (+0.04 vs +0.07 z)."),
             L("same-day drop: per mouse, descriptive; level gap: Mann-Whitney p = 0.01, neurons pooled, Day 2 only",
               "같은 날 감소: 마우스별 기술 통계; 수준 차이: Mann-Whitney p = 0.01, 뉴런 합산, Day 2만")),
            (L("Q3  Does it change the responses to the pain stimuli?", "Q3  통증 자극 반응을 바꾸나?"),
             L("Population response: no change from baseline. Pin UP rose after vehicle (12 -> 22%, mostly 5939, Day 2), not after SBI-553 (11 -> 10%"
               "); the drug gap is Day 2 only (12 vs 42%; Day 1 8 vs 2%, baseline 7 vs 1%). Heat UP rose after both, light UP after SBI-553 only. Baseline "
               "movement-onset neurons grew after SBI-553 but weakened after vehicle (***).",
               "집단 반응: 기준 대비 변화 없음. pin UP은 vehicle 후 증가 (12 -> 22%, 주로 5939, Day 2), SBI-553 후는 그대로 (11 -> 10%); "
               "약물 차이는 Day 2에만 (12 vs 42%; Day 1 8 vs 2%, 기준 7 vs 1%). heat UP은 둘 다 증가, light UP은 SBI-553 후만. 기준 movement onset 뉴런은 SBI-553 후 강해지고 "
               "vehicle 후 약해짐 (***)."),
             L("time-shift test per neuron, McNemar / Fisher, Mann-Whitney (neurons pooled); pin: smaller SBI-day change in 3/4 mice, "
               "vehicle rise * with max-T (p = 0.02)",
               "뉴런별 시간 이동 검정, McNemar / Fisher, Mann-Whitney (뉴런 합산); pin: SBI 날 변화가 더 작은 마우스 3/4, max-T로 vehicle 증가 * (p = 0.02)")),
            (L("Q4  Stimulus or movement?", "Q4  자극인가 움직임인가?"),
             L("Movement onset is the strongest BLA signal; its responders rose after both injections (SBI-553 4/4 mice, vehicle 2/4). 55% "
               "of neuron-stimulus pairs keep >= half their response without a withdrawal, yet only 2-7% of neurons follow the stimulus alone.",
               "가장 강한 BLA 신호는 movement onset; 그 반응 뉴런은 두 주사 후 모두 증가 (SBI-553 4/4마리, vehicle 2/4). 뉴런-자극 쌍의 55%는 "
               "withdrawal 없이도 반응 절반 이상 유지, 그래도 자극만 따르는 뉴런은 2-7%뿐."),
             L("BH per neuron (max-T 3/4 mice); Wilcoxon; encoding model (2-7% is a lower bound, check 3); Spearman over 16 sessions: ns except % DOWN vs movement (rho -0.60, p = 0.01)",
               "뉴런별 BH (max-T 3/4마리); Wilcoxon; encoding model (2-7%는 하한값, 검증 3); 16개 세션 Spearman: % DOWN vs 움직임 (rho -0.60, p = 0.01) 외 모두 ns")),
            (L("Q5  Sex?", "Q5  성별?"),
             L("Candidates only, both females vs both males: bigger firing drop on the SBI-553 day (stimulus blocks, log2 -1.58 / -0.53 vs -0.22 / 0.00), "
               "light-touch neurons after SBI-553 overlap (40 / 15 vs 29 / 0%), more withdrawal after vehicle (61 / 70 vs 54 / 55%).",
               "후보일 뿐, 두 암컷 vs 두 수컷: SBI-553 날 활동 감소 더 큼 (자극 블록, log2 -1.58 / -0.53 vs -0.22 / 0.00), SBI-553 후 light touch 뉴런 "
               "겹침 (40 / 15 vs 29 / 0%), vehicle 후 withdrawal 더 많음 (61 / 70 vs 54 / 55%)."),
             L("2 vs 2 mice: descriptive only", "2 vs 2마리: 기술 통계만"))]
    tb = s.shapes.add_table(len(rows) + 1, 3, In(0.45), In(1.45), In(12.4), In(5.1)).table
    for j, (w_, h_) in enumerate(((2.7, L("question", "질문")), (6.6, L("what we found", "결과")), (3.1, L("how sure (n, test)", "확실성 (n, 검정)")))):
        tb.columns[j].width = In(w_)
        c = tb.cell(0, j); c.text = h_
        for pp in c.text_frame.paragraphs:
            pp.font.size = P_(12.5); pp.font.bold = True; pp.font.name = dk.font
    tb.rows[0].height = In(0.35)
    for i, r in enumerate(rows, 1):
        tb.rows[i].height = In(1.05 if i == 3 else 0.9)
        for j, t_ in enumerate(r):
            c = tb.cell(i, j); c.text = t_
            c.margin_top = c.margin_bottom = In(0.03)
            for pp in c.text_frame.paragraphs:
                pp.font.size = P_((9.5 if j else 10.5) if kr else (10.5 if j else 11.5)); pp.font.bold = (j == 0); pp.font.name = dk.font   # 10-24: the longer timefix texts
    dk.box(s, 0.45, 6.72, 12.4, 0.7, L("n = 4 mice (2 F, 2 M), crossover; with 4 mice no test across mice can reach p < 0.05 (min 0.125): "
                                       "every neural result is exploratory. Next: more mice, then population dynamics.",
                                       "n = 4마리 (암 2, 수 2), crossover; 4마리로는 마우스 간 검정이 p < 0.05에 닿을 수 없음 (최소 0.125): "
                                       "뉴런 결과는 모두 탐색적. 다음: 마우스 추가, 그다음 집단 동역학."), 13)
    s.notes_slide.notes_text_frame.text = L(
        "Sources (re-checked 10-23): behaviour per delivery / per stimulus (crossover_per_session.csv; CMH: P6_behaviour_tests.csv), "
        "activity of all neurons (stimulus-block rate per neuron), population response to each stimulus, neurons UP / DOWN per stimulus "
        "(cells_classes.csv), check 4 (per mouse, max-T), stimulus vs withdrawal (O2 / O3), brain and behaviour (P7), sex.",
        "출처 (10-23 다시 확인): 자극당 / 자극별 행동 (crossover_per_session.csv; CMH: P6_behaviour_tests.csv), 모든 뉴런 활동 (뉴런별 "
        "자극 블록 rate), 자극별 집단 반응, 자극별 UP / DOWN 뉴런 (cells_classes.csv), 검증 4 (마우스별, max-T), 자극 vs withdrawal "
        "(O2 / O3), 뇌와 행동 (P7), 성별.")
    return s


def limits2(dk, lang):
    """what else limits the results (10-22 story judge; 10-23: Day-2 methods and the automatic detection added, CMH wording fixed)"""
    L = lambda en, kr: kr if lang == "kr" else en
    s = cards6(dk, L("LIMITS", "한계"), L("What the data can and cannot say yet (2/2)", "지금 데이터로 말할 수 있는 것과 없는 것 (2/2)"),
               L("Drug order, Day-2 methods, a lenient responder rule and pooled deliveries can make effects look larger than they are",
                 "투여 순서, Day 2 방법 차이, 너그러운 반응 판정, 자극 합산은 효과를 실제보다 크게 보이게 할 수 있음"),
               [(L("day", "날"), L("drug order", "투여 순서"),
                 L("SBI-553 on Day 1 = 5941, 5939; Day 2 = 5937, 1191. The strong movement drop follows this order; the pin gap is "
                   "Day 2 only (Day 1: 8 vs 2%, baseline 7 vs 1%).",
                   "Day 1 SBI-553 = 5941, 5939; Day 2 = 5937, 1191. 큰 움직임 감소가 이 순서를 따름; pin prick 차이는 Day 2에만 (Day 1: "
                   "8 vs 2%, 기준 7 vs 1%).")),
                (L("Day 2", "Day 2"), L("method and level differ", "방법과 수준이 다름"),
                 L("Day 1: key marks, Day 2: hand clicker (both re-timed 10-23). The Day-2 SBI-553 mice fired 0.80 vs 3.9 /min at baseline, so fewer "
                   "responders can be detected in them, whatever the drug.",
                   "Day 1 = 키 표시, Day 2 = 손 클리커 (둘 다 10-23 재정렬). Day 2 SBI-553 마우스는 기준 활동이 분당 0.80 vs 3.9라 약물과 상관없이 반응 뉴런이 "
                   "덜 잡힐 수 있음.")),
                (L("q 0.20", "q 0.20"), L("lenient responder call", "너그러운 반응 판정"),
                 L("BH q <= 0.20 per session; the strict max-T keeps ~8% of neurons as responsive (check 1).",
                   "세션마다 BH q <= 0.20; 엄격한 max-T는 뉴런 ~8%만 반응으로 남김 (검증 1).")),
                (L("CMH", "CMH"), L("deliveries pooled", "자극 합산"),
                 L("Behaviour stars treat the deliveries inside each mouse as independent (mouse = stratum); with mice as the unit the "
                   "smallest p is 0.125.",
                   "행동 별표는 마우스 안 자극을 독립으로 취급 (마우스 = 층); 마우스를 단위로 하면 최소 p는 0.125.")),
                (L("12", "12"), L("clicker presses checked", "확인한 클리커 누름"),
                 L("The clicker indicator stayed on > 1 s at 12 Day-2 deliveries; checked on the pressure trace (10-23): 1 removed, 3 presses added, 8 kept.",
                   "Day 2 자극 12번에서 로거 표시가 1초 넘게 켜짐; 10-23 압력 신호로 확인: 1개 제거, 누름 3개 추가, 8개 유지.")),
                (L("auto", "자동"), L("movement / grooming", "움직임 / grooming"),
                 L("Detected automatically and checked by eye on example clips only; no frame-by-frame agreement with hand scoring yet.",
                   "자동 검출, 예시 클립만 눈으로 확인; 아직 직접 채점과 프레임 단위 비교 없음."))],
               L("n = 4 mice (crossover); descriptive summary of the limits (no test).", "n = 4마리 (crossover); 한계의 기술 요약 (검정 없음)."),
               L("Also: pin responders are higher on Day 2 even at baseline (3/4 mice; heat 2/4), so day and drug order "
                 "overlap for every responder result.",
                 "또한: pin 반응 뉴런은 기준에서도 Day 2가 높음 (3/4마리; heat 2/4) - 반응 뉴런 결과마다 날과 투여 순서가 겹침."))
    return s


def divider(dk, lang, title, sub, cards=None, six=False, notes=""):
    L = lambda en, kr: kr if lang == "kr" else en
    bottom = L("Mini1p Ca imaging of the BLA and pain behaviour; n = 4 mice (2 F, 2 M), crossover. " + sub + ". Tests, named on every "
               "result slide: Wilcoxon, Mann-Whitney, McNemar, Fisher, CMH, time-shift test.",
               "BLA mini1p Ca 이미징과 통증 행동; n = 4마리 (암 2, 수 2), crossover. " + sub + ". 검정 (결과 슬라이드마다 표시): "
               "Wilcoxon, Mann-Whitney, McNemar, Fisher, CMH, 시간 이동 검정.")
    if six:
        s = cards6(dk, L("CONTENTS", "목차"), title, L(*HEAD), cards or [], bottom, notes)
    else:
        s = dk.cards(L("CONTENTS", "목차"), title, L(*HEAD), cards or [], bottom, notes)
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == L(*HEAD):
            for r_ in sh.text_frame.paragraphs[0].runs:
                r_.font.size = Pt(key_pt(L(*HEAD), lang == "kr"))
    return s


def glance2(dk, lang):
    """RESULTS AT A GLANCE, one card per question (10-23 round 3: Q4 and Q5 added; Q2 = the same-day drop; the pin numbers are
    both days pooled, the gap is Day 2); numbers checked 10-23"""
    L = lambda en, kr: kr if lang == "kr" else en
    s = cards6(dk, L("RESULTS AT A GLANCE", "결과 한눈에"), L("Day 1 + Day 2: the answers at a glance", "Day 1 + Day 2: 답을 한눈에"),
               L("SBI-553: less movement (3/4 mice), bigger same-day firing drop (3/4); pin UP 10 vs 22% after vehicle (both days; gap on Day 2)",
                 "SBI-553: 움직임 감소 (3/4마리), 같은 날 활동 감소 더 큼 (3/4); pin UP 10 vs 22% (두 날 합산; 차이는 Day 2뿐)"),
               [("3/4", L("Q1 less movement after SBI-553 than vehicle", "Q1 SBI-553 후 움직임이 vehicle보다 적음"),
                 L("0.56 vs 0.76 per delivery; strong only in the SBI-553-first mice (5941, 5939). Over 4 mice: p = 0.25.",
                   "자극당 0.56 vs 0.76; 크게 줄어든 것은 SBI-553을 먼저 받은 두 마리 (5941, 5939). 4마리: p = 0.25.")),
                ("3/4", L("Q2 bigger same-day firing drop after SBI-553", "Q2 같은 날 활동 감소가 SBI-553 후 더 큼"),
                 L("Day 1 -54 vs -37%, Day 2 -68 vs -36%. The level 1.3 vs 2.7 /min (p = 0.01, neurons pooled) is Day-2 driven.",
                   "Day 1 -54 vs -37%, Day 2 -68 vs -36%. 수준 1.3 vs 2.7/분 (p = 0.01, 뉴런 합산)은 Day 2가 주도.")),
                ("ns", L("Q3 average population response", "Q3 집단 평균 반응"),
                 L("No stimulus changed the mean response of all neurons from baseline (Wilcoxon: all ns); pin SBI-553 < vehicle (Mann-Whitney p = 0.04) also at baseline (p = 0.06).",
                   "어느 자극에서도 전체 뉴런 평균 반응이 기준 대비 변화 없음 (Wilcoxon: 모두 ns); pin SBI-553 < vehicle (Mann-Whitney p = 0.04)은 기준에서도 (p = 0.06).")),
                ("10%", L("Q3 pin UP neurons after SBI-553 (22% after vehicle)", "Q3 SBI-553 후 pin UP 뉴런 (vehicle 후 22%)"),
                 L("Both days pooled; the gap is Day 2 only (12 vs 42%; Day 1 8 vs 2%, baseline 7 vs 1%). Smaller SBI-day change in 3/4 mice; max-T: vehicle rise * (p = 0.02).",
                   "두 날 합산; 차이는 Day 2에만 (12 vs 42%; Day 1 8 vs 2%, 기준 7 vs 1%). SBI 날 변화가 더 작은 마우스 3/4; max-T: vehicle 증가 * (p = 0.02).")),
                ("2-7%", L("Q4 neurons that follow the stimulus alone", "Q4 자극만 따르는 뉴런"),
                 L("Movement onset is the strongest signal; movement alone explains 38-54% of neurons; 2-7% is a lower bound (check 3).",
                   "movement onset이 가장 강한 신호; 움직임만으로 뉴런 38-54%; 2-7%는 하한값 (검증 3).")),
                ("2+2", L("Q5 sex: candidates only", "Q5 성별: 후보만"),
                 L("Both females vs both males: bigger firing drop on the SBI-553 day, more withdrawal after vehicle; 2 vs 2 mice cannot be tested.",
                   "두 암컷 vs 두 수컷: SBI-553 날 활동 감소 더 큼, vehicle 후 withdrawal 더 많음; 2 vs 2마리는 검정 불가."))],
               L("n = 4 mice (crossover): 245 neurons on Day 1, 217 on Day 2 (185 followed on both days). Tests on each result slide; "
                 "neurons pooled = exploratory.",
                 "n = 4마리 (crossover): Day 1 뉴런 245, Day 2 217 (두 날 모두 추적 185). 검정은 각 결과 슬라이드; 뉴런 합산 = 탐색적."),
               L("One card per question. Pin numbers: 24 / 246 vs 47 / 216 neurons (both days); Day 2 13 / 110 vs 45 / 107. Firing = median "
                 "stimulus-block rate. The earlier glance card (185 neurons, pin +0.002) is in the supplement of the FULL deck.",
                 "질문마다 카드 하나. pin 수: 뉴런 24 / 246 vs 47 / 216 (두 날); Day 2 13 / 110 vs 45 / 107. 활동 = 자극 블록 rate 중앙값. "
                 "이전 카드 (185 뉴런, pin +0.002)는 FULL 덱 부록에."))
    return s


DAYC = {
    1: dict(
        title=("Day 1 alone (2 vs 2 different mice): SBI-553 vs vehicle", "Day 1만 따로 (서로 다른 2 vs 2마리): SBI-553 vs vehicle"),
        key=("Day 1: the SBI-553 mice moved less (0.40 vs 0.76) and their firing fell more (-54 vs -37%); UP heat 26 vs 11%, light 16 vs 1%; pin 8 vs 2% (baseline 7 vs 1%)",
             "Day 1: SBI-553 마우스가 덜 움직이고 (0.40 vs 0.76) 활동이 더 감소 (-54 vs -37%); UP heat 26 vs 11%, light 16 vs 1%; pin 8 vs 2% (기준 7 vs 1%)"),
        cards=[("0.40", ("movement per delivery after SBI-553 (vehicle mice 0.76)", "SBI-553 후 자극당 움직임 (vehicle 마우스 0.76)"),
                ("Same-day baselines 0.54 / 0.67. SBI-553 = 5941, 5939; vehicle = 1191, 5937. 2 vs 2 different mice: not testable.",
                 "같은 날 기준 0.54 / 0.67. SBI-553 = 5941, 5939; vehicle = 1191, 5937. 서로 다른 2 vs 2마리: 검정 불가.")),
               ("3.5/min", ("median firing after SBI-553 (vehicle mice 2.7)", "SBI-553 후 활동 중앙값 (vehicle 마우스 2.7)"),
                ("Baselines 7.7 vs 4.4: same-day drop -54% vs -37%. Neurons pooled: Mann-Whitney p = 0.56.",
                 "기준 7.7 vs 4.4: 같은 날 감소 -54% vs -37%. 뉴런 합산: Mann-Whitney p = 0.56.")),
               ("8%", ("pin-prick UP neurons after SBI-553 (vehicle mice 2%)", "SBI-553 후 pin prick UP 뉴런 (vehicle 마우스 2%)"),
                ("11 / 136 vs 2 / 109; baselines 7% / 1%. Fisher p = 0.04 (neurons pooled).",
                 "11 / 136 vs 2 / 109; 기준 7% / 1%. Fisher p = 0.04 (뉴런 합산).")),
               ("26%", ("heat UP neurons after SBI-553 (vehicle mice 11%)", "SBI-553 후 heat UP 뉴런 (vehicle 마우스 11%)"),
                ("35 / 136 vs 12 / 109; baselines 1% / 3%. Fisher p = 0.005 (neurons pooled).",
                 "35 / 136 vs 12 / 109; 기준 1% / 3%. Fisher p = 0.005 (뉴런 합산).")),
               ("16%", ("light-touch UP neurons after SBI-553 (vehicle mice 1%)", "SBI-553 후 light touch UP 뉴런 (vehicle 마우스 1%)"),
                ("22 / 136 vs 1 / 109, mostly 5941 (14 / 67); baselines 1% / 18%. Fisher p < 0.001 (neurons pooled).",
                 "22 / 136 vs 1 / 109, 주로 5941 (14 / 67); 기준 1% / 18%. Fisher p < 0.001 (뉴런 합산).")),
               ("-0.026", ("pin-prick population response after SBI-553 (z)", "SBI-553 후 pin prick 집단 반응 (z)"),
                ("vs +0.004 after vehicle and -0.011 at baseline; all neurons pooled, p = 0.07.",
                 "vehicle 후 +0.004, 기준 -0.011; 전체 뉴런 합산, p = 0.07."))],
        bottom=("Day 1 alone: n = 2 + 2 mice (245 neurons); different mice per drug, so this checks the drug order. Test: Fisher exact, "
                "Mann-Whitney U (neurons pooled = exploratory).",
                "Day 1만: n = 2 + 2마리 (뉴런 245); 약물마다 다른 마우스라 투여 순서 확인용. 검정: Fisher exact, Mann-Whitney U (뉴런 합산 = 탐색적)."),
        notes=("Firing = median stimulus-block rate of all neurons (transients / min). Day 1 deliveries = hand marks. The light-touch and "
               "population cards are the earlier Day-1 cards, kept (10-23).",
               "활동 = 모든 뉴런의 자극 블록 rate 중앙값 (분당 transient). Day 1 자극 = 직접 표시. light touch·집단 반응 카드는 이전 Day 1 "
               "카드를 그대로 둔 것 (10-23).")),
    2: dict(
        title=("Day 2 alone (2 vs 2 different mice): SBI-553 vs vehicle", "Day 2만 따로 (서로 다른 2 vs 2마리): SBI-553 vs vehicle"),
        key=("Day 2: pin UP rose after vehicle (22 -> 42%, mostly 5939), not after SBI-553 (16 -> 12%); heat UP rose in both (14 -> 35%, 3 -> 22%)",
             "Day 2: pin UP은 vehicle 후만 증가 (22 -> 42%, 주로 5939; SBI-553 16 -> 12%); heat UP은 둘 다 증가"),
        cards=[("0.71", ("movement per delivery after SBI-553 (vehicle mice 0.75)", "SBI-553 후 자극당 움직임 (vehicle 마우스 0.75)"),
                ("Same-day baselines 0.80 / 0.78. SBI-553 = 1191, 5937; vehicle = 5939, 5941. 2 vs 2 different mice: not testable.",
                 "같은 날 기준 0.80 / 0.78. SBI-553 = 1191, 5937; vehicle = 5939, 5941. 서로 다른 2 vs 2마리: 검정 불가.")),
               ("0.26/min", ("median firing after SBI-553 (vehicle mice 2.5)", "SBI-553 후 활동 중앙값 (vehicle 마우스 2.5)"),
                ("Baselines 0.80 vs 3.9 (the SBI-553 mice were already low): same-day drop -68% vs -36%. Mann-Whitney p < 0.001.",
                 "기준 0.80 vs 3.9 (SBI-553 마우스는 이미 낮음): 같은 날 감소 -68% vs -36%. Mann-Whitney p < 0.001.")),
               ("12%", ("pin-prick UP neurons after SBI-553 (vehicle mice 42%)", "SBI-553 후 pin prick UP 뉴런 (vehicle 마우스 42%)"),
                ("13 / 110 vs 45 / 107; baselines 16% / 22%; the vehicle rise is mostly 5939 (2 -> 21 / 47). Fisher p < 0.001.",
                 "13 / 110 vs 45 / 107; 기준 16% / 22%; vehicle 증가는 주로 5939 (2 -> 21 / 47). Fisher p < 0.001.")),
               ("22%", ("heat UP neurons after SBI-553 (vehicle mice 35%)", "SBI-553 후 heat UP 뉴런 (vehicle 마우스 35%)"),
                ("24 / 110 vs 37 / 107; baselines 3% / 14%. Fisher p = 0.049 (neurons pooled).",
                 "24 / 110 vs 37 / 107; 기준 3% / 14%. Fisher p = 0.049 (뉴런 합산).")),
               ("0.08", ("attending per delivery after SBI-553 (vehicle mice 0.09)", "SBI-553 후 자극당 attending (vehicle 마우스 0.09)"),
                ("Same-day baselines 0.09 / 0.03. 2 vs 2 different mice: not testable.",
                 "같은 날 기준 0.09 / 0.03. 서로 다른 2 vs 2마리: 검정 불가.")),
               ("+0.038", ("pin-prick population response after SBI-553 (z)", "SBI-553 후 pin prick 집단 반응 (z)"),
                ("vs +0.062 after vehicle and +0.042 at baseline; all neurons pooled, p = 0.10.",
                 "vehicle 후 +0.062, 기준 +0.042; 전체 뉴런 합산, p = 0.10."))],
        bottom=("Day 2 alone: n = 2 + 2 mice (217 neurons); different mice per drug, so this checks the drug order. Test: Fisher exact, "
                "Mann-Whitney U (neurons pooled = exploratory).",
                "Day 2만: n = 2 + 2마리 (뉴런 217); 약물마다 다른 마우스라 투여 순서 확인용. 검정: Fisher exact, Mann-Whitney U (뉴런 합산 = 탐색적)."),
        notes=("Day 2 deliveries come from the hand clicker (Day 1: key marks); both re-timed 10-23. Firing = median stimulus-block rate of all neurons (transients / "
               "min). The attending and population cards are the earlier Day-2 cards, kept (10-23).",
               "Day 2 자극은 손 클리커 (Day 1: 키 표시); 둘 다 10-23 재정렬. 활동 = 모든 뉴런의 자극 블록 rate 중앙값 (분당 transient). attending·집단 반응 카드는 "
               "이전 Day 2 카드를 그대로 둔 것 (10-23).")),
}


def daycard(dk, lang, d):
    """HOW SURE?: Day 1 alone / Day 2 alone with the SAME four cards (10-23 round 2: the old cards picked different events and hid
    that the pin-prick gap is Day 2 only); numbers checked 10-23 (cells_classes.csv, crossover_per_session.csv, whole-session rates)"""
    k = 1 if lang == "kr" else 0
    D = DAYC[d]
    s = cards6(dk, "얼마나 확실한가?" if k else "HOW SURE?", D["title"][k], D["key"][k],
               [(big, h[k], t[k]) for big, h, t in D["cards"]], D["bottom"][k], D["notes"][k])
    return s


def deck(path):
    return Presentation(path) if os.path.isfile(path) else None


# 10-24 (user screenshots): slide kinds the PI asked to SEE in the PI deck (not hidden): the all-mice population response per drug day
# (stimuli), responsive neurons only, up / down counts and their baseline-normalized change, behaviour per stimulus normalized
PI_SHOW = (r"^(All mice, (vehicle|SBI-553) day: population response, .*\(stimuli\)|All mice, responsive neurons only: .*\(stimu"
           r"|All mice: how many neurons go up or down|All mice: change in % UP|Behaviour per stimulus, normalized)")


def mark_backup(slide, lang, hide=True):
    """10-23 (user): a backup slide of the PI deck - light grey background, a grey 'BACKUP' tag, hidden in the slide show
    (right-click > Unhide Slide in PowerPoint to show it)"""
    from pptx.dml.color import RGBColor as RGB
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = RGB(0xD9, 0xD9, 0xD9)            # 10-24: visibly grey (was EEEFF1)
    tb = slide.shapes.add_textbox(Inches(10.9), Inches(7.08), Inches(2.3), Inches(0.3))
    pp = tb.text_frame.paragraphs[0]
    pp.text = "BACKUP (hidden)" if lang != "kr" else "보조 슬라이드 (숨김)"
    pp.font.size = Pt(11); pp.font.bold = True; pp.font.color.rgb = RGB(0x80, 0x86, 0x90)
    pp.font.name = "Malgun Gothic" if lang == "kr" else "Arial"
    if hide:
        slide._element.set("show", "0")


def _isbk(t):
    """the [BACKUP] tag of a notes text: at the start, or after the figure key moved into the notes (10-24)"""
    return bool(re.search(r"(?:^|\n)\[BACKUP\]", t or ""))


def _notes(slide):
    try:
        return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    except Exception:
        return ""


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--lang", default="en", choices=["en", "kr"])
    ap.add_argument("--kinds", nargs="+", default=["FULL", "PI"], choices=["FULL", "PI"])
    a = ap.parse_args(argv)
    lg = a.lang
    L = lambda en, kr: kr if lg == "kr" else en
    sub, suf = ("KR", "_KR") if lg == "kr" else ("EN", "")
    p = lambda st, name: os.path.join(V, st, sub, f"{name}{suf}.pptx")
    NSD = os.environ.get("NEW_STYLE_DIR", "NEW_STYLE")              # 10-23: NEW_STYLE_v3 (a deck of NEW_STYLE was open in PowerPoint)
    src = {k: deck(p(NSD, k)) for k in ("Day1+2_FULL", "Day1+2_PI", "Day1_FULL", "Day2_FULL", "Day1_PI", "Day2_PI", "PerMouse_Day1+2")}
    srcEN = {k: deck(os.path.join(V, NSD, "EN", f"{k}.pptx")) for k in src}     # same slides, English titles (story keys)
    print("new-style decks from", NSD)
    old = {k: deck(p("OLD_STYLE", k)) for k in ("Day1_FULL", "Day2_FULL", "Day1+2_longitudinal")}
    oldEN = {k: deck(os.path.join(V, "OLD_STYLE", "EN", f"{k}.pptx")) for k in old}
    # the same earlier decks built from the best-window re-run (RUN_bestwin_old_analyses_20261022.sh), when present
    bw = {d: deck(p("OLD_STYLE_BW", f"Day{d}_FULL_bestwindow")) for d in (1, 2)}
    bwEN = {d: deck(os.path.join(V, "OLD_STYLE_BW", "EN", f"Day{d}_FULL_bestwindow.pptx")) for d in (1, 2)}
    print("best-window re-run decks:", {d: v is not None for d, v in bw.items()})
    missing = [k for k, v in {**src, **old}.items() if v is None]
    print("missing inputs:", missing or "none")
    log = []
    for kind in a.kinds:
        dk = ND.Deck(lg, "")
        import winwords as WW                                  # 10-24: the windows the analysis used (per event in WIN_MODE=mix)
        rule = (L("Marks and response window chosen per event (notes)", "사건별 표시·반응 창 선택 (노트)") if WW.mix4() else
                L("Response window per event: 0-0.5 or 0-1 s (notes)", "사건별 반응 창: 0-0.5 또는 0-1 s (노트)") if WW.mixed() else
                L("Best window for every neuronal number: " + WW.en() + "; every star shows its pair",
                  "모든 뉴런 수치 = 최적 창: " + WW.kr() + "; 별표마다 비교 쌍 표시"))
        RULE_NOTES = L(WW.prefix_en()[0].upper() + WW.prefix_en()[1:] + ": " + WW.en() + ". No correction for the choice of window"
                       + (" or marks (the original marks are 0.40 / 0.37 s late against the video)" if WW.mix4() else "")
                       + ". Every star shows its pair.",
                       WW.prefix_kr() + ": " + WW.kr() + ". 창" + (" · 표시" if WW.mix4() else "") + " 선택에 대한 보정 없음"
                       + (" (옛 표시는 영상 대비 0.40 / 0.37 s 늦음)" if WW.mix4() else "") + ". 별표마다 비교 쌍 표시.")
        if kind == "FULL":
            parts = [("1", L("Day 1 + Day 2: the story", "Day 1 + Day 2: 이야기"), L("question -> behaviour -> activity -> responses (every "
                                                                              "event) -> checks -> one-page summary",
                                                                              "질문 -> 행동 -> 활동 -> 반응 (모든 사건) -> 검증 -> 한 페이지 요약")),
                     ("2", L("every mouse", "마우스별"), L("Day 1 and Day 2, best window", "Day 1과 Day 2, 최적 창")),
                     ("3 / 4", L("Day 1 alone, Day 2 alone", "Day 1 단독, Day 2 단독"),
                      L("new style + earlier slides (window-independent; the others re-run with the best window)",
                        "새 스타일 + 이전 슬라이드 (창과 무관한 것; 나머지는 최적 창으로 다시 계산)")),
                     ("5 + A", L("earlier Day 1 + 2 slides; archive", "이전 Day 1+2 슬라이드; 기록"), L("archive = earlier slides with the OLD 2-s window, marked", "기록 = 옛 2 s 창의 이전 슬라이드, 표시함"))]
            divider(dk, lg, L("SBI-553 mini1p, BLA: everything so far (FULL)", "SBI-553 mini1p, BLA: 지금까지 전부 (FULL)"), rule, parts, notes=RULE_NOTES)
        else:
            divider(dk, lg, L("SBI-553 mini1p, BLA: the key results (PI)", "SBI-553 mini1p, BLA: 핵심 결과 (PI)"), rule,
                    [("Q1", L("Does SBI-553 change pain behaviour?", "SBI-553가 통증 행동을 바꾸나?"),
                      L("reflexes, affective behaviour, movement per delivery", "반사, 정서 행동, 자극당 움직임")),
                     ("Q2", L("Does it change BLA activity?", "BLA 활동을 바꾸나?"), L("firing of every neuron, whole sessions",
                                                                                  "모든 뉴런의 활동, 세션 전체")),
                     ("Q3", L("Does it change the responses to pain stimuli?", "통증 자극 반응을 바꾸나?"),
                      L("population first, then single neurons, every stimulus and behaviour", "집단 먼저, 그다음 단일 뉴런, 모든 자극과 행동")),
                     ("Q4", L("Stimulus or movement?", "자극인가 움직임인가?"),
                      L("what the responsive neurons follow", "반응 뉴런이 무엇을 따르나")),
                     ("Q5", L("Is there a pattern by sex?", "성별 패턴이 있나?"), L("2 females vs 2 males: candidates only",
                                                                              "암컷 2 vs 수컷 2: 후보만")),
                     ("?", L("How sure are we?", "얼마나 확실한가?"), L("4 checks, Day 1 / Day 2 alone, limits, then the one-page summary",
                                                                   "검증 4개, Day 1 / Day 2 따로, 한계, 그리고 한 페이지 요약"))],
                    six=True, notes=RULE_NOTES)
        import glob as _g
        import make_pipeline_figs as PF                         # HOW WE DID IT: pipeline, 2-min video, logger, behaviours
        vids = sorted(_g.glob(os.path.join(PF.VID, "*behaviour_ca_*_small.mp4")), key=os.path.getmtime)
        vid = vids[-1] if vids else None
        post = vid.replace("_small.mp4", "_frame.png") if vid else None
        vid2 = os.path.join(PF.VID, "behaviour_examples_all_small.mp4")
        post2 = os.path.join(PF.VID, "behaviour_examples_all_frame.png")
        ENT = ["COVER"]
        ENT += PF.add_pipeline_slides(dk, lg, vid, post if post and os.path.isfile(post) else None,
                                      vid2 if os.path.isfile(vid2) else None, post2 if os.path.isfile(post2) else None)
        prs = dk.prs
        n_cover = len(prs.slides)
        cnt = {}

        NT = {"added": 0}
        SEEN = set()

        def add(name, D, filt=None, tag=None, label=None, ref=None, story=False, take=False, split=True, keyfix=None, backup=None):
            """ref = the EN version of D: the slide decisions are made on its English titles (same slides, same order);
            story = these slides are ordered by STORY; take = the take-home messages of TAKE; split = small figures -> parts"""
            if D is None:
                return 0
            R_ = ref if ref is not None else D
            dec = [filt(*title_of(s)) if filt else True for s in R_.slides]
            k = 0
            for i, s in enumerate(D.slides):
                if i >= len(dec) or not dec[i]:
                    continue
                if kind == "PI":                                 # 10-23: the PI deck never shows the same slide twice
                    et0 = title_of(R_.slides[i])[0]
                    if et0 in SEEN:
                        continue
                    SEEN.add(et0)
                news = clone_split(prs, s, tag, lg) if split else [clone(prs, s, tag, lg)]   # small figure: one slide per part
                NT["split"] = NT.get("split", 0) + (len(news) - 1)
                et, es = title_of(R_.slides[i])
                is_bk = (_isbk(_notes(s)) or _isbk(_notes(R_.slides[i]))
                         or (backup(et, es) if callable(backup) else bool(backup)))
                if is_bk and (re.search(r": where the responsive neurons are$", et) or re.match(PI_SHOW, et)):
                    is_bk = False                                # 10-24 (user): these kinds are asked for in the PI deck, shown
                for ns_ in news:
                    fit_big_numbers(ns_)
                    fix_texts(ns_, lg)
                    if lg == "kr" and ref is not None:
                        kr_notes(ns_, R_.slides[i])
                    if lg == "kr":
                        kr_notes_fix(ns_)
                    ENT.append(((keyfix(f"{es}|{et}") if keyfix else f"{es}|{et}")) if story else "OTHER")
                    if tag is None and trim_bottom(ns_, lg):           # first shorten the bottom, then the per-slide overrides
                        NT["trim"] = NT.get("trim", 0) + 1
                    if take and apply_take(ns_, et, lg):
                        NT["take"] = NT.get("take", 0) + 1
                    resize_key(ns_, lg)
                    fit_title(ns_)
                    notes_fix(ns_)
                    if tag is None and ensure_n_test(ns_, lg):    # only the new-style slides (no tag)
                        NT["added"] += 1
                    elif tag is not None and old_nt(ns_, title_of(R_.slides[i])[0], lg):   # earlier slides: n / test by title
                        NT["old"] = NT.get("old", 0) + 1
                    nt_ = _notes(ns_)
                    if _isbk(nt_):                               # 10-24: the moved figure key can come before the tag
                        ns_.notes_slide.notes_text_frame.text = re.sub(r"(^|\n)\[BACKUP\]\s*", lambda m_: m_.group(1) + (
                            "Backup slide. " if lg != "kr" else "보조 슬라이드. "), nt_, count=1)
                    if is_bk and kind == "PI":
                        mark_backup(ns_, lg)
                        NT["backup"] = NT.get("backup", 0) + 1
                    k += 1
            cnt[label or name] = cnt.get(label or name, 0) + k
            return k

        if kind == "FULL":
            add("Day1+2_FULL", src["Day1+2_FULL"], label="PART 1 Day 1 + Day 2 (new)", story=True, take=True, ref=srcEN["Day1+2_FULL"])
            glance2(dk, lg); ENT.append("GLANCE2")
            for d in (1, 2):
                daycard(dk, lg, d); ENT.append(f"DAYCARD{d}")
            one_page_summary(dk, lg); ENT.append("SUMMARY1")
            limits2(dk, lg); ENT.append("LIMITS2")
            add("PerMouse", src["PerMouse_Day1+2"], label="PART 2 every mouse (best window)", ref=srcEN["PerMouse_Day1+2"])
            for d in (1, 2):
                add(f"Day{d}_FULL", src[f"Day{d}_FULL"], label=f"PART {d + 2} Day {d} (new)", ref=srcEN[f"Day{d}_FULL"])
                keep = lambda t, sec: (any(re.search(r, t) for r in KEEP_OLD) and not any(re.search(r, t) for r in DEP_OLD) and not is_mouse_slide(sec)
                                       and not any(re.search(r, t) or re.search(r, sec) for r in DROP_OLD))
                add(f"old Day{d}_FULL", old[f"Day{d}_FULL"], keep, L(f"earlier Day-{d} deck (window-independent; marks before the 10-23 fix)",
                                                                       f"이전 Day {d} 덱 (창과 무관; 10-23 보정 전 표시)"), label=f"PART {d + 2} Day {d} (earlier, kept)",
                    ref=oldEN[f"Day{d}_FULL"])
                if bw[d] is not None:                            # the window-dependent earlier analyses, re-run with the best window
                    archf = lambda t, sec: ((not any(re.search(r, t) for r in KEEP_OLD) or any(re.search(r, t) for r in DEP_OLD))
                                            and not is_mouse_slide(sec) and not any(re.search(r, t) or re.search(r, sec) for r in DROP_OLD)
                                            and not re.match(r"^TOY ", t))
                    add(f"bw Day{d}", bw[d], archf, L(f"earlier Day-{d} analysis, best window of 10-22, marks BEFORE the 10-23 fix",
                                                      f"이전 Day {d} 분석, 10-22 최적 창, 10-23 보정 전 표시"),
                        label=f"PART {d + 2} Day {d} (earlier analyses, best window)", ref=bwEN[d])
            keepL = lambda t, sec: any(re.search(r, t) for r in KEEP_OLD) and not any(re.search(r, t) for r in DEP_OLD)
            add("old long", old["Day1+2_longitudinal"], keepL, L("earlier Day 1 + 2 deck (window-independent)", "이전 Day 1+2 덱 (창과 무관)"),
                label="PART 5 earlier Day 1 + 2 (kept)", ref=oldEN["Day1+2_longitudinal"])
            arch = lambda t, sec: ((not any(re.search(r, t) for r in KEEP_OLD) or any(re.search(r, t) for r in DEP_OLD)) and not is_mouse_slide(sec)
                                   and not any(re.search(r, t) or re.search(r, sec) for r in DROP_OLD))
            tagA = L("OLD 2-s WINDOW, marks before the 10-23 fix - superseded", "옛 2 s 창, 10-23 보정 전 표시 - 대체됨")
            archive_index(dk, lg, redone=all(v is not None for v in bw.values())); ENT += ["OTHER", "OTHER"]
            cnt["ARCHIVE index"] = 2
            for d in (1, 2):
                add(f"arch Day{d}", old[f"Day{d}_FULL"], arch, tagA, label=f"ARCHIVE Day {d} (old window)", ref=oldEN[f"Day{d}_FULL"],
                    split=False)
            add("arch long", old["Day1+2_longitudinal"], lambda t, sec: not keepL(t, sec), tagA, label="ARCHIVE Day 1 + 2 (old window)",
                ref=oldEN["Day1+2_longitudinal"], split=False)
        else:
            # the earlier glance card is replaced by glance2 (it stays in the FULL deck's supplement)
            add("Day1+2_PI", src["Day1+2_PI"], lambda t, sec: not re.match(r"^Day 1 \+ Day 2: SBI-553 vs vehicle in the same mice", t),
                label="Day 1 + Day 2 PI", story=True, take=True, ref=srcEN["Day1+2_PI"])
            # 10-23 round 2 (equal treatment, checks 1-2-3-4): three slides of the FULL deck
            add("PI extra", src["Day1+2_FULL"],
                lambda t, sec: bool(re.match(r"^(Check 3:|Population response to each behaviour|Pin prick: every neuron in the three conditions)", t)),
                label="from FULL: behaviour population, pin-prick neurons, check 3", story=True, take=True, ref=srcEN["Day1+2_FULL"],
                keyfix=lambda k_: "PIX|" + k_)
            # 10-23 evening (user): time spent in each behaviour, UP / DOWN traces, stability across days, spatial clustering, and
            # the one-event slides of the OTHER events as backup (grey, hidden), so every event is in the deck if the PI asks
            PIX2 = (r"^(Time spent in each behaviour after a delivery|UP, DOWN and the other neurons|Stable across days, shared between events"
                    r"|Are the responsive neurons clustered|Summary: responsive neurons per event)")
            add("PI extra 2", src["Day1+2_FULL"], lambda t, sec: bool(re.match(PIX2, t)),
                label="from FULL: time in behaviour, UP / DOWN traces, stability, clustering", story=True, take=True,
                ref=srcEN["Day1+2_FULL"], keyfix=lambda k_: "PIX|" + k_)
            EVB = r"^(Light touch|Mild touch|Heat|withdrawal|flinch|attending onset|movement onset|grooming onset|escape/rearing onset)"
            add("PI backup events", src["Day1+2_FULL"],
                lambda t, sec: bool(re.match(EVB + r": (every neuron in the three conditions|UP, DOWN)", t) or
                                    re.match(r"^Where the (?!Pin prick)", t)),
                label="from FULL: the other events (backup, hidden)", story=True, take=True, ref=srcEN["Day1+2_FULL"],
                keyfix=lambda k_: "PIX|" + k_, backup=True)
            if old.get("Day1_FULL") is not None:                  # methods of the earlier deck (window-independent)
                add("PI methods", old["Day1_FULL"],
                    lambda t, sec: bool(re.match(r"^(Design: one field of view|From the raw movie to neurons|Neurons per mouse)", t)),
                    label="methods of the earlier deck", story=True, ref=oldEN["Day1_FULL"], keyfix=lambda k_: "PIM|" + k_)
            glance2(dk, lg); ENT.append("GLANCE2")
            for d in (1, 2):                                     # the Day-alone cards: the same four cards on both days
                daycard(dk, lg, d); ENT.append(f"DAYCARD{d}")
            one_page_summary(dk, lg); ENT.append("SUMMARY1")
            limits2(dk, lg); ENT.append("LIMITS2")
        groups = reorder(prs, ENT, lg)                       # the story order (see STORY)
        NT["mix"] = sum(mix_fix(sl) for sl in prs.slides)        # 10-24: one-window wording -> window per event
        NT["face"] = sum(face_fix(sl, lg) for sl in prs.slides)   # 10-24 audit: n per mouse, named tests, titles, limits
        NT["over"] = apply_overrides(prs, lg)                       # 10-24 audit: per-variant key / bottom / notes, checked
        ban = os.environ.get("DECK_BANNER", "")              # 10-24: e.g. the old-marks variant, on every slide
        if ban:
            ban = ban.split("|")[1] if (lg == "kr" and "|" in ban) else ban.split("|")[0]
            for sl in prs.slides:
                tb = sl.shapes.add_textbox(Inches(5.6), Inches(0.02), Inches(7.5), Inches(0.26))   # 10-25: top right
                pp = tb.text_frame.paragraphs[0]
                pp.alignment = 3                                                              # right
                pp.text = ban; pp.font.size = Pt(10); pp.font.bold = True; pp.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
                pp.font.name = "Malgun Gothic" if lg == "kr" else "Arial"
        print("  story:", " > ".join(dict.fromkeys(g for g in groups if g != "OTHER")))
        os.makedirs(os.path.join(V, os.environ.get("OUT_DIR", "ALL_IN_ONE")), exist_ok=True)
        out = os.path.join(V, os.environ.get("OUT_DIR", "ALL_IN_ONE"), f"ALL_{kind}_{'KR' if lg == 'kr' else 'EN'}.pptx")
        prs.save(out)
        log.append(f"{out}: {len(prs.slides)} slides (cover {n_cover}) " + ", ".join(f"{k} {v}" for k, v in cnt.items())
                   + f"; n / test line added from the notes on {NT['added']} slides, by title on {NT.get('old', 0)} earlier slides; "
                   f"{NT.get('split', 0)} extra slides from split figures; take-home set on {NT.get('take', 0)}, bottom shortened on "
                   f"{NT.get('trim', 0)}")
        print(log[-1], flush=True)
    return log


if __name__ == "__main__":
    main()
