"""make_longitudinal_deck.py - the Day 1 + Day 2 deck (crossover: every mouse is its own control), PI rules.

  SUMMARY      what the two days show together (4 cards)
  DESIGN       the crossover: which mouse got what on which day, neurons and deliveries per session
  TRACKING     the same neurons on both days: maps per mouse, how each pair was decided, the transform
  CALCIUM      tracked neurons: baseline (mean of the two days) vs vehicle vs SBI-553 (ca_longitudinal.py L1-L4);
               every neuron of every session, tracked or not (your rule: single-day neurons are SHOWN, statistics use
               the tracked ones only)
  BEHAVIOUR    per delivery, vehicle day vs SBI-553 day within each mouse; spontaneous movement and grooming
  SUPPLEMENT   limits, the annotation quality checks
Every slide: title, a 'how to read' line, the figure, the test / n note and ONE plain result sentence (numbers read
from the analysis files). Day-2 neuron selections are provisional until you confirm them (red mark on every slide).

  CA_DAY=2 python make_longitudinal_deck.py --base C:/Users/hsollim/Documents/HEAL_mini1p_SBI553 --out <pptx>
      [--long-dir analysis/ca_longitudinal]
"""
import argparse
import os
import sys

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import safe_io  # noqa: E402,F401  (a figure save retries on a transient Windows write error)
import make_day1_ca_slides as S1                                    # noqa: E402
import make_day1_full_deck as FD                                    # noqa: E402
import make_day1_complete_deck as CD                                # noqa: E402

FW, FH = S1.FW, S1.FH
ARMS = {1: {"5941": "SBI-553", "5939": "SBI-553", "5937": "vehicle", "1191": "vehicle"},
        2: {"5937": "SBI-553", "1191": "SBI-553", "5941": "vehicle", "5939": "vehicle"}}
MO = ["5941", "5939", "5937", "1191"]
SEX = {"5941": "F", "5939": "M", "5937": "M", "1191": "F"}
CC = {"baseline": "#7f8c8d", "vehicle": "#2471a3", "SBI-553": "#c0392b"}
DATE = {1: "20261019", 2: "20261021"}


def pstr(p):
    return S1.pstr(p)


def kept(B, day, m):
    R = pd.read_csv(os.path.join(B, "mini1p_out", DATE[day], f"{m}_d{day}", "4_cell_review", f"{m}_d{day}_cell_review.csv"))
    return R


def fig_design(B, out, lang="en"):
    S = pd.read_csv(os.path.join(B, "analysis", "behavior_crossover", "crossover_per_session.csv")); S["mouse"] = S.mouse.astype(str)
    rows = []
    for m in MO:
        n1 = int((kept(B, 1, m).keep == 1).sum())
        try:
            n2 = int((kept(B, 2, m).keep == 1).sum())
        except FileNotFoundError:
            n2 = "-"
        pf = os.path.join(B, "analysis", "crossday", f"{m}_pairs.csv")
        nt = len(pd.read_csv(pf)) if os.path.isfile(pf) else "-"
        dl = []
        for day in (1, 2):
            for ph in ("baseline", "injection"):
                x = S[(S.mouse == m) & (S.day == day) & (S.phase == ph)]
                dl.append(str(int(x.deliveries.iloc[0])) if len(x) else "-")
        rows.append([m, SEX[m], ARMS[1][m], ARMS[2][m], str(n1), str(n2), str(nt), " / ".join(dl[:2]), " / ".join(dl[2:])])
    kr = lang == "kr"                                          # 10-23 review: the KR deck showed the English figure
    rc_ = {"font.family": "Malgun Gothic"} if kr else {}
    plt.rcParams.update(rc_)
    fig = plt.figure(figsize=(FW, FH))
    ax = fig.add_axes([0.03, 0.70, 0.94, 0.2]); ax.set_axis_off()
    x0 = 0.0
    for day, lab in ((1, "Day 1 (2026-10-19)"), (2, "Day 2 (2026-10-21)")):
        for k, (ph, col) in enumerate(((("기준 세션" if kr else "baseline session"), CC["baseline"]),
                                       (("주사 세션" if kr else "injection session"), "#d68910"))):
            ax.add_patch(plt.Rectangle((x0 + k * 0.22, 0.2), 0.2, 0.5, color=col, alpha=0.85))
            ax.text(x0 + k * 0.22 + 0.1, 0.45, ph, ha="center", va="center", fontsize=10, color="white", fontweight="bold")
        ax.text(x0 + 0.21, 0.85, lab, ha="center", fontsize=11, fontweight="bold")
        x0 += 0.52
    ax.set_xlim(0, 0.98); ax.set_ylim(0, 1)
    ax.text(0.49, 0.02, ("같은 시야, 같은 뉴런 추적; Day 2에는 주사 약물을 바꿈" if kr else
                         "same field of view, same neurons tracked; the treatment of the injection session is swapped on Day 2"),
            ha="center", fontsize=9.5)
    tb = fig.add_axes([0.03, 0.02, 0.94, 0.62]); tb.set_axis_off()
    cl_ = (["마우스", "성별", "Day 1 주사", "Day 2 주사", "뉴런 Day 1", "뉴런 Day 2", "추적", "자극 수 Day 1 (기준 / 주사)",
            "자극 수 Day 2 (기준 / 주사)"] if kr else
           ["mouse", "sex", "Day 1 injection", "Day 2 injection", "neurons Day 1", "neurons Day 2", "tracked",
            "deliveries Day 1 (base / inj.)", "deliveries Day 2 (base / inj.)"])
    T = tb.table(cellText=rows, colLabels=cl_,
                 loc="upper left", cellLoc="left", colLoc="left", colWidths=[0.07, 0.05, 0.11, 0.11, 0.1, 0.1, 0.08, 0.19, 0.19])
    T.auto_set_font_size(False); T.set_fontsize(10); T.scale(1, 2.1)
    for (r, c), cell in T.get_celld().items():
        cell.set_edgecolor("0.85")
        if r == 0:
            cell.set_text_props(fontweight="bold"); cell.set_facecolor("#f3f5f8")
        elif c in (2, 3):
            cell.get_text().set_color(CC[rows[r - 1][c]])
    fig.savefig(out, dpi=160); plt.close(fig)
    if kr:
        plt.rcParams.update({"font.family": plt.rcParamsDefault["font.family"]})
    return rows, ("Each mouse: a baseline session, then the injection session, on each of two days 2 days apart; SBI-553 on one "
                  "day and vehicle on the other (crossover, 2 mice each order). Deliveries differ between sessions, so behaviour "
                  "is always divided by deliveries.")


def tracking_summary(B):
    import json
    rows = []
    for m in MO:
        f = os.path.join(B, "analysis", "crossday", f"{m}_tracking.json")
        if os.path.isfile(f):
            J = json.load(open(f))
            # your three categories from the FINAL (confirmed) files: on both days = the pairs, one day only = the rest
            pf = os.path.join(B, "analysis", "crossday", f"{m}_pairs.csv")
            nb = len(pd.read_csv(pf)) if os.path.isfile(pf) else 0
            n1 = int((pd.read_csv(os.path.join(B, "mini1p_out", "20261019", f"{m}_d1", "4_cell_review", f"{m}_d1_cell_review.csv")).keep == 1).sum())
            n2 = int((pd.read_csv(os.path.join(B, "mini1p_out", "20261021", f"{m}_d2", "4_cell_review", f"{m}_d2_cell_review.csv")).keep == 1).sum())
            rows.append(dict(mouse=m, **{k: J[k] for k in ("day1_confirmed", "A_same_green", "A_same_day1_orange", "B_check",
                                                             "C_day2_deleted", "day1_not_found", "new_kept", "day2_kept")},
                             n_both=nb, n_d1=n1, n_d2=n2, n_d1_only=n1 - nb, n_d2_only=n2 - nb,
                             rotation=J["transform"]["rotation_deg"], scale=J["transform"]["scale"],
                             r=J["transform"]["image_reg"]["r"], patch_null=J["patch_null95"]))
    return pd.DataFrame(rows)


def fig_tracking_counts(TS, out):
    """your three categories per mouse (2026-10-22): green = the neuron exists on both days, magenta = Day 1 only,
    cyan = Day 2 only; B = how much the field of view turned between the days"""
    fig, ax = plt.subplots(1, 2, figsize=(FW, FH), gridspec_kw=dict(width_ratios=[1.6, 1]))
    fig.subplots_adjust(left=0.07, right=0.99, top=0.9, bottom=0.3, wspace=0.3)
    parts = [("n_both", "on both days", "#2ecc40"), ("n_d1_only", "Day 1 only", "#f012be"), ("n_d2_only", "Day 2 only", "#00d0ff")]
    y = np.arange(len(TS))[::-1]
    for yy, (_, r) in zip(y, TS.iterrows()):
        left = 0
        for k, lab, col in parts:
            ax[0].barh(yy, r[k], left=left, color=col, height=0.6, label=lab if yy == y[0] else None)
            if r[k] > 2:
                ax[0].text(left + r[k] / 2, yy, str(int(r[k])), ha="center", va="center", fontsize=10, fontweight="bold")
            left += r[k]
        ax[0].text(left + 1, yy, f"Day 1: {int(r.n_d1)}, Day 2: {int(r.n_d2)}", va="center", fontsize=9.5)
    ax[0].set_yticks(y, [f"{m}" for m in TS.mouse]); ax[0].set_xlabel("neurons")
    ax[0].set_xlim(0, (TS.n_both + TS.n_d1_only + TS.n_d2_only).max() * 1.45)
    ax[0].legend(fontsize=10, frameon=False, loc="upper center", bbox_to_anchor=(0.45, -0.2), ncol=3)
    ax[0].set_title("A  every neuron: on both days, Day 1 only, Day 2 only", loc="left")
    ax[1].bar(range(len(TS)), TS.rotation, color="0.5", width=0.5)
    for i, r in TS.iterrows():
        ax[1].text(i, r.rotation + (0.05 if r.rotation >= 0 else -0.05), f"{r.rotation:+.2f} deg\nscale {r.scale:.3f}",
                   ha="center", va="bottom" if r.rotation >= 0 else "top", fontsize=8.5)
    ax[1].axhline(0, color="0.4", lw=0.8)
    ax[1].set_xticks(range(len(TS)), TS.mouse); ax[1].set_ylabel("rotation Day 1 -> Day 2 (deg)")
    lim = max(1.0, TS.rotation.abs().max() * 1.6); ax[1].set_ylim(-lim, lim)
    ax[1].set_title("B  field-of-view transform", loc="left")
    for a in ax:
        for sp in ("top", "right"):
            a.spines[sp].set_visible(False)
    fig.savefig(out, dpi=160); plt.close(fig)
    return ("Each Day-1 neuron was followed to Day 2 and judged on Day-2 evidence (local Cn / std pattern, the Day-2 movie proof, the "
            "Day-2 cell rules); the Day-1 neurons not found by the Day-2 EXTRACT were recovered with their Day-1 footprint. "
            "Your confirmation 2026-10-22. Green = on both days, magenta = Day 1 only, cyan = Day 2 only.")


def fig_within_day(LT, out):
    """A: within-day change of the transient rate (injection - that day's baseline), vehicle day vs SBI-553 day, per mouse;
    B: per event, SBI-553-day change minus vehicle-day change of the event response (z), per mouse"""
    fig, ax = plt.subplots(1, 2, figsize=(FW, FH), gridspec_kw=dict(width_ratios=[1, 2.2]))
    fig.subplots_adjust(left=0.07, right=0.98, top=0.86, bottom=0.2, wspace=0.22)
    MC = {"5941": "#e67e22", "5939": "#8e44ad", "5937": "#16a085", "1191": "#2c3e50"}
    R = LT[LT.measure == "transient rate"].set_index("mouse")
    for m in MO:
        if m not in R.index:
            continue
        r = R.loc[m]
        ax[0].plot([0, 1], [r.med_d_vehicle, r.med_d_sbi], "-o", color=MC[m], lw=1.5, ms=6)
        ax[0].text(1.06, r.med_d_sbi, f"{m} (SBI-553 on Day {1 if ARMS[1][m] == 'SBI-553' else 2})\nn={int(r.n)}, {pstr(r.p_dsbi_vs_dveh)}",
                   va="center", fontsize=8, color=MC[m])
    ax[0].axhline(0, color="0.4", lw=0.8)
    ax[0].set_xticks([0, 1], ["vehicle day", "SBI-553 day"]); ax[0].set_xlim(-0.2, 1.9)
    ax[0].set_ylabel("within-day change of the transient rate\n(injection - baseline, transients/min;\nmedian over tracked neurons)")
    ax[0].set_title("A  transient rate", loc="left")
    E = LT[LT.measure == "event response"].copy()
    evs = [e for e in ["Light touch", "Mild touch", "Pin prick", "Heat", "withdrawal", "flinch", "movement onset",
                       "escape/rearing onset"] if e in set(E.event)]
    for i, m in enumerate(MO):
        X = E[E.mouse == m].set_index("event")
        for k, ev in enumerate(evs):
            if ev not in X.index:
                continue
            r = X.loc[ev]
            d = r.med_d_sbi - r.med_d_vehicle
            x = k + (i - 1.5) * 0.16
            ax[1].plot(x, d, "o", color=MC[m], ms=6, label=m if k == 0 else None)
            if r.p_dsbi_vs_dveh < 0.05:
                ax[1].text(x, d, "*", ha="center", va="bottom" if d >= 0 else "top", fontsize=12)
    ax[1].axhline(0, color="0.4", lw=0.8)
    ax[1].set_xticks(range(len(evs)), evs, fontsize=8.5, rotation=15, ha="right")
    lo, hi = ax[1].get_ylim(); ax[1].set_ylim(lo, hi + 0.35 * (hi - lo))
    ax[1].set_ylabel("SBI-553-day change - vehicle-day change\n(event response, z [0, 2] s - [-2, 0] s;\nmedian over tracked neurons)")
    ax[1].set_title("B  event responses (below 0 = smaller change on the SBI-553 day)", loc="left")
    h_, l_ = ax[1].get_legend_handles_labels()
    h_.append(Line2D([], [], color="k", ls="", marker="$*$", ms=9)); l_.append("* = that mouse: SBI-553-day change vs vehicle-day change")
    ax[1].legend(h_, l_, fontsize=8.5, frameon=False, ncol=5, loc="upper left")
    for a in ax:
        for sp in ("top", "right"):
            a.spines[sp].set_visible(False)
    fig.savefig(out, dpi=160); plt.close(fig)
    return ("Within-day change = injection session minus the same day's baseline session, per tracked neuron. Test per mouse: Wilcoxon "
            "signed-rank over tracked neurons, SBI-553-day change vs vehicle-day change; * p < 0.05 (uncorrected). n = 4 mice: no "
            "test across mice is shown (minimum possible p = 0.125).")


def split_rows(png, out_top, out_bottom):
    """a 4-mouse figure cut between mouse rows 2 and 3, so each half can be shown large enough to read"""
    from PIL import Image
    im = Image.open(png)
    W, H = im.size
    cut = int(H * 0.535)
    im.crop((0, 0, W, cut)).save(out_top)
    im.crop((0, cut, W, H)).save(out_bottom)
    return out_top, out_bottom


def fig_all_neurons(B, out):
    """every neuron of every session (tracked or not): transient rate in the stimulus blocks, still frames"""
    fig, ax = plt.subplots(1, 4, figsize=(FW, FH), sharey=False)
    fig.subplots_adjust(left=0.06, right=0.99, top=0.82, bottom=0.17, wspace=0.28)
    stim = S1.STIM
    res = []
    for i, m in enumerate(MO):
        a = ax[i]
        pf = os.path.join(B, "analysis", "crossday", f"{m}_pairs.csv")
        P = pd.read_csv(pf) if os.path.isfile(pf) else pd.DataFrame(columns=["day1_id", "day2_id"])
        trk = {1: set(P.day1_id.astype(int)), 2: set(P.day2_id.astype(int))}
        x = 0
        lab = []
        for day in (1, 2):
            f = os.path.join(B, "analysis", f"ca_day{day}", f"ca_d{day}_cells_whole_session.csv")
            if not os.path.isfile(f):
                continue
            W = pd.read_csv(f); W["mouse"] = W.mouse.astype(str); W = W[W.mouse == m]
            num = sum(W[f"still rate | {s}"].fillna(0) * W[f"still min | {s}"].fillna(0) for s in stim)
            den = sum(W[f"still min | {s}"].fillna(0) for s in stim)
            W = W.assign(v=num / den.replace(0, np.nan))
            for ph in ("pre", "post"):
                X = W[W.phase == ph]
                cond = "baseline" if ph == "pre" else ARMS[day][m]
                t = X.cell.astype(int).isin(trk[day]).values
                rng = np.random.default_rng(day * 10 + (ph == "post"))
                jit = rng.uniform(-0.15, 0.15, len(X))
                a.scatter(x + jit[~t], X.v.values[~t], s=7, facecolor="none", edgecolor=CC[cond], lw=0.6, alpha=0.7)
                a.scatter(x + jit[t], X.v.values[t], s=7, color=CC[cond], alpha=0.85, lw=0)
                a.plot([x - 0.3, x + 0.3], [np.nanmedian(X.v)] * 2, color="k", lw=2)
                lab.append((x, f"D{day}\n{'base' if ph == 'pre' else cond}"))
                res.append(dict(mouse=m, day=day, phase=ph, condition=cond, n_all=len(X), n_tracked=int(t.sum()),
                                median_all=float(np.nanmedian(X.v)), median_tracked=float(np.nanmedian(X.v.values[t])) if t.any() else np.nan))
                x += 1
        a.set_xticks([q[0] for q in lab], [q[1] for q in lab], fontsize=8)
        a.set_title(f"{m}", loc="left", fontsize=11)
        if i == 0:
            a.set_ylabel("Ca transients per min\n(stimulus blocks, still frames)")
        a.set_ylim(bottom=0)
        for sp in ("top", "right"):
            a.spines[sp].set_visible(False)
    fig.legend(handles=[Line2D([], [], marker="o", ls="", color="0.3", label="neuron found on both days (tracked)"),
                        Line2D([], [], marker="o", ls="", mfc="none", mec="0.3", label="neuron of this day only"),
                        Patch(color=CC["baseline"], label="baseline"), Patch(color=CC["vehicle"], label="vehicle"),
                        Patch(color=CC["SBI-553"], label="SBI-553"), Line2D([], [], color="k", lw=2, label="median, all neurons")],
               loc="upper center", ncol=6, frameon=False, fontsize=9)
    fig.savefig(out, dpi=160); plt.close(fig)
    return pd.DataFrame(res), ("Every neuron of every session (open = found on one day only, filled = tracked on both days). "
                               "Statistics on the next slides use the tracked neurons only (your rule).")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--long-dir", default="analysis/ca_longitudinal")
    a = ap.parse_args(argv)
    from pptx import Presentation
    from pptx.util import Inches
    B = a.base
    L = os.path.join(B, a.long_dir)
    od = os.path.splitext(a.out)[0] + "_figures"
    os.makedirs(od, exist_ok=True)
    f = lambda nm: os.path.join(od, nm)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(S1.SW), Inches(S1.SH)
    pvf = lambda: f(f"preview_slide{len(prs.slides) + 1:02d}.png")
    toy = "TOY" in a.long_dir.upper()
    # ------------------------------------------------------------------ numbers
    BX = os.path.join(B, "analysis", "behavior_crossover")
    BT = pd.read_csv(os.path.join(BX, "crossover_tests.csv")); BT["mouse"] = BT.mouse.astype(str)
    BA = pd.read_csv(os.path.join(BX, "crossover_across_mice.csv"))
    TS = tracking_summary(B)
    LT = pd.read_csv(os.path.join(L, "long_tests.csv")); LT["mouse"] = LT.mouse.astype(str)
    LM = pd.read_csv(os.path.join(L, "long_cell_matching.csv")); LM["mouse"] = LM.mouse.astype(str)
    lr = LT[LT.measure == "transient rate"].set_index("mouse")
    n_sbi_lower_ca = int(sum(lr.at[m, "med_d_sbi"] < lr.at[m, "med_d_vehicle"] for m in MO if m in lr.index))
    ba = BA.set_index("behaviour")
    # ------------------------------------------------------------------ SUMMARY
    ntr = int(LM.n_matched.sum()); n1 = int(LM.n_day1.sum())
    cards = [(f"{ntr}", "neurons tracked from Day 1 to Day 2",
              "Of " + f"{n1} Day-1 neurons: " + ", ".join(f"{r.mouse} {r.n_matched}/{r.n_day1}" for r in LM.itertuples())
              + ". Day-2 selection provisional."),
             (f"{n_sbi_lower_ca}/4", "mice: within-day change of the tracked neurons lower on the SBI-553 day",
              "Median within-day change of the transient rate, SBI-553 vs vehicle day: " + ", ".join(
                  f"{m} {lr.at[m, 'med_d_sbi']:+.1f} vs {lr.at[m, 'med_d_vehicle']:+.1f}" for m in MO if m in lr.index) + " /min."),
             (f"{int(ba.at['withdrawal', 'n_sbi_lower'])}/4", "mice: withdrawal per delivery rose less after SBI-553 than after vehicle",
              f"Flinch {int(ba.at['flinch', 'n_sbi_lower'])}/4, movement after a delivery {int(ba.at['general movement', 'n_sbi_lower'])}/4; "
              f"Wilcoxon over 4 mice p = {ba.at['withdrawal', 'p_wilcoxon_mice']:.3f} (the minimum possible)."),
             (f"{int(TS.n_d1_only.sum() + TS.n_d2_only.sum()) if len(TS) else 0}", "neurons on one day only",
              ("Day 1 only " + ", ".join(f"{r.mouse} {int(r.n_d1_only)}" for r in TS.itertuples()) + "; Day 2 only "
               + ", ".join(f"{r.mouse} {int(r.n_d2_only)}" for r in TS.itertuples()) + "; shown, not in the paired statistics.")
              if len(TS) else "tracking not run yet")]
    FD.summary_slide(prs, cards, pvf(), "Day 1 + Day 2 (crossover, 4 mice): SBI-553 vs vehicle within each mouse",
                     "Each mouse was imaged on two days with SBI-553 on one and vehicle on the other; the same neurons were "
                     "followed and every behaviour was divided by the deliveries of its session.",
                     howto="How to read: four results of Day 1 + Day 2 together, one per card; the big number is the key value, the "
                           "text says what it is (the slides that follow show each one).")
    # ------------------------------------------------------------------ DESIGN
    rows, nd = fig_design(B, f("s02_design.png"))
    FD.add_slide(prs, "DESIGN", "The crossover: SBI-553 on one day, vehicle on the other", f("s02_design.png"), nd,
                 "Each mouse had a baseline and an injection session on both days. Day 1: " + ", ".join(
                     f"{m} {ARMS[1][m]}" for m in MO) + "; Day 2 the other treatment.", pvf(),
                 howto="How to read: top = one day = baseline session then injection session; table = each mouse's treatment on each "
                       "day, its neurons and how many stimuli it received.")
    # ------------------------------------------------------------------ TRACKING
    maps = [os.path.join(B, "analysis", "crossday", f"{m}_tracking_map.png") for m in MO]
    if all(os.path.isfile(p) for p in maps):
        FD.add_slide(prs, "TRACKING", "The same neurons on Day 1 and Day 2 (one mouse per row pair)",
                     CD.grid(maps, 2, f("s03_tracking_maps.png")), "",
                     "Numbers = Day-1 neuron ids, written at the same neuron on both days. " + ", ".join(
                         f"{r.mouse}: {int(r.n_both)} on both days, {int(r.n_d1_only)} Day 1 only, {int(r.n_d2_only)} Day 2 only"
                         for r in TS.itertuples()) + ".", pvf(),
                     howto="How to read: per mouse, Day 1 (left) and Day 2 (right) max dF/F; green = the neuron is there on both days "
                           "(same number on both sides), magenta = Day 1 only, cyan = Day 2 only.")
    if len(TS):
        nt = fig_tracking_counts(TS, f("s04_tracking_counts.png"))
        FD.add_slide(prs, "TRACKING", "How each Day-1 neuron was decided on Day 2", f("s04_tracking_counts.png"), nt,
                     "On both days: " + ", ".join(f"{r.mouse} {int(r.n_both)}" for r in TS.itertuples())
                     + "; rotation between the days " + ", ".join(
                         f"{r.mouse} {r.rotation:+.1f} deg" for r in TS.itertuples()) + ".", pvf(),
                     howto="How to read: A = one bar per mouse, split by what happened to its Day-1 neurons on Day 2; B = how much "
                           "the field of view turned between the days.")
    # ------------------------------------------------------------------ CALCIUM
    R_all, na = fig_all_neurons(B, f("s05_all_neurons.png"))
    R_all.round(4).to_csv(f("s05_all_neurons.csv"), index=False)
    FD.add_slide(prs, "CALCIUM", "Every neuron of every session, tracked or not", f("s05_all_neurons.png"), na,
                 "Median transients per min, all neurons (D1 base / D1 inj / D2 base / D2 inj): " + "; ".join(
                     f"{m} " + " / ".join(f"{v:.1f}" for v in R_all[R_all.mouse == m].median_all) for m in MO) + ".", pvf(),
                 howto="How to read: one panel per mouse, the 4 sessions left to right; dots = neurons (filled = also found on the "
                       "other day), colour = condition, black bar = median.")
    for png, title, how, res in (
            ("long_L1_rate_three_conditions.png", "Tracked neurons: baseline vs vehicle vs SBI-553, transient rate",
             "How to read: one grey line per tracked neuron across the three conditions (baseline = mean of the two baseline "
             "sessions); coloured bar = median.",
             "Median transients per min, baseline / vehicle / SBI-553: " + ", ".join(
                 f"{m} {lr.at[m, 'med_baseline']:.1f} / {lr.at[m, 'med_vehicle']:.1f} / {lr.at[m, 'med_sbi']:.1f}"
                 f" ({pstr(lr.at[m, 'p_sbi_vs_veh'])})" for m in MO if m in lr.index) + "."),
            ):
        p = os.path.join(L, png)
        if os.path.isfile(p):
            FD.add_slide(prs, "CALCIUM", title, p, "Tests per mouse over tracked neurons: Friedman (3 conditions) and Wilcoxon "
                         "signed-rank (SBI-553 vs vehicle; within-day changes). Across mice n = 4 (Wilcoxon minimum p = 0.125).",
                         res, pvf(), howto=how)
    nw = fig_within_day(LT, f("s06_within_day_change.png"))
    floor = [m for m in MO if len(R_all[(R_all.mouse == m) & (R_all.day == 2)]) and (R_all[(R_all.mouse == m) & (R_all.day == 2)].median_all == 0).all()]
    ev_sig = LT[(LT.measure == "event response") & (LT.p_dsbi_vs_dveh < 0.05)]
    FD.add_slide(prs, "CALCIUM", "Four mice, each its own control: the within-day change", f("s06_within_day_change.png"), nw,
                 f"The SBI-553-day change of the transient rate was below the vehicle-day change in {n_sbi_lower_ca}/4 mice"
                 + (" (" + ", ".join(f"{m}: Day-2 median 0 transients/min, floor" for m in floor) + ")" if floor else "") + ". Event "
                 "responses with p < 0.05: " + (", ".join(f"{r.mouse} {r.event.lower()}" for r in ev_sig.itertuples()) or "none") + ".",
                 pvf(), howto="How to read: A = one line per mouse from its vehicle-day change to its SBI-553-day change; B = per event "
                              "and mouse, how much the SBI-553 day differed from the vehicle day (0 = the same).")
    # ------------------------------------------------------------------ BEHAVIOUR
    w = BT[BT.behaviour == "withdrawal"].set_index("mouse"); fl = BT[BT.behaviour == "flinch"].set_index("mouse")
    FD.add_slide(prs, "BEHAVIOUR", "Behaviour per delivery: vehicle day vs SBI-553 day, every mouse",
                 os.path.join(BX, "crossover_per_delivery.png"),
                 "Events per delivery (reflexes within 2 s, other behaviours: bout starts within 3 s). * = the within-day change "
                 "differs between the SBI-553 and the vehicle day (Poisson GLM, events ~ session x drug + log deliveries).",
                 "Withdrawal per delivery, vehicle day base -> inj. vs SBI-553 day base -> inj.: " + ", ".join(
                     f"{m} {w.at[m, 'veh_base']:.2f}->{w.at[m, 'veh_inj']:.2f} vs {w.at[m, 'sbi_base']:.2f}->{w.at[m, 'sbi_inj']:.2f}"
                     for m in MO) + ".", pvf(),
                 howto="How to read: per mouse two short lines: blue = vehicle day, red = SBI-553 day, each from its baseline (open) "
                       "to the injection session (filled).")
    FD.add_slide(prs, "BEHAVIOUR", "Within each mouse: change after SBI-553 vs after vehicle",
                 os.path.join(BX, "crossover_change_paired.png"),
                 "Change = events per delivery, injection minus that day's baseline. n = 4 mice; Wilcoxon signed-rank over mice "
                 "(minimum possible p = 0.125 with 4 mice).",
                 "SBI-553 change below vehicle change: " + ", ".join(
                     f"{b} {int(ba.at[b, 'n_sbi_lower'])}/4" for b in ("withdrawal", "flinch", "attending", "escape/rearing",
                                                                       "general movement") if b in ba.index) + ".", pvf(),
                 howto="How to read: one line per mouse from its vehicle-day change (left) to its SBI-553-day change (right); below 0 = "
                       "less after the injection than at that day's baseline.")
    S = pd.read_csv(os.path.join(BX, "crossover_per_session.csv")); S["mouse"] = S.mouse.astype(str)
    mv = lambda m, drug, ph: float(S[(S.mouse == m) & (S.drug_day == drug) & (S.phase == ph)].movement_pct_nostim.iloc[0])
    FD.add_slide(prs, "BEHAVIOUR", "Spontaneous movement and grooming, both days", os.path.join(BX, "crossover_movement_grooming.png"),
                 "Movement = % of frames with automatic general movement in the first 3 min (no stimulus); grooming = automatic "
                 "grooming, seconds per session (Day-2 maybe-clips not yet judged by you).",
                 "Time moving before any stimulus, base -> inj., vehicle day vs SBI-553 day: " + ", ".join(
                     f"{m} {mv(m, 'vehicle', 'baseline'):.0f}->{mv(m, 'vehicle', 'injection'):.0f} vs "
                     f"{mv(m, 'SBI-553', 'baseline'):.0f}->{mv(m, 'SBI-553', 'injection'):.0f}%" for m in MO) + ".", pvf(),
                 howto="How to read: per mouse, blue = vehicle day, red = SBI-553 day, open = baseline, filled = injection session.")
    # ------------------------------------------------------------------ SUPPLEMENT
    p2 = os.path.join(L, "long_L2_event_responses.png")
    if os.path.isfile(p2):
        halves = split_rows(p2, f("s20_event_responses_5941_5939.png"), f("s21_event_responses_5937_1191.png"))
        for hp, ms in zip(halves, (("5941", "5939"), ("5937", "1191"))):
            sig = LT[(LT.measure == "event response") & (LT.p_sbi_vs_veh < 0.05) & LT.mouse.isin(ms)]
            FD.add_slide(prs, "SUPPLEMENT", f"Tracked neurons: event responses in three conditions ({ms[0]}, {ms[1]})", hp,
                         "Response = z [0, 2] s minus [-2, 0] s around the event, per tracked neuron. Per panel: Friedman over the 3 "
                         "conditions (F p) and Wilcoxon signed-rank SBI-553 vs vehicle over tracked neurons; uncorrected.",
                         "SBI-553 vs vehicle differed (p < 0.05) for: " + (", ".join(f"{r.mouse} {r.event.lower()}" for r in sig.itertuples())
                                                                        or "none") + ".", pvf(),
                         howto="How to read: one row per mouse, one panel per event; grey line = one tracked neuron across baseline "
                               "(mean of the 2 days), vehicle and SBI-553; bar = median.")
    for png, title, how, res in (
            ("long_L3_population_psth.png", "Tracked neurons: mean response around each event",
             "How to read: mean over the tracked neurons of each mouse, around each event; grey = baseline, blue = vehicle, red = SBI-553.",
             "The tests are on the previous two slides; this view shows the time course."),
            ("long_L5_cell_matching.png", "Cell matching used by the longitudinal analysis",
             "How to read: Day-1 and Day-2 footprints of the tracked neurons per mouse.",
             "Tracked neurons used: " + ", ".join(f"{r.mouse} {r.n_matched}" for r in LM.itertuples()) + ".")):
        p = os.path.join(L, png)
        if os.path.isfile(p):
            FD.add_slide(prs, "SUPPLEMENT", title, p, ("Mean over the tracked neurons of each mouse; n = tracked neurons ("
                         + ", ".join(f"{r.mouse} {r.n_matched}" for r in LM.itertuples()) + ").") if "L3" in png else "",
                         res, pvf(), howto=how)
    q = os.path.join(B, "analysis", "behavior_quality", "behavior_quality.png")
    if os.path.isfile(q):
        Q = pd.read_csv(os.path.join(B, "analysis", "behavior_quality", "behavior_quality.csv"))
        FD.add_slide(prs, "SUPPLEMENT", "Annotation quality checks, all 16 sessions", q,
                     "Orphan = a reflex tap more than 5 s after the last delivery inside a stimulus block (a missed delivery or a "
                     "spontaneous reflex). Escape frames without movement were rearing in place (checked on the video frames).",
                     f"Orphan reflex taps: {int(Q.reflex_orphan_in_blocks.sum())} of {int(Q.reflex_taps.sum())} "
                     f"({100 * Q.reflex_orphan_in_blocks.sum() / Q.reflex_taps.sum():.1f}%); automatic grooming overlapping your "
                     f"lick / bite, guarding or attending: {Q[['grooming_on_lickbite_s', 'grooming_on_guarding_s', 'grooming_on_attending_s']].sum().sum():.1f} s.",
                     pvf(), howto="How to read: one row per session (Day, mouse, phase); A = reflex taps by time since the last "
                                  "delivery; B-D = automatic movement and grooming against your own scoring.")
    prs.save(a.out)
    print("->", a.out, "|", len(prs.slides), "slides | figures + previews in", od)


if __name__ == "__main__":
    main()
