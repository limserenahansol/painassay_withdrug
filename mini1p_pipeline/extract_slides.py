"""extract_slides.py - 10-26 (user rule: send only the changed / added slides, never the whole deck).

  python extract_slides.py <deck.pptx> <out.pptx> 1 3 5     keep only those slides (1-based), notes included
"""
import sys

from pptx import Presentation


def extract(src, out, keep):
    prs = Presentation(src)
    ids = prs.slides._sldIdLst
    for i, sid in reversed(list(enumerate(list(ids), 1))):
        if i not in keep:
            prs.part.drop_rel(sid.rId)
            ids.remove(sid)
    prs.save(out)
    return len(prs.slides)


if __name__ == "__main__":
    print("->", sys.argv[2], extract(sys.argv[1], sys.argv[2], {int(x) for x in sys.argv[3:]}), "slides")
