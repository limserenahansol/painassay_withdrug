"""make_auc_basenorm.py - 10-26 (user, on the population AUC slide: 'we might need normalized by baseline? as it might be baseline
change?'): population AUC normalized to each day's baseline session.

Per neuron (same ROI in the baseline and the injection session of one day):
    dAUC = AUC after the injection - AUC in the same day's baseline session      (z*s)
where AUC = area of (trial-mean z - its mean in the 0.5 s before the event) over the response window - the same AUC as on the
population AUC slide (make_pi5_deck.auc_cells). Tests: dAUC vs 0 = Wilcoxon (same neurons, = the blue / red brackets of the old
slide); vehicle day vs SBI-553 day = Mann-Whitney on dAUC (all neurons, different sets) and paired Wilcoxon on the 185 tracked
neurons; mice = paired Wilcoxon over the 4 mouse means (smallest possible p 0.125).

  python make_auc_basenorm.py   -> <PI5 figures>\\s3a_auc_basenorm.png + .json, and the 1-slide pptx + png in
                                   OneDrive\\Desktop\\1111112121212_FINAL_compact_added_slides
"""
import json
import os

import numpy as np
import pandas as pd
from scipy.stats import mannwhitneyu, wilcoxon

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402

B = r"C:\Users\hsollim\Documents\HEAL_mini1p_SBI553"
A = B + r"\analysis\_variants\mix4"
FIG = B + r"\lab_meeting\decks_Day1_Day2\v2_20261022\FINAL_20261026_PI5_mixmarks\figures"
DST = r"C:\Users\hsollim\OneDrive\Desktop\1111112121212_FINAL_compact_added_slides"
os.makedirs(DST, exist_ok=True)

plt.rcParams.update({"font.family": "Arial", "font.size": 12, "axes.spines.top": False, "axes.spines.right": False, "savefig.dpi": 200})
MICE = ["5941", "5939", "5937", "1191"]
SBI = {1: {"5941", "5939"}, 2: {"5937", "1191"}}
SBIDAY = {"5941": 1, "5939": 1, "5937": 2, "1191": 2}
SEX = {"5941": "F", "5939": "M", "5937": "M", "1191": "F"}
MC = {"5941": "#e67e22", "5939": "#8e44ad", "5937": "#16a085", "1191": "#2c3e50"}
MK = {"5941": "o", "5939": "s", "5937": "^", "1191": "D"}
OFF = {"5941": -0.21, "5939": -0.07, "5937": 0.07, "1191": 0.21}
CC = {"vb": "#c4c4c4", "vi": "#2471a3", "sb": "#7a7a7a", "si": "#c0392b"}
STIM = ["Light touch", "Mild touch", "Pin prick", "Heat"]
STARS = "* p < 0.05, ** p < 0.01, *** p < 0.001, ns p ≥ 0.05"


def st(p):
    return "ns" if p is None or not np.isfinite(p) or p >= 0.05 else "***" if p < 0.001 else "**" if p < 0.01 else "*"


def pf(p):
    if p is None or not np.isfinite(p):
        return "n.a."
    return "p < 0.001" if p < 0.001 else f"p = {p:.4f}" if p < 0.01 else f"p = {p:.3f}" if p < 0.1 else f"p = {p:.2f}"


cc = pd.read_csv(A + r"\event_analysis\cells_classes.csv")
PSTH = np.load(B + r"\analysis\new_decks\psth_cache_plainz_mix4.npz", allow_pickle=True)["D"].item()
TRK = pd.read_csv(A + r"\tracked_crossover\tracked_responses.csv"); TRK["mouse"] = TRK.mouse.astype(str)


def win(ev):
    e = cc[cc.event == ev].iloc[0]
    return float(e.pre), float(e.a), float(e.b)


def base_sub(Y, tt, ev):
    """trial-mean z minus each neuron's mean in the test's baseline window (0.5 s before the event) - as make_pi5_deck.base_sub"""
    pre, a, b = win(ev)
    a0 = min(a, 0)
    m = (tt >= a0 - pre - 1e-9) & (tt < a0 - 1e-9)
    return Y - np.nanmean(Y[m], 0, keepdims=True)


def auc_session(m, d, ph, ev):
    """per neuron of one session: AUC (z*s) of the baseline-subtracted trial-mean trace over the response window, and that trace"""
    v = PSTH[(m, d, ph, ev)]
    tt = v["tt"]
    Yb = base_sub(v["mean"], tt, ev)
    pre, a, b = win(ev)
    k0, k1 = np.searchsorted(tt, a - 1e-9), np.searchsorted(tt, b - 1e-9)
    auc = np.trapezoid(Yb[k0:k1 + 1], tt[k0:k1 + 1], axis=0)
    ids = [int(i) for i in v["ids"]]
    return pd.Series(auc, index=ids), pd.DataFrame(Yb, index=tt, columns=ids)


def analyse(ev):
    rows, tr = [], {"veh": [], "sbi": []}
    for m in MICE:
        for d in (1, 2):
            dr = "sbi" if m in SBI[d] else "veh"
            a0, y0 = auc_session(m, d, "pre", ev)
            a1, y1 = auc_session(m, d, "post", ev)
            ids = [i for i in a0.index if i in set(a1.index)]                  # same ROI in both sessions of the day
            for i in ids:
                rows.append(dict(mouse=m, day=d, drug=dr, cell=i, base=a0[i], after=a1[i], d=a1[i] - a0[i]))
            tr[dr].append(y1[ids].values - y0[ids].values)                    # change trace (after - baseline), same neurons
            tt = y0.index.values
    df = pd.DataFrame(rows)
    V, S = df[df.drug == "veh"], df[df.drug == "sbi"]
    # tracked neurons: the same cell on both days
    t = TRK[TRK.event == ev]
    lut = {(r.mouse, r.day, r.cell): r.d for r in df.itertuples()}
    tk = pd.DataFrame([dict(mouse=r.mouse, dv=lut.get((r.mouse, int(r.veh_day), int(r.veh_id))),
                            ds=lut.get((r.mouse, int(r.sbi_day), int(r.sbi_id)))) for r in t.itertuples()]).dropna()
    mv = V.groupby("mouse").d.mean(); ms = S.groupby("mouse").d.mean()
    mdiff = (ms - mv).reindex(MICE)
    R = dict(ev=ev, w=win(ev), n_veh=len(V), n_sbi=len(S), n_trk=len(tk),
             veh_mean=V.d.mean(), veh_sem=V.d.std(ddof=1) / np.sqrt(len(V)), sbi_mean=S.d.mean(), sbi_sem=S.d.std(ddof=1) / np.sqrt(len(S)),
             p_veh=wilcoxon(V.d)[1], p_sbi=wilcoxon(S.d)[1], p_between=mannwhitneyu(S.d, V.d)[1],
             p_trk=wilcoxon(tk.ds, tk.dv)[1], trk_veh=tk.dv.mean(), trk_sbi=tk.ds.mean(),
             trk_mice_lower=int(sum((g.ds - g.dv).mean() < 0 for _, g in tk.groupby("mouse"))),
             mouse_veh={m: float(mv[m]) for m in MICE}, mouse_sbi={m: float(ms[m]) for m in MICE},
             mice_lower=int((mdiff < 0).sum()), p_mice=wilcoxon(mdiff.values)[1],
             # the raw levels of the old slide, to show where its after-vs-after difference comes from
             base_veh=V.base.mean(), base_sbi=S.base.mean(), after_veh=V.after.mean(), after_sbi=S.after.mean(),
             p_base_vs_base=mannwhitneyu(S.base, V.base)[1], p_after_vs_after=mannwhitneyu(S.after, V.after)[1])
    trace = {k: np.hstack(x) for k, x in tr.items()}
    return R, trace, tt


RES = {ev: analyse(ev) for ev in STIM}


def bracket(a, x1, x2, y, txt, color="k", fs=11, h=None):
    h = h if h is not None else 0.015 * (a.get_ylim()[1] - a.get_ylim()[0])
    a.plot([x1, x1, x2, x2], [y, y + h, y + h, y], color=color, lw=1.2, clip_on=False)
    a.text((x1 + x2) / 2, y + h, txt, ha="center", va="bottom", color=color, fontsize=fs if txt != "ns" else fs - 1,
           fontweight="bold" if txt != "ns" else "normal", clip_on=False)


def fig(fn):
    fig, ax = plt.subplots(2, 4, figsize=(12.3, 4.75), gridspec_kw=dict(wspace=0.30, hspace=0.62, left=0.085, right=0.99, top=0.80,
                                                                         bottom=0.10, height_ratios=[1.15, 1.0]))
    for j, ev in enumerate(STIM):
        R, trace, tt = RES[ev]
        a = ax[0, j]
        vals = []
        for x, k, col in ((0, "veh", CC["vi"]), (1, "sbi", CC["si"])):
            mu, se = R[k + "_mean"], R[k + "_sem"]
            a.bar(x, mu, yerr=se, color=col, width=0.6, capsize=3, error_kw=dict(lw=1), zorder=2)
            mm = R["mouse_" + k]
            for m in MICE:
                a.scatter(x + OFF[m] * 0.9, mm[m], s=26, marker=MK[m], color=MC[m], edgecolor="white", lw=0.6, zorder=4)
            vals += [mu + se, mu - se] + list(mm.values())
        hi, lo = max(vals + [0]), min(vals + [0])
        span = hi - lo
        a.set_ylim(lo - 0.10 * span, hi + 0.80 * span)
        for x, k, col in ((0, "veh", CC["vi"]), (1, "sbi", CC["si"])):
            p = R["p_" + k]
            a.text(x, hi + 0.08 * span, "vs 0: " + st(p), ha="center", va="bottom", fontsize=10, color=col,
                   fontweight="bold" if p < 0.05 else "normal")
        bracket(a, 0, 1, hi + 0.38 * span, f"{st(R['p_between'])} ({pf(R['p_between'])})", fs=11, h=0.04 * span)   # 10-26 review: p shown
        a.axhline(0, color="0.4", lw=0.8)
        a.set_xticks([0, 1], ["vehicle\nday", "SBI-553\nday"], fontsize=10.5)
        for lab, c in zip(a.get_xticklabels(), (CC["vi"], CC["si"])):
            lab.set_color(c)
        a.tick_params(axis="y", labelsize=10)
        pre, a_, b_ = win(ev)
        a.set_title(f"{ev}  ({a_:g}-{b_:g} s)\n{R['n_trk']} tracked, day vs day: {st(R['p_trk'])}"
                    + (f" ({pf(R['p_trk'])})" if R["p_trk"] < 0.05 else ""), fontsize=11)
        b = ax[1, j]
        sel = (tt >= -3) & (tt <= 6)
        for k, col in (("veh", CC["vi"]), ("sbi", CC["si"])):
            Y = trace[k][sel]
            mu, se = Y.mean(1), Y.std(1, ddof=1) / np.sqrt(Y.shape[1])
            b.fill_between(tt[sel], mu - se, mu + se, color=col, alpha=0.2, lw=0)
            b.plot(tt[sel], mu, color=col, lw=1.3)
        b.axvspan(min(a_, 0) - pre, min(a_, 0), color="#d6eaf8", lw=0, zorder=0); b.axvspan(a_, b_, color="#e5e7e9", lw=0, zorder=0)
        b.axhline(0, color="0.5", lw=0.7); b.axvline(0, color="k", lw=0.8, ls="--")
        b.set_xlabel("time from event (s)", fontsize=10.5); b.tick_params(labelsize=10)
        b.set_title("change trace: injection − baseline session", fontsize=9.5, color="0.35")
    ax[0, 0].set_ylabel("ΔAUC (z·s): injection −\nbaseline session", fontsize=10.5)
    ax[1, 0].set_ylabel("Δz, mean ± SEM", fontsize=10.5)
    n = RES["Pin prick"][0]
    h = [Patch(color=CC["vi"], label=f"vehicle day ({n['n_veh']} neurons)"), Patch(color=CC["si"], label=f"SBI-553 day ({n['n_sbi']})"),
         Patch(color="#d6eaf8", label="0.5 s before event (subtracted first)"), Patch(color="#e5e7e9", label="response window = AUC")]
    fig.legend(handles=h, ncol=4, loc="upper center", bbox_to_anchor=(0.5, 1.0), frameon=False, fontsize=10.5, handlelength=1.2,
               columnspacing=1.2)
    hm = [Line2D([], [], ls="", marker=MK[m], color=MC[m], markeredgecolor="white", markersize=7,
                 label=f"{m} ({'female' if SEX[m] == 'F' else 'male'}, SBI-553 Day {SBIDAY[m]})") for m in MICE]
    fig.legend(handles=hm, ncol=4, loc="upper center", bbox_to_anchor=(0.5, 0.945), frameon=False, fontsize=10, handlelength=1.0,
               columnspacing=1.2)
    fig.savefig(fn); plt.close(fig)


F = FIG + r"\s3a_auc_basenorm.png"
fig(F)
J = {ev: {k: v for k, v in RES[ev][0].items()} for ev in STIM}
json.dump(J, open(F.replace(".png", ".json"), "w"), indent=1, default=float)
for ev in STIM:
    R = RES[ev][0]
    print(f"{ev:12s} dAUC veh {R['veh_mean']:+.4f} ({pf(R['p_veh'])})  sbi {R['sbi_mean']:+.4f} ({pf(R['p_sbi'])})  between {pf(R['p_between'])}"
          f" | tracked n={R['n_trk']} veh {R['trk_veh']:+.4f} sbi {R['trk_sbi']:+.4f} {pf(R['p_trk'])} lower {R['trk_mice_lower']}/4"
          f" | mice lower {R['mice_lower']}/4 {pf(R['p_mice'])}"
          f" | base veh {R['base_veh']:+.4f} sbi {R['base_sbi']:+.4f} {pf(R['p_base_vs_base'])}; after veh {R['after_veh']:+.4f} "
          f"sbi {R['after_sbi']:+.4f} {pf(R['p_after_vs_after'])}")
print("->", F)


# ------------------------------------------------------------------------------------------------ the 1-slide pptx (compact-deck style)
from lxml import etree  # noqa: E402
from matplotlib.font_manager import FontProperties  # noqa: E402
from matplotlib.textpath import TextPath  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

FOOT = "CeA-Ntsr1 mini1p  ·  SBI-553 vs vehicle crossover  ·  n = 4 mice  ·  exploratory, uncorrected (details in the notes)"


def fit(text, pt, bold, width, lo):
    prop = FontProperties(family="Arial", weight="bold" if bold else "normal")
    while pt > lo and TextPath((0, 0), text, size=pt, prop=prop).get_extents().width / 72 > width:
        pt -= 0.5
    if TextPath((0, 0), text, size=pt, prop=prop).get_extents().width / 72 > width:
        print("TOO LONG:", text)
    return pt


def tbox(s, x, y, w, h, text, size, bold=False, color=(0x22, 0x22, 0x22), align=None, italic=False):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = "Arial"; r.font.color.rgb = RGBColor(*color)
        if align:
            p.alignment = align
    return t


def arrow(s, x0, y0, x1, y1):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = RGBColor(0xC0, 0x39, 0x2B); c.line.width = Pt(3.5)
    te = etree.SubElement(c.line._get_or_add_ln(), "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")


P = RES["Pin prick"][0]
SECTION = "4  STIMULUS RESPONSES  ·  population (all neurons), change from each day's baseline session"
TITLE = "Change from each day's baseline session: vehicle and SBI-553 days do not differ"      # 10-26 after the review
KEY = (f"Pin after-vs-after gap ({pf(P['p_after_vs_after'])}) disappears once each neuron's own baseline session is subtracted "
       f"({pf(P['p_between'])})")
p4 = lambda k: ", ".join(f"{ev.split()[0].lower()} {pf(RES[ev][0][k])}" for ev in STIM)       # noqa: E731
BOTTOM = (f"ΔAUC = each neuron's AUC in the injection session − its AUC in the same day's baseline session (z·s), mean ± SEM; "
          f"markers = mouse means.\nVehicle vs SBI-553 day (bracket, Mann-Whitney over {P['n_veh']} / {P['n_sbi']} neurons): "
          f"{p4('p_between')}; vs 0 (Wilcoxon): all p ≥ {min(min(RES[ev][0]['p_veh'], RES[ev][0]['p_sbi']) for ev in STIM):.2f}.\n"
          f"4 mice, each on both days; mouse-level paired Wilcoxon all p ≥ {min(RES[ev][0]['p_mice'] for ev in STIM):.3f} (lowest possible "
          f"0.125). Pin raw AUC, baseline sessions: {P['base_veh']:.3f} vs {P['base_sbi']:.3f} z·s ({pf(P['p_base_vs_base'])}).")
NOTES = "\n".join([
    "WHY (user 10-26): on the previous slide the after-vehicle vs after-SBI-553 comparison mixes the drug with the SBI-553 day's "
    "lower baseline. Here every neuron is compared with itself in the same day's baseline session, and the two days' CHANGES are compared.",
    "MEASURE: per neuron-session AUC = area of (trial-mean z - its mean in the 0.5 s before the event) over the response window (pin "
    "0-0.5 s; light, mild, heat 0-1 s) - the same AUC as the previous slide. dAUC = AUC after the injection - AUC in the same day's "
    "baseline session (same ROI). Subtraction, not a ratio: the SBI-553-day pin baseline AUC is about 0 "
    f"({P['base_sbi']:+.4f} z·s), so % of baseline is undefined.",
    "TESTS: star over a bar = Wilcoxon signed-rank of dAUC vs 0 (= the blue / red brackets of the previous slide); black bracket = "
    "Mann-Whitney, vehicle-day dAUC vs SBI-553-day dAUC (all neurons, different sets); panel title = paired Wilcoxon of the two changes "
    "in the 185 neurons tracked on both days; mice = paired Wilcoxon over the 4 mouse means (smallest possible p 0.125). Exploratory, "
    "uncorrected.",
    "RESULTS: " + " | ".join(
        f"{ev}: dAUC vehicle day {R['veh_mean']:+.4f} ({pf(R['p_veh'])}), SBI-553 day {R['sbi_mean']:+.4f} ({pf(R['p_sbi'])}), "
        f"between {pf(R['p_between'])}; tracked {R['n_trk']}: {R['trk_veh']:+.4f} vs {R['trk_sbi']:+.4f} ({pf(R['p_trk'])}, "
        f"SBI-553-day change lower in {R['trk_mice_lower']}/4 mice); mouse means lower on the SBI-553 day in {R['mice_lower']}/4 "
        f"({pf(R['p_mice'])}); raw levels: baseline vs baseline {pf(R['p_base_vs_base'])}, after vs after {pf(R['p_after_vs_after'])}"
        for ev, R in ((ev, RES[ev][0]) for ev in STIM)) + ".",
    "READ: no stimulus changed from its own baseline on either day, and the change did not differ between days for any stimulus in all "
    "neurons. The pin after-vs-after difference of the previous slide (p = 0.041) is carried by the baselines (baseline vs baseline "
    f"{pf(P['p_base_vs_base'])}). Only light touch in the 185 tracked neurons differed ({pf(RES['Light touch'][0]['p_trk'])}, 1 of 4 "
    "stimuli, uncorrected; slide 'each neuron vs its own baseline' shows the same, p = 0.043 on R).",
    "WHY THIS DOES NOT CONTRADICT THE UP-NEURON SLIDES: the population mean adds UP, DOWN and non-responsive neurons together; the "
    "number of neurons that pass the shift test can change while the mean area of all neurons does not (see 'Responder numbers changed, "
    "response size did not').",
    "UP / DOWN ACCOUNTING (pin, vehicle day; from the independent review): 29 neurons became UP (mean dAUC +0.098) and 12 became DOWN "
    "(-0.120), so the mean stays at -0.002; the UP rise is mostly 5939, and neurons selected as UP after the injection have a positive "
    "change partly by selection.",
    "WHERE THE PIN GAP COMES FROM: the lower SBI-553-day baseline-session AUC comes from 5941 and 5939 (SBI-553 on Day 1): per mouse, "
    "SBI-553 day minus vehicle day, baseline / after / change = 5941 -0.042 / -0.040 / +0.002, 5939 -0.049 / -0.051 / -0.002, 5937 "
    "-0.003 / +0.026 / +0.028, 1191 +0.008 / +0.014 / +0.005. So it is a Day-1 vs Day-2 difference in 2 mice (drug order = day), present "
    "before any injection; in the 185 tracked neurons the pin baseline gap is paired p = 0.001. Neither the baseline nor the after gap "
    "survives a session random effect (p = 0.25 and 0.31). Baseline-session AUC was lower on the SBI-553 day for all 4 stimuli "
    "(mild p = 0.028).",
    "ROBUSTNESS (independent review): ANCOVA (after ~ baseline session + day) and mixed models with mouse / session random effects give "
    "the same answer, no stimulus significant (pin ANCOVA p = 0.61-0.63). Regression to the mean (lower SBI-553-day baselines) pushes the "
    "change score slightly toward the SBI-553 day; ANCOVA estimates are all slightly negative and none is significant. Light touch, "
    "tracked: Holm-corrected over 4 stimuli p = 0.14, paired t p = 0.063, driven by 5941 and 5939 (confounded with day order).",
    "WHICH BASELINE: here 'baseline session' = the session before the injection on the same day. The 0.5 s before each event is removed "
    "first, by subtraction, in every session (that is what the AUC already is).",
    "BOTTOM ROW: change trace = (injection - baseline session) of each neuron's trace (each already minus its own 0.5 s before), mean ± "
    "SEM over neurons; blue band = 0.5 s before, grey = response window (the AUC).",
    r"Made by mini1p_pipeline\make_auc_basenorm.py (numbers in figures\s3a_auc_basenorm.json); review scripts in the scratchpad "
    r"verify_basenorm folder."])

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
tbox(s, 0.45, 0.10, 11, 0.3, SECTION, 11, color=(0x55, 0x55, 0x55))
tbox(s, 0.45, 0.32, 12.4, 0.6, TITLE, fit(TITLE, 24, False, 12.15, 20))
tbox(s, 0.45, 0.92, 12.4, 0.5, KEY, fit(KEY, 15, True, 12.15, 13), bold=True, color=(0xC0, 0x39, 0x2B))
s.shapes.add_picture(F, Inches(0.5), Inches(1.45), Inches(12.3))
tbox(s, 0.45, 6.2, 12.4, 0.85, BOTTOM, 10.5)
tbox(s, 0.45, 7.1, 11.3, 0.3, FOOT, 9, color=(0x88, 0x88, 0x88), italic=True)
tbox(s, 12.2, 7.1, 0.9, 0.3, "9b / 14", 10, color=(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT)
arrow(s, 10.0, 2.38, 9.02, 2.60)                                  # at the pin change-vs-change bracket label
s.notes_slide.notes_text_frame.text = NOTES
OUTP = os.path.join(DST, "01_NEW_after_slide36_population_AUC_baseline_normalized.pptx")   # slide 36 of the user's deck = '9 / 14'
prs.save(OUTP)
print("->", OUTP)
