"""reorg_pi_deck.py - 10-26 (user): the mini1p part (slide 17 on) of the PI deck on the Desktop, re-ordered and made compact:
example -> analysis, overall -> specific, all neurons -> the same tracked neurons; one message per slide, arrows on the key
result, wordy text moved to the notes, less important slides kept as hidden BACKUP after the take-home (nothing deleted).
The original file is not touched; the result is written next to it.

  python reorg_pi_deck.py
"""
import copy
import io
import os

from lxml import etree
from matplotlib.font_manager import FontProperties
from matplotlib.textpath import TextPath
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SRC = r"C:\Users\hsollim\OneDrive\Desktop\1111112121212.pptx"
OUT = r"C:\Users\hsollim\OneDrive\Desktop\1111112121212_FINAL_compact.pptx"
ADD = r"C:\Users\hsollim\Documents\HEAL_mini1p_SBI553\lab_meeting\decks_Day1_Day2\v2_20261022\FINAL_20261025_3_mixmarks_mixwin\added_slides_20261026"
ABN = (r"C:\Users\hsollim\OneDrive\Desktop\1111112121212_FINAL_compact_added_slides"           # 10-26, make_auc_basenorm.py
       r"\01_NEW_after_slide36_population_AUC_baseline_normalized.pptx")
T33 = r"C:\Users\hsollim\AppData\Local\Temp\claude\C--Users-hsollim\e78b98bc-45eb-4eba-b6ff-ba7ac412ccce\scratchpad\reorg\T33.pptx"
E = 914400
RED = RGBColor(0xC0, 0x39, 0x2B)
STARS = "* p < 0.05, ** p < 0.01, *** p < 0.001, ns p ≥ 0.05"
FOOT = "CeA-Ntsr1 mini1p  ·  SBI-553 vs vehicle crossover  ·  n = 4 mice  ·  exploratory, uncorrected (details in the notes)"

prs = Presentation(SRC)
BL = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
for lay in prs.slide_layouts:
    if lay.name.lower() == "blank":
        BL = lay


def fit(text, pt, bold, width, lo):
    prop = FontProperties(family="Arial", weight="bold" if bold else "normal")
    while pt > lo and TextPath((0, 0), text, size=pt, prop=prop).get_extents().width / 72 > width:
        pt -= 0.5
    return pt


def set_text(sh, text, size=None, bold=None, color=None):
    """replace a text box's text, keeping the first run's formatting"""
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    r0 = p0.runs[0] if p0.runs else p0.add_run()
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs)[1:]:
        r._r.getparent().remove(r._r)
    r0.text = text
    if size:
        r0.font.size = Pt(size)
    if bold is not None:
        r0.font.bold = bold
    if color is not None:
        r0.font.color.rgb = color


def tbox(s, x, y, w, h, text, size, bold=False, color=(0x22, 0x22, 0x22), align=None, italic=False):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = "Arial"; r.font.color.rgb = RGBColor(*color)
        if align:
            p.alignment = align
    return t


def arrow(s, x0, y0, x1, y1, label=None, lx=None, ly=None, lw=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = RED; c.line.width = Pt(3.5)
    ln = c.line._get_or_add_ln()
    te = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    if label:
        tbox(s, lx, ly, lw or 2.2, 0.35, label, 12.5, bold=True, color=(0xC0, 0x39, 0x2B))


def roles(s):
    """the text boxes of a slide made by the deck builders, by position"""
    R = {}
    for sh in s.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        x, y = sh.left / E, sh.top / E
        if y < 0.06 and x > 4.5:
            R["banner"] = sh
        elif y < 0.25:
            R["section"] = sh
        elif y < 0.7:
            R["title"] = sh
        elif y < 1.4:
            R["key"] = sh
        elif y >= 7.0 and x > 11.5:
            R["num"] = sh
        elif y >= 7.0:
            R["footer"] = sh
        elif y >= 6.2:
            R["bottom"] = sh
    return R


def remove(sh):
    sh._element.getparent().remove(sh._element)


def edit(s, section, title, key, bottom, k, n):
    R = roles(s)
    moved = []
    for nm in ("key", "bottom", "footer", "banner"):
        if nm in R:
            moved.append(R[nm].text_frame.text)
    old = s.notes_slide.notes_text_frame.text
    s.notes_slide.notes_text_frame.text = "ON-SLIDE TEXT MOVED HERE (10-26):\n" + "\n".join(moved) + "\n\n" + old
    for nm in ("footer", "banner", "num"):
        if nm in R:
            remove(R[nm])
    set_text(R["section"], section, size=11)
    set_text(R["title"], title, size=fit(title, 24, False, 12.15, 20))
    if "key" in R:
        set_text(R["key"], key, size=fit(key, 15, True, 12.15, 13), bold=True, color=RED)
    else:
        tbox(s, 0.45, 0.92, 12.4, 0.5, key, fit(key, 15, True, 12.15, 13), bold=True, color=(0xC0, 0x39, 0x2B))
    if "bottom" in R:
        set_text(R["bottom"], bottom, size=11)
        pb = max([(sh.top + sh.height) / E for sh in s.shapes if sh.shape_type in (13, 16)] + [6.5])   # below the figure
        R["bottom"].top = Inches(min(max(6.55, pb + 0.02), 6.68))
        R["bottom"].height = Inches(0.42)
    footer(s, k, n)


def footer(s, k, n):
    tbox(s, 0.45, 7.1, 11.3, 0.3, FOOT, 9, color=(0x88, 0x88, 0x88), italic=True)
    tbox(s, 12.2, 7.1, 0.9, 0.3, f"{k} / {n}", 10, color=(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT)


def import_slide(path, idx=0):
    """copy slide idx of another deck (pictures + text boxes + lines) to the end of prs"""
    src = Presentation(path).slides[idx]
    s = prs.slides.add_slide(BL)
    for ph in list(s.placeholders):
        remove(ph)
    for sh in src.shapes:
        if sh.shape_type == 13:
            s.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
        else:
            s.shapes._spTree.insert_element_before(copy.deepcopy(sh._element), "p:extLst")
    s.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text if src.has_notes_slide else ""
    return s


def new_slide():
    s = prs.slides.add_slide(BL)
    for ph in list(s.placeholders):
        remove(ph)
    return s


S = {i: prs.slides[i - 1] for i in range(1, len(prs.slides) + 1)}      # original numbering
N_MAIN = 14
FIGD = r"C:\Users\hsollim\Documents\HEAL_mini1p_SBI553\lab_meeting\decks_Day1_Day2\v2_20261022\FINAL_20261026_PI5_mixmarks\figures"

# ------------------------------------------------------------------ the main sequence
ov = new_slide()
tbox(ov, 0.45, 0.10, 8, 0.3, "CeA-Ntsr1 mini1p  ·  OVERVIEW", 11, color=(0x55, 0x55, 0x55))
tbox(ov, 0.45, 0.32, 12.4, 0.6, "SBI-553 vs vehicle in CeA-Ntsr1 neurons: what we did and what we found", 24)
tbox(ov, 0.45, 0.95, 12.4, 0.5, "4 mice, crossover: each day = a baseline session, then an injection session (vehicle or SBI-553)", 15,
     bold=True, color=(0xC0, 0x39, 0x2B))
steps = [("1", "The data", "one example session:\nbehaviour + Ca together"),
         ("2", "Behaviour", "what the mice did\nafter each stimulus"),
         ("3", "Overall activity", "all neurons\n→ the same tracked neurons"),
         ("4", "Stimulus responses", "population → single neurons\n→ pin-prick close-up → tracked"),
         ("5", "Take-home", "three messages\n+ limits")]
x = 0.45
for i, (n_, h, d) in enumerate(steps):
    b = ov.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.3), Inches(2.2), Inches(2.6))
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0xF4, 0xF6, 0xF7) if i < 4 else RGBColor(0xFD, 0xED, 0xEC)
    b.line.color.rgb = RGBColor(0x2C, 0x3E, 0x50) if i < 4 else RED; b.line.width = Pt(1.5)
    tbox(ov, x + 0.1, 2.45, 2.05, 0.6, n_, 30, bold=True, color=(0x2C, 0x3E, 0x50) if i < 4 else (0xC0, 0x39, 0x2B), align=PP_ALIGN.CENTER)
    tbox(ov, x + 0.1, 3.1, 2.05, 0.75, h, 16, bold=True, align=PP_ALIGN.CENTER)
    tbox(ov, x + 0.1, 3.85, 2.05, 1.0, d, 12, color=(0x44, 0x44, 0x44), align=PP_ALIGN.CENTER)
    if i < 4:
        a_ = ov.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.23), Inches(3.4), Inches(0.26), Inches(0.4))
        a_.fill.solid(); a_.fill.fore_color.rgb = RGBColor(0x95, 0xA5, 0xA6); a_.line.fill.background()
    x += 2.52
tbox(ov, 0.45, 5.35, 12.4, 0.9,
     "SBI-553 on Day 1 in 5941 (F) and 5939 (M); on Day 2 in 5937 (M) and 1191 (F). Every comparison is made two ways: within a day "
     "(injection vs that day's baseline, same neurons) and between days (after vehicle vs after SBI-553).", 12.5, color=(0x33, 0x33, 0x33))
footer(ov, 1, N_MAIN)
ov.notes_slide.notes_text_frame.text = ("Order of the talk: example first, then the analyses; overall activity before stimulus responses; "
                                        "all neurons before the tracked (same-cell) neurons. Backup slides after the take-home are hidden.")

ht = new_slide()
tbox(ht, 0.45, 0.10, 8, 0.3, "HOW TO READ THE FIGURES", 11, color=(0x55, 0x55, 0x55))
tbox(ht, 0.45, 0.32, 12.4, 0.6, "Design and the colour code used on every slide", 24)
tbox(ht, 0.45, 0.92, 12.4, 0.5, "Each day: baseline session, then injection; every result is read within a day AND between days", 15, bold=True,
     color=(0xC0, 0x39, 0x2B))
ht.shapes.add_picture(FIGD + r"\s0_design.png", Inches(0.5), Inches(1.45), Inches(12.3))
ht.shapes.add_picture(FIGD + r"\s_howto.png", Inches(0.5), Inches(3.62), Inches(12.3))
footer(ht, 2, N_MAIN)
ht.notes_slide.notes_text_frame.text = (
    "Design: 4 mice (2 F, 2 M), crossover; SBI-553 on Day 1 for 5941 and 5939, on Day 2 for 5937 and 1191; vehicle on the other day. "
    "Each session has 4 stimulus blocks (light touch, mild touch, pin prick, heat) in random order.\n"
    "Response of a neuron = mean z (dF/F / robust SD of the session) in the response window minus the 0.5 s before the event.\n"
    "Tests: within a day the same neurons are compared (Wilcoxon / McNemar); between days all neurons are different sets (Mann-Whitney "
    "/ Fisher) unless they are the 185 tracked neurons (paired). Behaviour: CMH over deliveries with mouse as stratum. "
    "Neurons pooled over mice, exploratory, uncorrected; mouse-level tests with 4 mice cannot go below p = 0.125.")

edit(S[17], "1  THE DATA  ·  one example session", "One session: paw contacts, Ca movie and pin-prick neurons together",
     "Shade = 0-0.5 s after each pin prick: the response window used in every test",
     "5941, Day 1 baseline; n = 1 example session (descriptive)", 3, N_MAIN)

edit(S[19], "2  BEHAVIOUR  ·  every behaviour, every stimulus", "Pin prick: flinches rose after vehicle but fell after SBI-553",
     "Change from the same day's baseline differed between days only for pin flinch (p < 0.001) and pin escape (p < 0.05)",
     "Bars = mean of 4 mice of (% after injection - % at that day's baseline), markers = mice. Star over a bar = CMH vs baseline; "
     f"bracket = day difference (logistic over deliveries). {STARS}", 4, N_MAIN)
arrow(S[19], 8.55, 3.95, 7.75, 3.3, "flinch fell after SBI-553", 8.6, 3.85, 2.6)

edit(S[18], "2  BEHAVIOUR  ·  the behaviours that differed", "After SBI-553: fewer light-touch movements and pin flinches than after vehicle",
     "But not below the SBI-553-day baseline: part of the gap is a rise after vehicle",
     "n = 4 mice. Bars = mean of mice, markers = mice. CMH over deliveries (mouse = stratum): blue / red = vs the same day's baseline, "
     f"black = after vehicle vs after SBI-553. {STARS}", 5, N_MAIN)
arrow(S[18], 7.55, 3.55, 6.55, 3.68, "rise after vehicle", 7.6, 3.38, 2.0)

edit(S[23], "3  OVERALL ACTIVITY  ·  all neurons", "Firing fell after both injections; lower after SBI-553 than after vehicle",
     "After vs after p = 0.012, while the two baselines were equal (p = 0.64)",
     "Dot = one neuron (vehicle day 216, SBI-553 day 246), 4 mice. Wilcoxon within day (same neurons); Mann-Whitney between days. "
     f"{STARS}", 6, N_MAIN)
arrow(S[23], 7.25, 2.72, 5.6, 2.88)

edit(S[24], "3  OVERALL ACTIVITY  ·  the same tracked neurons, whole session", "Whole session, same 185 neurons: activity lower after SBI-553",
     "Each row = one mouse's mean z (1-s bins); the drop is quantified on the next slide",
     "Population mean z of the 185 tracked neurons, blocks rebuilt in one order (light, mild, pin, heat); bottom row of each group = "
     "mean of 4 mice. Descriptive.", 7, N_MAIN)

w2 = import_slide(os.path.join(ADD, "12_FIXED_slides27-28_tracked_activity_change_tests.pptx"))
edit(w2, "3  OVERALL ACTIVITY  ·  the same tracked neurons, quantified", "Same 185 neurons: activity fell more on the SBI-553 day",
     "Stimulus blocks -0.12 vs -0.03 z (p < 0.001), lower in 4/4 mice; but drug order = day",
     "Each tracked neuron: mean z after the injection minus its own same-day baseline. Star = Wilcoxon vs 0; bracket = paired Wilcoxon "
     f"between days; 'mice' = Wilcoxon over 4 mice (smallest possible p 0.125). {STARS}", 8, N_MAIN)
arrow(w2, 12.55, 5.25, 12.0, 4.6)

edit(S[26], "4  STIMULUS RESPONSES  ·  population (all neurons)",
     "Population responses: clear on the vehicle day, weak on the whole SBI-553 day",
     "Weak already at the SBI-553-day baseline, so not a drug effect by itself; after vs after only pin (p = 0.041)",
     "Bars = AUC of (z - 0.5 s before) in the response window, mean ± SEM of neurons (216 / 246); traces below; 4 mice. Star = Wilcoxon "
     f"vs 0; blue / red = within day; black = Mann-Whitney after vs after. {STARS}", 9, N_MAIN)
arrow(S[26], 10.1, 2.15, 9.35, 2.33)
abn = import_slide(ABN)              # 10-26 (user: 'normalized by baseline?'): the same AUC minus each day's baseline session, '9b'

edit(S[29], "4  STIMULUS RESPONSES  ·  how a neuron is called UP or DOWN",
     "UP / DOWN = the response beats 1000 time-shifted copies of the event times",
     "Response window minus the 0.5 s before, vs 1000 shifted event trains; BH q ≤ 0.20 → UP (R > 0) or DOWN (R < 0)",
     "Example: 5939, Day 2 after vehicle, pin prick: 47 neurons → 21 UP, 9 DOWN. Top: every trial; middle: mean ± SEM; bottom: grey = "
     "the 1000 shifted trains, line = the real response.", 10, N_MAIN)
arrow(S[29], 4.85, 4.55, 4.22, 4.92)

edit(S[30], "4  STIMULUS RESPONSES  ·  each neuron vs its own baseline", "Responder numbers changed, response size did not",
     "Touch responders rose on the SBI-553 day; pin-prick responders rose on the vehicle day",
     "Top: response change of each neuron (after - its baseline). Bottom: change in % UP / DOWN. Star over a bar = McNemar / Wilcoxon "
     f"vs baseline; bracket = Mann-Whitney between days. Neurons 216 / 246, 4 mice. {STARS}", 11, N_MAIN)

edit(S[34], "4  STIMULUS RESPONSES  ·  close-up: pin prick", "Pin prick: UP neurons stayed UP after vehicle (72%), not after SBI-553 (29%)",
     "Fisher p = 0.002 (neurons pooled); the vehicle-day side comes mostly from 5941",
     "Each baseline class = 100 %; all neurons (vehicle day 216, SBI-553 day 246), 4 mice; Fisher exact on the key transitions, "
     f"markers = mice. {STARS}", 12, N_MAIN)
arrow(S[34], 9.25, 2.85, 9.95, 3.38)

t33 = import_slide(T33)
edit(t33, "4  STIMULUS RESPONSES  ·  the same tracked neurons",
     "Same 185 neurons: fewer pin-prick UP neurons after SBI-553 (9%) than after vehicle (23%)",
     "Pin after vs after McNemar p < 0.001; light-touch and heat UP neurons rose after SBI-553 too",
     "Top: % UP (up) / DOWN (hatched) of the 185 neurons found on both days; markers = mice; all brackets McNemar (same neurons). "
     f"Bottom: mean trace of each session's UP neurons (z - 0.5 s before). 4 mice. {STARS}", 13, N_MAIN)
arrow(t33, 9.85, 2.55, 9.05, 2.45)

th = new_slide()
tbox(th, 0.45, 0.10, 8, 0.3, "5  TAKE-HOME", 11, color=(0x55, 0x55, 0x55))
tbox(th, 0.45, 0.32, 12.4, 0.6, "Take-home", 26)
tbox(th, 2.6, 0.42, 10.3, 0.5, "After vehicle, pin-prick responses rose in behaviour and CeA neurons; not on the SBI-553 day", 14.5, bold=True, color=(0xC0, 0x39, 0x2B))
cards = [("Behaviour", "After SBI-553: fewer light-touch movements and pin flinches than after vehicle",
          "but not below the SBI-553-day baseline: part is a rise after vehicle"),
         ("Overall activity", "Firing fell after both injections, more on the SBI-553 day in the same neurons",
          "-0.12 vs -0.03 z (p < 0.001); lower in 4/4 mice"),
         ("Stimulus responses", "Pin-prick UP neurons rose and stayed after vehicle, not after SBI-553",
          "same 185 neurons 9% vs 23% UP (p < 0.001); touch responders rose on the SBI-553 day")]
for i, (h, msg, sup) in enumerate(cards):
    x = 0.45 + i * 4.2
    b = th.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.2), Inches(3.95), Inches(3.5))
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0xF4, 0xF6, 0xF7); b.line.color.rgb = RGBColor(0x2C, 0x3E, 0x50); b.line.width = Pt(1.5)
    tbox(th, x + 0.2, 1.35, 3.55, 0.5, f"{i + 1}  {h}", 17, bold=True, color=(0x2C, 0x3E, 0x50))
    tbox(th, x + 0.2, 1.95, 3.55, 1.5, msg, 18, bold=True, color=(0xC0, 0x39, 0x2B))
    tbox(th, x + 0.2, 3.55, 3.55, 1.0, sup, 14, color=(0x33, 0x33, 0x33))
lim = th.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(5.0), Inches(12.4), Inches(1.3))
lim.fill.solid(); lim.fill.fore_color.rgb = RGBColor(0xFE, 0xF9, 0xE7); lim.line.color.rgb = RGBColor(0xF3, 0x9C, 0x12)
tbox(th, 0.65, 5.12, 12.0, 1.1, "Limits: n = 4 mice; the drug order is the day order (SBI-553 Day 1 in 2 mice, Day 2 in 2); population "
     "responses were already weak at the SBI-553-day baseline; neurons pooled, exploratory, uncorrected.", 13, color=(0x33, 0x33, 0x33))
footer(th, 14, N_MAIN)
th.notes_slide.notes_text_frame.text = ("Mouse-level tests with 4 mice cannot go below p = 0.125, so every p in this talk is neuron- or "
                                        "delivery-pooled and exploratory.")

bk = new_slide()
tbox(bk, 0.45, 2.8, 12.4, 1.0, "BACKUP", 40, bold=True, color=(0x7F, 0x8C, 0x8D), align=PP_ALIGN.CENTER)
tbox(bk, 0.45, 3.9, 12.4, 0.6, "hidden slides: shown only when asked", 16, color=(0x7F, 0x8C, 0x8D), align=PP_ALIGN.CENTER)

f70 = import_slide(os.path.join(ADD, "06_FIXED_slide70_UP_DOWN_all_tests.pptx"))

# ------------------------------------------------------------------ order + hidden flags
main = [ov, ht, S[17], S[19], S[18], S[23], S[24], w2, S[26], abn, S[29], S[30], S[34], t33, th]
backup = [S[20], S[21], S[22], S[25], S[27], f70, S[28], S[33], S[31], S[32], S[35]]
for s in backup:
    s._element.set("show", "0")
    R = roles(s)
    if "section" in R:
        tx = R["section"].text_frame.text
        if not tx.startswith("BACKUP"):
            set_text(R["section"], "BACKUP  ·  " + tx.replace("NTSR!", "Ntsr1"), size=11)
lst = prs.slides._sldIdLst
ids = {s.slide_id: el for el in list(lst) for s in [None] if False}
by_id = {el.get("id"): el for el in list(lst)}
order = [prs.slides.index(s) for s in [S[i] for i in range(1, 17)] + main + [bk] + backup]
els = list(lst)
for el in els:
    lst.remove(el)
for i in order:
    lst.append(els[i])
assert len(lst) == len(els), (len(lst), len(els))
prs.save(OUT)
print("->", OUT, len(prs.slides), "slides;", "main", len(main), "backup", len(backup))
