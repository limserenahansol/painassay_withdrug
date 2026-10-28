"""make_pi2_ca_summary.py - 10-28 (user, on the first 2-slide summary: 'your summary is wrong, we need to show some interesting Ca
imaging neuronal activity data!'): 2 PI slides built around the Ca imaging. Same data as ALL_FULL mixmarks (analysis\\_variants\\mix4).

  slide 1  CeA-Ntsr1 neurons respond to pain : one field of view with the pin-prick UP / DOWN neurons; every neuron's pin-prick
           response (heat map, baseline sessions); % UP / DOWN neurons for every stimulus and behaviour (baseline sessions)
  slide 2  SBI-553 : stimulus-block activity of the same 185 neurons per mouse; pin-prick responses of the same neurons in the
           4 sessions (heat maps, one row = one neuron); pin flinch (behaviour)

  python make_pi2_ca_summary.py -> OneDrive\\Desktop\\mini1p_SBI553_PI_2slides_Ca\\  (the earlier 2-slide file stays as it was)
"""
import json
import os

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = open(os.path.join(HERE, "make_pi5_deck.py"), encoding="utf-8").read().split("\n")


def find(prefix, start=1):
    for i in range(start - 1, len(SRC)):
        if SRC[i].startswith(prefix):
            return i + 1
    raise KeyError(prefix)


G = {"__name__": "make_pi2_ca_summary", "__file__": os.path.join(HERE, "make_pi5_deck.py")}
exec(compile("\n".join(SRC[:find("prs = Presentation()") - 1]), "make_pi5_deck.py", "exec"), G)
exec(compile("\n".join(SRC[find("SEGT = ") - 1:find("def fig_w2") - 1]), "make_pi5_deck.py", "exec"), G)
globals().update({k: v for k, v in G.items() if not k.startswith("__")})

import scipy.io as sio  # noqa: E402
from matplotlib.colors import TwoSlopeNorm  # noqa: E402

OUTD = r"C:\Users\hsollim\OneDrive\Desktop\mini1p_SBI553_PI_2slides_Ca"
os.makedirs(OUTD, exist_ok=True)
plt.rcParams.update({"font.size": 12, "axes.titlesize": 12.5, "xtick.labelsize": 11, "ytick.labelsize": 11})
RC = np.load(os.path.join(B, r"analysis\baseline_methods\cache.npz"), allow_pickle=True)["R"].item()   # trial-mean traces per method
TT = np.arange(int(-5 * 30), int(8 * 30) + 1) / 30.0
UPC, DNC, RSC = "#e74c3c", "#2e86c1", "#b0b7bb"
NOTE = {}


def trace(m, d, ph, ev, cid):
    v = RC.get((m, d, ph, ev))
    if v is None:
        return None
    w = np.where(np.asarray(v["ids"]).astype(int) == int(cid))[0]
    return v["tr"]["M1"][:, w[0]] if len(w) else None


def resp_window(ev):
    pre, a, b = win(ev)
    return (TT >= a - 1e-9) & (TT <= b + 1e-9)


# ============================================================================================ figure 1: neurons respond to pain
def fig1(fn):
    fig = plt.figure(figsize=(12.3, 4.6))
    gs = fig.add_gridspec(1, 3, width_ratios=[1.05, 1.15, 1.2], wspace=0.50, left=0.015, right=0.99, top=0.86, bottom=0.17)
    # A field of view, 5941 Day-1 baseline, pin-prick classes
    a = fig.add_subplot(gs[0])
    root = os.path.join(B, r"mini1p_out\20261019\5941_d1")
    S = np.asarray(sio.loadmat(os.path.join(root, "3_extract", "5941_d1_extract.mat"), variable_names=["spatial_weights"])
                   ["spatial_weights"], np.float32)
    bg = np.asarray(sio.loadmat(os.path.join(root, "3_extract", "5941_d1_projections.mat"), variable_names=["max_dff"])["max_dff"], float)
    Rv = pd.read_csv(os.path.join(root, "4_cell_review", "5941_d1_cell_review.csv"))
    cls = cc[(cc.mouse == "5941") & (cc.day == 1) & (cc.phase == "pre") & (cc.event == "Pin prick")].set_index("cell").cls
    lo, hi = np.percentile(bg, [1, 99.8])
    img = np.clip((bg - lo) / (hi - lo), 0, 1) ** 0.7
    rows, cols = np.where(np.any(S[:, :, (Rv.keep == 1).values] > 0, axis=2))
    r0, r1 = max(rows.min() - 15, 0), min(rows.max() + 15, img.shape[0]); c0, c1 = max(cols.min() - 15, 0), min(cols.max() + 15, img.shape[1])
    a.imshow(img[r0:r1, c0:c1], cmap="gray", vmin=0, vmax=1)
    n = {"activated": 0, "inhibited": 0, "non-responsive": 0}
    for j in np.where(Rv.keep.values == 1)[0]:
        f = S[r0:r1, c0:c1, j]
        k = cls.get(int(Rv.id.iloc[j]), "non-responsive")
        n[k] += 1
        col, lw = {"activated": (UPC, 1.6), "inhibited": (DNC, 1.6), "non-responsive": ("#dddddd", 0.7)}[k]
        if f.max() > 0:
            a.contour((f >= 0.3 * f.max()).astype(float), levels=[0.5], colors=[col], linewidths=lw)
    a.set_axis_off()
    a.set_title(f"A  One field of view (mouse 5941, baseline)\n{n['activated']} UP, {n['inhibited']} DOWN of {sum(n.values())} neurons "
                "to pin prick", loc="left", fontsize=11.5)
    NOTE["fov"] = n
    # B heat map: every neuron's pin-prick response, baseline sessions of both days
    b = fig.add_subplot(gs[1])
    Y, C = [], []
    for (m, d, ph, ev), v in RC.items():
        if ev != "Pin prick" or ph != "pre":
            continue
        cl = cc[(cc.mouse == m) & (cc.day == d) & (cc.phase == ph) & (cc.event == ev)].set_index("cell").cls.reindex(v["ids"]).values
        Y.append(v["tr"]["M1"]); C += list(cl)
    Y = np.hstack(Y); C = np.array(C)
    k = resp_window("Pin prick")
    order = np.argsort(-np.nanmean(Y[k], 0))
    sel = (TT >= -2) & (TT <= 4)
    im = b.imshow(Y[sel][:, order].T, aspect="auto", cmap="RdBu_r", norm=TwoSlopeNorm(0, -1.0, 1.0), interpolation="nearest",
                  extent=[TT[sel][0], TT[sel][-1], Y.shape[1], 0])
    b.axvline(0, color="k", lw=0.8, ls="--")
    b.set_xlabel("time from pin prick (s)"); b.set_ylabel(f"neurons ({Y.shape[1]}), sorted by response")
    b.set_title("B  Every neuron's pin-prick response\n(baseline sessions, 4 mice)", loc="left", fontsize=11.5)
    cb = fig.colorbar(im, ax=b, fraction=0.05, pad=0.02); cb.set_label("z − mean 0.5 s before", fontsize=10.5); cb.ax.tick_params(labelsize=9.5)
    side = np.array([{"activated": 1, "inhibited": -1}.get(x, 0) for x in C[order]])
    for i, s_ in enumerate(side):
        if s_:
            b.plot([4.08, 4.08], [i, i + 1], color=UPC if s_ > 0 else DNC, lw=3, clip_on=False, solid_capstyle="butt")
    b.set_xlim(-2, 4.15)
    NOTE["heat_n"] = int(Y.shape[1])
    # C % UP / DOWN for every event, baseline sessions
    c = fig.add_subplot(gs[2])
    EV = [("Light touch", "light"), ("Mild touch", "mild"), ("Pin prick", "pin"), ("Heat", "heat"), ("withdrawal", "with-\ndrawal"),
          ("flinch", "flinch"), ("movement onset", "move-\nment")]
    bb = cc[cc.phase == "pre"]
    up = [100 * (bb[bb.event == e].cls == "activated").mean() for e, _ in EV]
    dn = [100 * (bb[bb.event == e].cls == "inhibited").mean() for e, _ in EV]
    x = np.arange(len(EV))
    c.bar(x - 0.2, up, width=0.4, color=UPC, label="UP"); c.bar(x + 0.2, dn, width=0.4, color=DNC, label="DOWN")
    c.axvline(3.5, color="0.6", lw=0.8, ls=":")
    c.text(1.5, 27.5, "stimuli", ha="center", fontsize=11, color="0.3"); c.text(5, 27.5, "behaviours", ha="center", fontsize=11, color="0.3")
    c.set_ylim(0, 30)
    c.set_xticks(x, [l for _, l in EV], fontsize=10.5)
    c.set_ylabel("% of neurons (baseline sessions)")
    c.legend(frameon=False, fontsize=10.5, loc="upper left", bbox_to_anchor=(0.0, 0.93))
    c.set_title("C  Neurons UP / DOWN for each event\n(baseline sessions, 462 neuron-sessions)", loc="left", fontsize=11.5)
    NOTE["pct"] = {l.replace("\n", ""): (round(u, 1), round(d_, 1)) for (_, l), u, d_ in zip(EV, up, dn)}
    fig.canvas.draw()
    pos = b.get_position()
    arr = (0.5 + (pos.x0 + 0.03) * 12.3, 1.55 + (1 - pos.y1 + 0.05) * 4.6 * 12.3 / 12.3)
    fig.savefig(fn); plt.close(fig)
    return arr


# ============================================================================================ figure 2: SBI-553
def fig2(fn):
    fig = plt.figure(figsize=(12.3, 4.6))
    gs = fig.add_gridspec(1, 3, width_ratios=[1.0, 2.15, 0.95], wspace=0.42, left=0.07, right=0.99, top=0.78, bottom=0.17)
    T = TRK.drop_duplicates(["mouse", "day1_id", "day2_id"])
    W = SEGW.set_index(["mouse", "day", "phase", "cell"])["blocks"]
    # A stimulus-block activity per mouse, tracked neurons
    a = fig.add_subplot(gs[0])
    lev = {c: {} for c in ("vb", "vi", "sb", "si")}
    for m, g in T.groupby("mouse"):
        for c, dk, ph, ik in (("vb", "veh_day", "pre", "veh_id"), ("vi", "veh_day", "post", "veh_id"), ("sb", "sbi_day", "pre", "sbi_id"),
                              ("si", "sbi_day", "post", "sbi_id")):
            vals = [W.get((m, int(r[dk]), ph, int(r[ik])), np.nan) for _, r in g.iterrows()]
            lev[c][m] = float(np.nanmean(vals))
    X = [0, 1, 2.4, 3.4]
    for x, c in zip(X, ("vb", "vi", "sb", "si")):
        a.bar(x, np.mean(list(lev[c].values())), color=CC[c], width=0.72, zorder=2)
    for m_ in MICE:
        for xs, cs in ((X[:2], ("vb", "vi")), (X[2:], ("sb", "si"))):
            a.plot(xs, [lev[cs[0]][m_], lev[cs[1]][m_]], color=MC[m_], lw=1.2, zorder=3)
            for x_, c_ in zip(xs, cs):
                a.scatter(x_, lev[c_][m_], s=34, marker=MK[m_], color=MC[m_], edgecolor="white", lw=0.6, zorder=5)
    a.set_ylim(0, 1.45)
    D = tracked_change()
    p = wilcoxon(D["sbi_blocks"], D["veh_blocks"])[1]
    lower = sum(D[D.mouse == m_]["sbi_blocks"].mean() < D[D.mouse == m_]["veh_blocks"].mean() for m_ in MICE)
    a.plot([0.5, 0.5, 2.9, 2.9], [1.22, 1.26, 1.26, 1.22], color="k", lw=1.1)
    a.text(1.7, 1.27, f"bigger drop on the SBI-553 day\n{st(p)} ({pf(p)}), {lower}/4 mice", ha="center", va="bottom", fontsize=10.5,
           fontweight="bold")
    a.set_xticks(X, ["base", "after", "base", "after"])
    a.text(0.5, -0.17, "vehicle day", transform=a.get_xaxis_transform(), ha="center", va="top", color=CC["vi"], fontsize=11)
    a.text(2.9, -0.17, "SBI-553 day", transform=a.get_xaxis_transform(), ha="center", va="top", color=CC["si"], fontsize=11)
    a.set_ylabel("activity during stimulus blocks\n(mean z of the tracked neurons)")
    a.set_title("A  Activity, same 185 neurons", loc="left", fontsize=11.5)
    NOTE["levels"] = {c: {m_: round(v, 3) for m_, v in lev[c].items()} for c in lev}
    NOTE["block_change"] = dict(veh=round(D.veh_blocks.mean(), 3), sbi=round(D.sbi_blocks.mean(), 3), p=p, lower=int(lower))
    # B heat maps: pin-prick response of the same neurons in the 4 sessions
    sub = gs[1].subgridspec(1, 5, width_ratios=[1, 1, 1, 1, 0.07], wspace=0.08)
    t = TRK[TRK.event == "Pin prick"].drop_duplicates(["mouse", "day1_id", "day2_id"])
    M = {c: [] for c in ("vb", "vi", "sb", "si")}
    rows_mouse = []
    for m_ in MICE:
        g = t[t.mouse == m_]
        blocks = {}
        for c, dk, ph, ik in (("vb", "veh_day", "pre", "veh_id"), ("vi", "veh_day", "post", "veh_id"), ("sb", "sbi_day", "pre", "sbi_id"),
                              ("si", "sbi_day", "post", "sbi_id")):
            blocks[c] = np.column_stack([trace(m_, int(r[dk]), ph, "Pin prick", int(r[ik])) for _, r in g.iterrows()])
        k = resp_window("Pin prick")
        o = np.argsort(-np.nanmean(blocks["vi"][k], 0))
        for c in M:
            M[c].append(blocks[c][:, o])
        rows_mouse.append((m_, g.shape[0]))
    sel = (TT >= -2) & (TT <= 4)
    for i, c in enumerate(("vb", "vi", "sb", "si")):
        ax_ = fig.add_subplot(sub[i])
        Y = np.hstack(M[c])[sel]
        im = ax_.imshow(Y.T, aspect="auto", cmap="RdBu_r", norm=TwoSlopeNorm(0, -1.0, 1.0), interpolation="nearest",
                        extent=[TT[sel][0], TT[sel][-1], Y.shape[1], 0])
        ax_.axvline(0, color="k", lw=0.7, ls="--")
        y0 = 0
        for m_, nr in rows_mouse:
            if y0:
                ax_.axhline(y0, color="k", lw=0.9)
            if i == 0:
                ax_.text(-2.15, y0 + nr / 2, m_, ha="right", va="center", fontsize=10, color=MC[m_], fontweight="bold")
            y0 += nr
        ax_.set_title(CL[c].replace("baseline, ", "baseline\n").replace("after ", "after\n"), fontsize=10.5,
                      color=CC[c] if c in ("vi", "si") else "0.35")
        ax_.set_yticks([])
        ax_.set_xticks([0, 2]); ax_.tick_params(labelsize=10)
        if i == 0:
            ax_.set_ylabel("same 185 neurons (grouped by mouse)", labelpad=34)
        ax_.set_xlabel("s from pin" if i else "s from pin prick", fontsize=10)
        if i == 1:
            hb = ax_
    cax = fig.add_subplot(sub[4])
    cb = fig.colorbar(im, cax=cax); cb.set_label("z − mean 0.5 s before", fontsize=10); cb.ax.tick_params(labelsize=9)
    fig.text(0.315, 0.875, "B  Pin-prick response of each neuron in the 4 sessions (rows sorted by the after-vehicle response)", fontsize=11.5)
    # C behaviour: pin flinch
    c_ = fig.add_subplot(gs[2])
    pm = {c: beh_pct("flinch", "Pin prick", c)[0] for c in ("vb", "vi", "sb", "si")}
    for x, c in zip(X, ("vb", "vi", "sb", "si")):
        c_.bar(x, np.mean(list(pm[c].values())), color=CC[c], width=0.72, zorder=2)
    for m_ in MICE:
        for xs, cs in ((X[:2], ("vb", "vi")), (X[2:], ("sb", "si"))):
            c_.plot(xs, [pm[cs[0]][m_], pm[cs[1]][m_]], color=MC[m_], lw=1.0, zorder=3)
    pa = cmh("flinch", "Pin prick", "SBI-553 vs vehicle")
    c_.set_ylim(0, 125)
    c_.plot([1, 1, 3.4, 3.4], [106, 110, 110, 106], color="k", lw=1.1)
    c_.text(2.2, 111, f"{st(pa)} ({pf(pa)})", ha="center", va="bottom", fontsize=10.5, fontweight="bold")
    c_.set_xticks(X, ["base", "after", "base", "after"])
    c_.text(0.5, -0.17, "vehicle", transform=c_.get_xaxis_transform(), ha="center", va="top", color=CC["vi"], fontsize=11)
    c_.text(2.9, -0.17, "SBI-553", transform=c_.get_xaxis_transform(), ha="center", va="top", color=CC["si"], fontsize=11)
    c_.set_ylabel("pin pricks with a flinch (%)")
    c_.set_title("C  Behaviour", loc="left", fontsize=11.5)
    NOTE["flinch"] = {c: round(float(np.mean(list(pm[c].values()))), 1) for c in pm}; NOTE["flinch_p"] = pa
    t2 = t.copy()
    NOTE["pinUP"] = {c: {m_: int(t2[t2.mouse == m_][f"cls_{c}"].eq("activated").sum()) for m_ in MICE} for c in ("vb", "vi", "sb", "si")}
    fig.legend(handles=[Line2D([], [], ls="", marker=MK[m_], color=MC[m_], markeredgecolor="white", markersize=8,
                               label=f"{m_} ({SEX[m_]}, SBI-553 Day {SBIDAY[m_]})") for m_ in MICE], ncol=4, loc="upper center",
               bbox_to_anchor=(0.5, 1.0), frameon=False, fontsize=11, handlelength=1.0, columnspacing=1.6)
    fig.canvas.draw()
    p1 = hb.get_position()
    fig.savefig(fn); plt.close(fig)
    return dict(A=(0.5 + 0.30 * 12.3, 1.55 + (1 - 0.42) * 4.6), B=(0.5 + p1.x0 * 12.3, 1.55 + (1 - p1.y1) * 4.6))


F1, F2 = os.path.join(OUTD, "fig1_neurons_respond_to_pain.png"), os.path.join(OUTD, "fig2_sbi553.png")
A1 = fig1(F1)
A2 = fig2(F2)
json.dump(NOTE, open(os.path.join(OUTD, "numbers.json"), "w"), indent=1, default=float)
print(json.dumps(NOTE, default=lambda x: round(float(x), 4)))

# ============================================================================================ slides
from lxml import etree  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402

RED = (0xC0, 0x39, 0x2B)


def tb(s, x, y, w, h, text, size, bold=False, color=(0x22, 0x22, 0x22), italic=False, align=None):
    t_ = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t_.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p_ = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p_.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = "Arial"; r.font.color.rgb = RGBColor(*color)
        if align:
            p_.alignment = align
    return t_


def arrow(s, x0, y0, x1, y1):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = RGBColor(*RED); c.line.width = Pt(3.5)
    te = etree.SubElement(c.line._get_or_add_ln(), "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")


FOOT = "CeA-Ntsr1 mini1p (Ca imaging)  ·  SBI-553 vs vehicle crossover  ·  n = 4 mice  ·  exploratory, uncorrected"
pc = NOTE["pct"]
bc = NOTE["block_change"]
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

s = prs.slides.add_slide(prs.slide_layouts[6])
tb(s, 0.45, 0.12, 12.4, 0.3, "CeA-Ntsr1 mini1p  ·  1 / 2: what the neurons do", 12, color=(0x55, 0x55, 0x55))
tb(s, 0.45, 0.36, 12.4, 0.6, "CeA-Ntsr1 neurons respond to pain within 0.5 s, some UP and some DOWN", 26)
tb(s, 0.45, 0.98, 12.4, 0.45, f"Pin prick recruited the most neurons of the 4 stimuli ({pc['pin'][0]:.0f}% UP, {pc['pin'][1]:.0f}% DOWN); "
   f"movement onset even more ({pc['move-ment'][0]:.0f}% / {pc['move-ment'][1]:.0f}%)", 15, bold=True, color=RED)
s.shapes.add_picture(F1, Inches(0.5), Inches(1.55), Inches(12.3))
tb(s, 0.45, 6.25, 12.4, 0.75,
   "UP / DOWN = response window (pin 0-0.5 s; others 0-1 s) vs the time just before, tested against 1000 time-shifted event trains, "
   "BH q ≤ 0.20. A: outlines on the max ΔF/F projection, red = UP, blue = DOWN. B: one row = one neuron (bar at right = UP / DOWN).", 12)
tb(s, 0.45, 7.1, 11.6, 0.3, FOOT, 9.5, color=(0x88, 0x88, 0x88), italic=True)
tb(s, 12.2, 7.1, 0.9, 0.3, "1 / 2", 10, color=(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT)
arrow(s, 11.25, 3.55, 10.78, 4.02)                                   # panel C: pin-prick bars
s.notes_slide.notes_text_frame.text = (
    f"FOV 5941 Day 1 baseline, pin prick: {NOTE['fov']}. Heat map: {NOTE['heat_n']} neuron-sessions (baseline sessions of both days, 4 "
    f"mice), trial-mean z minus its mean in the 0.5 s before. % UP / DOWN in baseline sessions: {pc}. Pin UP / DOWN per mouse "
    "(baseline sessions, both days): 5941 32 / 30 of 127, 1191 18 / 14 of 171, 5939 2 / 0 of 116, 5937 1 / 0 of 48, so most responders "
    "are in 5941 and 1191. Same pattern with baseline z (-2..0 s) and dF/F baselines (analysis\\baseline_methods).")

s = prs.slides.add_slide(prs.slide_layouts[6])
tb(s, 0.45, 0.12, 12.4, 0.3, "CeA-Ntsr1 mini1p  ·  2 / 2: SBI-553 vs vehicle (crossover, same neurons)", 12, color=(0x55, 0x55, 0x55))
tb(s, 0.45, 0.36, 12.4, 0.6, "SBI-553 lowered CeA-Ntsr1 activity during stimuli, and pin flinches", 26)
tb(s, 0.45, 0.98, 12.4, 0.45, f"Same 185 neurons: activity during stimuli dropped more on the SBI-553 day ({bc['sbi']:.2f} vs {bc['veh']:.2f} z, "
   f"{bc['lower']}/4 mice); pin flinch {NOTE['flinch']['si']:.0f} vs {NOTE['flinch']['vi']:.0f}%", 15, bold=True, color=RED)
s.shapes.add_picture(F2, Inches(0.5), Inches(1.55), Inches(12.3))
tb(s, 0.45, 6.2, 12.4, 0.4, "A: lines = mice, bars = mean; paired Wilcoxon on each neuron's change from its own baseline session. "
   "B: same neuron order in all 4 panels. C: CMH over deliveries.", 12)
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(6.58), Inches(12.4), Inches(0.46))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xFE, 0xF9, 0xE7); box.line.color.rgb = RGBColor(0xF3, 0x9C, 0x12); box.shadow.inherit = False
tb(s, 0.6, 6.6, 12.1, 0.42, "Not yet firm: the extra pin UP neurons after vehicle are almost all in 5939 (B, 2 → 19); responses normalized "
   "to each day's baseline did not differ by day. 4 mice, exploratory.", 12, color=(0x33, 0x33, 0x33))
tb(s, 0.45, 7.1, 11.6, 0.3, FOOT, 9.5, color=(0x88, 0x88, 0x88), italic=True)
tb(s, 12.2, 7.1, 0.9, 0.3, "2 / 2", 10, color=(0x88, 0x88, 0x88), align=PP_ALIGN.RIGHT)
arrow(s, 4.05, 3.35, 3.48, 3.88)                                     # panel A: after-SBI-553 bar
s.notes_slide.notes_text_frame.text = (
    f"A: stimulus-block mean z of the tracked neurons per mouse (vb / vi / sb / si): {NOTE['levels']}. Change from the same day's baseline "
    f"session, mean of 185 neurons: vehicle day {bc['veh']:+.3f}, SBI-553 day {bc['sbi']:+.3f} z, paired Wilcoxon {pf(bc['p'])}; lower on "
    f"the SBI-553 day in {bc['lower']}/4 mice (both drug orders). No-stimulus periods did not change (p = 0.58). B: pin UP counts of the "
    f"185 neurons per mouse {NOTE['pinUP']}. C: pin flinch % (mean of mice) {NOTE['flinch']}, after vs after CMH {pf(NOTE['flinch_p'])}; "
    "lower after SBI-553 in 3 mice, equal in 5939. Day order: every mouse had more pin UP neurons after its Day-2 injection.")
OUT = os.path.join(OUTD, "mini1p_SBI553_PI_2slides_Ca.pptx")
prs.save(OUT)
print("->", OUT)
